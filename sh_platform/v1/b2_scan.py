"""B2 scan: per KPI of the signed SOW, are the fields its catalogue rows need present in a landing file?

available = every needed field found in a landing file with at least one data row; partly = some; not = none.
Writes `build/scan/<iid>.json` per V1-BUILD.md. Headers and row counts only: no model call, nothing guessed."""
from __future__ import annotations

import datetime as dt
import logging
import re
from pathlib import Path

from sh_platform.jsonfile import lock_for, read_json_strict, write_json_atomic
from sh_platform.warehouse import _read_tabular  # the landing reader the bronze loader already uses

log = logging.getLogger(__name__)
_IID = re.compile(r"[A-Za-z0-9][A-Za-z0-9_-]{0,79}")


class ScanError(ValueError):
    def __init__(self, status: int, code: str, message: str):
        super().__init__(message)
        self.status, self.code, self.message = status, code, message


def key(s) -> str:
    return re.sub(r"[^a-z0-9]", "", str(s or "").lower())


def now() -> str:
    return dt.datetime.now(dt.timezone.utc).isoformat(timespec="seconds")


def scan_file(plane, iid: str) -> Path:
    if not isinstance(iid, str) or not _IID.fullmatch(iid):
        raise ScanError(404, "not_found", "There is no such intent.")
    return plane.path("build", "scan", f"{iid}.json")


def read_sow(plane, iid: str) -> dict:
    scan_file(plane, iid)  # validates the id
    sow = read_json_strict(plane.path("intents", iid, "sow.json"), dict)
    if sow is None:
        raise ScanError(404, "not_found", "This intent has no intent of work yet.")
    if sow.get("state") != "signed":
        raise ScanError(409, "not_signed", "Sign the intent of work first; the scan reads the signed version.")
    return sow


def read_scan(plane, iid: str) -> dict | None:
    return read_json_strict(scan_file(plane, iid), dict)


# ---- landing files: one entry per CSV and per XLSX sheet -------------------------------------------------------------

def _xlsx_sheets(p: Path) -> list[tuple[str, list[str], int]]:
    from openpyxl import load_workbook
    wb = load_workbook(p, read_only=True, data_only=True)
    try:
        out = []
        for ws in wb.worksheets:
            rows = [r for r in ws.iter_rows(values_only=True) if any(v not in (None, "") for v in r)]
            if rows:
                out.append((ws.title, [str(h).strip() for h in rows[0] if h not in (None, "")], len(rows) - 1))
        return out
    finally:
        wb.close()


def landing_files(plane) -> list[dict]:
    root = plane.path("landing")
    out = []
    for p in sorted(root.rglob("*")) if root.is_dir() else []:
        suf = p.suffix.lower()
        if not p.is_file() or suf not in (".csv", ".xlsx") or p.name.startswith("_b4_"):  # _b4_: Flow's staged copies
            continue
        rel = p.relative_to(root).as_posix()
        try:
            if suf == ".csv":
                header, rows = _read_tabular(p)
                sheets = [(None, header, sum(1 for r in rows if any(str(c).strip() for c in r)))]
            else:
                sheets = _xlsx_sheets(p)
        except (OSError, ValueError, UnicodeDecodeError):
            log.warning("landing file %s could not be read for the scan", rel, exc_info=True)
            continue
        for sheet, header, n in sheets:
            out.append({"path": f"{rel}#{sheet}" if sheet else rel, "folder": rel.split("/")[0] if "/" in rel else "",
                        "name": p.name, "columns": [h for h in header if h], "rows": n})
    return out


def _files_of(src: dict, files: list[dict]) -> list[dict]:
    """The landing files that belong to a source (folder or file stem matches its id or name, or its sample's file name)."""
    names = {key(src.get("id")), key(src.get("name"))} - {""}
    sample = Path(str(src.get("sample_path") or "")).name
    return [f for f in files if key(f["folder"]) in names or key(Path(f["name"]).stem) in names or (sample and f["name"] == sample)]


# ---- the scan ---------------------------------------------------------------------------------------------------------

def _needs(sow: dict, kpi: dict) -> list[tuple[str, str]]:
    """(source_id, field) pairs the KPI needs: its landing catalogue rows, else every catalogue row feeding it, else its sources' fields."""
    rows = [c for c in sow.get("catalogue") or [] if kpi.get("code") in (c.get("feeds_kpis") or [])]
    rows = [c for c in rows if c.get("layer") == "landing"] or rows
    pairs = [(c.get("source_id") or "", str(f.get("source_field") or f.get("name") or "").split(".")[-1])
             for c in rows for f in c.get("fields") or []]
    if not pairs:
        srcs = {s.get("id"): s for s in sow.get("sources") or []}
        pairs = [(sid, f.get("name") or "") for sid in kpi.get("source_ids") or [] for f in (srcs.get(sid) or {}).get("fields") or []]
    return list(dict.fromkeys((s, f) for s, f in pairs if f))


def scan_kpi(sow: dict, kpi: dict, files: list[dict]) -> dict:
    srcs = {s.get("id"): s for s in sow.get("sources") or []}
    used, missing, needs = [], [], _needs(sow, kpi)
    for sid, field in needs:
        src = srcs.get(sid) or {"id": sid}
        own = _files_of(src, files)
        found = any(key(field) in {key(c) for c in f["columns"]} and f["rows"] > 0 for f in own or files)
        if found:
            used.append(sid)
        else:
            where = ", ".join(f"landing/{f['path']}" for f in own) or (f"landing/ (the {src.get('name') or sid or 'source'} file)")
            missing.append({"what": f"{field} ({src.get('name') or sid or 'no source named'})", "where_expected": where})
    if not needs:
        missing.append({"what": "the fields this KPI needs: no catalogue row or source names them", "where_expected": "the intent of work catalogue"})
    status = "not" if not used else ("available" if not missing else "partly")
    return {"code": kpi.get("code"), "status": status, "sources_used": sorted({s for s in used if s}), "missing": missing}


def run_scan(plane, iid: str) -> dict:
    sow, path = read_sow(plane, iid), scan_file(plane, iid)
    files = landing_files(plane)
    rec = {"intent_id": iid, "at": now(), "kpis": [scan_kpi(sow, k, files) for k in sow.get("kpis") or [] if k.get("code")], "agreed": None}
    with lock_for(path):
        write_json_atomic(path, rec)
    return rec


def agree(plane, iid: str, scope, by: str) -> dict:
    path = scan_file(plane, iid)
    if not isinstance(scope, list) or not scope or not all(isinstance(c, str) for c in scope):
        raise ScanError(422, "invalid", "Name at least one KPI code to agree.")
    with lock_for(path):
        rec = read_scan(plane, iid)
        if rec is None:
            raise ScanError(409, "no_scan", "Run the scan first.")
        status = {k["code"]: k["status"] for k in rec["kpis"]}
        bad = [c for c in scope if status.get(c) in (None, "not")]
        if bad:
            raise ScanError(422, "not_buildable", f"These cannot be built from the files in hand: {', '.join(bad)}.")
        rec["agreed"] = {"by": by, "at": now(), "scope": list(dict.fromkeys(scope))}
        write_json_atomic(path, rec)
    return rec


def send_back(plane, iid: str) -> dict:
    path = scan_file(plane, iid)
    with lock_for(path):
        rec = read_scan(plane, iid)
        if rec is None:
            raise ScanError(409, "no_scan", "Run the scan first.")
        rec["agreed"] = None
        write_json_atomic(path, rec)
    return rec
