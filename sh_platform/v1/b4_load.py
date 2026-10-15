"""B4 Flow loads: the mapping sheet's rows (build/mapping/<iid>.json) become pipeline contracts (SHAPES 3), one per source, run through the
built-in runner (sh_platform.build.runner.run_pipeline) into the company warehouse and logged in warehouse/runs.json. Flow writes a note per
run (sh_platform.build.agents.flow.explain_run). Failed checks and missing files are posted to the intent's room.

The runner reads one file from the landing root and takes its checks and report tables from the blueprint. A source here lands as many
monthly files under landing/<source>/, and its checks and report tables live in the mapping sheet, so this module
1. stacks every landed file of a source into landing/_b4_<source>.csv (headers united, so a renamed column shows up as empty cells),
2. runs the runner on that file, then
3. runs the mapping's checks the runner listed as "could not run" (it found no blueprint fields for them) and builds the mapping's gold
   tables, with the runner's own transforms, and writes the results back onto the same run record."""
from __future__ import annotations

import csv
import datetime as dt
import logging
import os
import re
from pathlib import Path

from sh_platform.build import contracts, runner, transforms
from sh_platform.build.contracts import NOT_YET, slug
from sh_platform.jsonfile import lock_for, write_json_atomic
from sh_platform.warehouse import open_company, refresh_readonly_copy

log = logging.getLogger(__name__)
IID = re.compile(r"[A-Za-z0-9][A-Za-z0-9_-]{0,79}")
PREFIX = "v1_"  # the runner creates or replaces silver.<source>: the prefix keeps the warehouse's existing tables untouched
TIMES = {"expected": "WD5", "sla": "WD8"}  # month-end files; a guess, so every contract here is marked "guessed"
STAGED = "_b4_"


def mapping(plane, iid: str) -> dict:
    return contracts._read(plane, "build", "mapping", f"{iid}.json") if IID.fullmatch(iid or "") else {}


def contracts_for(plane, iid: str) -> tuple[list[dict], int]:
    """(contracts, rows left out). A row that is not "mapped" or names no source is left out. Each contract also carries the mapping's
    source id and table (to find the files), the field each check reads (`fields_of`) and the gold measures (`gold`)."""
    sow = contracts._read(plane, "intents", iid, "sow.json")
    named = {s.get("id"): s for s in sow.get("sources") or [] if isinstance(s, dict)}
    out: dict[str, dict] = {}
    left = 0
    for r in mapping(plane, iid).get("rows") or []:
        sid = str(r.get("source_id") or "")
        if r.get("status") != "mapped" or not slug(sid):
            left += 1
            continue
        s = named.get(sid, {})
        src = PREFIX + slug(sid) + (f"__{slug(str(r.get('source_table') or ''))}" if len(s.get("paths") or []) > 1 and r.get("source_table") else "")
        c = out.setdefault(src, {"source": src, "title": s.get("name") or sid, "lands_at": "landing", "pattern": f"{STAGED}{src}.csv", **TIMES,
                                 "schema": [], "rules": [], "depends_on": [], "owner": s.get("owner") or NOT_YET, "marts": [], "mark": "guessed",
                                 "source_id": sid, "source_table": r.get("source_table") or "", "fields_of": {}, "gold": {}})
        fields = [x.strip() for x in str(r.get("source_field") or "").split(",") if x.strip()]
        field = fields[0] if fields else ""
        for x in fields:
            if x not in [y["field"] for y in c["schema"]]:
                c["schema"].append({"field": x, "type": "auto"})
        for text in (r.get("test"), r.get("rule")):
            if field and text and transforms._kind(str(text).lower()):
                rid = f"{r.get('target_field') or field}:{text}"
                if rid not in c["rules"]:
                    c["rules"].append(rid)
                c["fields_of"][rid] = field
        op = (r.get("op") if r.get("op") in ("sum", "latest") else None) or transforms.op_of(str(r.get("transformation") or ""))
        table = slug(str(r.get("target_table") or "").removeprefix("gold."))
        if r.get("layer") == "gold" and op and field and table:
            c["gold"].setdefault(table, []).append((slug(r.get("target_field") or field), field, op))
            if f"gold.{table}" not in c["marts"]:
                c["marts"].append(f"gold.{table}")
    return sorted(out.values(), key=lambda c: c["source"]), left


# ---- the landed files -----------------------------------------------------------------------------------------------------------------

def files_of(plane, c: dict) -> list[Path]:
    """Every CSV or XLSX landed for a source: under a landing folder named like the source (or its table), or at the landing root with
    a name that starts like it. Oldest first."""
    land = plane.area("landing")
    keys = {k for k in (slug(c["source_id"]), slug(c["source_table"])) if k}
    if not land.is_dir() or not keys:
        return []
    found = [p for d in land.iterdir() if d.is_dir() and slug(d.name) in keys for p in d.rglob("*")]
    found += [p for p in land.iterdir() if p.is_file() and not p.name.startswith(STAGED) and any(slug(p.stem).startswith(k) for k in keys)]
    return sorted({p for p in found if p.is_file() and p.suffix.lower() in (".csv", ".xlsx")}, key=lambda p: (p.stat().st_mtime, p.name))


def _cell(v) -> str:
    if v is None:
        return ""
    if isinstance(v, dt.datetime):
        return v.date().isoformat() if v.time() == dt.time() else v.isoformat(sep=" ")
    if isinstance(v, dt.date):
        return v.isoformat()
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v)


def _table(p: Path, sheet: str) -> list[list[str]]:
    """Header plus rows as text. An XLSX gives the sheet named like the source table, else its first sheet."""
    if p.suffix.lower() == ".csv":
        with p.open(newline="", encoding="utf-8-sig") as fh:
            return [r for r in csv.reader(fh) if any(x.strip() for x in r)]
    from openpyxl import load_workbook
    wb = load_workbook(p, read_only=True, data_only=True)
    try:
        ws = next((w for w in wb.worksheets if slug(sheet) and slug(w.title) == slug(sheet)), wb.worksheets[0])
        return [[_cell(v) for v in r] for r in ws.iter_rows(values_only=True) if any(v not in (None, "") for v in r)]
    finally:
        wb.close()


def stage(plane, c: dict) -> Path | None:
    """landing/_b4_<source>.csv: every landed file of the source stacked under one header (the union of theirs, first seen first). Its
    time is the newest landed file's, so the runner's on-time check reads the real arrival. None when nothing has landed."""
    files = files_of(plane, c)
    if not files:
        return None
    header: list[str] = []
    parts = []
    for p in files:
        t = _table(p, c["source_table"])
        if not t:
            continue
        h = [str(x).strip() for x in t[0]]
        header += [x for x in h if x and x not in header]
        parts.append((h, t[1:]))
    if not header:
        return None
    dest = plane.path("landing", f"{STAGED}{c['source']}.csv")
    with dest.open("w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh)
        w.writerow(header)
        for h, rows in parts:
            for r in rows:
                d = dict(zip(h, r))
                w.writerow([d.get(x, "") for x in header])
    newest = max(p.stat().st_mtime for p in files)
    os.utime(dest, (newest, newest))
    return dest


# ---- one run ----------------------------------------------------------------------------------------------------------------------------

def _finish(plane, c: dict, run: dict) -> dict:
    """The mapping's checks the runner could not run (no blueprint fields) and the mapping's gold tables, onto the same run record."""
    src, rid_run = c["source"], run["id"]
    con = open_company(plane)
    try:
        typed = [(n, str(t).lower()) for n, t in con.execute("select column_name, data_type from information_schema.columns where table_schema = 'silver' "
                                                              "and table_name = ? order by ordinal_position", [src]).fetchall()]
        rules, held, still = list(run.get("rules") or []), list(run.get("quarantined") or []), []
        for rid in run.get("not_run") or []:
            f = c["fields_of"].get(rid)
            res = transforms.check_rule(con, src, rid.partition(":")[2], [f], [f], typed) if f else None
            if res is None:
                still.append(rid)
                continue
            rules.append({"rule": rid, "passed": res[0], "count": res[1]})
            if not res[0]:  # every mapping check blocks: its rows are set aside like the runner's own
                held += transforms.quarantine(con, src, rid, rid.partition(":")[2], [f], [f], typed, rid_run)
        for table, ms in c["gold"].items():
            if not transforms.build_mart(con, src, table, ms, typed):
                rules.append({"rule": f"gold.{table}: left as it was, no period could be worked out ({NOT_YET})", "passed": False, "count": 0})
    finally:
        con.close()
    failed = run["status"] == "failed" or any(not r["passed"] for r in rules)
    run = {**run, "rules": rules, "not_run": still, "quarantined": held, "status": "failed" if failed else run["status"],
           "marts_ready": None if failed else run.get("marts_ready") or dt.datetime.now(dt.timezone.utc).isoformat(),
           "sla_met": None if failed else run.get("sla_met")}
    refresh_readonly_copy(plane)
    path = plane.path("warehouse", "runs.json")
    with lock_for(path):
        write_json_atomic(path, {"runs": [run if r.get("id") == rid_run else r for r in runner.read_runs(plane)]})
        try:  # the company map follows every write that changes what it shows
            from sh_platform.v1.m1_map import touch; touch(plane)
        except ImportError:
            pass
    return run


def load(plane, c: dict) -> dict:
    """Stage, run, finish. Raises FileNotFoundError when nothing has landed, runner.Blocked, or ValueError/OSError for a file that will
    not load (the message holds no row)."""
    if stage(plane, c) is None:
        raise FileNotFoundError(f"No file has landed for {c['title']} yet. It is expected under landing/{slug(c['source_id'])}/.")
    return _finish(plane, c, runner.run_pipeline(plane, c))
