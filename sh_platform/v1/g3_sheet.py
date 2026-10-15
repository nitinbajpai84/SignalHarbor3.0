"""G3 The benchmark sheet: a signed spreadsheet document (`intents/<iid>/benchmarks.json`, artefact kind `benchmarks`).

Rows sit at the lowest granularity the owner reports at: one row per KPI and per combination of the dimensions the owner names (period,
entity, account, ...), with the reported value, the expected value after the load, a tolerance, the source of the figure and who stated it.
Every figure is copied by this program: from the person's own words (P4's parser, clause by clause) or from their uploaded report through
X5's reading (`x5_extract._lines`, `parse`, `extract`: the model may name the line, the program copies the value and keeps the line).
A row without a reported value is `needs figure` and blocks signing. The raiser signs, and the owner of any KPI that is not the raiser's.
The old per-KPI `benchmark` on the SOW becomes a pointer: {"sheet": true, "rows": [row ids], ...the first row's figures}.
Flow's check after the load (p4_benchmarks.verify) compares gold with each row at the same granularity (`got_by`).

Record: {intent_id, version, state: draft|signed, rows: [{id, code, kpi, dims, reported, expected, tolerance, source, source_doc, line,
stated_by, proposed, dropped, status}], signature, signatures: [{name, role, at, version}], updated_at}.
"""
from __future__ import annotations

import io
import logging
import re
import secrets
from datetime import date, datetime, timezone

from sh_platform.jsonfile import lock_for, read_json_strict, write_json_atomic
from sh_platform.v1 import p4_benchmarks as P4
from sh_platform.v1 import wh

log = logging.getLogger(__name__)
NEEDS = "needs figure"
PERIOD_WORDS = re.compile(r"^(?:period|month|date|as at|as of|year|time)$", re.I)
BY = re.compile(r"\bby\s+([a-z]+(?:\s[a-z]+)?)(?=\s+by\b|\s+for\b|[,.;:\n]|$)", re.I)
FILLER = re.compile(r"\b(?:for|was|were|is|are|reported|report|we|the|at|of|should|read|reads|after|load|expected|expect|figure|figures|number|numbers|"
                    r"in|on|and|a|an|i|my|our|last|month|period|within|about)\b", re.I)
EDITABLE = ("reported", "expected", "tolerance")


class SheetError(Exception):
    def __init__(self, status: int, message: str):
        super().__init__(message)
        self.status, self.message = status, message


def _now() -> str:
    return datetime.now(timezone.utc).isoformat(timespec="seconds")


def _norm(v) -> str:
    return " ".join(str(v or "").lower().replace("_", " ").split())


def _path(plane, iid: str):
    return plane.path("intents", iid, "benchmarks.json")


def read(plane, iid: str) -> dict:
    try:
        rec = read_json_strict(_path(plane, iid), dict)
    except (OSError, ValueError):
        raise SheetError(409, "The benchmark sheet could not be read. Please ask your administrator.") from None
    return rec or {"intent_id": iid, "version": 1, "state": "draft", "rows": [], "signature": None, "signatures": [], "updated_at": None}


def _write(plane, iid: str, sheet: dict) -> None:
    for r in sheet["rows"]:
        r["status"] = "proposed" if r.get("proposed") else NEEDS if r.get("reported") is None else "ready"
    sheet["updated_at"] = _now()
    _path(plane, iid).parent.mkdir(parents=True, exist_ok=True)
    write_json_atomic(_path(plane, iid), sheet)


def label(row: dict) -> str:
    dims = [v for d, v in (row.get("dims") or {}).items() if d != "period" and v]
    return f"{row.get('kpi') or row.get('code')}" + (f" ({', '.join(dims)})" if dims else "") + f", {(row.get('dims') or {}).get('period') or ''}".rstrip(", ")


def _key(r: dict) -> tuple:
    return r.get("code"), tuple(sorted((d, _norm(v)) for d, v in (r.get("dims") or {}).items()))


def upsert(plane, iid: str, new: list[dict], *, reopen: bool = True) -> dict:
    """Add or replace rows (same KPI and dimensions = same row). A signed sheet that changes is reopened: version + 1, signature cleared,
    the history of signatures stays."""
    with lock_for(_path(plane, iid)):
        sheet = read(plane, iid)
        have = {_key(r): r for r in sheet["rows"]}
        changed = False
        for n in new:
            old = have.get(_key(n))
            if old is None:
                row = {"id": "br_" + secrets.token_hex(4), "proposed": False, "dropped": False, "source_doc": None, "line": None, **n}
                sheet["rows"].append(row)
                have[_key(row)] = row
                changed = True
            else:
                if any(old.get(f) != n.get(f) for f in (*EDITABLE, "source", "source_doc", "line")):
                    changed = True
                old.update({k: v for k, v in n.items() if k != "id"})
        if changed and sheet["state"] == "signed" and reopen:
            sheet.update(state="draft", signature=None, version=int(sheet.get("version") or 1) + 1)
        _write(plane, iid, sheet)
    return sheet


def set_figure(plane, iid: str, rid: str, body: dict, by: str) -> dict:
    """A person types a figure into a row of the sheet (their own words in a cell). Refused on a signed sheet."""
    with lock_for(_path(plane, iid)):
        sheet = read(plane, iid)
        if sheet["state"] == "signed":
            raise SheetError(409, "The benchmark sheet is signed. A change goes through a deviation.")
        row = next((r for r in sheet["rows"] if r.get("id") == rid), None)
        if row is None:
            raise SheetError(404, "There is no such row on the sheet.")
        for f in EDITABLE:
            if body.get(f) in (None, ""):
                continue
            try:
                row[f] = float(str(body[f]).replace(",", ""))
            except ValueError:
                raise SheetError(400, f"{f.capitalize()} must be a number.") from None
        if row.get("reported") is not None:
            row["expected"] = row["expected"] if row.get("expected") is not None else row["reported"]
            row["tolerance"] = row["tolerance"] if row.get("tolerance") is not None else _tol(row["expected"])
            row.update(source=f"said by {by}", stated_by=by, source_doc=None, line=None, proposed=False)
        _write(plane, iid, sheet)
    return sheet


def _tol(expected) -> float:
    return round(abs(expected) * P4.DEFAULT_TOLERANCE_PCT / 100, 6)


# ---- the SOW pointer ------------------------------------------------------------------------------------------------------------------

def apply_pointers(sow: dict, sheet: dict) -> None:
    """Each KPI's `benchmark` points at its rows on the sheet (and carries the first row's figures for readers of the old field)."""
    for k in sow.get("kpis") or []:
        rows = [r for r in sheet.get("rows") or [] if r.get("code") == k.get("code") and not r.get("dropped")]
        if rows:
            r0 = rows[0]
            k["benchmark"] = {"sheet": True, "rows": [r["id"] for r in rows], "period": (r0.get("dims") or {}).get("period"),
                              "reported": r0.get("reported"), "expected": r0.get("expected"), "tolerance": r0.get("tolerance"),
                              "source": r0.get("source")}


def _sync_sow(plane, iid: str, sheet: dict) -> None:
    sp = plane.path("intents", iid, "sow.json")
    with lock_for(sp):
        try:
            sow = read_json_strict(sp, dict) or {}
        except (OSError, ValueError):
            return
        if sow and sow.get("state") != "signed":  # a signed SOW is frozen; the sheet is the truth either way
            apply_pointers(sow, sheet)
            write_json_atomic(sp, sow)


# ---- rows from the person's words ----------------------------------------------------------------------------------------------------

def _kpis(sow: dict) -> list[dict]:
    return [k for k in sow.get("kpis") or [] if isinstance(k, dict) and k.get("code")]


def _clean(words) -> list[str]:
    out = []
    for w in words if isinstance(words, list) else []:
        w = re.sub(r"[^a-z ]", "", str(w).lower()).strip()[:30]
        if w and w not in out:
            out.append(w)
    return out[:3]


def dims_named(seg: str, hint=None) -> list[str]:
    return _clean(list(hint or []) + [m.group(1) for m in BY.finditer(seg)])


def _dim_name(named: list[str]) -> str:
    return next((d for d in named if not PERIOD_WORDS.match(d)), "item")


def _label_of(clause: str, kname: str) -> str:
    s = re.sub(re.escape(kname), " ", clause, flags=re.I) if kname else clause
    s = P4.ISO.sub(" ", s)
    s = P4.MONTH.sub(" ", s)
    s = P4.TOL.sub(" ", s)
    s = re.sub(P4.NUM, " ", s)
    s = FILLER.sub(" ", re.sub(r"[%±:;()]|\+/-", " ", s))
    return " ".join(s.split()).strip(" -,.")[:60]


STATED = re.compile(r"expect|should (?:read|be)|within|tolerance|after the load|reads", re.I)


def _figures(line: str) -> list[str]:
    """One line is one row, unless it says two things: "A 12,400,000 and B 8,100,000" is two rows (a line that says what it expects stays one)."""
    if STATED.search(line) or P4.TOL.search(line):
        return [line]
    parts = re.split(r";|,(?=\s*[^\d\s])|\band\b", line)
    ok = len(parts) > 1 and all(re.search(P4.NUM, P4.ISO.sub(" ", p)) for p in parts)
    return parts if ok else [line]



def rows_from_text(kpi: dict, seg: str, by: str, named: list[str], today: date | None = None) -> list[dict]:
    """Rows for one KPI from its part of the person's words: one per line that holds a figure."""
    t = today or wh.today()
    explicit = P4.ISO.search(seg) or P4.MONTH.search(seg)
    seg_period = P4._period(seg, t)[0]
    out = []
    for clause in (c for line in re.split(r"\n+", seg) for c in _figures(line)):
        got = P4.parse(clause, by, t)
        if not got:
            continue
        has_p = P4.ISO.search(clause) or P4.MONTH.search(clause)
        period = got["period"] if has_p or not explicit else seg_period
        who = _label_of(clause, kpi.get("name") or "")
        dims = {"period": period}
        if who:
            dims[_dim_name(named)] = who
        out.append({"code": kpi["code"], "kpi": kpi.get("name"), "dims": dims, "reported": got["reported"], "expected": got["expected"],
                    "tolerance": got["tolerance"], "source": got["source"], "stated_by": by})
    return out


def from_talk(plane, iid: str, sow: dict, patch_kpis, text: str, by: str, asked: bool, today: date | None = None) -> None:
    """P4's reading of the person's words, then the sheet: figures they gave become rows (copied by the program), a dimension named with
    no figure becomes a `needs figure` row. `sow` is the caller's in-memory SOW; its KPIs' `benchmark` end as pointers."""
    kpis = _kpis(sow)
    before = {k["code"]: k.get("benchmark", "unset") for k in kpis}
    P4.from_talk(sow, patch_kpis, text, by, asked, today)
    hints = {P4._norm(i.get("name")): i.get("benchmark_by") for i in patch_kpis or [] if isinstance(i, dict)}
    t = today or wh.today()
    last = P4._period("", t)[0]
    new = []
    for k in kpis:
        b = k.get("benchmark", "unset")
        if isinstance(before[k["code"]], dict) and before[k["code"]].get("sheet") and b is None:
            k["benchmark"] = before[k["code"]]  # "none" for a KPI that already has rows on the sheet leaves them
            continue
        seg = P4.segment(kpis, k, text)
        named = dims_named(seg, hints.get(P4._norm(k.get("name")))) if (b is not before[k["code"]] or hints.get(P4._norm(k.get("name")))) else []
        if isinstance(b, dict) and not b.get("sheet") and b is not before[k["code"]]:
            new += rows_from_text(k, seg, by, named, t) or [{"code": k["code"], "kpi": k.get("name"), "dims": {"period": b["period"]}, "reported": b["reported"],
                                                           "expected": b["expected"], "tolerance": b["tolerance"], "source": b["source"], "stated_by": by}]
        elif named and not any(r.get("code") == k["code"] for r in read(plane, iid)["rows"]):
            new.append({"code": k["code"], "kpi": k.get("name"), "dims": {"period": last, **{_dim_name(named): ""}}, "reported": None,
                        "expected": None, "tolerance": None, "source": "", "stated_by": by})
    if new:
        sheet = upsert(plane, iid, new)
        apply_pointers(sow, sheet)


# ---- rows from an uploaded report (X5) ------------------------------------------------------------------------------------------------

def _per(cell: str) -> str | None:
    m = P4.ISO.search(cell)
    if m:
        return f"{m.group(1)}-{m.group(2)}"
    m = P4.MONTH.search(cell)
    if m and m.group(2):
        return f"{m.group(2)}-{P4.MONTHS[m.group(1).lower()]:02d}"
    return None


def _matches(header: str, kpi: dict) -> bool:
    h, n = _norm(header), _norm(kpi.get("name"))
    return bool(h and n and (h == n or n in h or (len(h) >= 4 and h in n)))


def _num(cell: str) -> float | None:
    from sh_platform.v1 import x5_extract as X
    got = X.parse("percent" if "%" in cell else "number", cell)
    try:
        return float(got) if got != "" else None
    except ValueError:
        return None


def _doc_row(kpi: dict, dims: dict, value: float, name: str, i: int, line: str, rel: str, by: str) -> dict:
    # the line's text is not kept: the viewer page shows it only to whoever may read the document
    return {"code": kpi["code"], "kpi": kpi.get("name"), "dims": dims, "reported": value, "expected": value, "tolerance": _tol(value),
            "source": f"{name}, line {i + 1}", "source_doc": rel, "line": i + 1, "stated_by": by}


def rows_from_lines(lines: list[str], kpis: list[dict], rel: str, by: str, model=None, plane=None, today: date | None = None) -> list[dict]:
    """Tables first (KPI as a column, or KPI as a row with periods across), then `label: value` lines through X5's locate. The value is
    always copied from the line by `x5_extract.parse`; the model only points at the line."""
    from sh_platform.v1 import x5_extract as X
    name = rel.rsplit("/", 1)[-1]
    doc_period = next((p for p in map(_per, lines) if p), None) or P4._period("", today or wh.today())[0]
    out, i = [], 0
    while i < len(lines):
        if not lines[i].startswith("|"):
            i += 1
            continue
        j = i
        while j < len(lines) and lines[j].startswith("|"):
            j += 1
        block = [[c.strip() for c in ln.strip("|").split("|")] for ln in lines[i:j]]
        head = block[0]
        cols = {}
        for c, h in enumerate(head):
            k = next((k for k in kpis if _matches(h, k)), None)
            if k:
                cols[c] = k
        if cols:  # KPI as a column; the other text columns are the dimensions
            for r, cells in enumerate(block[1:], 1):
                cells += [""] * (len(head) - len(cells))
                dims = {}
                for c, h in enumerate(head):
                    if c in cols or not cells[c] or _num(cells[c]) is not None and not _per(cells[c]):
                        continue
                    dims["period" if PERIOD_WORDS.match(h) or _per(cells[c]) else _norm(h)] = _per(cells[c]) or cells[c]
                dims.setdefault("period", doc_period)
                for c, k in cols.items():
                    v = _num(cells[c]) if cells[c] else None
                    if v is not None:
                        out.append(_doc_row(k, dict(dims), v, name, i + r, lines[i + r], rel, by))
        else:  # KPI as a row; the header cells after the first are periods (or entities)
            for r, cells in enumerate(block[1:], 1):
                k = next((k for k in kpis if _matches(cells[0], k)), None)
                if not k:
                    continue
                for c in range(1, min(len(cells), len(head))):
                    v = _num(cells[c]) if cells[c] else None
                    if v is None:
                        continue
                    per = _per(head[c])
                    dims = {"period": per or doc_period, **({} if per else {"item": head[c]})}
                    out.append(_doc_row(k, dims, v, name, i + r, lines[i + r], rel, by))
        i = j
    if not out:  # plain lines: "Net worth: 12,400,000" through X5's own locate, with the model's pointer when there is a model
        t = {"name": "benchmark report", "examples": [], "fields": [{"name": k["name"], "format": "number", "where_usually": k["name"], "meaning": ""}
                                                                    for k in kpis if k.get("name")]}
        got = X.extract(plane, lines, t, rel, model) if plane is not None else {}
        for k in kpis:
            g = got.get(k.get("name")) or {}
            if g.get("value") not in (None, "") and g.get("line") is not None:
                out.append(_doc_row(k, {"period": doc_period}, float(g["value"]), name, g["line"], lines[g["line"]], rel, by))
    return out


def from_document(plane, iid: str, rel: str, by: str, *, sow: dict | None = None, model=None) -> list[dict]:
    """Read last month's report from a document of the company's files into rows on the sheet. Raises SheetError when it cannot be read."""
    from sh_platform.v1 import x5_extract as X
    try:
        lines = X._lines(plane, rel)
    except X.Refused as e:
        raise SheetError(404, e.message) from None
    except Exception:  # noqa: BLE001 - an unreadable file is a plain sentence, not a crash
        raise SheetError(400, "That file could not be read.") from None
    if sow is None:
        sow = read_json_strict(plane.path("intents", iid, "sow.json"), dict) or {}
    rows = rows_from_lines(lines, _kpis(sow), rel, by, model, plane)
    if not rows:
        raise SheetError(409, "I could not find any of the KPI figures in that document. Type the figures, or name the line they are on.")
    sheet = upsert(plane, iid, rows)
    if sow.get("state") != "signed":
        _sync_sow(plane, iid, sheet)
    return rows


def from_files(plane, iid: str, sow: dict, files: list[dict], by: str, model=None) -> int:
    """Files dropped in the talk while the benchmarks are asked for: each readable one is read for the KPI figures. Count of rows read."""
    from sh_platform.v1 import x5_extract as X
    n = 0
    for f in files or []:
        if isinstance(f, dict) and f.get("path"):
            try:
                rows = rows_from_lines(X._lines(plane, f["path"]), _kpis(sow), f["path"], by, model, plane)
            except Exception:  # noqa: BLE001 - a file that is not a report is just a file
                continue
            if rows:
                apply_pointers(sow, upsert(plane, iid, rows))
                n += len(rows)
    return n


def missing(plane, iid: str) -> list[dict]:
    return [r for r in read(plane, iid)["rows"] if r.get("status") == NEEDS and not r.get("dropped")]


def ask_line(rows: list[dict]) -> str:
    """North's follow-up for rows with no figure: words only, no digits of its own."""
    names = list(dict.fromkeys(str(r.get("kpi") or r.get("code")) for r in rows))
    return (f"I have the way you report {', '.join(names)}, but not the figures yet. Give me last month's numbers for each, "
            "or drop last month's report in and I will read them from it.")


# ---- signing -------------------------------------------------------------------------------------------------------------------------

def _people(plane, iid: str) -> tuple[str | None, set[str]]:
    raiser, known = None, set()
    try:
        from sh_platform.v1 import p1_project as P1
        p = P1.project_of(plane, iid)
        if p:
            raiser = (p.get("people") or {}).get("raiser")
            known |= {n for n in (p.get("people") or {}).values() if n}
    except Exception:  # noqa: BLE001
        pass
    try:
        t = read_json_strict(plane.path("intents", iid, "thread.json"), dict) or {}
        known |= {x.get("name") for x in t.get("turns") or [] if isinstance(x, dict) and x.get("by") == "person" and x.get("name")}
    except (OSError, ValueError):
        pass
    return raiser, known


def required(plane, iid: str, sheet: dict | None = None) -> list[str]:
    """The raiser, and the owner of any KPI on the sheet that is not the raiser's (when the owner is a person we know)."""
    sheet = sheet or read(plane, iid)
    raiser, known = _people(plane, iid)
    sow = read_json_strict(plane.path("intents", iid, "sow.json"), dict) or {}
    onsheet = {r["code"] for r in sheet["rows"] if not r.get("proposed") and not r.get("dropped")}
    owners = {str(k.get("owner")).strip() for k in _kpis(sow) if k["code"] in onsheet and k.get("owner")}
    return [n for n in dict.fromkeys([raiser, *sorted(o for o in owners if o in known)]) if n]


def sign_block(plane, iid: str, by: str | None = None) -> str | None:
    sheet = read(plane, iid)
    gaps = missing(plane, iid)
    if gaps:
        return (f"{len(gaps)} row{'s have' if len(gaps) > 1 else ' has'} no reported figure yet ({', '.join(dict.fromkeys(label(r) for r in gaps))}). "
                "Give the figures or drop last month's report in before the sheet is signed.")
    need = required(plane, iid, sheet)
    if by is not None and not need:
        return "Nobody is named to sign this sheet yet; sign the intent of work first."
    if by is not None and by not in need:
        return f"Only {' and '.join(need)} can sign the benchmark sheet."
    return None


def sign(plane, iid: str, by: str) -> dict:
    with lock_for(_path(plane, iid)):
        sheet = read(plane, iid)
        if sheet["state"] == "signed":
            raise SheetError(409, "The benchmark sheet is already signed.")
        block = sign_block(plane, iid, by)
        if not block:
            try:
                from sh_platform.v1 import p1_project as P1
                block = P1.sign_block(plane, iid, "benchmarks", by)
            except ImportError:
                pass
        if block:
            raise SheetError(409, block)
        v = int(sheet.get("version") or 1)
        if not any(s.get("name") == by and s.get("version") == v for s in sheet["signatures"]):
            sheet["signatures"].append({"name": by, "role": "raiser" if by == (_people(plane, iid)[0]) else "owner", "at": _now(), "version": v})
        done = {s["name"] for s in sheet["signatures"] if s.get("version") == v}
        if set(required(plane, iid, sheet)) <= done:
            sheet.update(state="signed", signature=sheet["signatures"][-1])
        _write(plane, iid, sheet)
    _sync_sow(plane, iid, sheet)
    return sheet


# ---- the loaded figure at the row's own granularity -------------------------------------------------------------------------------------

def got_by(plane, sow: dict, code: str, period: str, dims: dict) -> tuple[float | None, str | None]:
    """(gold figure for the KPI in the period and the dimension values, why it cannot be read). A dimension with no column in the gold
    table cannot be compared: that is said, never guessed."""
    from sh_platform.v1.c1_cards import table_of
    from sh_platform.warehouse import open_readonly
    table = table_of(sow, code)
    if not table:
        return None, "the catalogue names no gold table for this figure"
    try:
        con = open_readonly(plane)
    except Exception:  # noqa: BLE001
        return None, "the warehouse could not be opened"
    try:
        cols = [c[0] for c in con.execute(f"describe gold.{table}").fetchall()]
        if len(cols) < 2:
            return None, "the gold table has no figure column"
        where, args = ['substr(cast("%s" as varchar), 1, 7) = ?' % cols[0]], [period]
        for d, v in dims.items():
            col = next((c for c in cols[1:-1] if _norm(c) == _norm(d) or _norm(d) in _norm(c)), None)
            if col is None:
                return None, f"the gold table has no {d} column to compare with"
            where.append('lower(trim(cast("%s" as varchar))) = ?' % col)
            args.append(_norm(v))
        row = con.execute(f'select sum(try_cast("{cols[-1]}" as double)), count(*) from gold.{table} where ' + " and ".join(where), args).fetchone()
    except Exception:  # noqa: BLE001 - duckdb.Error and friends
        return None, "the gold table could not be read"
    finally:
        con.close()
    return (float(row[0]), None) if row and row[1] and row[0] is not None else (None, "no loaded row matches this line")


# ---- the document ----------------------------------------------------------------------------------------------------------------------

def columns(sheet: dict) -> list[str]:
    dims = ["period"] + [d for d in dict.fromkeys(d for r in sheet["rows"] for d in (r.get("dims") or {})) if d != "period"]
    return dims


def lines(sheet: dict) -> list[str]:
    out = []
    for r in sheet["rows"]:
        if r.get("dropped"):
            continue
        d = ", ".join(f"{k} {v}" for k, v in (r.get("dims") or {}).items() if v)
        fig = ("needs figure" if r.get("reported") is None else
               f"reported {r['reported']:g}, expected {r['expected']:g} after the load, within {r['tolerance']:g}")
        out.append(f"**{r.get('kpi') or r.get('code')}** ({d}): {fig}" + (f". Source: {r['source']}" if r.get("source") else "")
                   + (f". Stated by {r['stated_by']}" if r.get("stated_by") else "") + (" (proposed)" if r.get("proposed") else ""))
    return out


def xlsx(sheet: dict) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    wb = Workbook()
    ws = wb.active
    ws.title = "Benchmarks"
    dims = columns(sheet)
    head = ["KPI", *[d.capitalize() for d in dims], "Reported", "Expected after load", "Tolerance", "Source", "Stated by", "Status"]
    ws.append(head)
    for c in range(1, len(head) + 1):
        ws.cell(1, c).font = Font(bold=True)
        ws.cell(1, c).fill = PatternFill("solid", fgColor="E8EEF4")
        ws.column_dimensions[ws.cell(1, c).column_letter].width = 22
    for r in sheet["rows"]:
        if r.get("dropped"):
            continue
        ws.append([r.get("kpi") or r.get("code"), *[(r.get("dims") or {}).get(d, "") for d in dims], r.get("reported"), r.get("expected"),
                   r.get("tolerance"), r.get("source") or "", r.get("stated_by") or "", r.get("status") or ""])
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def view(plane, iid: str, viewer_name: str | None = None) -> dict:
    sheet = read(plane, iid)
    return {"sheet": sheet, "columns": columns(sheet), "sign_block": sign_block(plane, iid, viewer_name), "required": required(plane, iid, sheet),
            "missing": [label(r) for r in missing(plane, iid)]}
