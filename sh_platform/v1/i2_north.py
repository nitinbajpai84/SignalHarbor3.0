"""I2 North's thread: one decision maker talks to North; North keeps `intents/<iid>/thread.json` and drafts `intents/<iid>/sow.json`.

North asks one question per unmet need (goal, process, inputs, kpis, rules, unclear), one follow-up at most, and updates the
intent of work after every turn. It never guesses: what is not known becomes an `unclear` entry. Pack KPIs it suggests are
marked `pack_suggested` and stay unclear until the decision maker settles them. Nothing here signs (I4 does, by a person).

Reused, not copied: North's role and rules (`sh_platform/intents/north.py` rules(), north_rules.yaml role/rules) and the document
rules of the private-room North (`modules/m1_intent/agent/north.py` load_rules()["documents_rules"]); the model call is the kernel
runtime (`sh_platform.build.agents.runtime.get_model`, `sh_platform.agents.loop.run_agent`: egress, number guard, run log).
The drafting view follows 2.0 `drafting_view` (a fixed list of fields North may see). With no model (SH_MODEL_PROVIDER=off) a
rule-based fallback asks a fixed question per need and parses the person's text; its turns are named "North (rules only)".
"""
from __future__ import annotations

import csv
import io
import json
import logging
import re
import secrets
from datetime import datetime, timezone

from sh_platform.agents.loop import run_agent
from sh_platform.agents.registry import ToolRegistry
from sh_platform.agents.runs import RunRegistry
from sh_platform.build.agents.runtime import OpenGate, register_names
from sh_platform.context_docs import reader
from sh_platform.intents import lifecycle as L
from sh_platform.intents import store as S
from sh_platform.jsonfile import lock_for, read_json_strict, write_json_atomic

log = logging.getLogger(__name__)

NEEDS = ("goal", "process", "inputs", "kpis", "rules", "unclear")
NORTH, FALLBACK_NAME = "North", "North (rules only)"
SAMPLE_EXT = {"csv", "tsv", "xlsx"}
MAX_TEXT = 4000
MAX_PACK = 3
QUESTIONS = {  # the fallback's fixed question per need; the model words its own
    "goal": "Why does this matter now, who will read the result, which decisions will it support, and by when do you need it?",
    "process": "Walk me through how this is done today, step by step: who does what, with which inputs, and what comes out of each step?",
    "inputs": "Which systems or files hold the data for this? Drop sample files or documents here if you have them.",
    "kpis": "Which figures do you want to see? For each, give a name and how it is worked out, if you know.",
    "rules": "Are there rules on this data: who may see it, how long it is kept, and who must approve it?",
    "unclear": "Is anything still unsure, or something only someone else can settle? Say \"nothing\" if not.",
}
FOLLOW_UPS = {
    "goal": "Who reads it, and by when do you need it?",
    "process": "Who does each step, and in which system?",
    "inputs": "Who owns each of those sources?",
    "kpis": "How is each of those figures worked out, and for which period?",
    "rules": "Who may see this data, and who approves it?",
    "unclear": "Who could settle that?",
}
READY = "The intent of work is ready: read it and sign when you are content."
DONT_KNOW = re.compile(r"\b(?:don'?t know|do not know|not sure|unsure|no idea)\b", re.I)
NOTHING = re.compile(r"^\s*(?:no|none|nothing|nope|n/a|nothing else|no,? that'?s all)\b[\s.!]*$", re.I)
_SPLIT = re.compile(r"(?:\r?\n)+|;\s*|(?<=[.!?])\s+")
_STOP = {"which", "their", "there", "about", "with", "from", "share", "total", "value", "versus", "known", "method"}

SOW_PROMPT = (
    "You are drafting an intent of work with the person who raised it, one need at a time. The need you asked about is in `need`; "
    "the next need is in `next_need` (null when this was the last). Read the person's latest message and reply with ONLY one JSON "
    "object {{\"section\": {{...}}, \"met\": true|false, \"unclear\": [{{\"question\", \"about\", \"who_can_settle\"}}], "
    "\"follow_up\": \"...\", \"next_question\": \"...\"}}. section holds only what the person said, for the need: {shape}. "
    "Never put a bare number outside a string. met is true when the need is answered well enough to write down. A point the person "
    "does not know, or that someone else must settle, goes in unclear, never guessed. follow_up is one short question when met is "
    "false, else \"\". next_question is one plain question for next_need, else \"\". No text outside the JSON."
)
SHAPES = {
    "goal": "{\"why\", \"readers\": [], \"decisions\": [], \"needed_by\"}",
    "process": "{\"steps\": [{\"name\", \"who\", \"does\", \"inputs\": [], \"outputs\": [], \"system\"}], "
               "\"handovers\": [{\"from_step\", \"to_step\", \"what\"}]} with step names as from_step and to_step",
    "inputs": "{\"sources\": [{\"name\", \"kind\": \"system\"|\"file\"|\"document\", \"holds\", \"owner\"}]}",
    "kpis": "{\"kpis\": [{\"name\", \"definition\", \"formula\", \"unit\", \"period\", \"owner\"}]} with \"\" for what they did not say",
    "rules": "{\"governance\": [{\"rule\", \"kind\": \"access\"|\"retention\"|\"approval\", \"applies_to\": []}]}",
    "unclear": "{} (put each open point in unclear)",
}


class Refused(Exception):
    def __init__(self, status: int, message: str):
        super().__init__(message)
        self.status, self.message = status, message


def _now() -> str:
    return datetime.now(timezone.utc).isoformat(timespec="seconds")


def _id(prefix: str) -> str:
    return f"{prefix}_{secrets.token_hex(6)}"


# ---- records ---------------------------------------------------------------------------------------------------------------

def _path(plane, iid: str, name: str):
    if not S.ID_RE.fullmatch(str(iid)):
        raise Refused(404, S.NOT_FOUND)
    return plane.path("intents", iid, name)


def _read(plane, iid: str, name: str) -> dict:
    import time
    for attempt in range(6):  # ER-14, RC-11: on Windows a read that meets a writer's os.replace (the background scan) fails for a moment
        try:
            data = read_json_strict(_path(plane, iid, name), dict)
            break
        except PermissionError:
            if attempt == 5:
                raise Refused(500, S.UNREADABLE) from None
            time.sleep(0.05)
        except (OSError, ValueError):
            raise Refused(500, S.UNREADABLE) from None
    if data is None:
        raise Refused(404, S.NOT_FOUND)
    return data


def read_thread(plane, iid: str) -> dict:
    return _read(plane, iid, "thread.json")


def read_sow(plane, iid: str) -> dict:
    return _read(plane, iid, "sow.json")


def empty_sow(iid: str, line: str) -> dict:
    return {"id": f"sow-{iid}", "intent_id": iid, "title": line, "state": "draft", "version": 1,
            "goal": {"line": line, "why": "", "readers": [], "decisions": [], "needed_by": ""},
            "process": {"steps": [], "handovers": []}, "kpis": [], "sources": [], "catalogue": [], "governance": [],
            "unclear": [], "signature": None}


def _turn(by: str, name: str, text: str, need: str | None, files: list | None = None) -> dict:
    return {"id": _id("t"), "by": by, "name": name, "text": text, "at": _now(), "need": need, "files": files or []}


def _next_need(thread: dict) -> str | None:
    return next((n for n in NEEDS if n not in thread["needs_done"]), None)


# ---- what North may see (2.0 drafting_view shape: a fixed list of fields) -----------------------------------------------------

def _company_page(plane, iid: str | None = None, viewer=None) -> dict:
    """C4: what North may read about the company: its name and industry (control plane) and, as `document_passages`, one line per
    context-library item ticked for this conversation that the viewer may see. The I1 company page is no longer read."""
    page: dict = {}
    try:
        from sh_platform.entitlements import EntitlementStore
        co = EntitlementStore(plane.root.parent).company(plane.company_id) or {}
        page = {k: str(co[k]) for k in ("name", "industry") if co.get(k)}
    except Exception:  # noqa: BLE001 - a view without the company's name still runs
        log.warning("the company record could not be read for North")
    if iid:
        try:
            from sh_platform.v1 import c4_library
            page["document_passages"] = c4_library.lines(plane, iid, viewer)
        except Exception:  # noqa: BLE001 - a broken library never breaks the thread
            log.warning("the context library could not be read for %s", iid)
    return page


def file_preview(plane, iid: str, f: dict, viewer=None, check: bool = False) -> str:
    """Samples: the columns and five example rows; documents: the first part of the text. "" when it cannot be read, or when a library
    path has no index row (or, with `check`, none this viewer may read)."""
    try:
        from sh_platform.v1 import c4_library
        if not c4_library.path_ok(plane, f["path"], viewer, check):
            return ""
        p = plane.path(*f["path"].split("/"))
        ext = reader.ext_of(f["name"])
        if ext in ("csv", "tsv"):
            rows = list(csv.reader(io.StringIO(p.read_text(encoding="utf-8-sig", errors="replace")), delimiter="\t" if ext == "tsv" else ","))[:6]
            return "\n".join(", ".join(r) for r in rows)
        if ext == "xlsx":
            from openpyxl import load_workbook
            wb = load_workbook(p, read_only=True, data_only=True)
            out = []
            for ws in wb.worksheets:
                out.append(f"Sheet {ws.title}:")
                out += [", ".join("" if c is None else str(c) for c in row) for row in ws.iter_rows(max_row=6, values_only=True)]
            wb.close()
            return "\n".join(out)
        return "\n".join(t for _, t in reader.read_segments(f["name"], p.read_bytes()))[:3000]
    except Exception:  # noqa: BLE001 - a file North cannot read is only left out of its view
        log.warning("a thread file could not be read for North")
        return ""


def drafting_view(plane, iid: str, thread: dict, sow: dict, need: str | None, viewer=None) -> dict:
    files = [f for t in thread["turns"] for f in t.get("files") or []][-6:]
    return {"title": sow["title"], "need": need, "next_need": None,
            "statement_of_work": {k: sow[k] for k in ("goal", "process", "kpis", "sources", "governance", "unclear")},
            "earlier_turns": [{"by": t["name"], "text": t["text"]} for t in thread["turns"][-8:]],
            "files": [{"name": f["name"], "kind": f["kind"], "preview": file_preview(plane, iid, f, viewer, True)} for f in files],
            "company_page": _company_page(plane, iid, viewer)}


def fit_view(view: dict) -> dict:
    """One JSON string through the gateway: records flattened to lines, then trimmed (see egress_fit)."""
    from sh_platform.v1.egress_fit import fit
    return fit(view, (("company_page", "document_passages"), ("earlier_turns", None), ("files", None)))


# ---- the model (kernel runtime only) ----------------------------------------------------------------------------------------

def _system() -> str:
    from modules.m1_intent.agent.north import load_rules
    from sh_platform.intents.north import rules
    r = rules()
    docs = load_rules()["documents_rules"]
    return "\n\n".join([r["role"].strip(), "Rules:\n" + "\n".join(f"- {x}" for x in r["rules"]),
                        "Company documents and files:\n" + "\n".join(f"- {x}" for x in docs)])


_PROPOSED = {"goal": "goal", "kpi": "kpis", "source": "inputs", "governance": "rules", "open_question": "unclear"}


def _proposed(proposals: list) -> list[tuple[str, dict, list]]:
    """W05: the drafts North's propose tools returned as (need, section, unclear) triples that `apply` takes; none is written here."""
    out = []
    for p in proposals:
        if p.get("kind") == "process_step":
            out.append(("process", {"steps": [p["step"]]}, []))
        elif p.get("kind") == "sow" and p.get("section") in _PROPOSED:
            patch = p.get("sow_patch") or {}
            need = _PROPOSED[p["section"]]
            if need == "goal":
                out.append((need, patch.get("goal") or {}, []))
            elif need == "unclear":
                out.append((need, {}, patch.get("unclear") or []))
            else:
                key = {"kpis": "kpis", "inputs": "sources", "rules": "governance"}[need]
                out.append((need, {key: patch.get(key) or []}, []))
    return out


def _ask_model(model, plane, view: dict, need: str, next_need: str | None, latest: str, people: list[str],
               viewer=None, iid: str | None = None) -> dict | None:
    """The model's reading of the latest message as {section, met, unclear, follow_up, next_question, proposed}, or None.
    W05: a tool loop (shared tools plus North's own, LOOP-N); what its propose tools drafted comes back in `proposed`."""
    if model is None:
        return None
    from modules.m1_intent.domain.private_room import figures
    view = {**view, "need": need, "next_need": next_need,
            "facts": figures(latest, *(t["text"] if isinstance(t, dict) else str(t) for t in view["earlier_turns"]))}
    register_names(plane, "", people)
    reg, line = ToolRegistry(OpenGate()), ""
    if viewer is not None:  # no person, no access filter to apply: the loop then runs without tools, as before
        from sh_platform.v1 import north_tools as NT
        from sh_platform.v1.agent_tools import tools_for
        tools_for(plane, viewer, "north", module_id="m1_intent", iid=iid, registry=reg)
        NT.register(reg, plane, viewer, iid or "")
        line = " " + NT.PROMPT_LINE
    res = run_agent(model, reg, plane.company_id, "m1_intent", NORTH,
                    _system() + "\n\n" + SOW_PROMPT.format(shape=SHAPES[need]) + line, view,
                    f"LATEST MESSAGE FROM THE PERSON: {latest}", recorder=RunRegistry(plane))
    if res.status != "done":
        return None
    a, b = res.text.find("{"), res.text.rfind("}")
    try:
        got = json.loads(res.text[a:b + 1]) if 0 <= a < b else None
    except ValueError:
        return None
    if not (isinstance(got, dict) and isinstance(got.get("section", {}), dict)):
        return None
    got["proposed"] = _proposed(res.proposals)
    return got


# ---- the rule-based fallback --------------------------------------------------------------------------------------------------

def _parts(text: str) -> list[str]:
    return [p.strip(" -*•\t") for p in _SPLIT.split(text) if p and p.strip(" -*•\t.")]


def _labelled(text: str) -> dict[str, str]:
    """"key: value" lines, keys lower-cased."""
    out = {}
    for line in text.splitlines():
        k, sep, v = line.partition(":")
        if sep and v.strip() and len(k.split()) <= 3:
            out[k.strip().lower()] = v.strip()
    return out


def _list(v: str) -> list[str]:
    return [x.strip() for x in re.split(r",|\band\b", v or "") if x.strip()]


def parse(need: str, text: str) -> dict:
    """A structured parse of the person's text for one need (the fallback's `section`)."""
    kv, parts = _labelled(text), _parts(text)
    if need == "goal":
        return {"why": kv.get("why") or (parts[0] if parts else ""), "readers": _list(kv.get("readers") or kv.get("who reads") or ""),
                "decisions": _list(kv.get("decisions") or ""), "needed_by": kv.get("needed by") or kv.get("by when") or ""}
    if need == "process":
        steps = []
        for p in parts:
            who, _, does = p.partition(" ")
            named = who[:1].isupper() and who.lower() not in ("we", "then", "the", "i", "it", "every", "each")
            steps.append({"name": p[:80], "who": who if named else "", "does": p, "inputs": [], "outputs": [], "system": ""})
        return {"steps": steps, "handovers": []}
    if need == "inputs":
        return {"sources": [{"name": n, "kind": "system", "holds": "", "owner": ""} for n in _list(" and ".join(parts))]}
    if need == "kpis":
        kpis = []
        for p in parts:
            name, sep, definition = p.partition(":")
            if not sep:
                name, sep, definition = p.partition(" is ")
            kpis.append({"name": name.strip()[:120], "definition": definition.strip(), "formula": "", "unit": "", "period": "", "owner": ""})
        return {"kpis": kpis}
    if need == "rules":
        def kind(s: str) -> str:
            s = s.lower()
            return "retention" if re.search(r"\b(keep|kept|retain|delete|years?|months?)\b", s) else \
                "approval" if re.search(r"\b(approv|sign)", s) else "access"
        return {"governance": [{"rule": p, "kind": kind(p), "applies_to": []} for p in parts]}
    return {}


def _fallback(need: str, next_need: str | None, text: str) -> dict:
    unsure = bool(DONT_KNOW.search(text))
    met = unsure or bool(NOTHING.match(text)) or len(text.split()) >= 4
    unclear = [{"question": p, "about": "unclear", "who_can_settle": ""} for p in _parts(text)] \
        if need == "unclear" and not NOTHING.match(text) else []
    return {"section": {} if unsure else parse(need, text), "met": met, "unclear": unclear,
            "follow_up": "" if met else FOLLOW_UPS[need], "next_question": QUESTIONS.get(next_need or "", "")}


# ---- applying a section to the intent of work -----------------------------------------------------------------------------

def _s(v, limit: int = 600) -> str:
    return str(v or "").strip()[:limit]


def _ls(v) -> list[str]:
    return [_s(x, 200) for x in v if _s(x)] if isinstance(v, list) else _list(_s(v))


def apply(sow: dict, need: str, section: dict, unclear: list, by: str) -> None:
    sec = section if isinstance(section, dict) else {}
    if need == "goal":
        g = sow["goal"]
        g["why"] = _s(sec.get("why")) or g["why"]
        g["readers"] = _ls(sec.get("readers")) or g["readers"]
        g["decisions"] = _ls(sec.get("decisions")) or g["decisions"]
        g["needed_by"] = _s(sec.get("needed_by"), 80) or g["needed_by"]
    elif need == "process":
        steps = sow["process"]["steps"]
        for st in sec.get("steps") or []:
            if isinstance(st, dict) and _s(st.get("name") or st.get("does")):
                steps.append({"n": len(steps) + 1, "name": _s(st.get("name") or st.get("does"), 120), "who": _s(st.get("who"), 120),
                              "does": _s(st.get("does")), "inputs": _ls(st.get("inputs")), "outputs": _ls(st.get("outputs")),
                              "system": _s(st.get("system"), 120), "control": _s(st.get("control")), "exception": _s(st.get("exception"))})
        num = {s["name"].lower(): s["n"] for s in steps}
        for h in sec.get("handovers") or []:
            if isinstance(h, dict) and num.get(_s(h.get("from_step")).lower()) and num.get(_s(h.get("to_step")).lower()):
                sow["process"]["handovers"].append({"from_step": num[_s(h["from_step"]).lower()], "to_step": num[_s(h["to_step"]).lower()],
                                                    "what": _s(h.get("what"))})
    elif need == "inputs":
        have = {s["name"].lower() for s in sow["sources"]}
        for src in sec.get("sources") or []:
            name = _s(src.get("name"), 120) if isinstance(src, dict) else ""
            if name and name.lower() not in have:
                have.add(name.lower())
                kind = src.get("kind") if src.get("kind") in ("file", "system", "document") else "system"
                sow["sources"].append({"id": _id("src"), "name": name, "kind": kind, "holds": _s(src.get("holds")), "sample_path": "",
                                       "owner": _s(src.get("owner"), 120), "fields": []})
    elif need == "kpis":
        for k in sec.get("kpis") or []:
            name = _s(k.get("name"), 120) if isinstance(k, dict) else ""
            if name and name.lower() not in {x["name"].lower() for x in sow["kpis"]}:
                definition = _s(k.get("definition"))
                sow["kpis"].append({"code": f"K{len(sow['kpis']) + 1:02d}", "name": name, "definition": definition,
                                    "formula": _s(k.get("formula")), "unit": _s(k.get("unit"), 40), "period": _s(k.get("period"), 40),
                                    "owner": _s(k.get("owner"), 120), "source_ids": [], "status": "defined" if definition else "unclear",
                                    "pack_suggested": False})
                if not definition:
                    unclear = unclear + [{"question": f"How is {name} worked out?", "about": name, "who_can_settle": ""}]
    elif need == "rules":
        for r in sec.get("governance") or []:
            if isinstance(r, dict) and _s(r.get("rule")):
                sow["governance"].append({"rule": _s(r["rule"]), "kind": r.get("kind") if r.get("kind") in ("access", "retention", "approval") else "access",
                                          "applies_to": _ls(r.get("applies_to")), "stated_by": by})
    asked = {u["question"].lower() for u in sow["unclear"]}
    for u in unclear or []:
        q = _s(u.get("question"), 300) if isinstance(u, dict) else ""
        if q and q.lower() not in asked:
            asked.add(q.lower())
            sow["unclear"].append({"id": _id("u"), "question": q, "about": _s(u.get("about"), 120) or need,
                                   "who_can_settle": _s(u.get("who_can_settle"), 120), "settled": False, "answer": ""})


def pack_suggestions(sow: dict, pack, said: str) -> None:
    """At most three pack KPIs whose name shares a word with what the person said; marked pack_suggested and unclear."""
    try:
        kpis = pack.section("kpi_tree").get("kpis") or []
    except (KeyError, AttributeError, OSError):
        return
    words = {w for w in re.findall(r"[a-z]{5,}", said.lower())} - _STOP
    have = {k["name"].lower() for k in sow["kpis"]}
    added = 0
    for k in kpis:
        name = str(k.get("name") or "")
        if added >= MAX_PACK or not name or name.lower() in have or not (set(re.findall(r"[a-z]{5,}", name.lower())) - _STOP) & words:
            continue
        added += 1
        sow["kpis"].append({"code": f"K{len(sow['kpis']) + 1:02d}", "name": name, "definition": str(k.get("definition") or ""),
                            "formula": str(k.get("formula") or ""), "unit": str(k.get("unit") or ""), "period": "", "owner": "",
                            "source_ids": [], "status": "unclear", "pack_suggested": True})
        sow["unclear"].append({"id": _id("u"), "question": f"Industry practice (suggestion): do you want {name}, defined as the industry usually does?",
                               "about": name, "who_can_settle": "The person who raised it", "settled": False, "answer": ""})


# ---- the acts -----------------------------------------------------------------------------------------------------------------

def create(plane, viewer, line: str, *, model=None) -> dict:
    line = " ".join(str(line or "").split())[:300]
    if not line:
        raise Refused(400, "Say in one line what you want to achieve.")
    rec = L.new_record("initiative", plane.company_id, L.person(viewer.account_id, viewer.name), S.now_iso(), title=line, purpose=line)
    rec["state"] = "open"
    S.create_intent(plane, rec)
    iid = rec["id"]
    thread = {"intent_id": iid, "turns": [_turn("person", viewer.name, line, None)], "needs_done": [], "state": "open"}
    thread["turns"].append(_turn("north", NORTH if model else FALLBACK_NAME, QUESTIONS["goal"], "goal"))
    sow = empty_sow(iid, line)
    write_json_atomic(_path(plane, iid, "sow.json"), sow)
    write_json_atomic(_path(plane, iid, "thread.json"), thread)
    with __import__("contextlib").suppress(Exception):  # P1: whoever starts it is its raiser
        from sh_platform.v1.p1_project import create_project; create_project(plane, iid, line, viewer.name, viewer.account_id)
    return {"id": iid, **thread, "sow": sow}


def _place(name: str, kind: str | None) -> tuple[str, str]:
    """(kind, plane-relative path): a sample lands in `landing/<source>/`, a document in `documents/` (the folders I1, I5 and B2 read)."""
    ext = reader.ext_of(name)
    if kind not in ("sample", "document"):
        kind = "sample" if ext in SAMPLE_EXT else "document"
    if kind == "sample" and ext not in SAMPLE_EXT:
        kind = "document"  # only a table can be a sample
    if kind == "sample":
        source = re.sub(r"[^a-z0-9]+", "_", name.rsplit(".", 1)[0].lower()).strip("_") or "source"
        return kind, f"landing/{source}/{name}"
    return kind, f"documents/{name}"


def checked_name(filename: str, data: bytes, limit: int = reader.MAX_UPLOAD_BYTES) -> str:
    """The safe file name, after the type, empty and size checks (shared with C4's library upload)."""
    name = reader.safe_name(filename)
    ext = reader.ext_of(name)
    if ext not in SAMPLE_EXT and ext not in reader.PICTURE_TYPES and ext not in reader.allowed_extensions():
        raise Refused(422, f"Cannot accept .{ext or '(no extension)'} files here. Allowed: {', '.join(sorted(SAMPLE_EXT | reader.PICTURE_TYPES | set(reader.allowed_extensions())))}.")
    if not data:
        raise Refused(400, "That file is empty.")
    if len(data) > limit:
        raise Refused(413, f"That file is too large. The limit is {limit // 1_000_000} MB.")
    return name


def save_file(plane, iid: str, filename: str, data: bytes, kind: str | None = None) -> dict:
    read_thread(plane, iid)  # the intent's thread must exist
    name = checked_name(filename, data)
    kind, rel = _place(name, kind)
    p = plane.path(*rel.split("/"))
    if p.is_file() and p.read_bytes() != data:  # never overwrite a file the company page, a scan or a load already reads
        raise Refused(409, "A different file with that name is already there. Rename yours and add it again.")
    p.parent.mkdir(parents=True, exist_ok=True)
    p.write_bytes(data)
    return {"name": name, "path": rel, "kind": kind}


def stored_file(plane, entry: dict) -> dict:
    """A file record the screen sends back after uploading it: it must already be stored where `_place` puts it."""
    name = reader.safe_name(str(entry.get("name") or ""))
    kind, rel = _place(name, entry.get("kind"))
    if not plane.path(*rel.split("/")).is_file():
        raise Refused(400, "Upload the file first, then send it.")
    return {"name": name, "path": rel, "kind": kind}


def _scan(plane, iid: str, viewer=None) -> None:
    try:
        from sh_platform.v1 import catalogue_scan
    except ImportError:
        return  # I5 not landed yet
    try:
        catalogue_scan.scan(plane, iid, viewer=viewer)
    except Exception:  # noqa: BLE001 - a failed scan never loses the person's turn
        log.exception("catalogue scan failed for %s", iid)


def _room(plane, iid: str, text: str, sow_id: str) -> None:
    try:
        from sh_platform.v1 import room
        room.post(plane, iid, "north", "handover", text, to=["people"], refs=[{"kind": "sow", "id": sow_id}])
    except Exception:  # noqa: BLE001 - a missing or broken room never breaks the thread
        log.warning("room post skipped for %s", iid)


def turn(plane, iid: str, viewer, text: str, files: list[dict] | None = None, *, model=None, pack=None) -> dict:
    """The person's turn and North's reply. `files` are records already stored (save_file / stored_file).
    Returns the thread and the intent of work as they now stand."""
    text = str(text or "").strip()[:MAX_TEXT]
    thread_path = _path(plane, iid, "thread.json")
    with lock_for(thread_path):
        thread = read_thread(plane, iid)
        sow = read_sow(plane, iid)
        if sow["state"] != "draft":
            raise Refused(409, "The intent of work is signed. It can no longer change here.")
        files = list(files or [])
        if not text and not files:
            raise Refused(400, "Write something or add a file.")
        need = _next_need(thread)
        thread["turns"].append(_turn("person", viewer.name, text, need, files))
        if files:
            _scan(plane, iid, viewer)  # I5 writes sow.sources and sow.catalogue from the files
        nxt = None
        reply = READY
        if need:
            asked = sum(1 for t in thread["turns"] if t["by"] == "north" and t["need"] == need)
            later = [n for n in NEEDS if n not in thread["needs_done"] and n != need]
            nxt = later[0] if later else None
            view = drafting_view(plane, iid, thread, read_sow(plane, iid), need, viewer)
            said = text or "I have added files."
            got = _ask_model(model, plane, fit_view(view), need, nxt, said, [viewer.name], viewer, iid) if model else None
            name = NORTH if got else FALLBACK_NAME
            got = got or _fallback(need, nxt, said)
            met = bool(got.get("met")) or asked >= 2 or (need == "inputs" and bool(files))  # one follow-up at most
            unclear = list(got.get("unclear") or [])
            if met and not got.get("met") and not files:  # still not answered after the follow-up: never guessed, written down as unclear
                unclear.append({"question": FOLLOW_UPS[need], "about": need, "who_can_settle": ""})
            with lock_for(_path(plane, iid, "sow.json")):  # re-read: the scan may have written to it
                sow = read_sow(plane, iid)
                apply(sow, need, got.get("section") or {}, unclear, viewer.name)
                for pneed, psec, punc in got.get("proposed") or []:  # W05: sections the model proposed through its tools
                    if pneed == "process":
                        have = {x["name"].lower() for x in sow["process"]["steps"]}
                        psec = {"steps": [x for x in psec["steps"] if x["name"].lower() not in have]}
                    apply(sow, pneed, psec, punc, viewer.name)
                if need == "kpis" and pack is not None:
                    pack_suggestions(sow, pack, " ".join([sow["goal"]["line"], sow["goal"]["why"], text]))
                write_json_atomic(_path(plane, iid, "sow.json"), sow)
            if met:
                thread["needs_done"].append(need)
                reply = (_s(got.get("next_question"), 600) or QUESTIONS[nxt]) if nxt else READY
            else:
                reply = _s(got.get("follow_up"), 600) or FOLLOW_UPS[need]
                nxt = need
        else:
            name = NORTH if model else FALLBACK_NAME
        thread["turns"].append(_turn("north", name, reply, nxt if need else None))
        became_ready = thread["state"] == "open" and _next_need(thread) is None
        if became_ready:
            thread["state"] = "sow_ready"
            _scan(plane, iid, viewer)  # the landing folders and documents go into the catalogue before the person reads the SOW
        write_json_atomic(thread_path, thread)
    if became_ready:
        _room(plane, iid, READY, sow["id"])
    return {**thread, "sow": read_sow(plane, iid)}
