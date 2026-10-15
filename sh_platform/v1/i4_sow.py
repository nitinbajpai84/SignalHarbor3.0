"""Card I4: the intent of work (`intents/<iid>/sow.json`, V1-BUILD.md "Records"): read, edit a line while draft, sign, send
back, download. Signing freezes the record (CLAUDE.md rule 7) and hands the signed content to the 2.0 stores: every defined KPI
becomes a confirmed company term (sh_platform/intents/glossary.add_term), every source and catalogue row becomes a row of a module 2
catalogue whose baseline is then marked signed (modules/m2_baseline/store/catalogue.add_row, store/baseline.set_state).
Sign pattern copied from 2.0 sh_platform/intents/charter.py post_charter_sign (confirm flag, version guard, signed is read-only)."""
from __future__ import annotations

import hashlib
import io
import logging
import re
from datetime import datetime, timezone

from sh_platform.jsonfile import lock_for, read_json_strict, write_json_atomic

log = logging.getLogger(__name__)

_IID = re.compile(r"^[A-Za-z0-9][A-Za-z0-9_-]{0,79}$")
MAX_TEXT = 2000
SECTIONS = ("goal", "process", "kpis", "sources", "catalogue", "governance", "unclear")
# What a person may change in place, per section (the contract's own field names; ids, codes and links stay as drafted).
EDITABLE = {
    "title": None,
    "goal": {"line", "why", "readers", "decisions", "needed_by"},
    "process": {"name", "who", "does", "inputs", "outputs", "system", "control", "exception"},  # one step of process.steps
    "kpis": {"name", "definition", "formula", "unit", "period", "owner"},
    "sources": {"name", "holds", "owner"},
    "catalogue": {"table", "meaning"},
    "governance": {"rule", "applies_to", "stated_by"},
    "unclear": {"answer"},
    "template_fields": None,  # any field key of the pinned template (a custom must-fill field has no other slot)
}
LISTS = {"readers", "decisions", "inputs", "outputs", "applies_to"}


class SowError(Exception):
    """A plain-words refusal: `status` for the route, `code` for the screen."""

    def __init__(self, status: int, code: str, message: str):
        super().__init__(message)
        self.status, self.code, self.message = status, code, message


def _now() -> str:
    return datetime.now(timezone.utc).isoformat(timespec="seconds")


def path_of(plane, iid: str):
    if not isinstance(iid, str) or not _IID.match(iid):
        raise SowError(404, "not_found", "We could not find that intent of work.")
    return plane.path("intents", iid, "sow.json")


def read(plane, iid: str) -> dict:
    try:
        sow = read_json_strict(path_of(plane, iid), dict)
    except (OSError, ValueError):
        raise SowError(409, "unreadable", "The intent of work could not be read. Please ask your administrator.") from None
    if sow is None:
        raise SowError(404, "not_found", "We could not find that intent of work. Start one with North first.")
    return sow


def open_unclear(sow: dict) -> list[dict]:
    return [u for u in sow.get("unclear") or [] if isinstance(u, dict) and not u.get("settled")]


def sign_block(sow: dict) -> str | None:
    """Why it cannot be signed yet, or None."""
    if sow.get("state") == "signed":
        return "This intent of work is signed."
    return None  # open unclear points do not block: the raiser may sign and leave them marked for the architect


def t1_block(plane, iid: str) -> str | None:
    """T1: a must-fill field of the pinned template is empty. Fails closed: a template that cannot be read blocks the signature."""
    from sh_platform.v1 import t1_templates as T1
    from sh_platform.v1 import i2_north as N
    try:
        return T1.sign_block(plane, iid)
    except N.Refused as e:
        if e.status == 404:  # not an intent North started: it has no template to check
            return None
        return "The intent of work cannot be signed yet: its template could not be read. Please ask your administrator."
    except Exception:  # noqa: BLE001
        log.exception("T1 sign check failed")
        return "The intent of work cannot be signed yet: its template could not be read. Please ask your administrator."


def _text(value, what: str):
    if isinstance(value, list):
        out = [" ".join(str(v).split()) for v in value if isinstance(v, (str, int, float)) and str(v).strip()]
        if any(len(v) > MAX_TEXT for v in out):
            raise SowError(422, "invalid", f"The {what} is too long.")
        return out
    if not isinstance(value, str) or len(value) > MAX_TEXT:
        raise SowError(422, "invalid", f"Please write the {what} as a short piece of text.")
    return value.strip()


def _guard(sow: dict, version) -> None:
    if sow.get("state") == "signed":
        raise SowError(409, "read_only", "This intent of work is signed, so it can no longer be changed.")
    if version is not None and version != sow.get("version"):
        raise SowError(409, "stale", "The intent of work changed since you opened it. Please reload and try again.")


def edit(plane, iid: str, body: dict, by: str) -> dict:
    """One line: {section, index?, field, value, version?}. Settling an unclear point is section "unclear", field "answer"."""
    section, field, index = body.get("section"), body.get("field"), body.get("index")
    if section not in EDITABLE or (section not in ("title", "template_fields") and field not in EDITABLE[section])             or (section == "template_fields" and not (isinstance(field, str) and _IID.match(field))):
        raise SowError(422, "invalid", "That part of the intent of work cannot be changed here.")
    path = path_of(plane, iid)
    with lock_for(path):
        sow = read(plane, iid)
        _guard(sow, body.get("version"))
        if section == "title":
            value = _text(body.get("value"), "title")
            if not value:
                raise SowError(422, "invalid", "Please write a title.")
            sow["title"] = value
        else:
            value = _text(body.get("value"), field.replace("_", " "))
            if field in LISTS and isinstance(value, str):
                value = [v.strip() for v in value.split(",") if v.strip()]
            if section == "template_fields":
                if not isinstance(sow.get("template_fields"), dict):
                    sow["template_fields"] = {}
                target = sow["template_fields"]
            elif section == "goal":
                if not isinstance(sow.get("goal"), dict):
                    sow["goal"] = {}
                target = sow["goal"]
            else:
                rows = (sow.get("process") or {}).get("steps") if section == "process" else sow.get(section)
                if not isinstance(rows, list) or not isinstance(index, int) or isinstance(index, bool) or not 0 <= index < len(rows) \
                        or not isinstance(rows[index], dict):
                    raise SowError(404, "not_found", "We could not find that line.")
                target = rows[index]
            if section == "unclear":
                if not value:
                    raise SowError(422, "invalid", "Please write the answer that settles this point.")
                target["settled"] = True
            target[field] = value
        write_json_atomic(path, sow)
    return sow


def send_back(plane, iid: str, reason, by: str) -> dict:
    """The raiser returns the draft to North: a person turn in the thread (contract shape) that reopens it."""
    why = _text(reason, "reason")
    if not why:
        raise SowError(422, "invalid", "Please say what North should change.")
    sow = read(plane, iid)
    _guard(sow, None)
    tpath = plane.path("intents", iid, "thread.json")
    with lock_for(tpath):
        try:
            thread = read_json_strict(tpath, dict) or {"intent_id": iid, "turns": [], "needs_done": [], "state": "open"}
        except (OSError, ValueError):
            raise SowError(409, "unreadable", "North's thread could not be read. Please ask your administrator.") from None
        thread.setdefault("turns", []).append({"id": "t_" + hashlib.sha1(f"{iid}{_now()}{why}".encode()).hexdigest()[:12], "by": "person",
                                               "name": by, "text": f"Sent back from the intent of work: {why}", "at": _now(),
                                               "need": None, "files": []})
        thread["state"] = "open"
        write_json_atomic(tpath, thread)
    _room(plane, iid, "person", "question", f"{by} sent the intent of work back to North: {why}", ["north"], sow, name=by)
    return sow


# ---- signing ----------------------------------------------------------------------------------------------------------

def baseline_id_of(iid: str) -> str:
    """The module 2 baseline that holds this intent's signed catalogue: derived from the intent id, so no new field is needed."""
    return hashlib.md5(f"v1-sow:{iid}".encode()).hexdigest()


def sign(plane, iid: str, name: str, role: str, version, note: str = "") -> tuple[dict, dict]:
    """(the signed record, what was written where). Freezes sow.json, then writes glossary and catalogue and posts the hand-over."""
    _scan(plane, iid)  # I5 refreshes sources and catalogue from the files first (no-op when it has not landed)
    path = path_of(plane, iid)
    with lock_for(path):
        sow = read(plane, iid)
        _guard(sow, version)
        with __import__("contextlib").suppress(ImportError): from sh_platform.v1 import p1_project as P1; P1.check_sign(plane, iid, "intent", name, lambda m: SowError(409, "comments_open", m))  # P1
        block = sign_block(sow) or t1_block(plane, iid)
        if block:
            raise SowError(409, "not_ready", block)
        sow["state"] = "signed"
        sow["signature"] = {"name": name, "role": role, "at": _now(), **({"note": " ".join(str(note).split())[:500]} if str(note or "").strip() else {})}  # V6: the signer's note to the raiser (design B3)
        # history (orchestrator ruling, 7 Oct): C2 clears `signature` when it reopens the record, so every signature is also kept here
        sow.setdefault("signatures", []).append({**sow["signature"], "version": sow.get("version")})
        write_json_atomic(path, sow)
        try:  # the company map follows every write that changes what it shows
            from sh_platform.v1.m1_map import touch; touch(plane)
        except ImportError:
            pass
    written = {"glossary": _glossary(plane, iid, sow, name), "catalogue": _catalogue(plane, iid, sow, name)}
    _room(plane, iid, "north", "handover",
          f"{name} signed the intent of work (version {sow.get('version')}): {len(sow.get('kpis') or [])} KPIs, "
          f"{len(sow.get('sources') or [])} sources. Span, please draft the three layers from it.", ["span"], sow)
    return sow, written


def _scan(plane, iid: str) -> None:
    try:
        from sh_platform.v1 import catalogue_scan
    except ImportError:
        return
    fn = getattr(catalogue_scan, "scan", None)
    if fn is None:
        return
    try:
        fn(plane, iid)
    except Exception:  # the scan helps; it never stops a signature
        log.exception("the catalogue scan before signing failed for %s", iid)


def _glossary(plane, iid: str, sow: dict, by: str) -> dict:
    from sh_platform.intents import glossary as G
    out = {"created": 0, "reused": 0, "proposed": 0, "skipped": 0}
    for k in sow.get("kpis") or []:
        definition = " ".join(x for x in (str(k.get("definition") or "").strip(),
                                          f"Formula: {k['formula']}" if k.get("formula") else "") if x)
        if not k.get("name") or not definition:
            out["skipped"] += 1
            continue
        try:
            _, how = G.add_term(plane, str(k["name"]), definition, k.get("owner") or None, by, iid,
                                aliases=[str(k["code"])] if k.get("code") else None)
            out[how] += 1
        except Exception:  # one term that cannot be written does not undo the signature
            log.exception("a KPI definition of %s could not be written to the glossary", iid)
            out["skipped"] += 1
    return out


def _catalogue(plane, iid: str, sow: dict, by: str) -> dict:
    """The SOW's sources and catalogue as rows of a module 2 catalogue, then marked signed. A later signed version (after C2 reopened
    the record) leaves every row already written untouched and adds only rows that are new or changed (a changed row is a new row)."""
    from modules.m2_baseline.domain import catalogue as DC
    from modules.m2_baseline.store import baseline as B
    from modules.m2_baseline.store import catalogue as MC
    bid = baseline_id_of(iid)
    folder = B.root(plane) / bid
    try:
        with lock_for(folder / "baseline.json"):
            base = B.read_baseline(plane, bid)
            if base is None:
                B.write_file(folder / "baseline.json", {"id": bid, "charter_outcome_id": None, "state": "draft", "created_at": _now()})
        if base is not None and base.get("state") == "signed":
            B.set_state(plane, bid, "review")  # opened only for the length of this write
        have = B.read_file(folder / "catalogue.json", dict) or {}
        added = {"sources": 0, "tables": 0, "fields": 0}

        def row(tab: str, cells: dict) -> str:
            cells = _cells(cells)
            same = next((r for r in have.get(tab) or [] if all(DC.norm(DC.text_of(r.get(h))) == DC.norm(v) for h, v in cells.items())), None)
            if same is not None:
                return DC.row_id(tab, same)
            added[tab] += 1
            return MC.add_row(plane, bid, tab, cells, by)

        src_rows = {}
        for s in sow.get("sources") or []:
            src_rows[s.get("id")] = row("sources", {
                "Source name": s.get("name"), "Source type": s.get("kind"),
                "Structured or unstructured": "Unstructured" if s.get("kind") == "document" else "Structured",
                "Where it lands or lives": s.get("sample_path"), "Data owner": s.get("owner"), "Notes": s.get("holds")})
        for c in sow.get("catalogue") or []:
            parent = src_rows.get(c.get("source_id"))
            if parent is None:  # a silver or gold table built from several sources: it hangs off the SOW itself
                parent = row("sources", {"Source name": f"Intent of work: {sow.get('title') or iid}", "Notes": "Tables built inside the warehouse"})
            feeds = ", ".join(str(x) for x in c.get("feeds_kpis") or [])
            item = row("tables", {
                "Source ID": parent, "Kind": c.get("layer"), "Name of table, file or document type": c.get("table"),
                "What one row or record means": c.get("meaning"), "Notes": f"Feeds KPIs: {feeds}" if feeds else None})
            for f in c.get("fields") or []:
                row("fields", {"Item ID": item, "Field name as in the source": f.get("name"), "Business meaning": f.get("meaning"),
                               "Data type": f.get("type"), "Notes": f"From {f['source_field']}" if f.get("source_field") else None})
        B.set_state(plane, bid, "signed")
        return {"baseline_id": bid, **added}
    except Exception:  # the record stays signed; the screen says the catalogue was not written
        log.exception("the signed catalogue of %s could not be written", iid)
        return {"baseline_id": bid, "error": "The catalogue could not be written. Please ask your administrator."}


def _cells(d: dict) -> dict:
    return {k: str(v)[:MAX_TEXT] for k, v in d.items() if v not in (None, "", [])}


def _room(plane, iid: str, by: str, kind: str, text: str, to: list[str], sow: dict, name: str | None = None) -> None:
    try:
        from sh_platform.v1.room import post
    except ImportError:
        return
    try:
        post(plane, iid, by, kind, text, to=to, refs=[{"kind": "sow", "id": str(sow.get("id") or iid)}], name=name)
    except Exception:  # a missing or broken room never breaks a card
        log.exception("the room post for %s failed", iid)


# ---- download ---------------------------------------------------------------------------------------------------------

def _j(xs) -> str:
    return ", ".join(str(x) for x in xs or []) or "Not stated"


def _end(x) -> str:
    """One full stop at the end of a sentence, never two (UC9-03: a definition often ends with its own)."""
    x = str(x).rstrip()
    return x if x.endswith((".", "!", "?", ":")) else x + "."


def kind_word(kind) -> str:
    """UC9-04: a governance rule's kind as a reader writes it ('access' -> 'Access', 'retention_period' -> 'Retention period')."""
    return str(kind or "Rule").replace("_", " ").strip().capitalize()


def _blocks(sow: dict) -> list[tuple[str, list[str]]]:
    """The eight sections as (heading, lines); Markdown and Word are written from the same lines. UC9-03: sources and KPIs are named as
    the intent names them (no source ids, no KPI codes), and a sentence ends with one full stop."""
    g = sow.get("goal") or {}
    steps = (sow.get("process") or {}).get("steps") or []
    hand = (sow.get("process") or {}).get("handovers") or []
    sig = sow.get("signature")
    src = {s.get("id"): s.get("name") for s in sow.get("sources") or [] if s.get("id") and s.get("name")}
    kpi = {k.get("code"): k.get("name") for k in sow.get("kpis") or [] if k.get("code") and k.get("name")}
    return [
        ("1. Goal", [f"**{g.get('line') or 'Not stated'}**", f"Why: {g.get('why') or 'Not stated'}", f"Readers: {_j(g.get('readers'))}",
                     f"Decisions: {_j(g.get('decisions'))}", f"Needed by: {g.get('needed_by') or 'Not stated'}"]),
        ("2. Process", [" ".join(x for x in [f"{s.get('n')}. {s.get('name')} ({s.get('who') or 'who not stated'}):", _end(s['does']) if s.get('does') else "",
                                            _end(f"In: {_j(s.get('inputs'))}"), _end(f"Out: {_j(s.get('outputs'))}"),
                                            _end(f"System: {s.get('system') or 'Not stated'}")] if x) for s in steps]
         + [f"Hand-over from step {h.get('from_step')} to {h.get('to_step')}: {h.get('what')}" for h in hand]),
        ("3. KPIs", [f"**{k.get('name')}** ({k.get('status')}{', suggested by the pack' if k.get('pack_suggested') else ''}): "
                     f"{_end(k.get('definition') or 'No definition yet')} {_end('Formula: ' + str(k.get('formula') or 'Not stated'))} "
                     f"Unit: {k.get('unit') or '-'}, period: {k.get('period') or '-'}, owner: {k.get('owner') or 'Not named'}. "
                     f"Sources: {_j(src.get(i, i) for i in k.get('source_ids') or [])}."
                     for k in sow.get("kpis") or []]),
        ("4. Sources", [f"**{s.get('name')}** ({s.get('kind')}): {_end(s['holds']) + ' ' if s.get('holds') else ''}Owner: {s.get('owner') or 'Not named'}."
                        + (f" Fields: {_j(f.get('name') for f in s.get('fields') or [])}." if s.get("fields") else "") for s in sow.get("sources") or []]),
        ("5. Data catalogue", [f"**{str(c.get('table') or '').replace('_', ' ').capitalize()}** ({c.get('layer')} table {c.get('table')}): "
                               f"{_end(c['meaning']) + ' ' if c.get('meaning') else ''}From {src.get(c.get('source_id'), c.get('source_id')) or 'several sources'}; "
                               f"feeds {_j(kpi.get(x, x) for x in c.get('feeds_kpis') or [])}." for c in sow.get("catalogue") or []]),
        ("6. Governance", [f"{kind_word(r.get('kind'))}: {_end(r.get('rule') or 'Not stated')} Applies to {_j(r.get('applies_to'))}. Stated by {r.get('stated_by') or 'Not named'}."
                           for r in sow.get("governance") or []]),
        ("7. Unclear", [f"{u.get('question')} ({'settled: ' + str(u.get('answer')) if u.get('settled') else 'open; ' + str(u.get('who_can_settle') or 'who can settle not named')})"
                        for u in sow.get("unclear") or []]),
        ("8. Signature", [f"Signed by {sig.get('name')} ({sig.get('role')}) on {str(sig.get('at'))[:10]}." if sig else "Not signed yet. This is a draft."]),
    ]


def filename(sow: dict, ext: str) -> str:
    stem = re.sub(r"[^A-Za-z0-9]+", "-", str(sow.get("title") or "statement-of-work")).strip("-")[:60] or "statement-of-work"
    return f"{stem}-v{sow.get('version') or 1}.{ext}"


def layout_of(plane, iid: str):
    """T1b: the documents laid out from the intent's templates, or None (then the eight plain sections are written)."""
    try:
        from sh_platform.v1 import t1_layout
        if not plane.path("intents", iid, "template.json").is_file():  # began before templates: the plain sections
            return None
        return t1_layout.layout(plane, iid, read(plane, iid))
    except Exception:  # noqa: BLE001
        log.warning("the template layout could not be made for %s", iid)
        return None


def _note(sow: dict) -> str:
    sig = sow.get("signature")
    return f"Version {sow.get('version')}, {sow.get('state')}." + (f" Signed by {sig.get('name')} ({sig.get('role')}) on {str(sig.get('at'))[:10]}." if sig else "")


def to_markdown(sow: dict, layout: dict | None = None) -> str:
    if layout:
        from sh_platform.v1 import t1_layout
        return t1_layout.to_markdown(layout, f"Intent of work: {sow.get('title') or ''}", _note(sow))
    out = [f"# Intent of work: {sow.get('title') or ''}", f"Version {sow.get('version')}, {sow.get('state')}.", ""]
    for head, lines in _blocks(sow):
        out += [f"## {head}", ""] + ([f"- {x}" for x in lines] or ["- Nothing yet."]) + [""]
    return "\n".join(out)


def docx_available() -> bool:
    try:
        import docx  # noqa: F401
    except ImportError:
        return False
    return True


def to_docx(sow: dict, layout: dict | None = None) -> bytes:
    if layout:
        from sh_platform.v1 import t1_layout
        return t1_layout.to_docx(layout, f"Intent of work: {sow.get('title') or ''}", _note(sow))
    import docx
    d = docx.Document()
    d.add_heading(f"Intent of work: {sow.get('title') or ''}", 0)
    d.add_paragraph(f"Version {sow.get('version')}, {sow.get('state')}.")
    for head, lines in _blocks(sow):
        d.add_heading(head, 1)
        for x in lines or ["Nothing yet."]:
            d.add_paragraph(x.replace("**", ""), style="List Bullet")
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()
