"""K10 Consume reading room: the Morning Note, Daily P&L and NAV, Market signals, the fund league and the reports library.

Nothing here is worked out by a model. Every figure comes from a function that already exists and already filters by what the viewer may see
(k9_daily, k4_office, k5_managers, k7_benchmark, e3_calc, e1_reports, e5_brief); this file only puts them side by side, adds the index and
benchmark comparisons that market_observations allow, and words the sentences from those figures. Where a figure cannot be worked out the
answer says "Not available yet." and (where useful) why; a missing figure is never shown as zero.

Outside material (index levels, exchange rates, manager letters and newsletters) is always marked `outside: true` with its source and date, so the
screen can label it and so nobody mistakes it for the company's own figures.

V6: which index a fund is compared with comes from a signed fund-to-index mapping (the head of data signs it, `sign_index_map`), never from the fund's
name. The old asset-class and name rule survives only as the PROPOSAL shown to the head of data before signing (`_propose`). A fund with no signed
row shows "No index is signed for this fund yet." and no comparison.
ponytail: a fund's own return needs confirmed statements at both month ends and its calls and distributions; until that exists `fund_return` and
`alpha` are None for every fund and the screen shows "Not available yet".
"""
from __future__ import annotations

import calendar
import html
import io
import logging
import re
from datetime import date, datetime, timedelta, timezone
from decimal import ROUND_HALF_UP, Decimal

from sh_platform.v1.scope_memo import memo
from sh_platform.jsonfile import lock_for, read_json_strict, write_json_atomic
from sh_platform.v1 import e1_reports, e3_calc, e4_cases, e5_brief, k4_office as O, k5_managers as K5, k7_benchmark, wh
from sh_platform.v1 import k9_daily as D
from sh_platform.v1.i2_north import Refused

log = logging.getLogger(__name__)
ZERO = Decimal(0)
NOT_YET = "Not available yet."
PERIODS = ("month", "ytd", "inception")
LIVE_REPORTS = [
    ("manager_scorecard", "Manager scorecard", "Statements on time, latest values and what changed, for every manager.", "/centre/consume/managers"),
    ("liquidity_ladder", "Liquidity ladder, 30 / 60 / 90 days", "Calls and distributions against cash, by entity.", "/centre/consume/office#liquidity"),
    ("exposure_limits", "Exposure and limits", "The signed limits, checked every morning.", "/centre/action/cases"),
]


# ---- small helpers ------------------------------------------------------------------------------------------------------------

def _dec(v) -> Decimal | None:
    return None if v is None else Decimal(str(v))


def _money(v) -> str:
    n = _dec(v)
    if n is None:
        return NOT_YET
    return f"{n:,.0f}" if n == n.to_integral_value() else f"{n:,.2f}"


def _signed(v) -> str:
    n = _dec(v)
    return NOT_YET if n is None else ("+" if n > 0 else "-" if n < 0 else "") + _money(abs(n))


def _millions(v) -> str:
    n = _dec(v)
    return NOT_YET if n is None else f"{(n / 1000000).quantize(Decimal('0.01'), ROUND_HALF_UP):,.2f} million"


def _pct(v, plus: bool = False) -> str:
    n = _dec(v)
    return NOT_YET if n is None else f"{'+' if plus and n > 0 else ''}{n:.2f}%"


def _pts(v) -> str:
    n = _dec(v)
    return NOT_YET if n is None else f"{'+' if n > 0 else ''}{n:.2f}"


def _d(v) -> str:
    """7 Oct (no leading zero), from a date or an ISO text."""
    x = v if isinstance(v, date) and not isinstance(v, datetime) else _day(v)
    return f"{x.day} {x:%b}" if x else ""


def _dy(v) -> str:
    x = v if isinstance(v, date) and not isinstance(v, datetime) else _day(v)
    return f"{x.day} {x:%b %Y}" if x else ""


def _day(v) -> date | None:
    if v is None or v == "":
        return None
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    try:
        return date.fromisoformat(str(v)[:10])
    except ValueError:
        return None


def _tid(*p) -> str:
    return "k10:" + ":".join(str(x) for x in p)


def _word(s) -> str:
    return str(s or "").replace("_", " ").capitalize()


def _safe(name: str, fn, default):
    try:
        return fn()
    except Exception:  # one part failing must not blank the page; the reason goes to the log, the person gets plain words
        log.warning("consume reading room: %s could not be read", name, exc_info=True)
        return default


def _month_name(d: date | None) -> str:
    return calendar.month_name[d.month] if d else ""


def _period_end(issued: dict, day: date) -> date:
    return _day(issued.get("period_end")) if issued.get("available") and issued.get("period_end") else day.replace(day=1) - timedelta(1)


@memo
def _names(plane) -> dict[str, str]:
    with wh.connect(plane, read_only=True) as con:
        return {str(r[0]): r[1] for r in con.execute("select id, display_name from users").fetchall()}


@memo
def _company(plane) -> str:
    with wh.connect(plane, read_only=True) as con:
        r = con.execute("select name from organizations limit 1").fetchone()
    return r[0] if r else ""


# ---- cases (the "needs you" part) ---------------------------------------------------------------------------------------------

@memo
def _cases(plane, viewer) -> list[dict]:
    rows = _safe("cases", lambda: e4_cases.listing(plane, viewer, {"status": "active"}), [])
    return [{"number": r.get("number"), "title": r["title"], "kind": "Over limit" if r.get("exception_type") == "limit_breach" else "Open case",
             "limit": r.get("exception_type") == "limit_breach", "material": r.get("materiality_status") == "material",
             "href": f"/centre/action/cases/{r.get('number') or r['id']}"} for r in rows]


# ---- reports library ----------------------------------------------------------------------------------------------------------

def _stamp(v) -> str | None:
    return None if v is None else str(v)


def reports_library(plane, viewer) -> dict:
    names, rows, cases = _names(plane), [], _cases(plane, viewer)
    for r in e1_reports.listing(plane):
        vs = list(reversed(r["versions"]))  # oldest first, so version 1 is the first one issued
        versions = [{"id": v["id"], "number": i + 1, "status": v["status"], "state": "Signed" if v["status"] == "issued" else "Awaiting signature",
                     "period_end": _stamp(v["period_end"]), "issued_at": _stamp(v["issued_at"]), "signed_by": names.get(str(v["approved_by"])) if v["status"] == "issued" else None,
                     "link": f"/centre/consume/brief/{v['id']}", "pdf": f"/api/v1/consume/reports/{v['id']}/pdf" if v["status"] == "issued" else None,
                     "pptx": f"/api/v1/consume/reports/{v['id']}/pptx" if v["status"] == "issued" else None}
                    for i, v in enumerate(vs)]
        last = versions[-1] if versions else None
        rows.append({"id": r["id"], "code": r["report_code"], "name": r["name"], "audience": r["audience"], "live": False,
                     "summary": f"{' and '.join(str(r['currencies'] or '').split(','))}." + (f" Issued {_d(last['issued_at'])}, {_stamp(last['issued_at'])[11:16]}." if last and last["issued_at"] else ""),
                     "state": last["state"] if last else "Awaiting signature", "version": last["number"] if last else None, "signed_by": last and last["signed_by"],
                     "signed_at": last and last["issued_at"], "link": last and last["link"], "pdf": last and last["pdf"], "pptx": last and last["pptx"], "versions": versions,
                     "trail_id": _tid("report", r["id"])})
    over = [c for c in cases if c["limit"]]
    for code, name, summary, link in LIVE_REPORTS:
        if code == "exposure_limits":
            summary = f"{len(over)} limit case{'s' if len(over) != 1 else ''} open." if over else "No limit is over."
        rows.append({"id": code, "code": code, "name": name, "audience": None, "live": True, "summary": summary, "state": "Live", "version": None, "signed_by": None,
                     "signed_at": None, "link": link, "pdf": None, "pptx": None, "versions": [], "trail_id": _tid("report", code)})
    waiting = [r for r in rows if r["state"] == "Awaiting signature"]
    return {"reports": rows, "waiting": len(waiting), "pdf_available": pdf_available(), "pptx_available": pptx_available(),
            "delivery": [{"text": "The Morning Note arrives by email at 07:00.", "href": "/centre/operations#deliveries", "action": "See the email"}]}


# ---- report PDF ---------------------------------------------------------------------------------------------------------------

def pdf_available() -> bool:
    try:
        import reportlab  # noqa: F401
        return True
    except ImportError:
        return False


def pptx_available() -> bool:
    try:
        import pptx  # noqa: F401
        return True
    except ImportError:
        return False


def _issued(plane, viewer, report_version_id: str, currency: str | None) -> tuple[dict, str | None, list[list[str]]]:
    """A signed version as issued: (the brief, who signed it, table rows with a header). Refused for a version nobody has signed."""
    try:
        b = e5_brief.brief(plane, viewer, report_version_id, currency)
    except LookupError:
        raise Refused(404, "We could not find that report.") from None
    if b["version_status"] != "issued":
        raise Refused(409, "Only a signed version can be exported. This one is still waiting for a signature.")
    with wh.connect(plane, read_only=True) as con:
        sv = con.execute("select approved_by from report_versions where id = ?", [report_version_id]).fetchone()
    who = _names(plane).get(str(sv[0])) if sv else None
    ccy = b.get("currency") or ""
    return b, who, [["Line", f"As issued ({ccy})", f"Known now ({ccy})", "Status now"]] + [
        [x["label"], _money(x["issued"]), _money(x["known_now"]), _word(x["status"])] for x in b["lines"]]


def _file_name(b: dict, ext: str) -> str:
    return f"{re.sub(r'[^a-z0-9]+', '-', b['report_name'].lower()).strip('-')}-{b['period_end']}.{ext}"


def report_pptx(plane, viewer, report_version_id: str, currency: str | None = None) -> tuple[bytes, str]:
    """E5: (bytes, file name) of a signed report version as a PowerPoint: a title slide, then the lines as issued, 10 to a slide. Same rules as the PDF."""
    b, who, rows = _issued(plane, viewer, report_version_id, currency)
    if not pptx_available():
        raise Refused(404, "A PowerPoint file cannot be made on this server. Open the report on screen instead.")
    from pptx import Presentation
    from pptx.util import Pt
    deck = Presentation()
    first = deck.slides.add_slide(deck.slide_layouts[0])
    first.shapes.title.text = b["report_name"]
    first.placeholders[1].text = f"Period ending {_dy(b['period_end'])}. Issued {_dy(b['issued_at'])}" + (f", signed by {who}" if who else "") + f". Data cutoff {_dy(b['cutoff_at'])}."
    body = rows[1:] or [["No lines you may see", "", "", ""]]
    for n in range(0, len(body), 10):
        sl = deck.slides.add_slide(deck.slide_layouts[5])
        sl.shapes.title.text = b["report_name"] + (f" ({n // 10 + 1})" if len(body) > 10 else "")
        chunk = [rows[0], *body[n:n + 10]]
        tbl = sl.shapes.add_table(len(chunk), 4, deck.slide_width // 20, deck.slide_height // 4, deck.slide_width * 9 // 10, deck.slide_height // 2).table
        for i, r in enumerate(chunk):
            for j, v in enumerate(r):
                tbl.cell(i, j).text = str(v)
                tbl.cell(i, j).text_frame.paragraphs[0].font.size = Pt(12)
    notes = deck.slides[-1].shapes.add_textbox(deck.slide_width // 20, deck.slide_height * 9 // 10, deck.slide_width * 9 // 10, deck.slide_height // 20)
    notes.text_frame.text = "A signed version never changes. As issued is what was signed; known now is what the company knows today and is not part of the signed report."
    notes.text_frame.paragraphs[0].font.size = Pt(10)
    buf = io.BytesIO()
    deck.save(buf)
    return buf.getvalue(), _file_name(b, "pptx")


def report_pdf(plane, viewer, report_version_id: str, currency: str | None = None) -> tuple[bytes, str]:
    """(bytes, file name) of a signed report version as issued. Refused for a version nobody has signed. Only the lines the viewer may see are printed."""
    b, who, rows = _issued(plane, viewer, report_version_id, currency)
    if not pdf_available():
        raise Refused(404, "A PDF cannot be made on this server. Open the report on screen instead.")
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
    st = getSampleStyleSheet()
    table = Table(rows, colWidths=[230, 100, 100, 80], repeatRows=1)
    table.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E6EEF0")), ("ALIGN", (1, 1), (2, -1), "RIGHT"),
                               ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#9AA5A8")), ("FONTSIZE", (0, 0), (-1, -1), 9)]))
    out = [Paragraph(html.escape(b["report_name"]), st["Title"]),
           Paragraph(html.escape(f"Period ending {_dy(b['period_end'])}. Issued {_dy(b['issued_at'])}" + (f", signed by {who}" if who else "") +
                                 f". Data cutoff {_dy(b['cutoff_at'])}."), st["Normal"]), Spacer(1, 12), table, Spacer(1, 12)]
    if b.get("banner"):
        out.append(Paragraph(html.escape(b["banner"]), st["Normal"]))
    if b.get("partial"):
        out.append(Paragraph("Some lines are left out because you do not have access to them.", st["Normal"]))
    out.append(Paragraph("A signed version never changes. The as issued column is what was signed; the known now column is what the company knows today and is not part of the signed report.", st["Italic"]))
    buf = io.BytesIO()
    SimpleDocTemplate(buf, pagesize=A4, title=b["report_name"], author=who or "Signal Harbor").build(out)
    return buf.getvalue(), _file_name(b, "pdf")


# ---- market data --------------------------------------------------------------------------------------------------------------

@memo
def _observations(plane, day: date) -> dict[str, list]:
    cut = datetime.combine(day, datetime.max.time(), tzinfo=timezone.utc)
    with wh.connect(plane, read_only=True) as con:
        rows = e3_calc._live(con, plane, "market_observations", "supersedes_observation_id", wh.now(),
                             "t.series_key in (select series_key from market_observations where observation_type in ('fx','index_level','yield','macro')) and t.observed_at <= ?", [cut])
    by: dict[str, list] = {}
    for r in rows:
        by.setdefault(r["series_key"], []).append((r["observed_at"].astimezone(timezone.utc).date(), Decimal(str(r["value"])), r))
    for obs in by.values():
        obs.sort(key=lambda x: x[0])
    return by


def _level(obs: list, on: date, within: int = 45):
    """(date, level) of the latest observation on or before `on` that is not older than `within` days; None otherwise."""
    got = [o for o in obs if o[0] <= on]
    return (got[-1][0], got[-1][1]) if got and (on - got[-1][0]).days <= within else None


def _bounds(day: date, which: str, first: date | None) -> tuple[date, date]:
    end = day.replace(day=1) - timedelta(1)
    start = {"month": end.replace(day=1) - timedelta(1), "ytd": date(end.year - 1, 12, 31), "inception": first or end}[which]
    return start, end


def _change(obs: list, start: date, end: date, spread_bps: Decimal = ZERO) -> dict:
    a, b = _level(obs, start), _level(obs, end)
    if not a or not b or a[1] == 0:
        return {"start": start, "end": end, "start_level": None, "end_level": None, "change_pct": None}
    months = (end.year - start.year) * 12 + end.month - start.month if (start + timedelta(1)).day == 1 and (end + timedelta(1)).day == 1 else 0
    spread = spread_bps / 10000 * (Decimal(months) / 12 if months else Decimal((end - start).days) / 365)  # as e3_calc.benchmark_return does
    return {"start": start, "end": end, "start_level": a[1], "end_level": b[1], "start_date": a[0], "end_date": b[0], "change_pct": ((b[1] / a[1] - 1) + spread) * 100,
            "change_abs": b[1] - a[1]}


def _direction(pct) -> str | None:
    n = _dec(pct)
    return None if n is None else "Rising" if n > 0 else "Falling" if n < 0 else "Flat"


def _comparators(plane, by: dict, day: date) -> dict:
    """The yardsticks, all from the signed benchmark and the index levels: the blend itself, its liquid sleeve (60/40) and the private markets proxy."""
    sig = D._signed(plane)
    d = e3_calc.benchmark_definition(plane, sig["benchmark_code"], day) if sig else None
    if not d:
        return {}
    comps = d["components"]
    plain = sorted((c for c in comps if not c["spread_bps"]), key=lambda c: -c["weight"])
    spread = next((c for c in comps if c["spread_bps"]), None)
    out: dict = {"code": d["benchmark_code"], "name": d["benchmark_name"], "version": d["version_number"], "components": comps}
    wsum = sum((c["weight"] for c in plain), ZERO)
    if plain and wsum:
        out["sleeve"] = {"label": "/".join(f"{round(c['weight'] / wsum * 100)}" for c in plain) + " " + " and ".join(
            re.sub(r"^global |\s*index$", "", D.SERIES.get(c["series_key"], c["series_key"]).lower()) for c in plain), "parts": [(c["series_key"], c["weight"] / wsum) for c in plain]}
    if spread:
        out["proxy"] = {"series_key": spread["series_key"], "bps": spread["spread_bps"],
                        "label": f"Private markets proxy ({D.SERIES.get(spread['series_key'], spread['series_key']).lower()} + {float(spread['spread_bps']) / 100:g}% a year)"}
    return out


def _sleeve_change(by: dict, parts, start: date, end: date):
    got = [(w, _change(by.get(k, []), start, end)["change_pct"]) for k, w in parts]
    return None if any(c is None for _, c in got) else sum((w * c for w, c in got), ZERO)


def _series_name(key: str) -> str:
    """Our own name for a series, else the name the live-API connector gave it, else the key."""
    if key in D.SERIES:
        return D.SERIES[key]
    from sh_platform.v1.c5_connect.api import PROVIDERS
    return next((s["label"] for p in PROVIDERS.values() for s in p["series"].values() if s["key"] == key), key)


def _provenance(plane, by: dict) -> dict[str, dict]:
    """Where the newest figure of each series came from, read from the warehouse: the source system it was loaded by, how it was connected
    (a file, or a live API), and the date. A figure loaded from the demo seed says so."""
    last = {k: o[-1] for k, o in by.items() if o}
    ids = [str(v[2]["source_record_id"]) for v in last.values() if v[2].get("source_record_id")]
    got: dict[str, tuple] = {}
    if ids:
        with wh.connect(plane, read_only=True) as con:
            for r in con.execute("select r.id, sy.source_name, sy.connection_method, b.storage_uri from source_records r join source_batches b on b.id = r.source_batch_id "
                                 f"join source_systems sy on sy.id = b.source_system_id where r.id in ({','.join('?' * len(ids))})", ids).fetchall():
                got[str(r[0])] = r[1:]
    out = {}
    for k, (d, _, row) in last.items():
        src = got.get(str(row.get("source_record_id")))
        if not src:
            out[k] = {"label": None, "tag": None, "date": d, "demo": False}
            continue
        name, method, uri = src
        demo = str(uri or "").startswith("seed://")
        out[k] = {"label": f"{name} (demo data)" if demo else name, "tag": "LIVE API" if method == "live_api" else "FILE", "date": d, "demo": demo}
    return out


def _headlines(plane) -> list[dict]:
    """News the live-API connector brought in: title, where it is from and its date. Outside material, never a figure."""
    from sh_platform.consume import market_store
    items = [b for b in market_store.read(plane).get("added", []) if isinstance(b, dict) and b.get("live_api") and b.get("title")]
    return [{"title": b["title"], "source": b.get("source") or "Live API", "date": b.get("date"), "tag": "LIVE API", "outside": True, "trail_id": _tid("headline", b.get("id", b["title"]))}
            for b in sorted(items, key=lambda b: str(b.get("date") or ""), reverse=True)[:5]]


def _index_rows(plane, by: dict, comp: dict, day: date, sources: bool = False) -> list[dict]:
    rows = []
    order = list(D.SERIES)  # our own series first (equity, bonds, exchange rate), then any others by key
    for key, obs in sorted(by.items(), key=lambda kv: (order.index(kv[0]) if kv[0] in order else len(order), kv[0])):
        kind = obs[0][2]["observation_type"]
        fx = kind == "fx"
        last = obs[-1]
        per = {w: _change(obs, *_bounds(day, w, obs[0][0])) for w in PERIODS}
        rows.append({"series_key": key, "name": _series_name(key), "kind": "fx" if fx else "index" if kind == "index_level" else kind, "outside": True,
                     "source": ("WMR 4pm fixing" if "WMR" in key else "Currency fixing") if fx else "Index provider",
                     "latest": {"date": last[0], "level": last[1]}, "periods": per, "direction": _direction(per["month"]["change_pct"]),
                     "trail_id": _tid("market", key),
                     "where": {"text": f"{_series_name(key)}: levels recorded in the warehouse from {obs[0][0]:%d %b %Y} to {last[0]:%d %b %Y}; the change is the closing level over the opening level."}})
    if "proxy" in comp:
        eq = by.get(comp["proxy"]["series_key"], [])
        per = {w: _change(eq, *_bounds(day, w, eq[0][0] if eq else None), spread_bps=comp["proxy"]["bps"]) for w in PERIODS}
        rows.append({"series_key": "proxy", "name": comp["proxy"]["label"], "kind": "proxy", "outside": True, "source": "Signed rule", "latest": None, "periods": per,
                     "direction": _direction(per["month"]["change_pct"]), "trail_id": _tid("market", "proxy"),
                     "where": {"text": f"The index plus the signed spread of {float(comp['proxy']['bps']):g} basis points a year, as in the office benchmark {comp['code']} version {comp['version']}."}})
    rows.append({"series_key": "regional", "name": "Regional equity (Asia ex-Japan)", "kind": "placeholder", "outside": True, "source": None, "latest": None,
                 "periods": {w: {"change_pct": None, "start_level": None, "end_level": None} for w in PERIODS}, "direction": None,
                 "trail_id": _tid("market", "regional"), "where": {"text": "No feed is loaded for this market yet."}, "message": NOT_YET})
    if sources:
        prov = _provenance(plane, by)
        for r in rows:
            r["provenance"] = prov.get(r["series_key"])
        proxy = next((r for r in rows if r["kind"] == "proxy"), None)
        eq = prov.get(comp["proxy"]["series_key"]) if proxy else None
        if proxy:
            proxy["provenance"] = {"label": "Calculated from the equity index and the signed spread", "tag": None, "date": eq and eq["date"], "demo": bool(eq and eq["demo"])}
    return rows


def _fx_sentence(rows: list[dict]) -> str:
    fx = next((r for r in rows if r["kind"] == "fx" and r["series_key"].startswith("FX:USDSGD")), None)
    ch = fx and fx["periods"]["month"]["change_pct"]
    if ch is None:
        return ""
    if ch == 0:
        return " The US dollar did not move against the Singapore dollar."
    return f" The US dollar {'rose' if ch > 0 else 'fell'} against the Singapore dollar, so every Singapore dollar figure is a little {'higher' if ch > 0 else 'lower'} this month."


def _news(plane, viewer, managers: list[str]) -> list[dict]:
    """Letters, newsletters and questionnaires from the library that name a manager, as written. A locked file shows no text."""
    out = []
    for name in managers:
        shown, raw = K5._library(plane, viewer, name)
        for x in K5._documents_with_passages(plane, name, shown, raw):
            word = {"newsletter": "Manager newsletter", "letter": "Manager letter", "due diligence questionnaire": "Manager questionnaire"}.get(x["kind"])
            if not word:
                continue  # meeting notes are the office's own words, not outside material
            out.append({"manager": name, "source": word, "title": (t := re.sub(r"\.\w+$", "", x["name"])), "date": _stamp(x["at"])[:10], "text": (x.get("passage") or "").removeprefix(t).strip() or None,
                        "locked": x.get("passage") is None, "outside": True, "href": "/centre/library", "trail_id": _tid("news", x["id"])})
    return sorted(out, key=lambda n: n["date"], reverse=True)[:8]


def _status_weakest(statuses) -> str | None:
    return next((s.capitalize() for s in D.WORST if s in (statuses or [])), None)


@memo
def _funds(plane, viewer, day: date) -> list[dict]:
    """Every fund the viewer may see with its last value, status word and the note under it (from k5_managers.inbox)."""
    box = K5.inbox(plane, viewer, day)
    ids = [f["fund_id"] for m in box["managers"] for f in m["funds"]]
    with wh.connect(plane, read_only=True) as con:
        cls = {str(r[0]): r[1] for r in con.execute(f"select id, asset_class from instruments where id in ({','.join('?' * len(ids))})", ids).fetchall()} if ids else {}
    out = []
    for m in box["managers"]:
        for f in m["funds"]:
            lv, sts = f.get("last_value"), f.get("statements") or []
            recv = max((_day(s["received_at"]) for s in sts if s.get("received_at")), default=None)
            vd = _day(lv and lv["valuation_date"])
            word = (lv or {}).get("status_word")
            late = bool(f.get("earlier_missing"))
            if late:
                word = "Late"
            rev = next((s["revision"] for s in sts if s.get("revision")), None)
            st = (lv or {}).get("status")
            if not lv:
                note = "No value is recorded yet."
            elif st == "revised" and rev:
                note = f"Final {_d(recv)} replaced the estimate ({_signed(rev['change'])})."
            elif st == "estimated":
                note = f"Estimate {_d(recv)}" + (f"; final due {_d(f['due'])}." if f.get("due") else ".")
            elif st == "unverified":
                note = f"Read {_d(recv)}, not confirmed."
            elif st == "stale" or late:
                note = f"Last valued {_d(vd)} ({(day - vd).days} days)." + (f" {len(f['earlier_missing'])} earlier statement{'s' if len(f['earlier_missing']) != 1 else ''} overdue." if late else
                                                                          f" Next statement not due until {_d(f['due'])}." if f.get("state") == "not_due" and f.get("due") else "")
            else:
                note = f"Statement {_d(recv)}." if recv else f"Valued {_d(vd)}."
            out.append({"fund_id": f["fund_id"], "fund": f["fund_name"], "manager": f["manager"], "manager_id": m["manager_id"], "asset_class": cls.get(f["fund_id"]),
                        "value": lv and lv["amount"], "currency": lv and lv["currency"], "valuation_date": vd, "status": word or NOT_YET, "status_key": (word or "").lower(),
                        "note": note, "trail_id": _tid("fund", f["fund_id"]),
                        "where": {"text": f"Latest value of record for {f['fund_name']}" + (f", dated {_dy(vd)}" if vd else "") + ", from the manager's statements in the warehouse."}})
    return out


def _propose(asset_class, name: str) -> str:
    """The proposal the head of data starts from (asset class and name). Never used to compare anything until it is signed."""
    ac, n = (asset_class or "").lower(), name.lower()
    if ac in ("listed_equity", "equity"):
        return "equity"
    if ac in ("bond", "private_credit"):
        return "bond"
    if ac == "hedge_fund":
        return "equity" if "equity" in n else "bond"
    return "proxy"


def _pick(rows: list[dict], which: str) -> dict | None:
    if which == "proxy":
        return next((r for r in rows if r["kind"] == "proxy"), None)
    return next((r for r in rows if r["kind"] == "index" and which in r["name"].lower()), None)


def _fund_index(rows: list[dict], f: dict, period: str, signed: dict | None = None) -> dict:
    key = (signed or {}).get(f["fund_id"])
    ix = next((r for r in rows if r["series_key"] == key), None) if key else None
    out = {"compared_with": ix and ix["name"], "index_return": ix and ix["periods"][period]["change_pct"], "index_direction": ix and ix["direction"],
           "index_signed": ix is not None, "fund_return": None, "alpha": None,
           "message": "Not available yet: a fund return needs confirmed statements at both month ends."}
    if ix is None:
        out["message"] = NO_INDEX
    return out


# ---- the signed fund-to-index mapping (V6) ------------------------------------------------------------------------------------

NO_INDEX = "No index is signed for this fund yet."


def _map_path(plane):
    return plane.path("consume", "fund_index_map.json")


def _signed_map(plane) -> dict:  # not @memo: signing reads it under the lock and answers with what it wrote, in the same request
    d = read_json_strict(_map_path(plane))
    return d if isinstance(d, dict) and isinstance(d.get("rows"), dict) else {}


def index_map(plane, viewer) -> dict:
    """The mapping as signed (who, when) and, for each fund the viewer may see, the signed index or the proposal to sign."""
    from sh_platform.v1.seats import HEAD_OF_DATA, seat_of
    day = O._day(None)
    by = _observations(plane, day)
    comp = _safe("benchmarks", lambda: _comparators(plane, by, day.replace(day=1) - timedelta(1)), {})
    rows = _index_rows(plane, by, comp, day)
    choices = [{"series_key": r["series_key"], "name": r["name"]} for r in rows if r["kind"] in ("index", "proxy")]
    signed = _signed_map(plane)
    mine = signed.get("rows") or {}
    out = []
    for f in _safe("funds", lambda: _funds(plane, viewer, day), []):
        p = _pick(rows, _propose(f["asset_class"], f["fund"]))
        out.append({"fund_id": f["fund_id"], "fund": f["fund"], "series_key": mine.get(f["fund_id"]), "proposed_series_key": p and p["series_key"]})
    return {"signed": bool(mine), "signed_by": signed.get("signed_by"), "signed_at": signed.get("signed_at"), "version": signed.get("version", 0),
            "funds": out, "choices": choices, "can_sign": HEAD_OF_DATA in seat_of(viewer)}


def sign_index_map(plane, viewer, picks: dict) -> dict:
    """The head of data signs which index each fund is read against. Picks are {fund_id: series_key}; a fund left out keeps its signed row."""
    from sh_platform.v1.seats import HEAD_OF_DATA, REFUSAL, seat_of
    if HEAD_OF_DATA not in seat_of(viewer):
        raise Refused(403, REFUSAL[HEAD_OF_DATA])
    view = index_map(plane, viewer)
    known, ok = {f["fund_id"] for f in view["funds"]}, {c["series_key"] for c in view["choices"]}
    picks = {str(k): str(v) for k, v in (picks or {}).items()}
    if not picks:
        raise Refused(400, "Choose an index for at least one fund.")
    if set(picks) - known or set(picks.values()) - ok:
        raise Refused(400, "Some of those funds or indexes are not known. Reload and try again.")
    plane.path("consume").mkdir(parents=True, exist_ok=True)
    with lock_for(_map_path(plane)):
        old = _signed_map(plane)
        write_json_atomic(_map_path(plane), {"version": int(old.get("version", 0)) + 1, "rows": {**(old.get("rows") or {}), **picks},
                                              "signed_by": viewer.name, "signed_at": datetime.now(timezone.utc).isoformat()})
    return index_map(plane, viewer)


# ---- Market signals -----------------------------------------------------------------------------------------------------------

def signals(plane, viewer, currency: str = "USD", as_of=None, period: str = "month") -> dict:
    cur, day = currency.upper(), O._day(as_of)
    if cur not in ("USD", "SGD") or period not in PERIODS:
        raise ValueError("bad_request")
    by = _observations(plane, day)
    comp = _safe("benchmarks", lambda: _comparators(plane, by, day.replace(day=1) - timedelta(1)), {})
    rows = _index_rows(plane, by, comp, day, sources=True)
    start, end = _bounds(day, period, min((o[0][0] for o in by.values() if o), default=None))
    pf_rows = _safe("portfolios", lambda: _portfolio_rows(plane, viewer, cur, start, end, period, by, comp, rows), [])
    funds = _safe("funds", lambda: _funds(plane, viewer, day), [])
    sm = _signed_map(plane).get("rows")
    fund_rows = [{**f, **_fund_index(rows, f, period, sm)} for f in funds]
    managers = sorted({f["manager"] for f in funds})
    lead = next((r for r in pf_rows if r["benchmark_label"] and "blend" in r["benchmark_label"].lower()), None) or next(iter(pf_rows), None)
    eq = _pick(rows, "equity")
    read = ""
    if lead and lead["return_pct"] is not None and eq and eq["periods"][period]["change_pct"] is not None:
        read = f"{lead['portfolio']} returned {_pct(lead['return_pct'])} while {eq['name'].lower()} {'rose' if eq['periods'][period]['change_pct'] >= 0 else 'fell'} {_pct(abs(eq['periods'][period]['change_pct']))}."
        lv = _safe("levers", lambda: D._levers(plane, viewer, cur, *_bounds(day, "month", None)), {})
        mv = [m for m in (lv.get("movers") or [])[:2]]
        if period == "month" and len(mv) == 2:
            read += f" The largest contributors were {mv[0]['name']} ({_signed(mv[0]['amount'])}) and {mv[1]['name']} ({_signed(mv[1]['amount'])})."
        read += _fx_sentence(rows) if period == "month" else ""
    note = None
    if comp:
        parts = ", ".join("private markets sleeve" + f" {round(c['weight'] * 100)}%" if c["spread_bps"] else f"{D.SERIES.get(c['series_key'], c['series_key']).lower()} {round(c['weight'] * 100)}%"
                          for c in sorted(comp["components"], key=lambda c: (bool(c["spread_bps"]), -c["weight"])))
        with wh.connect(plane, read_only=True) as con:
            pend = [r[0] for r in con.execute("select version_number from benchmark_definitions where benchmark_code = ? and approval_status <> 'approved' order by version_number", [comp["code"]]).fetchall()]
        note = f"{comp['name']}: {parts}. Approved version {comp['version']}" + (f"; version {pend[-1]} is still being agreed, so it is not used here." if pend else ".")
    return {"currency": cur, "as_of": day, "period": period, "period_start": start, "period_end": end, "refreshed": wh.now(),
            "indexes": rows, "portfolios": pf_rows, "benchmark_note": note, "funds": fund_rows, "read": read or None,
            "news": _safe("news", lambda: _news(plane, viewer, managers), []), "headlines": _safe("headlines", lambda: _headlines(plane), []),
            "outside_note": "Market figures are outside material: context for reading our numbers, never the source of them."}


def _portfolio_rows(plane, viewer, cur: str, start: date, end: date, period: str, by: dict, comp: dict, rows: list[dict]) -> list[dict]:
    out = []
    sleeve = comp.get("sleeve")
    sleeve_pct = sleeve and _sleeve_change(by, sleeve["parts"], start, end)
    proxy = _pick(rows, "proxy")
    for pf in O._portfolios(plane, viewer)[0]:
        r = e3_calc.portfolio_return(plane, pf["id"], start, end, currency=cur)
        ret = r["value"] * 100 if r["value"] is not None else None
        status = _status_weakest(r["statuses"]) if ret is not None else None
        liquid = "liquid" in f"{pf['portfolio_code']} {pf['portfolio_name']}".lower()
        yard = []
        if liquid and comp:
            b = e3_calc.benchmark_return(plane, comp["code"], start, end, None, cur)
            ok = b["value"] is not None and b["definition"] and b["definition"]["approval_status"] == "approved"
            yard.append(("Office blend", b["value"] * 100 if ok else None))
            if sleeve:
                yard.append((sleeve["label"], sleeve_pct))
        elif proxy:
            yard.append((proxy["name"].split(" (")[0], proxy["periods"][period]["change_pct"]))
        for label, bench in yard or [(None, None)]:
            out.append({"portfolio": pf["portfolio_name"], "portfolio_id": pf["id"], "entity": pf.get("entity_name"), "benchmark_label": label,
                        "return_pct": ret, "benchmark_pct": bench, "ahead_pts": ret - bench if ret is not None and bench is not None else None, "status": status,
                        "message": None if ret is not None else NOT_YET, "trail_id": _tid("portfolio", pf["id"], label or "none", period),
                        "where": {"text": f"{pf['portfolio_name']} return between {_dy(start)} and {_dy(end)} against {label or 'no benchmark'}: opening and closing values from the warehouse, "
                                          "index levels recorded there; flows are taken out of the return."}})
    return out


# ---- Fund league (E4) ---------------------------------------------------------------------------------------------------------

TRUST = {"official": 0, "revised": 1, "estimated": 2, "unverified": 3, "stale": 4, "late": 5}


def _fund_counts(funds: list[dict]) -> dict:
    count = lambda *k: sum(f["status_key"] in k for f in funds)  # noqa: E731
    return {"official": count("official"), "revised": count("revised"), "estimated": count("estimated"), "old": count("stale", "unverified", "late"), "total": len(funds)}


def _who(funds: list[dict], *keys: str) -> str:
    """The managers behind funds of those statuses, with a count where a manager has several: "Kestrel Ridge Advisors x2, Cedarpoint Ventures"."""
    c: dict[str, int] = {}
    for f in funds:
        if f["status_key"] in keys:
            c[f["manager"]] = c.get(f["manager"], 0) + 1
    return ", ".join(f"{m} \u00d7{n}" if n > 1 else m for m, n in c.items())


def league(plane, viewer, as_of=None, rank: str = "alpha") -> dict:
    day = O._day(as_of)
    if rank not in ("alpha", "value"):
        raise ValueError("bad_request")
    by = _observations(plane, day)
    comp = _safe("benchmarks", lambda: _comparators(plane, by, day.replace(day=1) - timedelta(1)), {})
    ix = _index_rows(plane, by, comp, day)
    sm = _signed_map(plane).get("rows")
    funds = [{**f, **_fund_index(ix, f, "month", sm)} for f in _funds(plane, viewer, day)]
    funds.sort(key=(lambda f: -(_dec(f["value"]) or ZERO)) if rank == "value" else
               (lambda f: (f["alpha"] is None, -(f["alpha"] or ZERO), TRUST.get(f["status_key"], 9), -(_dec(f["value"]) or ZERO))))
    for i, f in enumerate(funds, 1):
        f["rank"] = i
    n = _fund_counts(funds)
    worst = max((f for f in funds if f["valuation_date"] and f["status_key"] in ("stale", "late")), key=lambda f: (day - f["valuation_date"]).days, default=None)
    start, end = _bounds(day, "month", None)
    return {"as_of": day, "period_end": end, "funds": funds, "managers": len({f["manager"] for f in funds}), "rank": rank,
            "summary": {**n,
                        "ahead": sum(1 for f in funds if f["alpha"] is not None and f["alpha"] > 0) if any(f["alpha"] is not None for f in funds) else None,
                        "ahead_note": "Once September returns load" if not any(f["alpha"] is not None for f in funds) else None},
            "worst": worst and {"fund": worst["fund"], "text": f"{worst['fund']} cannot be ranked fairly. Its last value is from {_dy(worst['valuation_date'])}, {(day - worst['valuation_date']).days} days old.",
                                "href": "/centre/action/managers"},
            "note": "Index returns are for the month, from month-end levels. Fund returns and alpha show as Not available yet until each statement is confirmed."}


# ---- Daily P&L and NAV (E2) ---------------------------------------------------------------------------------------------------

def pnl(plane, viewer, currency: str = "USD", as_of=None) -> dict:
    cur, day = currency.upper(), O._day(as_of)
    if cur not in ("USD", "SGD"):
        raise ValueError("bad_request")
    d = D.daily(plane, viewer, cur, day)
    iv = _safe("issued", lambda: O.issued_vs_known(plane, viewer, cur), {"report": None, "message": NOT_YET})
    other = "SGD" if cur == "USD" else "USD"
    alt = _safe("other currency", lambda: O.issued_vs_known(plane, viewer, other), {})
    alt_total = next((l for l in alt.get("lines", []) if l["label"] == "Total"), None)
    nb = _safe("bridge", lambda: O.nav_bridge(plane, viewer, cur, day), {"portfolios": []})
    lv, rt = d["levers"], d["returns"]
    ytd = (rt.get("periods") or [None, None])[1] if rt.get("available") else None
    bridge = []
    for p in nb.get("portfolios", []):
        why = [x for x in p.get("parts", []) if _dec(x["amount"]) not in (None, ZERO) and x["code"] != "unexplained"]
        bridge.append({"portfolio": p["portfolio_name"], "entity": p.get("entity_name"), "administrator": p.get("administrator_nav"), "ours": p.get("investment_total"),
                       "difference": p.get("difference"), "available": p.get("available", True), "nav_date": p.get("nav_date"),
                       "why": "; ".join(f"{x['label']}: {x['evidence'][0]['label']}" if x.get("evidence") else x["label"] for x in why) or "Nothing to explain.",
                       "trail_id": _tid("bridge", p["portfolio_id"], cur, day), "where": p.get("where")})
    tot_a = sum((_dec(b["administrator"]) for b in bridge if b["administrator"] is not None), ZERO)
    tot_o = sum((_dec(b["ours"]) for b in bridge if b["ours"] is not None), ZERO)
    return {"currency": cur, "as_of": day, "issued": iv, "other_currency": other, "other_total": alt_total and {"issued": alt_total["issued"], "known_now": alt_total["known_now"]},
            "rate": _rate(plane, iv),
            "liquid_book": {"day": {"available": False, "message": f"Not available yet: {DAY_PNL_WAIT[0].lower()}{DAY_PNL_WAIT[1:]}"},
                            "month": {"label": _month_name(_day(lv.get("period_end"))) if lv.get("available") else "", "change": lv.get("change"), "percent": lv.get("percent"),
                                      "available": bool(lv.get("available")), "message": lv.get("message"), "movers": lv.get("movers") or [], "trail_id": lv.get("trail_id"), "where": lv.get("where")},
                            "ytd": {"available": bool(ytd and ytd.get("portfolio") is not None), "percent": ytd and ytd.get("portfolio"), "message": (ytd or {}).get("message") or NOT_YET,
                                    "trail_id": (ytd or {}).get("trail_id")}},
            "bridge": bridge, "bridge_total": {"administrator": tot_a, "ours": tot_o, "difference": tot_o - tot_a} if bridge else None,
            "by_entity": (d["nav"].get("by_entity") or []) if d["nav"].get("available") else [], "total": d["nav"].get("total"), "partial": d["nav"].get("partial"),
            "next": None}


def _rate(plane, iv: dict) -> dict | None:
    rep = iv.get("report")
    pe = _day(rep and rep["period_end"])
    if not pe:
        return None
    r = e3_calc.fx(plane, "USD", "SGD", pe)
    return {"pair": "USD/SGD", "date": pe, "rate": r} if r else None


# ---- Morning Note (E1) --------------------------------------------------------------------------------------------------------

def _moved_rows(lines: list[dict], lv: dict, month: str) -> list[dict]:
    """What moved: report lines that changed since the report was issued, then the liquid book's month and its largest contributors."""
    moved = [{"label": f"{l['label']} changed since the report", "detail": f"Report {_money(l['issued'])}, now {_money(l['known_now'])}", "amount": l["difference"],
              "status": _word(l["status"]), "trail_id": _tid("moved", l["line_key"]), "case_id": l.get("case_id")} for l in lines]
    if lv.get("available"):
        moved.append({"label": f"The liquid book {'rose' if _dec(lv['change']) >= 0 else 'fell'} in {month}", "detail": f"{lv['mover_book'] or 'Liquid book'}: " + ", ".join(
            f"{p['label'].lower()} {_signed(p['amount'])}" for p in lv["parts"] if _dec(p["amount"])), "amount": lv["change"], "status": lv.get("status"), "trail_id": lv["trail_id"]})
        moved += [{"label": f"{'Largest contributor' if i == 0 else 'Next'}: {x['name']}", "detail": x["plain"], "amount": x["amount"], "status": None, "trail_id": x["trail_id"]} for i, x in enumerate(lv["movers"][:2])]
    return moved


def morning_note(plane, viewer, currency: str = "USD", as_of=None) -> dict:
    d = D.daily(plane, viewer, currency, as_of)
    cur, day = d["currency"], d["as_of"]
    iv, nav, rt, lv, mk, lq = d["issued"], d["nav"], d["returns"], d["levers"], d["markets"], d["liquidity"]
    month = _month_name(d["period_end"])
    cases = _cases(plane, viewer)
    limits = [c for c in cases if c["limit"]]
    lib = _safe("reports", lambda: reports_library(plane, viewer), {"reports": []})
    issued_row = next((r for r in lib["reports"] if r.get("versions") and any(v["id"] in (iv.get("brief_link") or "") for v in r["versions"])), None)
    signed = issued_row and next((v for v in issued_row["versions"] if v["id"] in (iv.get("brief_link") or "")), None)
    # headline and summary, worded from the figures
    diff = _dec(iv.get("difference")) if iv.get("available") else None
    if iv.get("available") and iv.get("known_now") is not None:
        head = f"The office is worth {cur} {_millions(iv['known_now'])}"
        head += (f", {_money(abs(diff))} {'more' if diff > 0 else 'less'} than {month}'s report said." if diff else ", the same as the report said.")
        short = f"the office is worth {cur} {_millions(iv['known_now'])}"
    elif nav.get("available") and nav.get("total") is not None:
        head, short = f"The office is worth {cur} {_millions(nav['total'])}.", f"the office is worth {cur} {_millions(nav['total'])}"
    else:
        head, short = NOT_YET, NOT_YET
    lines = _safe("lines", lambda: [l for l in O.issued_vs_known(plane, viewer, cur).get("lines", []) if l["label"] != "Total" and _dec(l["difference"])], [])
    parts = [f"{l['label']} now reads {_money(l['known_now'])} against {_money(l['issued'])} in the report issued {_d(iv.get('issued_at'))}." for l in lines[:2]]
    m = (rt.get("periods") or [None])[0] if rt.get("available") else None
    if m and m.get("portfolio") is not None and m.get("difference") is not None:
        ahead = _dec(m["difference"])
        parts.append(f"The {rt['portfolio']} returned {_pct(m['portfolio'])} in {month}, {abs(ahead):.2f} points {'ahead of' if ahead >= 0 else 'behind'} its benchmark.")
    watch_cases = [c for c in cases if not c["limit"]]
    funds = _safe("funds", lambda: _funds(plane, viewer, day), [])
    loose = [f for f in funds if f["status_key"] in ("estimated", "unverified")]
    unconfirmed = (_dec(lq.get("calls_30")) - (_dec(lq.get("cash")) - _dec(lq.get("cash_after_confirmed")))) if lq.get("available") else None
    tiles = [
        {"key": "value", "label": "Value, known now", "available": bool(iv.get("available")), "value": iv.get("known_now"), "kind": "money", "currency": cur,
         "note": f"{_signed(iv.get('difference'))} since the report" if iv.get("available") else iv.get("message"), "status": iv.get("status"), "trail_id": iv.get("trail_id"), "where": iv.get("where")},
        {"key": "return", "label": f"Liquid book, {month}", "available": bool(m and m.get("portfolio") is not None), "value": m and m.get("portfolio"), "kind": "percent",
         "note": f"Benchmark {_pct(m['benchmark'])} · {_pts(m['difference'])} pts" if m and m.get("portfolio") is not None else (m or {}).get("message") or rt.get("message"),
         "status": rt.get("status"), "trail_id": rt.get("trail_id"), "where": rt.get("where")},
        {"key": "calls", "label": "Capital calls, next 30 days", "available": bool(lq.get("available")), "value": lq.get("calls_30"), "kind": "money", "currency": cur,
         "note": (f"{_money(_dec(lq['cash']) - _dec(lq['cash_after_confirmed']))} confirmed · {_money(unconfirmed)} not yet") if lq.get("available") else lq.get("message"),
         "status": lq.get("status"), "trail_id": lq.get("trail_id"), "where": lq.get("where")},
        {"key": "limits", "label": "Limits over", "available": True, "value": str(len(limits)), "kind": "count", "note": "Open limit cases" if limits else "None over",
         "status": None, "trail_id": _tid("limits", day), "where": {"text": "Open cases raised by the signed limits when they were last checked."}},
    ]
    moved = _moved_rows(lines, lv, month)
    watch = []
    if loose:
        watch.append({"text": f"{len(loose)} fund value{'s' if len(loose) != 1 else ''} not confirmed yet: " + ", ".join(f["fund"] for f in loose) + ". They are shown with their status word.", "trail_id": _tid("watch", "unconfirmed")})
    watch += [{"text": c["title"], "trail_id": _tid("watch", c["number"]), "href": c["href"]} for c in watch_cases[:2]]
    return {"currency": cur, "as_of": day, "period_end": d["period_end"], "company": _company(plane), "dateline": f"{day:%A} {day.day} {day:%B %Y}", "headline": head, "short": short,
            "summary": " ".join(parts) or None, "tiles": tiles, "moved": moved, "watch": watch,
            "needs_you": {"count": len(cases), "items": cases[:3], "href": "/centre/action"},
            "markets": [{"name": s["name"], "kind": s["kind"], "level": s["level"], "month_change_pct": s["month_change_pct"], "day_change_pct": s["day_change_pct"],
                         "outside": True, "trail_id": s["trail_id"]} for s in (mk.get("series") or [])],
            "reports": [{"name": r["name"], "state": r["state"], "version": r["version"], "link": r["link"]} for r in lib["reports"]],
            "built": {"text": f"Built from the company warehouse as of {_dy(day)}." + (f" Report issued {_d(iv.get('issued_at'))}" + (f", signed by {signed['signed_by']}" if signed and signed["signed_by"] else "") + "." if iv.get("issued_at") else ""),
                      "report": iv.get("report")},
            "comment_target": {"kind": "brief_line", "id": f"company/morning-note-{day:%Y-%m-%d}"}}


# ---- Today and Financial health (Consume redesign) ----------------------------------------------------------------------------

DAY_PNL_WAIT = "Waits for the daily price file."


def _plain_index(key: str) -> str:
    return re.sub(r"^global |\s*index$", "", _series_name(key).lower()).strip()


def _markets_line(plane, day: date) -> dict | None:
    """The markets in one sentence from the month-end index changes and the newest exchange rate, with the source of each figure read from the warehouse."""
    by = _observations(plane, day)
    if not by:
        return None
    start, end = _bounds(day, "month", None)
    prov = _provenance(plane, by)
    idx, fxs = [], []
    order = list(D.SERIES)  # our own series first, in the order the office reads them: equity, bonds, the exchange rate
    for key, obs in sorted(by.items(), key=lambda kv: (order.index(kv[0]) if kv[0] in order else len(order), kv[0])):
        kind = obs[0][2]["observation_type"]
        if kind == "index_level" and len(idx) < 3 and (ch := _change(obs, start, end)["change_pct"]) is not None:
            idx.append((key, f"{_plain_index(key)} {_pct(ch, True)}"))
        elif kind == "fx" and not fxs:
            r = obs[-1][2]
            fxs.append((key, f"{r['base_currency']}/{r['quote_currency']} {obs[-1][1]:.4f}"))
    parts = idx + fxs
    if not parts:
        return None
    text = ", ".join(([f"{', '.join(t for _, t in idx)} in {_month_name(end)}"] if idx else []) + [t for _, t in fxs])
    text = text[:1].upper() + text[1:]
    seen, sources = set(), []
    for key, _ in parts:
        p = prov.get(key) or {}
        k = (p.get("label"), p.get("tag"), p.get("date"))
        if k not in seen:
            seen.add(k)
            sources.append({"label": p.get("label"), "tag": p.get("tag"), "date": p.get("date")})
    return {"text": text + ".", "month": _month_name(end), "sources": sources}


def _say(w: dict | None) -> dict | None:
    """A "Where is this from?" block as plain words (the long list of source rows stays with the pages that open it)."""
    return w and {"text": w.get("text")}


def today(plane, viewer, currency: str = "USD", as_of=None) -> dict:
    """The Today page: one sentence, four key numbers, what moved and the markets in a line. Only what the page shows is read."""
    d = D.daily(plane, viewer, currency, as_of, skip=("liquidity", "markets"))
    cur, day = d["currency"], d["as_of"]
    iv, nav, rt, lv = d["issued"], d["nav"], d["returns"], d["levers"]
    month = _month_name(d["period_end"])
    m = (rt.get("periods") or [None])[0] if rt.get("available") else None
    has_bench = bool(m and m.get("portfolio") is not None and m.get("difference") is not None)
    total = nav.get("total") if nav.get("available") else None
    sentence = f"The office is worth {cur} {_millions(total)}." if total is not None else NOT_YET
    if lv.get("available"):
        sentence += f" The liquid book {'rose' if _dec(lv['change']) >= 0 else 'fell'} {_millions(abs(_dec(lv['change'])))} in {month}"
        sentence += (f" and {'beat' if _dec(m['difference']) >= 0 else 'trailed'} its benchmark by {abs(_dec(m['difference'])):.2f} points." if has_bench else ".")
    n_pf, n_en = nav.get("portfolios", 0), nav.get("entities", 0)
    tiles = [
        {"key": "nav", "label": "Net asset value, today", "available": total is not None, "value": total, "kind": "money", "currency": cur,
         "note": f"{n_pf} portfolio{'' if n_pf == 1 else 's'} · {n_en} {'entity' if n_en == 1 else 'entities'}" if total is not None else "Not every portfolio can be valued yet.",
         "status": nav.get("status"), "trail_id": nav.get("trail_id"), "where": _say(nav.get("where"))},
        {"key": "day", "label": "P&L, yesterday", "available": False, "value": None, "kind": "change", "currency": cur, "note": DAY_PNL_WAIT, "status": None,
         "trail_id": _tid("day-pnl", day), "where": {"text": "A day's P&L needs a price for every holding on both days. Until the daily price file arrives only month ends are priced."}},
        {"key": "month", "label": f"P&L, {month}", "available": bool(lv.get("available")), "value": lv.get("change"), "kind": "change", "currency": cur,
         "note": f"Liquid book {_pct(lv.get('percent'))}" if lv.get("available") else lv.get("message"), "status": lv.get("status"), "trail_id": lv.get("trail_id"), "where": _say(lv.get("where"))},
        {"key": "bench", "label": "Against the benchmark", "available": has_bench, "value": m and m.get("difference"), "kind": "points",
         "note": f"Benchmark {_pct(m['benchmark'])}" if has_bench else (m or {}).get("message") or rt.get("message"), "status": rt.get("status"),
         "trail_id": rt.get("trail_id"), "where": _say(rt.get("where"))},
    ]
    lines = [l for l in iv.get("lines", []) if l["label"] != "Total" and _dec(l["difference"])] if iv.get("available") else []
    moved = _moved_rows(lines, lv, month)
    funds = _safe("funds", lambda: _funds(plane, viewer, day), [])
    for f in [f for f in funds if f["status_key"] in ("estimated", "unverified")][:3]:
        moved.append({"label": f["fund"], "detail": f["note"], "amount": f["value"] if f["currency"] == cur else None, "status": f["status"],
                      "trail_id": f["trail_id"], "unconfirmed": True})
    return {"currency": cur, "as_of": day, "period_end": d["period_end"], "dateline": f"{day:%A} {day.day} {day:%B}", "headline": sentence, "tiles": tiles, "moved": moved,
            "markets": _safe("markets line", lambda: _markets_line(plane, day), None),
            "built": f"Built from the company warehouse as of {_dy(day)}." + (f" Report issued {_d(iv.get('issued_at'))}." if iv.get("issued_at") else ""),
            "comment_target": {"kind": "brief_line", "id": f"company/morning-note-{day:%Y-%m-%d}"}}


def _trust(lines: list[dict]) -> dict:
    """How far to trust a value: the weakest status among its lines, and which holdings carry it."""
    st = next((w for w in D.WORST if any(l["data_status"] == w for l in lines)), None)
    if st is None:
        return {"status": "Not available yet", "detail": None}
    names = sorted({l["name"] for l in lines if l["data_status"] == st})
    return {"status": st.capitalize(), "detail": None if st == "official" else ", ".join(names[:2]) + (f" and {len(names) - 2} more" if len(names) > 2 else "")}


def _sum_or_none(vals) -> Decimal | None:
    vals = list(vals)
    return None if not vals or any(v is None for v in vals) else sum(vals, ZERO)


def health(plane, viewer, currency: str = "USD", as_of=None) -> dict:
    """Financial health: each portfolio and entity with value, month P&L, return, the lead over its benchmark and how far the value is trusted;
    the fund counts that the fund league (its own call) opens up."""
    cur, day = currency.upper(), O._day(as_of)
    if cur not in ("USD", "SGD"):
        raise ValueError("bad_request")
    start, end = _bounds(day, "month", None)
    h = O.holdings(plane, viewer, cur, day)
    by = _observations(plane, day)
    comp = _safe("benchmarks", lambda: _comparators(plane, by, end), {})
    ix = _index_rows(plane, by, comp, day)
    lead: dict[str, dict] = {}
    for r in _safe("portfolios", lambda: _portfolio_rows(plane, viewer, cur, start, end, "month", by, comp, ix), []):
        lead.setdefault(r["portfolio_id"], r)  # the first yardstick of a portfolio is its main one
    rows, ents = [], []
    for e in h["entities"]:
        pl = []
        for p in e["portfolios"]:
            r = e3_calc.portfolio_return(plane, p["portfolio_id"], start, end, currency=cur)
            pnl = r["closing"] - r["opening"] - r["flows"] if r["value"] is not None else None
            y = lead.get(p["portfolio_id"])
            row = {"portfolio_id": p["portfolio_id"], "portfolio": p["portfolio_name"], "entity": e["entity_name"], "value": p["total"], "pnl": pnl,
                   "return_pct": r["value"] * 100 if r["value"] is not None else None, "ahead_pts": y and y["ahead_pts"], "benchmark_label": y and y["benchmark_label"],
                   "trust": _trust(p["lines"]), "message": p["message"] or (None if r["value"] is not None else (r["problems"] or [NOT_YET])[0]),
                   "trail_id": _tid("health", p["portfolio_id"], cur, day),
                   "where": {"text": f"{p['portfolio_name']}: the latest value of each of its {len(p['lines'])} holdings on {_dy(day)} in {cur}, from manager statements and administrator files. "
                                     f"Month P&L is the value on {_dy(end)} less the value on {_dy(start)} less confirmed cash flows."}}
            rows.append(row)
            pl.append(row)
        ents.append({"entity_id": e["entity_id"], "entity": e["entity_name"], "portfolios": len(pl), "value": e["total"], "pnl": _sum_or_none(x["pnl"] for x in pl),
                     "trust": _trust([l for p in e["portfolios"] for l in p["lines"]]), "trail_id": _tid("health", "entity", e["entity_id"], cur, day),
                     "where": {"text": f"{e['entity_name']}: the sum of its {len(pl)} portfolios above."}})
    funds = _safe("funds", lambda: _funds(plane, viewer, day), [])
    br = _safe("bridge", lambda: O.nav_bridge(plane, viewer, cur, day), {"portfolios": []})
    open_ = [p for p in br.get("portfolios", []) if p.get("available", True)]
    unexplained = _sum_or_none(next((x["amount"] for x in p.get("parts", []) if x["code"] == "unexplained"), ZERO) for p in open_) if open_ else None
    return {"currency": cur, "as_of": day, "period": {"start": start, "end": end, "month": _month_name(end)}, "portfolios": rows, "entities": ents,
            "total": {"value": h["total"], "pnl": _sum_or_none(r["pnl"] for r in rows), "count": len(rows), "unexplained": unexplained},
            "funds": {**_fund_counts(funds), "estimated_who": _who(funds, "estimated"), "old_who": _who(funds, "stale", "unverified", "late")}}


# ---- the daily email (E6): the template package 4 sends --------------------------------------------------------------------------

def render_email(note: dict, base_url: str = "") -> dict:
    """{subject, html, text} for the Morning Note by email. Pure: takes the dict `morning_note` returns, escapes every value, uses inline styles only and
    no pictures or scripts, so it reads the same in any mail program. `base_url` prefixes the links to Signal Harbor (empty keeps them relative)."""
    e = lambda s: html.escape(str(s if s is not None else ""))  # noqa: E731
    day = note["as_of"] if isinstance(note["as_of"], date) else _day(note["as_of"])
    subject = f"Morning Note, {_d(day)}: {note['short']}"
    cell = lambda t: (f'<td style="padding:10px 12px;border:1px solid #d5dde0;vertical-align:top;width:25%"><div style="font-size:12px;color:#5b6b70">{e(t["label"])}</div>'  # noqa: E731
                      f'<div style="font-size:22px;font-weight:600;color:#14262b">{e(_tile_text(t))}</div><div style="font-size:12px;color:#5b6b70">{e(t.get("note") or "")}</div></td>')
    para = lambda s: f'<p style="margin:12px 0;font-size:14px;line-height:1.5;color:#14262b">{e(s)}</p>'  # noqa: E731
    link = lambda href, text: f'<a href="{e(base_url + href)}" style="color:#0b6b73">{e(text)}</a>'  # noqa: E731
    body = [para(note["summary"])] if note.get("summary") else []
    body += [para(w["text"]) for w in note.get("watch", [])]
    if note["needs_you"]["items"]:
        body.append('<p style="margin:12px 0;font-size:14px;color:#14262b"><strong>Needs you today</strong></p>')
        body += [f'<p style="margin:6px 0;font-size:14px">{e(c["kind"])}: {link(c["href"], c["title"])}</p>' for c in note["needs_you"]["items"]]
    doc = (f'<!doctype html><html><head><meta charset="utf-8"><title>{e(subject)}</title></head><body style="margin:0;background:#f3f6f7;font-family:Arial,Helvetica,sans-serif"><table role="presentation" width="100%" cellpadding="0" cellspacing="0"><tr><td align="center">'
           f'<table role="presentation" width="640" cellpadding="0" cellspacing="0" style="background:#ffffff;border:1px solid #d5dde0"><tr><td style="padding:20px 24px">'
           f'<div style="font-size:12px;color:#5b6b70">Signal Harbor · {e(note.get("company"))}</div>'
           f'<div style="font-size:12px;letter-spacing:1px;color:#0b6b73;margin-top:8px">MORNING NOTE · {e(note["dateline"].upper())}</div>'
           f'<h1 style="font-size:22px;margin:8px 0;color:#14262b">{e(note["headline"])}</h1>'
           f'<table role="presentation" width="100%" cellpadding="0" cellspacing="6"><tr>{"".join(cell(t) for t in note["tiles"])}</tr></table>'
           f'{"".join(body)}<p style="margin:16px 0">{link("/centre/consume", "Open Signal Harbor")} · {link("/centre/consume", "Where is this from?")}</p>'
           f'<p style="margin:16px 0 0;font-size:12px;color:#5b6b70">{e(note["built"]["text"])} Market figures are outside material. Reply to this email to ask a question.</p>'
           f'</td></tr></table></td></tr></table></body></html>')
    text = "\n".join([note["headline"], ""] + [f"{t['label']}: {_tile_text(t)} ({t.get('note') or ''})" for t in note["tiles"]] + [""] + ([note["summary"], ""] if note.get("summary") else []) +
                     [w["text"] for w in note.get("watch", [])] + [f"{c['kind']}: {c['title']}" for c in note["needs_you"]["items"]] + ["", note["built"]["text"]])
    return {"subject": subject, "html": doc, "text": text}


def _tile_text(t: dict) -> str:
    if not t.get("available"):
        return NOT_YET
    v = t["value"]
    return _pct(v) if t["kind"] == "percent" else str(v) if t["kind"] == "count" else f"{t.get('currency', '')} {_millions(v)}" if _dec(v) and abs(_dec(v)) >= 1000000 else f"{t.get('currency', '')} {_money(v)}"
