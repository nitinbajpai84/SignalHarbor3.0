"""T1b: the three documents of an intent laid out from the templates it follows, and their export.

`layout(plane, iid, sow)` gives the intent of work, the process map and the definitions as ordered sections: heading, grey instruction,
example, whether it is required, and the content the statement of work already holds in that section (`filled`). North's patch fills a
section by writing the SOW slot (or `template_fields[key]`) the section is bound to, so the layout needs no copy of its own. The
exports (Markdown, Word, spreadsheet) are written from the same layout, so the template's headings and tables survive. No model calls.
"""
from __future__ import annotations

import io

from sh_platform.v1 import t1_seeds as TS
from sh_platform.v1.i4_sow import kind_word

DOCS = ("intent_of_work", "process_map", "definitions")
TITLES = {"intent_of_work": "Intent of work", "process_map": "Process map", "definitions": "Definitions"}
FREE_HEADING = "Anything else"


def _j(xs) -> str:
    return ", ".join(str(x) for x in xs or [] if x)


def _steps(sow) -> list[dict]:
    return [{"n": s.get("n"), "name": s.get("name") or "", "who": s.get("who") or "", "does": s.get("does") or "", "inputs": _j(s.get("inputs")),
             "outputs": _j(s.get("outputs")), "system": s.get("system") or ""} for s in (sow.get("process") or {}).get("steps") or [] if isinstance(s, dict)]


def _kpi_rows(sow) -> list[dict]:
    names = {s.get("id"): s.get("name") for s in sow.get("sources") or [] if isinstance(s, dict)}
    return [{"name": k.get("name") or "", "definition": k.get("definition") or "", "formula": k.get("formula") or "", "unit": k.get("unit") or "",
             "period": k.get("period") or "", "owner": k.get("owner") or "", "source": _j(names.get(x) for x in k.get("source_ids") or []),
             "_status": k.get("status"), "_suggested": bool(k.get("pack_suggested")), "_edit": {"section": "kpis", "index": i, "field": "definition"}}
            for i, k in enumerate(sow.get("kpis") or []) if isinstance(k, dict)]


def _bench(sow) -> list[str]:
    try:
        from sh_platform.v1 import p4_benchmarks as P4
        return [f"{k.get('name')}: {P4.line(k['benchmark'])}" for k in sow.get("kpis") or [] if isinstance(k, dict) and k.get("benchmark")]
    except Exception:  # noqa: BLE001 - a figure that cannot be shown is left to the person's own words
        return []


def _raw(kind: str, key: str, sow: dict):
    """What the SOW holds for a section bound to it (str, list of str or list of rows), or None when the section has no SOW slot."""
    g = sow.get("goal") or {}
    if kind == "intent_of_work":
        if key == "goal":
            return [x for x in (g.get("line"), f"Why: {g['why']}" if g.get("why") else "") if x]
        if key == "readers":
            return [x for x in (_j(g.get("readers")), f"They decide: {_j(g.get('decisions'))}" if g.get("decisions") else "") if x]
        if key == "frequency":
            return g.get("needed_by") or ""
        if key == "process":
            return [f"{s['n']}. {s['name']}" + (f" ({s['who']})" if s["who"] else "") + (f": {s['does']}" if s["does"] else "") for s in _steps(sow)]
        if key == "kpis":
            return [{c: r[c] for c in ("name", "definition", "formula", "unit", "period", "owner")} for r in _kpi_rows(sow)]
        if key == "inputs":
            return [{"name": s.get("name") or "", "kind": s.get("kind") or "", "holds": s.get("holds") or "", "owner": s.get("owner") or ""}
                    for s in sow.get("sources") or [] if isinstance(s, dict)]
        if key == "rules":
            return [f"{kind_word(r.get('kind'))}: {r.get('rule')}" for r in sow.get("governance") or [] if isinstance(r, dict)]
        if key == "unclear":
            return [u["question"] + (f" ({u['who_can_settle']} can settle it)" if u.get("who_can_settle") else "")
                    for u in sow.get("unclear") or [] if isinstance(u, dict) and u.get("question") and not u.get("settled")]
        if key == "benchmarks":
            return _bench(sow) or None
    elif kind == "process_map":
        if key == "steps":
            return _steps(sow)
        if key == "handovers":
            return [{"from_step": h.get("from_step"), "to_step": h.get("to_step"), "what": h.get("what") or ""} for h in (sow.get("process") or {}).get("handovers") or []]
    elif kind == "definitions" and key == "kpis":
        return _kpi_rows(sow)
    return None


def _has(v) -> bool:
    return bool(v.strip()) if isinstance(v, str) else v not in (None, [], {})


def _rows(raw) -> list[dict]:
    return [r for r in raw if isinstance(r, dict)] if isinstance(raw, list) else []


def _lines(raw) -> list[str]:
    if isinstance(raw, str):
        return [x.strip() for x in raw.splitlines() if x.strip()]
    return [(_row_line(x) if isinstance(x, dict) else str(x)) for x in raw or [] if x not in (None, "")]


def _row_line(r: dict) -> str:
    vals = [str(v) for k, v in r.items() if not k.startswith("_") and v not in (None, "")]
    return (vals[0] + (": " + ", ".join(vals[1:]) if len(vals) > 1 else "")) if vals else ""


def _row_edits(kind: str, key: str, sow: dict) -> list[dict]:
    """Edits for a filled section bound to SOW rows: one per editable cell, with the row's index (the existing section edit route takes them)."""
    def e(label, section, field, index, value):
        return {"label": label, "section": section, "field": field, "index": index, "value": value if isinstance(value, str) else _j(value)}
    if kind == "intent_of_work" and key == "rules":
        return [e(kind_word(r.get("kind")), "governance", "rule", i, r.get("rule") or "") for i, r in enumerate(sow.get("governance") or []) if isinstance(r, dict)]
    if key == "kpis" and kind in ("intent_of_work", "definitions"):
        return [e(f"{r['name']}: {c}", "kpis", c, r["_edit"]["index"], r[c]) for r in _kpi_rows(sow) for c in ("name", "definition", "formula", "unit", "period", "owner")]
    if kind == "intent_of_work" and key == "inputs":
        return [e(f"{r.get('name')}: {c}", "sources", c, i, r.get(c) or "") for i, r in enumerate(sow.get("sources") or []) if isinstance(r, dict) for c in ("name", "holds", "owner")]
    return []


def render_section(kind: str, s: dict, sow: dict) -> dict:
    key, typ = s["key"], s.get("type") or "text"
    tf = sow.get("template_fields") or {}
    raw = _raw(kind, key, sow)
    if not _has(raw):
        raw = tf.get(key)  # what North or the person wrote under the section's key
    if typ == "table":
        rows = _rows(raw)
        cols = s.get("columns") or (list(rows[0]) if rows else [])
        value = {"columns": ["text"], "rows": [{"text": x} for x in _lines(raw)]} if not rows and _has(raw) else {"columns": cols, "rows": rows}
        filled = bool(value["rows"])
    elif typ in ("list",):
        value = _lines(raw)
        filled = bool(value)
    else:
        value = "\n".join(_lines(raw))
        filled = bool(value.strip())
    edits = []
    if kind == "intent_of_work" and key == "goal":
        edits = [{"label": "Goal", "section": "goal", "field": "line", "value": (sow.get("goal") or {}).get("line") or ""},
                 {"label": "Why", "section": "goal", "field": "why", "value": (sow.get("goal") or {}).get("why") or ""}]
    elif kind == "intent_of_work" and key == "frequency":
        edits = [{"label": "Needed by", "section": "goal", "field": "needed_by", "value": (sow.get("goal") or {}).get("needed_by") or ""}]
    elif _raw(kind, key, sow) is None and typ in ("text", "long_text", "list"):
        edits = [{"label": s["heading"], "section": "template_fields", "field": key, "value": value if isinstance(value, str) else chr(10).join(value)}]
    elif filled:
        edits = _row_edits(kind, key, sow)
    came = [x.get("name") for x in sow.get("sources") or [] if isinstance(x, dict) and x.get("kind") in ("file", "document") and x.get("name")] \
        if kind == "intent_of_work" and key == "inputs" and filled else []
    return {**{k: s.get(k) for k in ("key", "instruction", "example", "required", "type", "columns", "labels", "view")}, "heading": s["heading"],
            "filled": filled, "value": value, "edits": edits, "from": came}


def _choices(plane, kind: str) -> list[dict]:
    from sh_platform.v1 import t1_templates as T
    return [{"id": r["id"], "name": (T._snapshot(plane, r["id"], r["version"]) or r)["name"]} for r in T._all(plane)
            if r["kind"] == kind and r.get("state") == "adopted" and r.get("version")]


def layout(plane, iid: str, sow: dict | None = None) -> dict | None:
    """{documents: [{kind, title, template: {id, name, version}, choices, sections, map?}]} or None when no template can be read."""
    from sh_platform.v1 import t1_templates as T
    sow = sow if sow is not None else T.record_for_intent(plane, iid)
    docs = []
    for kind in DOCS:
        t = T.for_intent(plane, iid, kind)
        if not t:
            continue
        secs = [render_section(kind, s, sow) for s in t.get("sections") or []]
        free = t.get("free_text") or {}
        tf = (sow.get("template_fields") or {}).get(T.FREE_KEY)
        secs.append({"key": T.FREE_KEY, "heading": FREE_HEADING, "instruction": "Anything the template does not ask for.", "example": "",
                     "required": bool(free.get("required")), "type": "long_text", "columns": None, "labels": None, "view": None,
                     "filled": _has(tf), "value": "\n".join(_lines(tf)),
                     "edits": [{"label": FREE_HEADING, "section": "template_fields", "field": T.FREE_KEY, "value": tf if isinstance(tf, str) else ""}], "from": []})
        docs.append({"kind": kind, "title": TITLES[kind], "template": {"id": t["id"], "name": t["name"], "version": t["version"]},
                     "choices": _choices(plane, kind) if kind != "intent_of_work" else [], "sections": secs,
                     **({"map": sow.get("process") or {"steps": [], "handovers": []}} if kind == "process_map" else {})})
    return {"documents": docs} if docs else None


# ---- export: the same layout as Markdown, Word and a spreadsheet ------------------------------------------------------------------

def _head(c: str, labels: dict | None) -> str:
    return (labels or {}).get(c) or c.replace("_", " ").capitalize()


def _cells(rows: list[dict], cols: list[str]) -> list[list[str]]:
    return [[str(r.get(c) if r.get(c) is not None else "") for c in cols] for r in rows]


def to_markdown(lay: dict, title: str, note: str) -> str:
    out = [f"# {title}", note, ""]
    for d in lay["documents"]:
        out += [f"## {d['title']} ({d['template']['name']})", ""]
        for s in d["sections"]:
            out += [f"### {s['heading']}", ""]
            if not s["filled"]:
                out += [f"_{s['instruction']}_" + (f" _e.g. {s['example']}_" if s["example"] else ""), ""]
            elif s["type"] == "table":
                cols = s["value"]["columns"]
                out += ["| " + " | ".join(_head(c, s.get("labels")) for c in cols) + " |", "|" + " --- |" * len(cols)]
                out += ["| " + " | ".join(x.replace("|", "/") for x in r) + " |" for r in _cells(s["value"]["rows"], cols)] + [""]
            elif s["type"] == "list":
                out += [f"- {x}" for x in s["value"]] + [""]
            else:
                out += [s["value"], ""]
    return "\n".join(out)


def to_docx(lay: dict, title: str, note: str) -> bytes:
    import docx
    from docx.shared import RGBColor
    d = docx.Document()
    d.add_heading(title, 0)
    d.add_paragraph(note)
    for doc in lay["documents"]:
        d.add_heading(f"{doc['title']} ({doc['template']['name']})", 1)
        for s in doc["sections"]:
            d.add_heading(s["heading"], 2)
            if not s["filled"]:
                p = d.add_paragraph()
                r = p.add_run(s["instruction"] + (f" For example: {s['example']}" if s["example"] else ""))
                r.italic, r.font.color.rgb = True, RGBColor(0x80, 0x80, 0x80)
            elif s["type"] == "table":
                cols = s["value"]["columns"]
                t = d.add_table(rows=1, cols=len(cols))
                t.style = "Table Grid"
                for i, c in enumerate(cols):
                    t.rows[0].cells[i].text = _head(c, s.get("labels"))
                for r in _cells(s["value"]["rows"], cols):
                    cells = t.add_row().cells
                    for i, x in enumerate(r):
                        cells[i].text = x
            elif s["type"] == "list":
                for x in s["value"]:
                    d.add_paragraph(x, style="List Bullet")
            else:
                d.add_paragraph(s["value"])
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def to_xlsx(lay: dict, title: str) -> bytes:
    """One sheet per document: a heading row per section, its table (or lines) under it, empty sections show their grey instruction."""
    from openpyxl import Workbook
    from openpyxl.styles import Font
    wb = Workbook()
    wb.remove(wb.active)
    for doc in lay["documents"]:
        ws = wb.create_sheet(doc["title"][:31])
        ws.append([title]); ws["A1"].font = Font(bold=True, size=14)
        ws.append([f"Template: {doc['template']['name']}"])
        for s in doc["sections"]:
            ws.append([])
            ws.append([s["heading"]]); ws.cell(ws.max_row, 1).font = Font(bold=True)
            if not s["filled"]:
                ws.append([s["instruction"]]); ws.cell(ws.max_row, 1).font = Font(italic=True, color="808080")
            elif s["type"] == "table":
                cols = s["value"]["columns"]
                ws.append([_head(c, s.get("labels")) for c in cols])
                for c in ws[ws.max_row]:
                    c.font = Font(bold=True)
                for r in _cells(s["value"]["rows"], cols):
                    ws.append(r)
            else:
                for x in (s["value"] if s["type"] == "list" else s["value"].splitlines()):
                    ws.append([x])
        ws.column_dimensions["A"].width = 48
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
