"""Harbor 3.0 figures: every figure (KPI) of the signed brief ends in a FILLED ready (gold) table, or in a check line that names the figure and why.

Why this exists (docs/v3/FIGURES-DIAGNOSIS.md): the design pack was drafted before any file was mapped, so its ready tables read the standard model tables
(valuation observations), not the tables the load really filled (position snapshots); a file that did not load was still called loaded; a figure with no
table was a quiet note in the pack; and the answer path went to Lumen, who found empty tables. Nothing here invents a figure, a table or a rule: designs are
Span's own (b10_design, rules first, the model for what the rules cannot place), the numbers come from the ready tables, and a missing figure says so.

Hooks (v3_harbor): `load_file` in b_load; `cover` inside a_pack (before submit); chain B gets `b_figures` after the run on a copy
(`install`); Decision 3 gets one check line per figure; `figure_answer` runs in the ask path before Lumen. All of it needs the Ctx `x` of the chain only
for `x.api / x.plane / x.th / x.app`; this module imports v3_harbor lazily so the two never import each other at load time.
"""
from __future__ import annotations

import logging
import os
import re
from datetime import date
from decimal import Decimal

log = logging.getLogger(__name__)
DESIGN, ADDS_UP, ANSWER = "model.design-ready-table", "check.adds-up", "consume.answer-question"
LOADED = ("loaded", "already_loaded")
META = {"id", "organization_id"}
MONTHS = {m: i for i, m in enumerate(["january", "february", "march", "april", "may", "june", "july", "august", "september", "october", "november", "december"], 1)}
GENERIC = {"the", "and", "for", "portfolio", "fund", "entity", "account", "file", "total", "what", "was", "were", "how", "much", "our", "are", "worth", "family"}  # every entity of a family office is the family's: "each family member" names none (r3b uc-100)
CCY_COL = re.compile(r"^(?:currency|ccy)(?:_code)?$", re.I)  # a currency code is a unit of the figure, never a member to filter by or a name to print
UNIT_WORDS = {"total", "number", "value", "amount", "rate", "net", "gross", "average", "change", "count", "each", "per"}
LIMIT = 6  # figures designed in one step


# ---- the brief's figures and the warehouse ------------------------------------------------------------------------------------------

def _D():
    from sh_platform.v1 import b10_design
    return b10_design


def kpis_of(plane, iid: str | None) -> list[dict]:
    """The figures of the signed brief, as written there."""
    if not iid:
        return []
    from sh_platform.v1 import i4_sow
    try:
        return [k for k in i4_sow.read(plane, iid).get("kpis") or [] if isinstance(k, dict) and k.get("name")]
    except Exception:  # noqa: BLE001 - no brief, no figures
        log.info("no brief to read the figures from", exc_info=True)
        return []


def table_rows(plane, copy: bool = False) -> dict[str, int]:
    """Rows of every table of the warehouse's main schema (the live one; `copy`: the trial-run copy when there is one)."""
    from sh_platform.v1 import b9_verify as V, wh
    con = V._copy(plane) if copy else wh.connect(plane, read_only=True)
    try:
        names = [r[0] for r in con.execute("select table_name from information_schema.tables where table_schema = 'main'").fetchall()]
        return {n: con.execute(f'select count(*) from "{n}"').fetchone()[0] for n in names if re.fullmatch(r"[A-Za-z_][A-Za-z0-9_]*", n)}
    finally:
        con.close()


def fact_names() -> set[str]:
    """The tables a load writes facts into (the intake's own list); the rest of the model is master data that is there before any file."""
    from sh_platform.v1 import e0c_load
    return {*getattr(e0c_load, "_TABLES", {}), "transactions"}


def fact_tables(rows: dict[str, int]) -> set[str]:
    """The fact tables that hold rows now."""
    return {t for t in fact_names() if rows.get(t)}


# ---- which ready table covers a figure ----------------------------------------------------------------------------------------------

def _slug(s) -> str:
    return re.sub(r"[^a-z0-9]+", "_", str(s).lower()).strip("_")


def _tokens(s) -> set[str]:
    return _D()._toks(str(s).replace("_", " "))


def covers(kpi: dict, g: dict) -> bool:
    """A gold table (pack or plan form) holds the figure: it lists it, has a column named for it, or has words for most of what the figure is called."""
    name = str(kpi.get("name") or "")
    if name.lower() in {str(k).lower() for k in g.get("kpis") or []}:
        return True
    want = _tokens(name)
    if not want:
        return False
    cols = [str(c.get("name") if isinstance(c, dict) else c) for c in g.get("columns") or []] + list((g.get("col_defs") or {}))
    if _slug(name) in {_slug(c) for c in cols}:
        return True
    have = _tokens(" ".join([g.get("name", ""), g.get("purpose", ""), *cols]))
    key = {w for w in want if w not in UNIT_WORDS} or want
    return len(key & have) / len(key) >= .5 and len(want & have) >= 1


def covering(kpi: dict, gold: list[dict]) -> list[dict]:
    """The tables that hold the figure. A figure bound to a recipe is held by its recipe's table only, once there is one (c7: "Cash variance per entity"
    was read from the cash-by-entity table, whose words cover it)."""
    from sh_platform.v1 import v3_recipes as R
    r = R.bound(kpi) if "built_from" in kpi else None
    mine = [g for g in gold if r and str(g.get("purpose") or "").startswith(f"{R.MARK}{r['id']}:")]
    return mine or [g for g in gold if covers(kpi, g)]


def sources_of(g: dict, gold: list[dict], seen: set | None = None) -> set[str]:
    """The silver tables a gold table finally reads (through other gold tables of the pack)."""
    seen = seen or set()
    by, out = {t["name"]: t for t in gold}, set()
    for d in g.get("depends_on") or []:
        if d in seen:
            continue
        seen.add(d)
        out |= sources_of(by[d], gold, seen) if d in by else {d}
    return out


def dead(g: dict, gold: list[dict], rows: dict[str, int]) -> bool:
    """Every fact table the ready table reads is empty (master data such as portfolios does not make a table live)."""
    src = sources_of(g, gold) & fact_names()
    return bool(src) and not any(rows.get(s) for s in src)


def plain(name: str, gold: list[dict] | None = None) -> str:
    """A table as people say it: without 'gold_' and, when the pack's tables share one, the project's name that starts each of them."""
    names = [g["name"] for g in gold or []]
    cut = len(os.path.commonprefix(names).rpartition("_")[0]) + 1 if len(names) > 1 else 0
    return re.sub(r"^gold_(prj_[0-9a-f]{8}_)?", "", name[cut:] if cut > 5 and name.startswith(names[0][:cut]) else name).replace("_", " ")


# ---- designing a ready table for a figure --------------------------------------------------------------------------------------------

def dims_of(kpi: dict, goal: str) -> list[str]:
    """What the figure is split by, from its own words ("net worth by portfolio", "for each entity") or the goal line that names it; the 2.0 design only
    keeps the ones the company model can reach."""
    own = f"{kpi.get('name') or ''} {kpi.get('definition') or ''}".lower()
    found = [m.group(1) for m in re.finditer(r"\b(?:by|per|for each|each|across)\s+([a-z]+)", own)]
    if not found and _tokens(kpi.get("name")) & _tokens(goal):  # the goal says "net worth by portfolio"
        found = [m.group(1) for m in re.finditer(r"\b(?:by|per)\s+([a-z]+)", str(goal).lower())]
    skip = {"the", "month", "day", "week", "year", "quarter", "each", "all", "a", "an", "default"}
    return list(dict.fromkeys(w for w in found if w not in skip))


def viewer_of(x, who: str = "leo"):
    """The 2.0 viewer of the seat Harbor acts as (the data engineer), for the in-process calls that take one."""
    from sh_platform.profile.view import Viewer
    from sh_platform.v1 import v3_harbor as H
    st = x.app.state
    acct = st.account_store.get_by_email(f"{H.PEOPLE[who]}.{x.plane.company_id}@example.test") or st.account_store.get(x.account_id) or {}
    return Viewer(str(acct.get("display_name") or ""), frozenset(acct.get("roles") or []), False, acct.get("account_id"))


def projections(sql: str):
    """The select's columns as sqlglot expressions ([] when the text is not a readable select)."""
    import sqlglot
    from sqlglot import exp
    try:
        q = sqlglot.parse_one(sql, read="duckdb")
        return list((q.expression if isinstance(q, exp.Create) else q).expressions)
    except Exception:  # noqa: BLE001 - the plan check has already accepted it; nothing more to say here
        return []


def null_columns(sql: str) -> set[str]:
    """Columns a select fills with NULL or a fixed value ("cast(null as numeric) as drift", "12 as time_weighted_return"): a placeholder, not a
    design. r9: Span's model design gave a return of 12 and a drift of 100 to every row of the commitments; a figure is read from records, never typed."""
    from sqlglot import exp
    out = set()
    for e in projections(sql):
        inner = e.this if isinstance(e, exp.Alias) else e
        while isinstance(inner, exp.Cast):
            inner = inner.this
        if isinstance(inner, (exp.Null, exp.Literal)):
            out.add(e.alias_or_name)
    return out


AVG = re.compile(r"\bavg\s*\(", re.I)


def _close(text: str, i: int) -> int:
    """The index of the ")" that closes the "(" at i, or -1."""
    depth = 0
    for j in range(i, len(text)):
        if text[j] == "(":
            depth += 1
        elif text[j] == ")":
            depth -= 1
            if depth == 0:
                return j
    return -1


def avg_numbers(sql: str) -> str:
    """Every avg in a designed select reads its argument as a number. The warehouse has no avg of yes/no values: "avg(signatory_compliance) * 100" over
    a boolean column stopped the run on the copy (r3b uc-094, 'avg(BOOLEAN)'); a number's avg is the same either way. The text is unchanged when no
    avg needs it or it cannot be read. r5 fix: the SQL is edited where the avg is, never parsed and written out again: a parse-and-write round trip
    rewrote to_char(d, 'YYYY-MM') as CAST(d AS TEXT), which changes the result (the monthly keys of the tables)."""
    out, pos = [], 0
    for m in AVG.finditer(sql):
        if m.start() < pos:  # an avg inside the argument of an avg just read
            continue
        open_i = m.end() - 1
        close = _close(sql, open_i)
        if close < 0:
            return sql  # cannot be read: the plan check reads it and says what is wrong
        arg = sql[open_i + 1:close].strip()
        if re.match(r"(?i)cast\s*\(", arg) and _close(arg, arg.index("(")) == len(arg) - 1:
            continue  # already a cast of its argument
        prefix = "DISTINCT " if re.match(r"(?i)distinct\s", arg) else ""
        if prefix:
            arg = arg[8:].strip()
        out += [sql[pos:m.start()], "avg(", prefix, f"CAST({arg} AS DOUBLE)", ")"]
        pos = close + 1
    if pos == 0:
        return sql
    return "".join(out) + sql[pos:]


def has_column(kpi_name: str, cols: list[str]) -> bool:
    """A column of the table is named for the figure (most of its distinctive words are in the column names)."""
    want = {w for w in _tokens(kpi_name) if w not in UNIT_WORDS} or _tokens(kpi_name)
    return bool(want) and len(want & _tokens(" ".join(cols))) / len(want) >= .5


def holds(t: dict, figure: str) -> bool:
    """A designed table has a column for the figure. A table built from the figure's recipe (or from the design the office kept for it) holds it by what
    its columns mean, not by their names: "Time-weighted return net of fees" is portfolio_return_pct and benchmark_return_pct, "Single-name breaches and
    cost to cure" is concentration_pct, limit_pct and cost_to_cure (live r2, uc4 and uc5: the recipe's table was dropped for "no column")."""
    from sh_platform.v1 import v3_recipes as R
    cols = [e.alias_or_name for e in projections(t["sql"])]
    return not cols or str(t.get("purpose") or "").startswith((R.MARK, R.SAVED)) or has_column(figure, cols)


def design(x, todo: list[dict], filled: set[str]) -> tuple[list[dict], dict[str, str]]:
    """Span's design step for these figures (skill model.design-ready-table): the rules first, on the tables the load filled; what they cannot place goes
    to the model once. -> (new gold tables as the plan check passed them, {figure: why none}). Nothing is written yet."""
    from sh_platform.build.agents.runtime import get_model
    from sh_platform.v1 import b6_span as S, b10_pack as P, p1_project as P1
    D, plane, pid = _D(), x.plane, x.th["pid"]
    viewer = viewer_of(x)
    sow = S.signed_sow(plane, P1.read(plane, pid))
    goal = str((sow.get("goal") or {}).get("line") or "")
    grain = P._grain(sow)[0]
    names = [k["name"] for k in todo]
    kp = []
    for k in todo:
        dims = dims_of(k, goal)
        kp.append({**k, "benchmark_by": k.get("benchmark_by") or dims})
    new: dict[str, dict] = {}
    why: dict[str, str] = {}
    for model in (None, "model"):  # rules first; the model only for the figures the rules left
        left = [k for k in kp if not any(k["name"] in (t.get("kpis") or []) for t in new.values())]
        if not left:
            break
        base = S.draft_plan(plane, viewer, pid, None)
        before = {t["name"] for t in base["tables"]}
        base["notes"] = [n for n in base.get("notes") or [] if "could not place" not in n] + [f"Span could not place the KPI '{k['name']}' on a table that holds data." for k in left]
        with D.prefer(filled):
            got = D.design(plane, viewer, pid, {**sow, "kpis": left}, base, grain, get_model(x.app.state, plane) if model else None)
        for t in got["tables"]:
            if t.get("layer") == "gold" and t["name"] not in before:
                new[t["name"]] = {**t, "sql": avg_numbers(t["sql"]), **({"by_model": True} if model else {})}
        for n in got.get("notes") or []:
            m = re.search(r"Define the KPI '([^']+)': (.*)", n)
            if m and m.group(1) in names:
                why[m.group(1)] = m.group(2)[:300]
    for name, t in list(new.items()):  # a column of NULLs is not a figure; a table with no column for a figure does not hold it: the figures are named, the table is not kept
        empty = {_slug(n) for n in null_columns(t["sql"])}
        held = [n for n in t.get("kpis") or [] if _slug(n) not in empty and holds(t, n)]
        for n in t.get("kpis") or []:
            if n not in held:
                why[n] = ("Span could not work it out from the loaded data: the only design it found was an empty or fixed column." if _slug(n) in empty
                          else "Span could not work it out from the loaded data: its design had no column for it.")
        if held:
            new[name] = {**t, "kpis": held}
        else:
            del new[name]
    for t in new.values():
        for n in t.get("kpis") or []:
            why.pop(n, None)
    return list(new.values()), why


def defined(t: dict) -> dict[str, str]:
    """A definition for every column of a designed table: the figure's own from the design, the period and the split in plain words (nothing to settle)."""
    out = dict(t.get("col_defs") or {})
    try:
        names = [e.alias_or_name for e in projections(t["sql"])] or [str(c.get("name") if isinstance(c, dict) else c) for c in t.get("columns") or []]
    except Exception:  # noqa: BLE001 - the plan check reads it again; a column left without a definition is settled below
        names = [str(c.get("name") if isinstance(c, dict) else c) for c in t.get("columns") or []]
    for n in names:
        out.setdefault(n, "The first day of the period the figure belongs to." if n in ("month", "day", "week", "quarter", "year")
                       else "The currency the value is recorded in." if n == "currency" else f"The {n.replace('_', ' ')} the figure is for.")
    return out


def cover(x, phase: str = "draft") -> str:
    """Make sure each figure of the brief has a ready table in the design pack. phase "draft" (chain A, the pack not yet submitted): a figure with no table
    gets one designed. phase "loaded" (chain B, after the load): also a figure whose table reads only empty tables while the load filled others. The pack
    gets the new tables (a new version when it was approved: submitted and approved as the head of data), and x.th["figures"] says what was done."""
    from sh_platform.v1 import v3_recipes, v3_skills as SK
    th, A = x.th, f"/api/v1/build/p/{x.th['pid']}/architecture"
    kpis = kpis_of(x.plane, th.get("iid"))
    if not kpis:
        return "the brief names no figures"
    pk = x.api("leo", "GET", A + "/pack").get("pack") or {}
    gold, rows = pk.get("target", {}).get("gold") or [], table_rows(x.plane)
    filled = fact_tables(rows)
    todo, notes = [], {}
    for k in kpis:
        cov = covering(k, gold)
        if not cov or v3_recipes.stale(k, cov):  # A7: a figure with a recipe is built from it, also when the rules' table already covers it
            todo.append(k)
        elif phase == "loaded" and filled and not valued(x.plane, k, cov):  # built on the copy with the loaded data, and still no value for it
            todo.append(k)
    held = {n for h in (th.get("held_tables") or {}).values() if h.get("answered") != "used" for n in h.get("kpis") or []}
    todo = [k for k in todo if k["name"] not in held]  # r9: a table held for the person (or left out by them) is not designed again
    th["figures"] = [{"kpi": k["name"], "table": (covering(k, gold) or [{}])[0].get("name")} for k in kpis]
    if not todo:
        return f"all {len(kpis)} figures have a ready table"
    with SK.record(x.plane, th["id"], DESIGN) as rec:
        try:
            new, notes = design(x, todo[:LIMIT], filled)
        except Exception as e:  # noqa: BLE001 - a design that fails leaves the pack as it was; the figure is named at Decision 3
            log.warning("designing ready tables failed", exc_info=True)
            new, notes = [], {k["name"]: f"Span could not design it: {str(getattr(e, 'message', e))[:200]}" for k in todo}
        rec["detail"] = f"{len(new)} ready tables for {', '.join(k['name'] for k in todo[:LIMIT])}"
    have = {g["name"] for g in gold}
    for t in [t for t in new if t["name"] in have]:  # the design came back with a table the pack already has: that is not a new design
        for n in t.get("kpis") or []:
            notes[n] = "Span's design repeats a table the pack already has, which holds nothing."
    new = [t for t in new if t["name"] not in have]
    added = []
    th["model_tables"] = {**(th.get("model_tables") or {}), **{t["name"]: True for t in new if t.get("by_model")}}
    for t in new:
        try:
            x.api("leo", "PUT", A + "/pack/edit", {"section": "gold_add", "key": t["name"], "note": "Span designed a ready table for " + ", ".join(t.get("kpis") or []) + ".",
                  "changes": {"purpose": t.get("purpose") or "", "sql": t["sql"], "depends_on": list(t.get("depends_on") or []), "grain": t.get("grain") or "",
                              "keys": list(t.get("keys") or []), "definitions": defined(t)}})
            added.append(t)
        except Exception as e:  # noqa: BLE001 - the plan check refused it: the reason goes with the figure
            for n in t.get("kpis") or []:
                notes[n] = f"The design was refused: {str(getattr(e, 'text', e))[:200]}"
    new = added
    th["figure_notes"] = {**(th.get("figure_notes") or {}), **notes}
    if new and pk.get("state") == "approved":  # the edit forked a new draft: the same two steps as the first pack
        from sh_platform.v1 import v3_rows
        v3_rows.send(x)  # the new table's open rows are settled first (r9), else the submit refuses
        from sh_platform.v1 import v3_harbor as H
        H.b_run(x)  # the new pack version gets its own round (sources connected, built on the copy) before anyone can go live on it
        fresh(x)
    th["figures"] = [{"kpi": k["name"], "table": (covering(k, [*gold, *new]) or [{}])[0].get("name")} for k in kpis]
    return f"{len(new)} ready tables added" + (f"; no table for {', '.join(notes)}" if notes else "")


# ---- after the run: does each figure's table have rows? ---------------------------------------------------------------------------------

def measure_col(kpi: dict, g: dict) -> str | None:
    """The column of a ready table that holds the figure: named for it, else the number column sharing the most words with it."""
    cols = [c for c in g.get("columns") or [] if isinstance(c, dict)]
    exact = next((c["name"] for c in cols if _slug(c["name"]) == _slug(kpi["name"])), None)
    nums = [c["name"] for c in cols if str(c.get("type") or "").upper().startswith(("DECIMAL", "NUMERIC", "DOUBLE", "FLOAT", "BIGINT", "INTEGER", "HUGEINT"))]
    want = _tokens(kpi["name"])
    return exact or max((c for c in nums if want & _tokens(c)), key=lambda c: len(want & _tokens(c)), default=None)


def counts(con, table: str, col: str | None) -> tuple[int | None, int]:
    """(rows, rows with a value in `col`) of a table of the trial-run copy; rows None when the copy has no such table."""
    if not re.fullmatch(r"[A-Za-z_][A-Za-z0-9_]*", table) or (col and not re.fullmatch(r"[^\"]+", col)):
        return None, 0
    try:
        n, nn = con.execute(f'select count(*), count("{col}") from "{table}"' if col else f'select count(*), count(*) from "{table}"').fetchone()
        return n, nn
    except Exception:  # noqa: BLE001 - the table was not built
        return None, 0


def valued(plane, kpi: dict, cov: list[dict]) -> bool:
    """A ready table of the trial-run copy holds a value for the figure (check's own test). Rows alone are not enough: a table that joins every portfolio
    to an empty fact table has a row per portfolio and a NULL figure in each (test cycle 2, head_of_data: net worth was never redesigned on the loaded positions)."""
    from sh_platform.v1 import b9_verify as V
    try:
        con = V._copy(plane)
    except Exception:  # noqa: BLE001 - no copy to read: leave the design as it is; check() names the figure
        log.warning("the trial-run copy could not be read", exc_info=True)
        return True
    try:
        return any(counts(con, g["name"], measure_col(kpi, g))[1] for g in cov)
    finally:
        con.close()


def check(x) -> str:
    """Decision 3's figure lines (skill check.adds-up): for each figure of the brief, the rows of its ready table in the trial-run copy and whether the
    column that holds the figure has values; an empty, missing or all-empty table gets the figure's name and the reason. Stored in x.th["figures"]."""
    from sh_platform.v1 import b9_verify as V
    th = x.th
    kpis = kpis_of(x.plane, th.get("iid"))
    pk = x.api("leo", "GET", f"/api/v1/build/p/{th['pid']}/architecture/pack", soft=True).get("pack") or {}
    gold = pk.get("target", {}).get("gold") or []
    try:
        con, silver = V._copy(x.plane), table_rows(x.plane)
    except Exception:  # noqa: BLE001 - a busy copy: the lines say so
        log.warning("the trial-run copy could not be read", exc_info=True)
        con, silver = None, {}
    problems = [e.get("title", "") for e in th.get("exceptions") or []]
    out = []
    try:
        for k in kpis:
            got = [(g, *(counts(con, g["name"], measure_col(k, g)) if con else (None, 0))) for g in covering(k, gold)]
            best = max(got, key=lambda r: (r[2], r[1] or 0), default=None)
            reason, n, nn, g = "", 0, 0, None
            if best is None:
                reason = (th.get("figure_notes") or {}).get(k["name"]) or "No ready table was designed for it."
            else:
                g, n, nn = best[0], best[1], best[2]
                if n is None:
                    reason = f"The run did not build {plain(g['name'], gold)}."
                elif n == 0:
                    empty = sorted(s for s in sources_of(g, gold) if not silver.get(s))
                    reason = (f"{plain(g['name'], gold)} has no rows: it reads {', '.join(plain(s) for s in empty)}, which hold nothing yet." if empty
                              else f"{plain(g['name'], gold)} has no rows.")
                    reason += (" Not loaded: " + "; ".join(problems[:4]) + ".") if problems else ""
                    reason += (" No file yet for: " + ", ".join(th["left_out"]) + ".") if th.get("left_out") else ""
                elif nn == 0:
                    reason = f"{plain(g['name'], gold)} has {n} rows but none holds a value for {k['name']}."
            out.append({"kpi": k["name"], "table": g["name"] if g else None, "words": plain(g["name"], gold) if g else None, "rows": n or 0, "ok": bool(nn), "reason": reason,
                        **({"by_model": True} if g and (th.get("model_tables") or {}).get(g["name"]) else {})})
    finally:
        if con is not None:
            con.close()
    th["figures"] = out
    return f"{sum(f['ok'] for f in out)} of {len(out)} figures have rows"


def d3_lines(th: dict) -> list[dict]:
    return [{"label": f["kpi"], "value": f"{f['rows']} rows" if f["ok"] else "Not available yet", "ok": f["ok"],
             "note": (("Ready table " + (f.get("words") or plain(f["table"])) + f" has {f['rows']} rows.") + (" Span drafted it with the model from the brief's wording: check its definition." if f.get("by_model") else "")) if f["ok"] else f["reason"]} for f in th.get("figures") or [] if "ok" in f]


def install(H) -> None:
    """Called once by v3_harbor when it is loaded: the two chain B steps, their skills and the figure lines of Decision 3."""
    H.SKILLS_OF.update(b_figures=(ADDS_UP,))
    steps, final = H.CHAINS["B"]
    out = []
    for s in steps:
        out.append(s)
        if s[2].__name__ == "b_run":
            out.append(("Span", "Made sure each figure has a filled ready table", b_figures))
    H.CHAINS["B"] = (out, final)
    d3 = H.DECISIONS["decision3"]

    def decision3(x):
        d = d3(x)
        lines = d3_lines(x.th)
        if lines:
            d["checks"] = [*d.get("checks", []), *lines]
            bad = [l["label"] for l in lines if not l["ok"]]
            if bad:
                d["title"] = "Some figures have no numbers yet. Go live?"
                d["summary"] = (d.get("summary") or "") + f" {len(bad)} of {len(lines)} figures have no numbers yet: {', '.join(bad)}."
        return d
    H.DECISIONS["decision3"] = decision3
    from sh_platform.v1 import v3_recipes
    v3_recipes.install(H)  # A7: the tie-out lines of the figures built from a recipe, and the model designs that worked are kept


def fresh(x) -> None:
    """Build the pack on the trial-run copy again, now: the round Harbor reads at "Ran on a copy" may have run when the pack was approved, before a file
    was loaded."""
    x.api("leo", "POST", f"/api/v1/build/p/{x.th['pid']}/run?mode=sandbox", {}, soft=True)


def b_figures(x):
    """After the load and the run on a copy: the copy is rebuilt with the loaded data, each figure without rows gets a ready table designed on the tables the
    load filled (a new pack version, approved as the head of data, built again), and every figure is checked for rows."""
    fresh(x)
    done = cover(x, "loaded")
    return f"{done}; {check(x)}"


# ---- loading, honestly -----------------------------------------------------------------------------------------------------------------

def load_file(x, name: str, fid: str) -> tuple[bool, str]:
    """Load one intake file and say what really happened. A mapping proposed on the way (a second layout) is approved as the head of data, like the
    others; fund names that are exactly the instrument's own are approved in one click (2.0 approve-exact) and the file is loaded again. -> (loaded, why not)."""
    st, msg = "", ""
    for _ in range(4):
        out = x.api("leo", "POST", f"/api/v1/intake/files/{fid}/load", {})
        last = out.get("last_load") or out
        st, msg = str(last.get("status") or out.get("state") or ""), str(last.get("message") or "")
        if st in LOADED:
            return True, msg
        if st == "waiting_for_mapping" and last.get("proposal_id"):
            x.api("anita", "POST", f"/api/v1/model/mappings/{last['proposal_id']}/approve", {"note": "Approved by the head of data."}, soft=True)
            continue
        if (last.get("needs_approval") or {}).get("funds") and (x.api("leo", "POST", "/api/v1/evidence/aliases/approve-exact", soft=True) or {}).get("approved"):
            continue
        break
    held = last.get("held_back") or []
    return False, (msg or f"{name} did not load ({st or 'no answer'}).") + (f" Row {held[0].get('row')}: {held[0].get('reason')}." if held else "")


ACCOUNT = re.compile(r"account|acct|iban", re.I)


def account_questions(x) -> list[dict]:
    """Decision 2: one plain question per account number a file carries that no portfolio of the company has ("Which entity owns account X?", the
    entities as options). Without the answer every row of the bank statement is held back. Read from the files themselves, before anything loads."""
    from sh_platform.v1 import e0b_tools, e0c_load, wh
    plane, th = x.plane, x.th
    try:
        with wh.connect(plane, read_only=True) as con:
            known = {str(r[0]).strip().lower() for r in con.execute("select portfolio_code from portfolios").fetchall()}
            ents = [r[0] for r in con.execute("select entity_name from legal_entities order by entity_name").fetchall()]
        out: dict[str, str] = {}
        for mid, m in (th.get("maps") or {}).items():
            path = plane.path("dropbox", m["name"])
            if not ents or not path.exists() or path.suffix.lower() not in (".csv", ".xlsx", ".xlsm"):
                continue
            prof = e0b_tools.profile_sheet(path, None)
            _, cols, rows, _ = e0c_load._read(path, prof.get("sheet"), int(prof.get("header_row") or 1), {})
            col = next((c for c in cols if ACCOUNT.search(c)), None)
            for _, vals, _ in rows if col else []:
                code = str(vals.get(col) or "").strip()
                if code and code.lower() not in known:
                    out.setdefault(code, m["name"])
    except Exception:  # noqa: BLE001 - no account question is better than a broken Decision 2; the load still names the account it held back
        log.info("could not read the account numbers", exc_info=True)
        return []
    return [{"id": f"acct:{c}", "text": f"Which entity owns account {c}?", "options": ents, "assumed": None} for c in sorted(out)]


def tie_account(x, code: str, entity: str) -> str:
    """The person's answer: account `code` belongs to `entity`. It becomes a portfolio of that entity named by its account number (master data written
    once), and the proposed mappings that read the account column as anything else now read it as the portfolio code."""
    from sh_platform.v1 import e1_write, wh
    plane = x.plane
    with wh.connect(plane) as con:
        ent = con.execute("select id, base_currency from legal_entities where entity_name = ?", [entity]).fetchone()
        if not ent:
            return f"{entity} is not an entity of the company"
        if not con.execute("select 1 from portfolios where lower(portfolio_code) = lower(?)", [code]).fetchone():
            con.execute("insert into portfolios (id, organization_id, legal_entity_id, portfolio_code, portfolio_name, base_currency, active_from) values (?,?,?,?,?,?,?)",
                        [wh.new_id(), e1_write.org_id(plane), ent[0], code, f"Account {code} ({entity})", ent[1], date(2020, 1, 1)])
    for mid in list(x.th.get("maps") or {}):
        d = x.api("leo", "GET", f"/api/v1/model/mappings/{mid}", soft=True) or {}
        for c in d.get("columns") or []:
            if ACCOUNT.search(str(c.get("source"))) and f"{c.get('target_table')}.{c.get('target_field')}" != "portfolios.portfolio_code" and d.get("state") == "proposed":
                x.api("leo", "PUT", f"/api/v1/model/mappings/{mid}", {"columns": [{"source": c["source"], "target": "portfolios.portfolio_code"}]}, soft=True)
    return f"account {code} tied to {entity}"


def pick_source(plane, name: str, sources: list[str], th: dict) -> str | None:
    """Two sources of one brief fit a file name equally: the file's own words (title, sheet names, column headings) against what each source holds."""
    from sh_platform.v1 import i4_sow
    try:
        held = {s.get("name"): f"{s.get('name')} {s.get('holds') or ''}" for s in i4_sow.read(plane, th.get("iid")).get("sources") or [] if isinstance(s, dict)}
        path = plane.path("dropbox", name)
        from openpyxl import load_workbook
        if name.lower().endswith((".xlsx", ".xlsm")):
            wb = load_workbook(path, read_only=True, data_only=True)
            cells = [*wb.sheetnames] + [str(c) for ws in wb.worksheets for r in ws.iter_rows(max_row=12, values_only=True) for c in r if isinstance(c, str)]
        else:
            cells = path.read_text(encoding="utf-8-sig", errors="replace").splitlines()[:3]
    except Exception:  # noqa: BLE001 - no reading, no pick: the person is asked
        return None
    words = _tokens(" ".join(cells))
    scored = sorted(((len(_tokens(held.get(s, s)) & words), s) for s in sources), reverse=True)
    return scored[0][1] if scored and scored[0][0] > 0 and (len(scored) == 1 or scored[0][0] > scored[1][0]) else None


# ---- the answer path: the ready table of the figure the question names -------------------------------------------------------------------

def parse_when(text: str) -> tuple[int | None, int, int | None] | None:
    """(year, month, day) the question asks for: "30 September 2026", "September 2026", "2026-09", "8 October" (no year: None). None when it names no month."""
    t = text.lower()
    m = re.search(r"\b(\d{4})-(\d{2})(?:-(\d{2}))?\b", t)
    if m:
        return int(m[1]), int(m[2]), int(m[3]) if m[3] else None
    m = re.search(r"\b(?:(\d{1,2})(?:st|nd|rd|th)?\s+)?(" + "|".join(MONTHS) + r")\b(?:\s*,?\s*(\d{4}))?", t)
    return (int(m[3]) if m[3] else None, MONTHS[m[2]], int(m[1]) if m[1] else None) if m else None


def kpi_for(text: str, kpis: list[dict]) -> dict | None:
    """The figure the question names: one of its name's words is in the question, and of those the one whose name and definition share the most words with
    it, a word that only one figure uses counting most ("cash break" is Breaks, not Cash balance, though both have "cash")."""
    asked_by = _tokens(" ".join(re.findall(r"\b(?:which|by|per|each|every)\s+([a-z]+)", text.lower())))  # "which entity": a way to split a figure, not the figure
    q = _tokens(text) - asked_by
    words = {k["name"]: (_tokens(k["name"]), _tokens(f"{k['name']} {k.get('definition') or ''}")) for k in kpis}
    df: dict[str, int] = {}
    for _, full in words.values():
        for w in full:
            df[w] = df.get(w, 0) + 1
    best, top = None, 0.0
    for k in kpis:
        name, full = words[k["name"]]
        score = sum(1 / df[w] for w in q & full)
        if q & name and score > top:
            best, top = k, score
    return best


def aliases(plane) -> dict[str, str]:
    """code (and id) -> name for every master table of the company model that has an X_code and an X_name (portfolios, entities ...)."""
    from sh_platform.v1 import e0a_catalogue as E, wh
    out: dict[str, str] = {}
    try:
        with wh.connect(plane, read_only=True) as con:
            for t in E.current(plane)["tables"]:
                names = {c["name"] for c in t["columns"]}
                for c in sorted(n for n in names if n.endswith("_code") and n[:-5] + "_name" in names):
                    out.update({str(a): str(b) for a, b in con.execute(f'select "{c}", "{c[:-5]}_name" from "{t["name"]}"').fetchall() if a and b})
                    if "id" in names:  # a ready table may keep the master's id (a uuid) instead of its code
                        out.update({str(a): str(b) for a, b in con.execute(f'select "id", "{c[:-5]}_name" from "{t["name"]}"').fetchall() if a and b})
    except Exception:  # noqa: BLE001 - names are a convenience; the codes still match
        log.info("no names for the codes", exc_info=True)
    return out


def masters(plane) -> list[tuple[str, list[set[str]]]]:
    """(noun, one set of keys per row: code, id, name) for every master table of the company model that has an X_code and an X_name."""
    from sh_platform.v1 import e0a_catalogue as E, wh
    out = []
    try:
        with wh.connect(plane, read_only=True) as con:
            for t in E.current(plane)["tables"]:
                names = {c["name"] for c in t["columns"]}
                for c in sorted(n for n in names if n.endswith("_code") and n[:-5] + "_name" in names):
                    sel = f'"{c}", "{c[:-5]}_name"' + (', "id"' if "id" in names else "")
                    out.append((c[:-5].replace("_", " "), [{str(v) for v in r if v} for r in con.execute(f'select {sel} from "{t["name"]}"').fetchall()]))
    except Exception:  # noqa: BLE001 - no coverage line without the masters
        log.info("no master tables for coverage", exc_info=True)
    return out


def coverage(plane, members: set[str]) -> str:
    """"across the 3 portfolios loaded (of 5)" when the members a figure adds up are fewer than the master table's rows; "" when they are all, or unknown."""
    best = ""
    for noun, rows in masters(plane):
        hit = [r for r in rows if r & members]
        if hit and len(hit) < len(rows) and members <= set().union(*hit):
            best = f"across the {len(hit)} {noun}s loaded (of {len(rows)})"
    return best


CODE = re.compile(r"[^\s]*[\d_\-/.:#][^\s]*|[A-Z0-9]{2,}")  # no space and a digit or separator (EG-HOLD, FX-0007, a uuid), or all capitals (SGD)


def words_of(v, name=None) -> set[str]:
    """The words a member is known by: its name's, and its own when it is itself a name. A code or id has no words: "hold" in "we hold SGD" is not
    EG-HOLD (a code only matches whole, see named_in)."""
    return (set() if v is None or CODE.fullmatch(str(v)) else _tokens(v)) | (_tokens(name) if name else set())


def named_in(v, name, q: set[str], text: str) -> bool:
    """The question names this member: a word of its name (stop words aside: "all" is not "All funds"), or the whole code ("EG-HOLD", "SGD") standing on its
    own in the question."""
    whole = bool(v) and bool(CODE.fullmatch(str(v))) and not re.fullmatch(r"[0-9a-f-]{36}", str(v)) and bool(re.search(r"(?<![a-z0-9])" + re.escape(str(v).lower()) + r"(?![a-z0-9])", text.lower()))
    return bool(whole or (words_of(v, name) & q) - GENERIC)


def match_values(values: list[str], names: dict[str, str], q: set[str], text: str = "") -> list[str]:
    """The values of a dimension the question names (the value's own words or its name's); [] when it names none or all alike. A value whose whole
    name is in the question's own words ("Fund A") wins over a partial match: a one-letter name is no word."""
    scored = []
    for v in values:
        best = 0.0
        if any(c and re.search(r"(?<![a-z0-9])" + re.escape(str(c).lower()) + r"(?![a-z0-9])", text.lower()) for c in (None if re.fullmatch(r"[0-9a-f-]{36}", v) else v, names.get(v))):
            best = 2.0
        for c in (v, names.get(v)):
            t = words_of(c) if c else set()
            if t and (t & q) - GENERIC:
                best = max(best, len(t & q) / len(t))
        if best >= .5:
            scored.append((best, v))
    if not scored:
        return []
    top = max(s for s, _ in scored)
    got = [v for s, v in scored if s == top]
    return [] if len(got) == len(values) and len(values) > 1 else got


def fake_split(rows: list[dict], texts: list[str], when: str | None, meas: str) -> str | None:
    """The text column whose members all carry the very same figures (day by day), or None: a table whose design joined everything to everything (an
    entity column beside a total of the whole company) repeats one total under each name. Columns that follow one another one to one (an id and its name) count as one."""
    one = lambda c, d: len({(r.get(c), r.get(d)) for r in rows}) == len({r.get(c) for r in rows}) == len({r.get(d) for r in rows})  # noqa: E731
    for c in texts:
        if len({r.get(c) for r in rows}) < 2:
            continue
        rest = [d for d in texts if d != c and not one(c, d)]
        seen: dict = {}
        for r in rows:
            seen.setdefault(r.get(c), set()).add((str(r.get(when)) if when else "", tuple(str(r.get(d)) for d in rest), str(r.get(meas))))
        if len({frozenset(v) for v in seen.values()}) == 1 and len(next(iter(seen.values()))) > 1:  # one shared line could be chance
            return c
    return None


def _num(v) -> Decimal | None:
    return v if isinstance(v, Decimal) else (Decimal(str(v)) if isinstance(v, (int, float)) and not isinstance(v, bool) else None)


def _fmt(v: Decimal, money: bool) -> str:
    return f"{v:,.2f}" if money else f"{v:,.4f}".rstrip("0").rstrip(".")


def _month_words(d) -> str:
    return f"{list(MONTHS)[d.month - 1].capitalize()} {d.year}"


def figure_answer(plane, viewer, text: str, pid: str | None, th: dict | None = None) -> dict | None:
    """The deterministic answer: the figure the question names, read from its ready table, ahead of Lumen. None when the question names no figure of the
    brief or the table cannot be read (Lumen answers); an empty table gives "Not available yet" and why. Logged as a skill run."""
    if not pid or not th:
        return None
    from sh_platform.v1 import v3_figstore as FS, v3_skills as SK
    try:
        with SK.record(plane, th.get("id"), ANSWER) as rec:
            hit = _answer(plane, viewer, text, pid, th)
            rec["detail"] = (hit or {}).get("source", "not a question about a figure of the brief")
            if hit and (kpi := hit.pop("kpi", None)) and (link := FS.link_answer(plane, th, kpi)):
                hit["figure"] = {**(hit.get("figure") or {}), **link}  # "Where this number comes from": the figure record the answer was read through
    except Exception:  # noqa: BLE001 - any failure leaves the answer to Lumen
        log.info("figure answer failed", exc_info=True)
        return None
    return hit


CHANGE_Q = re.compile(r"\bwhy\b.*\b(?:move[ds]?|chang(?:e|ed|es)|differ\w*|shift\w*)\b|\bsince\b.*\b(?:report|issued|statement|last)\b|\bwhat (?:has |have )?chang(?:e|ed)\b|\b(?:move|chang)ed since\b", re.I)
CHANGE_COLS = (re.compile(r"original|previous|issued|reported|before"), re.compile(r"revis|amend|restat|updated|after|current"), re.compile(r"delta|change|differen"))


def change_columns(cols: dict[str, str]) -> tuple[str, str, str] | None:
    """(original, revised, delta) columns of a table that says what moved since the issued report, or None."""
    num = [c for c, ty in cols.items() if ty.lower().startswith(("decimal", "numeric", "double", "float", "bigint", "integer", "hugeint"))]
    got = [next((c for c in num if rx.search(c.lower())), None) for rx in CHANGE_COLS]
    return tuple(got) if all(got) and len(set(got)) == 3 else None  # type: ignore[return-value]


def _money(v: Decimal, sign: bool = False) -> str:
    return f"{v:+,.0f}" if sign and v == v.to_integral() else f"{v:+,.2f}" if sign else f"{v:,.0f}" if v == v.to_integral() else f"{v:,.2f}"


def _revisions(plane, text: str, trace: list[dict]) -> dict | None:
    """No ready table of original and revised values answers: the valuations themselves. An instrument with two different values for one date is a
    revision. When a row supersedes the other the original and the amended value are named, with the document the amended row cites; otherwise both
    values and their difference, the manager's documents of that date, and the word that the rows do not say which came later. The instrument is the one
    the question names (its name, or an identifier on record such as "Fund A - Evergreen Growth I" or a code), else every revised one. A question that
    names an instrument also gets its other valuations, newest first; "since the last meeting" starts at the date in the meeting note's file name."""
    from sh_platform.v1 import wh
    asked, q = parse_when(text), _tokens(text) - GENERIC
    with wh.connect(plane, read_only=True) as con:
        got = con.execute("select i.canonical_name, i.manager_name, v.valuation_date, v.value_amount, v.id, v.supersedes_valuation_id, d.original_filename "
                          "from valuation_observations v join instruments i on i.id = v.instrument_id left join documents d on d.id = v.evidence_document_id "
                          "order by 1, 3, v.value_amount").fetchall()
        ids = con.execute("select i.canonical_name, x.identifier_value from instrument_identifiers x join instruments i on i.id = x.instrument_id "
                          "where x.identifier_type <> 'security_id'").fetchall()
        docs = con.execute("select original_filename, valuation_date from documents where valuation_date is not null").fetchall()
        meet = con.execute("select original_filename from documents where regexp_matches(lower(original_filename), 'meeting|onsite') order by 1 desc").fetchall()
    tied = {n for n, v in ids if any(re.search(r"(?<![a-z0-9])" + re.escape(c.lower()) + r"(?![a-z0-9])", text.lower()) for c in {str(v), str(v).split(" - ")[0]} if len(c) > 2)}
    since = None
    if re.search(r"last (?:meeting|visit)|since we (?:met|last met)", text, re.I) and (m := next((re.search(r"\d{4}-\d{2}-\d{2}", f[0]) for f in meet if re.search(r"\d{4}-\d{2}-\d{2}", f[0])), None)):
        since = date.fromisoformat(m[0])
    groups: dict = {}
    for name, mgr, day, val, vid, sup, doc in got:
        if (asked and day.month != asked[1]) or (since and day < since):
            continue
        groups.setdefault((name, day), []).append((val, str(vid), str(sup) if sup else None, doc, mgr))
    revised = {k: v for k, v in groups.items() if len({x[0] for x in v}) > 1}
    named = {k: v for k, v in revised.items() if k[0] in tied or _tokens(k[0]) & q}
    explicit = {k[0] for k in groups if k[0] in tied or _tokens(k[0]) & q}
    revised = named or revised
    if not revised:
        return None
    lines, rows = [], []
    for (name, day), v in sorted(revised.items(), key=lambda kv: kv[0][1], reverse=True)[:6]:
        link = next(((a, b) for a in v for b in v if a[2] == b[1]), None)
        who = f"{name} at {day.day} {_month_words(day)}"
        if link:
            new, old = link
            line = f"{who}: original {_money(old[0])}, amended {_money(new[0])}, change {_money(new[0] - old[0], True)}" + (f" (amended per {new[3]})" if new[3] else "")
        else:
            lo, hi = min(x[0] for x in v), max(x[0] for x in v)
            mt = _tokens(v[0][4] or "") - GENERIC
            by = sorted({f for f, d in docs if d == day and mt & _tokens(re.sub(r"[_.-]", " ", f))})
            line = (f"{who}: two values on record, {_money(lo)} and {_money(hi)}, which differ by {_money(hi - lo)} (the rows do not say which came later"
                    + (f"; the manager's documents of that date: {', '.join(by)}" if by else "") + ")")
        lines.append((day, line))
        rows.append((day, {"label": f"{name} {day.isoformat()}", "value": line.split(": ", 1)[1]}))
    seen = {k for k in revised}
    for (name, day), v in sorted(groups.items(), key=lambda kv: kv[0][1], reverse=True):
        if name in explicit and (name, day) not in seen and len(lines) < 12:
            lines.append((day, f"{name} at {day.day} {_month_words(day)}: {_money(v[0][0])}"))
            rows.append((day, {"label": f"{name} {day.isoformat()}", "value": _money(v[0][0])}))
    lines, rows = [x for _, x in sorted(lines, key=lambda t: t[0], reverse=True)], [x for _, x in sorted(rows, key=lambda t: t[0], reverse=True)]
    return {"answer": "; ".join(lines) + ".", "figure": None, "rows": rows, "source": "valuations on record", "used": [],
            "trace": [*trace, {"kind": "tool", "label": "Compared the valuations on record for the same date"}]}


def _change_answer(plane, text: str, pid: str, th: dict, gold: list[dict], live: dict[str, dict]) -> dict | None:
    """"Why did X move since the issued report / what changed": the ready table that holds original, revised and delta values, read for the fund,
    portfolio and month the question names. An empty table says so and why; no such table leaves the question to Lumen."""
    from sh_platform.v1 import cx_dash as CD
    tabs = [(live[g["name"]], cc) for g in gold if g["name"] in live and (cc := change_columns({c["name"]: str(c["type"]) for c in live[g["name"]]["columns"]}))]
    if not tabs:
        return None
    trace = [{"kind": "tool", "label": "Read the ready table of original, revised and changed values"}, {"kind": "guard", "label": "Every figure comes from the table's rows"}]
    t, (o, r_, d) = max(tabs, key=lambda tc: tc[0]["rows"])
    if not t["rows"]:
        if hit := _revisions(plane, text, trace):
            return hit
        why = next((f["reason"] for f in th.get("figures") or [] if f.get("table") == t["name"] and f.get("reason")), "")
        return {"answer": f"Not available yet: the ready table of original and revised values ({plain(t['name'])}) has no rows." + (" " + why if why else _not_loaded(th)),
                "figure": None, "rows": [], "source": f"No rows in {t['name']}", "used": [], "trace": trace}
    cols = {c["name"]: str(c["type"]).lower() for c in t["columns"]}
    when = next((c for c, ty in cols.items() if ty.startswith(("date", "timestamp"))), None)
    names = aliases(plane)
    _, rows, _ = CD.run_sql(plane, f'select * from "{t["name"]}"' + (f' order by "{when}" desc' if when else ""))
    texts = [c for c, ty in cols.items() if ty.startswith(("varchar", "text", "char", "uuid")) and c not in META and not CCY_COL.match(c) and not c.endswith("run_id")
             and (not c.endswith("_id") or any(str(x.get(c)) in names for x in rows))]
    asked, q = parse_when(text), _tokens(text)
    if when and asked:
        rows = [x for x in rows if str(x[when])[5:7] == f"{asked[1]:02d}" and (not asked[0] or str(x[when])[:4] == str(asked[0]))] or rows
    for c in texts:  # the fund or manager the question names: its rows, whatever their date (uc7: Kestrel Ridge is July, the table's latest is September)
        pick = match_values(sorted({str(x[c]) for x in rows if x.get(c) is not None}), names, q, text)
        if pick:
            rows = [x for x in rows if str(x.get(c)) in pick]
    if when and rows and not asked:  # no month named: each fund's own latest row, not the latest month of the table
        key = (lambda x: str(x.get(texts[0]))) if texts else (lambda x: "")
        latest: dict[str, str] = {}
        for x in rows:
            latest[key(x)] = max(latest.get(key(x), ""), str(x[when])[:10])
        rows = [x for x in rows if str(x[when])[:10] == latest[key(x)]]
    rows = [x for x in rows if _num(x.get(o)) is not None and _num(x.get(r_)) is not None]
    if not rows:
        return {"answer": "Not available yet: the table of original and revised values has no row for what you named.", "figure": None, "rows": [], "source": t["name"], "used": [], "trace": trace}
    money = "decimal" in cols[r_]
    lines = []
    for x in rows[:12]:
        who = " · ".join(dict.fromkeys(names.get(str(x[c]), str(x[c])) for c in texts if x.get(c) is not None))
        lines.append((f"{who}{' ' + _month_words(date.fromisoformat(str(x[when])[:10])) if when and x.get(when) else ''}".strip(), x))
    fmt = lambda v: _fmt(_num(v), money)  # noqa: E731
    say = "; ".join(f"{a + ': ' if a else ''}issued {fmt(x[o])}, revised {fmt(x[r_])}, change {fmt(x[d]) if _num(x.get(d)) is not None else fmt(_num(x[r_]) - _num(x[o]))}" for a, x in lines)
    first = lines[0][1]
    ch = fmt(first[d]) if _num(first.get(d)) is not None else fmt(_num(first[r_]) - _num(first[o]))
    return {"answer": say + ".", "figure": {"label": (lines[0][0] or "Change since the issued report"), "value": fmt(first[r_]), "change": ch},
            "rows": [{"label": a, "value": f"{fmt(x[o])} to {fmt(x[r_])} ({fmt(x[d]) if _num(x.get(d)) is not None else fmt(_num(x[r_]) - _num(x[o]))})"} for a, x in lines],
            "source": plain(t["name"]), "used": [], "trace": trace}


MEMO_Q = re.compile(r"\b(?:support\w*|recommend\w*|memo)\b", re.I)  # "what figures support recommending X": every figure that has a row for X
BREACH_Q = re.compile(r"\b(?:over|above|exceed\w*|breach\w*|cure)\b", re.I)  # "over" is a stop word of the recipe match, so it is read here
RELEASE_Q = re.compile(r"\bwhich\b[^.?!]*\b(?:release|pay)\b", re.I)  # "which payments to release", "which to pay": every state of the batch
REPORT_Q = re.compile(r"\b(?:report|pack)\b", re.I)
TOTAL_Q = re.compile(r"\b(?:worth|total|net worth|nav|value)\b", re.I)


def report_total(plane, text: str) -> dict | None:
    """"What are we worth in the latest report": the report's own Total, read from the fields v3_loads stored (extracted_fields) for the latest report
    period, with the document named. Two documents of the same period (a workbook and its PDF) count once; two different totals in one currency are not
    answered (None: the ready table or Lumen answers). A currency code in the question picks that total; otherwise every total the report gives."""
    from sh_platform.v1 import wh
    if not (REPORT_Q.search(text) and TOTAL_Q.search(text)):
        return None
    with wh.connect(plane, read_only=True) as con:
        rows = con.execute("select d.original_filename, d.valuation_date, f.field_name, f.value_numeric, f.currency from extracted_fields f join documents d on d.id = f.document_id "
                           "where f.value_numeric is not null and d.valuation_date is not null and lower(f.field_name) like 'total%' "
                           "and coalesce(f.review_status, '') <> 'rejected' order by d.valuation_date desc, d.original_filename, f.page_number, f.line_number").fetchall()
    rows = [r for r in rows if re.fullmatch(r"total(?![_a-z]).*", str(r[2]).lower())]  # the report's own total ("Total", "Total (SGD) at WMR 4pm"), not a fund's total_commitment line
    if not rows:
        return None
    day = rows[0][1]
    rows = [r for r in rows if r[1] == day]
    want = {c for c in re.findall(r"\b[A-Z]{3}\b", text) if any(c == r[4] for r in rows)}
    tot: dict[str, Decimal] = {}
    for _, _, _, v, ccy in rows:
        if want and ccy not in want:
            continue
        if tot.setdefault(ccy or "", _num(v)) != _num(v):
            return None
    if not tot:
        return None
    docs = sorted({r[0] for r in rows})
    say = "; ".join(f"{c + ' ' if c else ''}{_fmt(v, False)}" for c, v in tot.items())
    when = f"{day.day} {_month_words(day)}"
    return {"answer": f"The latest report ({', '.join(docs)}, to {when}) gives a total of {say}.",
            "figure": {"label": f"Total in the latest report · {when}", "value": say.split("; ")[0], "change": ""}, "rows": [{"label": c or "Total", "value": _fmt(v, False)} for c, v in tot.items()],
            "source": f"{docs[0]} · the report's Total, as issued", "used": [],
            "trace": [{"kind": "tool", "label": "Read the totals stored from the latest report"}, {"kind": "guard", "label": "Every figure comes from the report's fields"}]}


HOLDING = _tokens("holding holdings position positions security securities issuer instrument fund")
NOT_WORDS = "missing overdue outstanding unpaid unmatched unreconciled failed breach"  # stemmed by _tokens where used
MISSING_Q = re.compile(r"\b(?:missing|overdue|never (?:came|arrived)|not (?:yet )?(?:received|arrived))\b.*\bstatements?\b"
                       r"|\bstatements?\b.*\b(?:missing|overdue|never (?:came|arrived)|not (?:yet )?(?:received|arrived))\b", re.I)


def missing_statements(plane, viewer, text: str) -> dict | None:
    """"Which statement is missing and how overdue is it?": each fund's statements that never arrived, named, with business days overdue, from the
    Managers inbox (k5_managers: the signed statement_due rule, weekdays). A fund with fewer than two statement dates of its own takes the cadence of
    its manager's statements together (Cedarpoint: one fund's March and the other's June make quarterly), else one March statement reads as monthly and
    April to August all show as missing. A viewer who may see every portfolio sees every fund, also one whose loaded values carry no portfolio."""
    if not MISSING_Q.search(text):
        return None
    from sh_platform.profile.view import Viewer
    from sh_platform.v1 import e1_access, k5_managers as K5, wh
    if not e1_access.is_admin(viewer):
        seen = {str(p) for p in e1_access.visible_portfolios(plane, viewer)}
        with wh.connect(plane, read_only=True) as con:
            every = {str(r[0]) for r in con.execute("select id from portfolios").fetchall()}
        if every and every <= seen:
            viewer = Viewer(viewer.name, frozenset(viewer.roles) | {"company_admin"}, False, getattr(viewer, "account_id", None))
    box = K5.inbox(plane, viewer, None)
    if not box["due_rule_signed"]:
        return {"answer": "Not available yet: " + box["rule_note"], "figure": None, "rows": [], "source": "Managers inbox", "used": [], "trace": []}
    docs = K5._statements(plane)
    out = []
    for m in box["managers"]:
        dates = lambda f: {K5._day(d["valuation_date"]) for d in docs.get(f["fund_id"], []) if d["valuation_date"]}  # noqa: E731
        mgr = K5._cadence([d for f in m["funds"] for d in dates(f)])
        for f in m["funds"]:
            cad = f["cadence"] if len(dates(f)) > 1 else mgr
            due = f["earlier_missing"] + ([{"period_end": f["period_end"], "due": f["due"], "days_late": f["days_late"]}] if f["state"] == "late" else [])
            out += [(f["fund_name"], e) for e in due if cad == "monthly" or date.fromisoformat(e["period_end"]).month in (3, 6, 9, 12)]
    trace = [{"kind": "tool", "label": "Read the managers' statements against the signed due rule"}, {"kind": "guard", "label": "Every date comes from the statements received"}]
    if not out:
        return {"answer": f"No statement is missing: every statement due by {K5._pretty(box['as_of'])} has arrived.", "figure": None, "rows": [],
                "source": "Managers inbox", "used": [], "trace": trace}
    out.sort(key=lambda p: -p[1]["days_late"])
    lines = [f"{n}, period {K5._pretty(e['period_end'])}, due {K5._pretty(e['due'])}, {e['days_late']} business days overdue" for n, e in out]
    return {"answer": f"Missing statement{'s' if len(out) > 1 else ''} at {K5._pretty(box['as_of'])}: " + "; ".join(lines) + ".",
            "figure": {"label": f"Missing statement · {out[0][0]}", "value": f"{out[0][1]['days_late']} business days overdue", "change": ""},
            "rows": [{"label": f"{n} · {K5._pretty(e['period_end'])}", "value": f"{e['days_late']} business days overdue"} for n, e in out][:12],
            "source": "Managers inbox · the signed statement due rule", "used": [], "trace": trace}


def _answer(plane, viewer, text: str, pid: str, th: dict) -> dict | None:
    from sh_platform.v1 import b10_pack as P, cx_dash as CD, talk
    kpis = kpis_of(plane, th.get("iid"))
    if pid not in {r["id"] for r in talk.reachable(plane, viewer)[0]}:
        return None
    if CHANGE_Q.search(text):  # why did X move, since the issued report, what changed: the table of original, revised and delta values, ahead of a figure by name
        gold = ((P.latest_approved(plane, pid) or {}).get("target") or {}).get("gold") or []
        hit = _change_answer(plane, text, pid, th, gold, {t["name"]: t for t in CD.gold_tables(plane, pid)})
        if hit:
            return hit
        if hit := _revisions(plane, text, [{"kind": "tool", "label": "No ready table of original and revised values: read the valuations"}, {"kind": "guard", "label": "Every figure comes from the valuation rows"}]):
            return hit
    if hit := report_total(plane, text):  # the report's own total, ahead of a ready table that may hold only part of it
        return hit
    if hit := missing_statements(plane, viewer, text):  # c4 fund_accountant: a count of statements received is not which one is missing
        return hit
    if RELEASE_Q.search(text) and (hit := recipe_rows(plane, text, pid, th, kpis, floor=2)):  # r1 uc-008: which to release is every state of the batch, release and hold
        return hit
    if MEMO_Q.search(text) and (hit := recipe_rows(plane, text, pid, th, kpis, floor=0, wide=True)):  # r3 uc-040: a memo's figures, each said for the entity named
        return hit
    from sh_platform.v1 import v3_consume as C, v3_rule_checks as RC
    calls = bool(C.CALLS.search(text)) and not C.OTHER.search(text)
    if calls and (hit := recipe_rows(plane, text, pid, th, kpis)):  # w6: can we pay the call, how much to convert: the recipes' rows, never the old forecast
        return hit
    k = RC.kpi_by_recipe(text, kpis) or kpi_for(text, kpis)  # the question's words name a recipe: the brief's figure built from it, whatever North called it
    if k is None:
        return recipe_rows(plane, text, pid, th, kpis, floor=2) or recipe_rows(plane, text, pid, th, kpis, floor=1)  # w6: no figure named ("which payments to release"): the recipes the question reads as; r3b uc-073: a bound table with rows is never left to the rules-only reply
    from sh_platform.v1 import v3_recipes as R
    k = _breach_figure(text, kpis, k) or k  # r3 uc5: "which holding is over its limit" is the breach table, not the concentration figure that shares its words
    if h := R.held(th, k["name"]):  # Decision 3 held it (a failed tie-out, a missing input): never stated as fact
        return {"answer": f"Not available yet: {k['name']}: {h['reason']}", "figure": None, "rows": [], "source": f"{k['name']} is held at Decision 3", "used": [],
                "trace": [{"kind": "guard", "label": "A figure held at Decision 3 is not stated"}]}
    mine = relevance(text, R.bound(k)) if R.bound(k) else 0
    named = named_relevance(text, R.bound(k)) if R.bound(k) else 0
    if max((relevance(text, r) for x in kpis if (r := R.bound(x))), default=0) > max(mine, 1) and max(
            (named_relevance(text, r) for x in kpis if (r := R.bound(x))), default=0) > named and (rows := recipe_rows(plane, text, pid, th, kpis, floor=2)):
        return rows  # w6 (uc-005 "Account number match ... 250,000"): the figure's name met a word, another figure's recipe fits the question's meaning
    pack = P.latest_approved(plane, pid) or {}
    gold = (pack.get("target") or {}).get("gold") or []
    live = {t["name"]: t for t in CD.gold_tables(plane, pid)}
    tabs = [live[g["name"]] for g in covering(k, gold) if g["name"] in live]
    ready = [t for t in tabs if t["rows"]]
    if ready and (hit := _every_row(plane, text, pid, th, k)):  # r2: every row of the set, from the recipe's own lines
        return hit
    if (asked := parse_when(text)) and asked[1] and any(_by_period(plane, t) for t in ready) and (rows := recipe_rows(plane, text, pid, th, [k], floor=1)):
        return rows  # r1 uc-029: a month named, over a table kept by month: the figure's own lines say it, for that month (not the flag's first row)
    run = next((s["run_id"] for s in reversed(th.get("steps") or []) if s.get("run_id")), None)
    trace = [{"kind": "tool", "label": f"Read the ready table for {k['name']}"}, {"kind": "guard", "label": "Every figure comes from the table's rows"}]
    if not ready and (rows := recipe_rows(plane, text, pid, th, kpis, floor=2)):
        return rows
    if not ready:
        why =next((f["reason"] for f in th.get("figures") or [] if f.get("kpi") == k["name"] and f.get("reason")), "") if tabs or gold else ""
        say = f"Not available yet: {k['name']} has no figure in the ready tables." + (" " + why.split(" Not loaded")[0] if why else "")
        return {"answer": say, "figure": None, "rows": [], "source": f"No rows for {k['name']}", "used": [], "trace": trace}
    meas = recipe_measures(k, gold)  # c6 ops_clerk: a recipe-built figure is read from its own table first, from the column the recipe names
    tabs = sorted(ready, key=lambda t: (t["name"] not in meas, -t["rows"]))
    says = any(isinstance(a, dict) and a.get("all") for a in (R.bound(k) or {}).get("answer") or []) and not (parse_when(text) or (None, 0))[1]
    for t, how in [*((t, RC.answer) for t in tabs), *([(None, "recipe")] if says else []), *((t, None) for t in tabs)]:  # a limit or a tolerance on record, on any table, before the recipe's own answer lines (r3: every key number said), before a plain read
        if how == "recipe":  # r3 uc-026, uc-033: the recipe's answer lines say every figure the question asks for, which the plain read leaves to the rows behind
            if (rows := recipe_rows(plane, text, pid, th, [k], floor=0)) and not rows["answer"].endswith("the figures are below."):
                return rows
            continue
        got = how(plane, k, t, text, th) if how else _read(plane, t, k, text, aliases(plane), th, meas.get(t["name"]))
        if got and str(got.get("answer") or "").startswith("Not available yet") and (rows := recipe_rows(plane, text, pid, th, kpis, floor=2)):
            return rows  # w6 (uc-001 "no split by email, when"): the figure named cannot say; the recipes the question reads as can
        if got:
            src = f"{plain(t['name'])}" + (f" · run {run}" if run else "")
            return {**got, "source": src, "used": [], "trace": trace, "kpi": k["name"]}
    return recipe_rows(plane, text, pid, th, kpis, floor=2) or recipe_rows(plane, text, pid, th, kpis, floor=1)


def _breach_figure(text: str, kpis: list[dict], k: dict) -> dict | None:
    """The brief's figure built from a breach recipe (its names say "breach") when the question asks what is over a limit and that recipe fits it at least
    as well as the figure chosen; None otherwise."""
    from sh_platform.v1 import v3_recipes as R
    if not BREACH_Q.search(text):
        return None
    mine = relevance(text, R.bound(k)) if R.bound(k) else 0
    cand = [(relevance(text, r), x) for x in kpis if (r := R.bound(x)) and re.search(r"breach", " ".join(r.get("names") or []), re.I)]
    best = max(cand, key=lambda t: t[0], default=None)
    return best[1] if best and best[1] is not k and best[0] >= max(3, mine) else None


def relevance(text: str, r: dict) -> int:
    """How many of the question's words the recipe's label, names, formula and columns use: the question's meaning, read against what each recipe holds."""
    from sh_platform.v1 import v3_recipes as R
    cols = r.get("columns") or {}
    return len(R._stems(text) & R._stems(" ".join([r["label"], *(r.get("names") or []), r.get("formula", ""), *cols, *map(str, cols.values())])))


def _by_period(plane, t: dict) -> bool:
    """A ready table kept by month: a text column whose values are "2026-09" (most of its rows)."""
    from sh_platform.v1 import cx_dash as CD
    text_cols = [c["name"] for c in t["columns"] if str(c["type"]).lower().startswith(("varchar", "text", "char"))]
    if not text_cols:
        return False
    _, rows, _ = CD.run_sql(plane, f'select * from "{t["name"]}"')
    return any(sum(bool(re.fullmatch(r"\d{4}-\d{2}", str(x.get(c) or ""))) for x in rows) * 2 >= len(rows) > 0 for c in text_cols)


def named_relevance(text: str, r: dict) -> int:
    """relevance over the recipe's label and names only, not its formula: a figure the question names by its own words ("close checklist") is not beaten by
    a figure whose formula happens to say "report" (r1 uc-029)."""
    from sh_platform.v1 import v3_recipes as R
    return len(R._stems(text) & R._stems(" ".join([r["label"], *(r.get("names") or [])])))


ROW_CAP = 100
SAID_WIDE = 4  # a memo ("what figures support ...") speaks for up to four figures
SAID_FIGS = 2  # the sentence speaks for the two figures the question fits best; the rest sit in the figures behind it
PERIOD_COLS = {"period_label", "period_name", "period", "quarter_label"}  # a table's own name for its period (r5 uc-044)
FLAGS = {"first_row", "is_total", "is_latest", "latest", "is_named", "is_opening"}  # helper columns of a recipe: they pick a row, a person never reads them


def _said(v, names: dict, iso: bool = False, money: bool | None = None, col: str = "") -> str | None:
    """A cell as a person says it: a whole amount without ".00", other amounts to the cent, rates to their own decimals (up to four), a date as
    "23 Oct 2026" (ISO in the figures behind the sentence, `iso`). `money`: the column is an amount (None: an amount is a number with two decimals)."""
    if v is None or v == "":
        return None
    if isinstance(v, bool):
        return "yes" if v else "no"
    if isinstance(v, str) and not iso and re.fullmatch(r"\d{4}-\d{2}-\d{2}", v):
        v = date.fromisoformat(v)
    if isinstance(v, date):
        return v.isoformat()[:10] if iso else f"{v.day} {v:%b %Y}"
    if (n := _num(v)) is not None:
        if re.search(r"(?:^|_)year$", col) and n == n.to_integral_value():  # a year is said as a year, never "2,021" (r3b uc-080)
            return str(int(n))
        return _money(n) if (n.as_tuple().exponent == -2 if money is None else money) else _fmt(n, False)
    return names.get(str(v), str(v))


def _holds(cond: str | None, x: dict) -> bool:
    """A line's `if`: "col" (truthy), "not col", "col = text", joined by " and "."""
    if not cond:
        return True
    if " and " in cond:
        return all(_holds(c, x) for c in cond.split(" and "))
    neg = cond.startswith("not ")
    c = cond[4:].strip() if neg else cond.strip()
    if "=" in c:
        col, want = (t.strip() for t in c.split("=", 1))
        ok = str(x.get(col)).lower() == want.lower()
    else:
        ok = bool(x.get(c)) and str(x.get(c)).lower() not in ("false", "0")
    return ok != neg


def _fill(tpl: str, x: dict, pool: list[dict], names: dict, money: set | None = None) -> str | None:
    """{col} from the row; {n}, {sum:col}, {max:col}, {min:col}, {list:col} over the rows the line holds for. None when a value is missing."""
    gone = False

    def one(m):
        nonlocal gone
        k = m.group(1)
        if k == "n":
            return str(len(pool))
        how, _, col = k.rpartition(":")
        cash = None if money is None else col in money
        if how == "list":
            seen = list(dict.fromkeys(s for y in pool if (s := _said(y.get(col), names))))
            got = (", ".join(seen[:5]) + (f" and {len(seen) - 5} more" if len(seen) > 5 else "")) if seen else None
        elif how in ("sum", "max", "min"):
            nums = [n for y in pool if (n := _num(y.get(col))) is not None]
            got = _said(sum(nums, Decimal(0)) if how == "sum" else (max if how == "max" else min)(nums), names, money=cash) if nums else None
        else:
            got = _said(x.get(col), names, money=cash, col=col)
        gone = gone or got is None
        return got or ""

    out = re.sub(r"\{([a-z0-9_:]+)\}", one, tpl)
    return None if gone else out


def say(r: dict, rows: list[dict], names: dict | None = None, money: set | None = None) -> str | None:
    """The recipe's own answer sentence (its `answer` lines in the pack) over the rows the question named, or None when it has none or no line holds.
    A line: `say` (the template), `if` (a condition on the row), `lead` (words before the list), `max` (rows said, default 3), `once` (said once over
    every row it holds for, for {n} and {sum:..}), `latest` (only the rows of the latest value of that column), `sort` (largest first), `stop` (no
    later line once it is said), `all` (every row said, never "and N more"; the question is then answered from these lines, not from a plain read of the table). `if`: "col", "not col", "col = text", joined by " and ". A row said by one line is not said again by a later one.
    `money`: the amount columns of the table (None: read from the values)."""
    names = names or {}
    out, used = [], set()
    for spec in r.get("answer") or []:
        spec = {"say": spec} if isinstance(spec, str) else spec
        pool = [x for x in rows if _holds(spec.get("if"), x)]
        if (c := spec.get("latest")) and (have := [x[c] for x in pool if x.get(c) is not None]):
            pool = [x for x in pool if x.get(c) == max(have)]
        if c := spec.get("sort"):  # largest first (a date: latest first)
            pool = sorted(pool, key=lambda x: n if (n := _num(x.get(c)) if not isinstance(x.get(c), date) else Decimal(x[c].toordinal())) is not None else Decimal("-Infinity"), reverse=True)
        items: list[str] = []
        for x in pool[:1] if spec.get("once") else pool:
            if id(x) in used and not spec.get("once"):
                continue
            if s := _fill(spec["say"], x, pool, names, money):
                items += [s] if s not in items else []
                used |= set() if spec.get("once") else {id(x)}
        if not items:
            continue
        cap = len(items) if spec.get("all") else int(spec.get("max", 3))  # `all`: the figures the question asks for are never cut to "and N more"
        line = "; ".join(items[:cap]) + (f"; and {len(items) - cap} more" if len(items) > cap else "")
        if spec.get("lead") and (lead := _fill(spec["lead"], pool[0], pool, names, money)):
            line = f"{lead} {line}"
        out.append(line[0].upper() + line[1:] + ("" if line.endswith((".", "?", "!")) else "."))
        if spec.get("stop"):
            break
    return " ".join(out) or None


def _plain_sentence(name: str, r: dict, rows: list[dict], who, names: dict, money: set) -> str | None:
    """No template: the figure, its subject and its measure: one row, or the total of a money measure in one currency, never a total of rates."""
    m = r.get("measure")
    rows = [x for x in rows if not x.get("is_total")] or rows
    vals = [n for x in rows if (n := _num(x.get(m))) is not None]
    if not vals:
        return None
    label = str(m).replace("_", " ")
    if len(rows) == 1 or (len(vals) == len(rows) and len(set(vals)) == 1):  # r3b uc-073: one figure repeated on every row is said once, never added up
        w = who(rows[0]) if len(rows) == 1 else ""
        return f"{name}{' for ' + w if w else ''}: {label} {_said(rows[0][m], names, money=m in money)}."
    ccy = {x.get("currency") for x in rows}
    if len(vals) == len(rows) and len(ccy) == 1 and m in money:
        unit = next(iter(ccy))
        return f"{name}: {label} {_said(sum(vals, Decimal(0)), names, money=True)}{' ' + str(unit) if unit else ''} in total over {len(rows)} rows."
    said = [f"{w + ' ' if (w := who(x)) else ''}{_said(x[m], names, money=m in money)}" for x in rows if _num(x.get(m)) is not None]
    return f"{name}, {label}: {'; '.join(said)}."  # r3b: a set is said row by row, never "one of N rows"


STATE_COL = re.compile(r"status|state|kind|reason|bucket|^(?:figure|figure_label|detail|source|note)$")  # a state of a row ("inside limit") is never the entity a question names; nor is the figure a line holds, its detail or its source (r6 uc-040: "commitment" and "Cedarpoint" in a document title kept only the memo's commitment lines)
RATE_COL = re.compile(r"rate|pct|percent|fixing|ratio|share|price")
ASKS_ALL = re.compile(r"\b(?:which|each|every|all)\b", re.I)  # "which filings ... are due": the rows of the set, never only those whose words the question shares (r3b uc-083)
SET_Q = re.compile(r"\b(?:each|every|all|which|across|whether|per)\b", re.I)  # a question about a set of rows, not one latest figure


def _every_row(plane, text: str, pid: str, th: dict, k: dict) -> dict | None:
    """r2: a question about every row ("each", "which", "across all funds", "whether to pay") is answered from the recipe's own answer lines over all
    its rows, not from the latest date's rows alone (uc-007, uc-013, uc-023, uc-050). None when the question is about one figure, names a month, the
    recipe has no such lines, or it reads the latest date (`latest`)."""
    from sh_platform.v1 import v3_recipes as R
    spec = [s for s in ((R.bound(k) or {}).get("answer") or []) if isinstance(s, dict)]
    if not spec or any("latest" in s for s in spec) or not SET_Q.search(text) or parse_when(text):  # a month named: the month's own rows (_read)
        return None
    return recipe_rows(plane, text, pid, th, [k], floor=0, every=True)


def recipe_rows(plane, text: str, pid: str, th: dict, kpis: list[dict], floor: int = 1, every: bool = False, wide: bool = False) -> dict | None:
    """The card's question answered from the brief's recipe-bound figures (w6): the figures whose recipe the question's words fit best (scored by
    `relevance`, at least `floor`, and within one of the best), the rows of their ready tables for what the question names. `answer` is a short
    sentence from each recipe's own `answer` lines (else the figure, its subject and its measure), in the pack's order; `rows` holds the figures behind
    it, a line per row with its readable columns, and a total per money column. Every number is a row's own. None when no such figure has rows.
    `wide` (a memo: "what figures support ..."): every bound figure with a row that names what the question names, up to SAID_WIDE sentences."""
    from sh_platform.v1 import b10_pack as P, cx_dash as CD, v3_recipes as R
    scored = [(relevance(text, r), k) for k in kpis if (r := R.bound(k)) and not R.held(th, k["name"])]
    top = max((sc for sc, _ in scored), default=0)
    figs = [k for sc, k in sorted(scored, key=lambda t: -t[0]) if wide or sc >= max(floor, top - 1)]
    if not figs:
        return None
    gold = ((P.latest_approved(plane, pid) or {}).get("target") or {}).get("gold") or []
    live = {t["name"]: t for t in CD.gold_tables(plane, pid)}
    names, q = aliases(plane), _tokens(text)
    order = {r["id"]: i for i, r in enumerate(R.recipes())}
    lines, said, seen, first = [], [], set(), None
    for k in figs:
        t = next((live[g["name"]] for g in covering(k, gold) if g["name"] in live and live[g["name"]]["rows"] and g["name"] not in seen), None)
        if t is None:
            continue
        seen.add(t["name"])
        r = R.bound(k)
        cols = {c["name"]: str(c["type"]).lower() for c in t["columns"]}
        texts = [c for c, ty in cols.items() if ty.startswith(("varchar", "text", "char")) and c not in META and not c.endswith("_id") and not c.startswith("run")]
        _, rows, _ = CD.run_sql(plane, f'select * from "{t["name"]}"')
        if str(cols.get("id")).startswith("varchar") and rows and not any(re.fullmatch(r"[0-9a-f-]{36}", str(x.get("id") or "")) for x in rows):
            texts = ["id", *texts]  # a recipe's own id that is a code ("TX-01") names its row; a uuid does not (r3b uc-083)
        period = [c for c in texts if rows and sum(bool(re.fullmatch(r"\d{4}-\d{2}", str(x.get(c) or ""))) for x in rows) * 2 >= len(rows)]
        if period and (asked := parse_when(text)) and asked[1]:  # a table kept by month, and the question names one: that month's rows, not its steps' titles (r1 uc-029)
            for c in period:
                hit = [x for x in rows if str(x.get(c))[5:7] == f"{asked[1]:02d}" and (not asked[0] or str(x.get(c))[:4] == str(asked[0]))]
                rows = hit or rows
        else:
            named = False
            for c in ([] if every or ASKS_ALL.search(text) else [c for c in texts if not STATE_COL.search(c)]):  # the entity, fund, account or currency the question names: every value that shares a word with it ("Cedarpoint": both funds); a question about every row names none
                hit = [x for x in rows if x.get(c) is not None and named_in(str(x[c]), names.get(str(x[c])), q, text)]
                rows, named = hit or rows, named or bool(hit)
            if wide and not named:
                seen.discard(t["name"])
                continue  # a memo is about what the question names: a figure with no row for it is left out
        money = [c for c, ty in cols.items() if ty.startswith("decimal") and ty.endswith(",2)") and c not in META and not RATE_COL.search(c)]  # never a rate

        def who(x, texts=texts):
            return " · ".join(names.get(str(x[c]), str(x[c])) for c in texts if x.get(c) not in (None, ""))

        said_by = [x for x in rows if any(_holds(sp.get("if"), x) for sp in (r.get("answer") or []) if isinstance(sp, dict) and sp.get("if"))]
        rows = said_by + [x for x in rows if not any(x is y for y in said_by)]  # the rows the answer names come first (a cap never drops them)
        for x in rows[:ROW_CAP]:
            vals = [f"{c.replace('_', ' ')} {v}" for c in cols if c not in texts and c not in META and c not in FLAGS and not c.endswith("_id")
                    and not c.startswith("run") and (v := _said(x.get(c), names, iso=True, money=c in money, col=c)) is not None]
            lines.append({"label": f"{k['name']} · {who(x)}" if who(x) else k["name"], "value": "; ".join(vals)})
            first = first or k["name"]
        body = [x for x in rows if not x.get("is_total")]  # a recipe's own total row is never added in again
        ccy = next((c for c in texts if CCY_COL.match(c)), None)
        parts = [c for c in money if len({x.get(c) for x in body}) > 1]  # a figure repeated on every row (the recipe's own total) is not a part to add (r3b uc-055)
        if len(body) > 1 and parts and len({x.get(ccy) for x in body} if ccy else {None}) == 1:  # "the total across all funds": one currency only
            tot = [f"{c.replace('_', ' ')} {_said(sum((_num(x.get(c)) or Decimal(0) for x in body), Decimal(0)), names, money=True)}" for c in parts]
            lines.append({"label": f"{k['name']} · total of {len(body)} rows", "value": "; ".join(tot)})
        if len(said) < (SAID_WIDE if wide else SAID_FIGS) and (s := say(r, rows, names, set(money)) or _plain_sentence(k["name"], r, rows, who, names, set(money))):
            said.append(((min(r.get("cards") or [999]), order.get(r["id"], 0)), s))  # the story's order: the notice before its payment, the short before the FX
    if not lines:
        return None
    sentence = " ".join(s for _, s in sorted(said, key=lambda t: t[0])) or f"{first}: the figures are below."
    return {"answer": sentence, "figure": {"label": lines[0]["label"], "value": lines[0]["value"].split("; ")[0], "change": ""},
            "rows": lines, "source": ", ".join(plain(n) for n in sorted(seen)), "used": [], "kpi": first,
            "trace": [{"kind": "tool", "label": "Read the ready tables of the figures the question is about"},
                      {"kind": "guard", "label": "Every figure comes from the tables' rows"},
                      {"kind": "tool", "label": "Said in one sentence from the recipes' own answer lines"}]}


def _not_loaded(th: dict | None) -> str:
    """What the load set aside, named, for an answer that cannot be given: the files that did not load and why (the thread's own record)."""
    ex = [f"{e['title'].removesuffix(' did not load')}: {e['text']}" for e in (th or {}).get("exceptions") or []
          if str(e.get("title") or "").endswith(" did not load") and e.get("text")]  # files only; a pending question is not a file that did not load
    return (" Not loaded: " + " | ".join(ex)[:700]) if ex else ""


def recipe_measures(k: dict, gold: list[dict]) -> dict[str, str]:
    """{ready table: the column its recipe names as the figure} for the pack's tables built from a recipe for this figure: the table lists the figure,
    or its recipe is the one the figure's name matches."""
    from sh_platform.v1 import v3_recipes as R
    r_k = (R.bound(k) or {}).get("id")
    out = {}
    for g in gold:
        p = str(g.get("purpose") or "")
        r = R.recipe(p[len(R.MARK):].split(":", 1)[0]) if p.startswith(R.MARK) else None
        if r and r.get("measure") and (k["name"] in (g.get("kpis") or []) or r["id"] == r_k):
            out[g["name"]] = r["measure"]
    return out


def _as_at(rows: list[dict], when: str, ent: str | None) -> str:
    """The as-at date a figure is read at when the question names none: the latest date on which the most entities have a row, so one entity that reported
    later (the family trust, 13 Oct) does not leave the family's figures on that one entity (r1 uc-017, uc-031). No entity column: the latest date."""
    by: dict[str, set] = {}
    for r in rows:
        by.setdefault(str(r[when])[:10], set()).add(str(r.get(ent)) if ent else "")
    most = max(len(v) for v in by.values())
    return max(d for d, v in by.items() if len(v) == most)


def _shares(k: dict, rows: list[dict], texts: list[str], meas: str, got, names: dict[str, str]) -> dict | None:
    """A share figure added up over each split the rows carry ("listed equity 36.37% by asset class"), with the total the rows give on the same as-at
    date. None when no dimension splits the rows, so the plain read answers."""
    parts = []
    for c in texts:
        sums: dict[str, Decimal] = {}
        for r in rows:
            n = _num(r.get(meas))
            if n is not None and r.get(c) is not None:
                key = names.get(str(r[c]), str(r[c])).replace("_", " ")
                sums[key] = sums.get(key, Decimal(0)) + n
        if len(sums) > 1:
            parts.append((c.replace("_", " "), sorted(sums.items(), key=lambda kv: -kv[1])))
    if not parts:
        return None
    total = next((_num(r.get("total_value")) for r in rows if _num(r.get("total_value")) is not None), None)
    head = f"{k['name']} on {got.day} {_month_words(got)}" if got else k["name"]
    say = (f"{head}: total {_money(total)}; " if total is not None else f"{head}: ") + "; ".join(
        f"by {d}: " + "; ".join(f"{a} {v:,.2f}%" for a, v in sums[:SAID_FIGS + 1]) + ("; and more in the figures" if len(sums) > SAID_FIGS + 1 else "")
        for d, sums in parts) + "."  # the sentence speaks for the largest shares; every share sits in the figures behind it
    return {"answer": say, "figure": {"label": head, "value": _money(total) if total is not None else "", "change": ""},
            "rows": [{"label": f"{k['name']} · by {d}", "value": "; ".join(f"{a} {v:,.2f}%" for a, v in sums)} for d, sums in parts]}


def _period_said(got: date | None, on_day: bool, month_words: str, labels: set[str]) -> str:
    """The period a plain read names: the month of its as-at date, or, when the table names its own period in one value ("2026Q3"), that name
    (r5 uc-044: "for 2026Q3 in June 2026" said the quarter by its date's month). A day is always said by its date."""
    if got and len(labels) == 1 and not on_day:
        return next(iter(labels))
    return month_words


def _read(plane, t: dict, k: dict, text: str, names: dict[str, str], th: dict | None = None, measure: str | None = None) -> dict | None:
    from sh_platform.v1 import cx_dash as CD
    cols = {c["name"]: str(c["type"]).lower() for c in t["columns"]}
    when = next((c for c, ty in cols.items() if ty.startswith(("date", "timestamp"))), None)
    nums = [c for c, ty in cols.items() if ty.startswith(("decimal", "numeric", "double", "float", "bigint", "integer", "hugeint")) and not c.endswith("_id")]
    want = _tokens(k["name"])
    if measure in cols:  # the recipe names the figure's column: no guessing by words
        meas = measure
    else:
        meas = max(nums, key=lambda c: (len(want & _tokens(c)), _slug(c) == _slug(k["name"])), default=None)
        if not meas or not (want & _tokens(meas) or len(nums) == 1):
            return None
        if (want & _tokens(NOT_WORDS)) - _tokens(meas):  # c4: "Missing statements" was read from statements_received_on_time; a count of what came is not what did not
            return None
    asked = parse_when(text)
    cond = f' where "{when}" is not null' if when else ""
    if when and asked and asked[0]:
        y, mth, d = asked
        end = date(y + (mth == 12), mth % 12 + 1, 1)
        cond += f" and cast(\"{when}\" as date) < date '{end.isoformat()}'"
    sql = f'select * from "{t["name"]}"{cond}' + (f' order by "{when}" desc' if when else "")
    _, rows, _ = CD.run_sql(plane, sql)
    texts = [c for c, ty in cols.items() if ty.startswith(("varchar", "text", "char", "uuid")) and c not in META and not CCY_COL.match(c)
             and (not c.endswith("_id") or any(str(r.get(c)) in names for r in rows))]  # a split by name; a run or batch id is not one
    q = _tokens(text)
    split_by = _tokens(" ".join(re.findall(r"(?:which|by|per|each|every)\s+([a-z]+)", text.lower()))) - {"how", "much", "many", "the", "our", "month", "day", "week", "year"}
    have = _tokens(" ".join(c for c in texts)) | {w for c in texts if any(str(r.get(c)) in names for r in rows) for w in _tokens(c)}
    if have & HOLDING:  # "which holding": an issuer, instrument or fund column is the holding
        have |= HOLDING
    if split_by - have and not (split_by & _tokens(k["name"])):  # "which entity": a table that has no split by it cannot say, and the nearest total is not the answer
        return {"answer": f"Not available yet: the ready table for {k['name']} has no split by {', '.join(sorted(split_by - have))}.", "figure": None, "rows": []}
    if (fake := fake_split(rows, [c for c in texts if c in cols], when, meas)):
        who = fake.replace("_name", "").replace("_id", "").replace("_", " ")
        return {"answer": f"Not available yet: the ready table for {k['name']} shows the same figure for every {who}, so it does not split by {who} (it adds up the whole company under each name)."
                          + _not_loaded(th), "figure": None, "rows": []}
    if when and asked and not asked[0]:  # a month and day without a year: the months of the table that match
        rows = [r for r in rows if int(str(r[when])[5:7]) == asked[1]]
        if asked[2] and any(str(r[when])[8:10] != "01" for r in rows):
            rows = [r for r in rows if int(str(r[when])[8:10]) == asked[2]]
        if not rows:
            return {"answer": f"Not available yet: {k['name']} has no figure for that day.", "figure": None, "rows": []}
    named = False
    for c in texts:  # the dimensions the question names
        vals = sorted({str(r[c]) for r in rows if r.get(c) is not None})
        pick = match_values(vals, names, q, text)
        if pick:
            rows = [r for r in rows if str(r.get(c)) in pick]
            named = True
    if not rows:
        return None
    register = bool(when and not asked and len(rows) > 1 and len({str(r[when])[:10] for r in rows}) == len(rows) and any(len({str(r.get(c)) for r in rows}) == len(rows) for c in texts))
    if when and not register:  # r3b uc-065, uc-066: a register (each member once, at its own date) is not cut to its latest date
        top = _as_at(rows, when, next((c for c in texts if "entity" in c), None))
        got = date.fromisoformat(top)
        if asked and asked[0] and (got.year, got.month) != (asked[0], asked[1]):
            return {"answer": f"Not available yet: {k['name']} has no figure for {_month_words(date(asked[0], asked[1], 1))}; the latest is {_month_words(got)}.",
                    "figure": None, "rows": []}
        rows = [r for r in rows if str(r[when])[:10] == top]
    else:
        got = None
    money = "decimal(18" in cols[meas] or "decimal(28" in cols[meas]
    vals = [(r, _num(r.get(meas))) for r in rows if _num(r.get(meas)) is not None]
    if not vals:
        return None
    if re.search(r"share|pct|percent", meas or "", re.I) and (sh := _shares(k, rows, texts, meas, got, names)):
        return sh
    ccy = next((c for c in cols if CCY_COL.match(c)), None)
    ccys = {str(r.get(ccy)) for r, _ in vals if r.get(ccy)} if ccy else set()
    unit = next(iter(ccys)) if len(ccys) == 1 and not re.search(rf"\b{re.escape(next(iter(ccys)))}\b", k["name"], re.I) else ""
    on_day = bool(got and asked and asked[2])  # a question about a day is answered for that day, not the month
    period = (f"{got.day} {_month_words(got)}" if on_day else _month_words(got)) if got else ""
    plab = next((c for c in texts if c in PERIOD_COLS), None)
    period = _period_said(got, on_day, period, {str(r[plab]) for r, _ in vals if r.get(plab) is not None} if plab else set())
    texts = [c for c in texts if c != plab]
    label = k["name"] + (f" · {period}" if got else "")
    split = [c for c in texts if len({str(r.get(c)) for r, _ in vals}) > 1] + ([ccy] if len(ccys) > 1 else [])  # never add one currency to another
    part = "" if named else coverage(plane, {str(r[c]) for r, _ in vals for c in texts if r.get(c) is not None})
    if len(vals) == 1 or not split:
        v = vals[0][1] if len(vals) == 1 else sum((v for _, v in vals), Decimal(0))
        who = " · ".join(dict.fromkeys(names.get(str(r[c]), str(r[c])) for r in [vals[0][0]] for c in texts if c in cols and r.get(c) is not None))
        num = _fmt(v, money) + (f" {unit}" if unit else "")
        return {"answer": f"{k['name']}{' for ' + who if who else ''}{' ' + part if part else ''}{(' on ' if on_day else ' in ') + period if got else ''}: {num}.",
                "figure": {"label": label + (f" · {who}" if who else ""), "value": num, "change": ""}, "rows": []}
    parts = [(" · ".join(dict.fromkeys(names.get(str(r[c]), str(r[c])) for c in split)), v) for r, v in vals]  # every member of the set, never cut (r3b)
    total = sum((v for _, v in vals), Decimal(0))
    add = money and len(ccys) <= 1 and not re.search(r"percent|ratio|rate|return|share", k["name"].lower())
    say = "; ".join(f"{a} {_fmt(v, money)}" for a, v in parts)
    owed = sum((_num(r.get("debt")) or Decimal(0) for r, _ in vals), Decimal(0)) if "debt" in cols and "debt" in k["name"].lower() else None
    after = f" after {_fmt(owed, money)} of debt" if owed else ""  # "net worth after debt": what the family owes is said with the total
    return {"answer": f"{k['name']}{' ' + part if part else ''}{(' on ' if on_day else ' in ') + period if got else ''}: {say}." + (f" Total {_fmt(total, money)}{after}." if add and len(parts) > 1 else ""),
            "figure": {"label": label + (" · total" if add else ""), "value": _fmt(total if add else parts[0][1], money), "change": ""},
            "rows": [{"label": a, "value": _fmt(v, money)} for a, v in parts]}


def note_figures(x, limit: int = 6) -> list[dict]:
    """The morning note's figures from the ready tables of the brief: each figure with rows, at its latest period (one line per split, up to `limit`)."""
    from sh_platform.v1 import b10_pack as P, cx_dash as CD
    out = []
    try:
        pack = P.latest_approved(x.plane, x.th["pid"]) or {}
        gold = (pack.get("target") or {}).get("gold") or []
        live = {t["name"]: t for t in CD.gold_tables(x.plane, x.th["pid"])}
        names = aliases(x.plane)
        from sh_platform.v1 import v3_recipes as R
        for k in [k for k in kpis_of(x.plane, x.th.get("iid")) if not R.held(x.th, k["name"])]:  # a held figure is not in the note
            t = max((live[g["name"]] for g in covering(k, gold) if g["name"] in live and live[g["name"]]["rows"]), key=lambda t: t["rows"], default=None)
            got = _read(x.plane, t, k, k["name"], names) if t else None
            if got and got.get("figure"):
                out += [{"label": f"{k['name']} {r['label']}", "value": r["value"], "change": ""} for r in got["rows"]] if got.get("rows") else [{**got["figure"]}]
    except Exception:  # noqa: BLE001 - a note without these figures is still a note
        log.info("no ready-table figures for the note", exc_info=True)
    return out[:limit]
