"""Harbor 3.0 loads: files whose home is the model's `extracted_fields`, not a table of rows (skills data.profile-file, data.read-pdf-fields).

A labelled-figures workbook (an issued report's summary sheet: "Period end | 2026-09-30", "Total | 57600000") is not rows of one kind: each
label and its value is kept as an extracted field of the workbook, registered as a document through the 2.0 document intake (page = the
sheet's place in the book, line = the row). A PDF whose lines hold labelled numbers but no amount written with a currency code (a report,
not a statement) goes the same way: v3_pdf marks it "figures" instead of asking which line holds the amount. Values are copied as written,
never worked out. State: thread["labelled"] = {name: {doc_id, n}}.
"""
from __future__ import annotations

import base64
import csv
import io
import logging
import re
from decimal import Decimal, InvalidOperation

log = logging.getLogger(__name__)
BOOK, PDF = "data.profile-file", "data.read-pdf-fields"
SHEETS = (".xlsx", ".xlsm", ".csv")
_NUM = re.compile(r"^(?:([A-Z]{3})\s*)?([-+]?\d[\d,]*(?:\.\d+)?)$")
_CCY = re.compile(r"\(([A-Z]{3})\)")
_DATEISH = re.compile(r"^\d{1,4}[/.-]\d{1,2}[/.-]\d{1,4}")


def _text(v) -> str:
    return "" if v is None else str(v).strip()


def _num(v) -> Decimal | None:
    if isinstance(v, bool):
        return None
    if isinstance(v, (int, float, Decimal)):
        return Decimal(str(v))
    m = _NUM.match(_text(v))
    try:
        return Decimal(m[2].replace(",", "")) if m else None
    except InvalidOperation:
        return None


def _ccy(*texts) -> str | None:
    """The currency a label or a column heading names, as '(USD)', or the code written before the amount."""
    for t in texts:
        m = _CCY.search(_text(t)) or (_NUM.match(_text(t)) if _text(t)[:3].isupper() else None)
        if m and m[1]:
            return m[1]
    return None


def _grid(name: str, data: bytes) -> list[tuple[str, list[list]]]:
    if name.lower().endswith(".csv"):
        return [(name, list(csv.reader(io.StringIO(data.decode("utf-8-sig", errors="replace")))))]
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    return [(ws.title, [list(r) for r in ws.iter_rows(values_only=True)]) for ws in wb.worksheets]


def _label(c) -> bool:
    t = _text(c)
    return bool(re.search(r"[A-Za-z]", t)) and not _DATEISH.match(t) and _num(t) is None


def sheet_figures(rows: list[list], page: int = 1) -> list[dict] | None:
    """The labelled figures of one sheet, or None when the sheet is a table of rows. A label row is a text label and one value (three or more
    such rows, at least half as many as the rows of a table, and their values not all of one kind, so a two-column list of funds and amounts stays a table).
    A row under a heading row is a figure too: its first cell is the label, its first number the value, the heading names the currency."""
    kv, wide, kinds, figs, head = 0, 0, set(), [], None
    for r, row in enumerate(rows, 1):
        cells = [(i, c) for i, c in enumerate(row) if _text(c)]
        if not cells:
            continue
        if len(cells) >= 3:
            wide += 1
            if all(_label(c) for _, c in cells) and head is None:
                head = {i: _text(c) for i, c in cells}
                continue
        if not _label(cells[0][1]) or len(cells) < 2:
            continue
        label, rest = _text(cells[0][1]), cells[1:]
        if len(cells) == 2:
            kv += 1
            if kv > 1:
                kinds.add("number" if _num(rest[0][1]) is not None else "text")
        i, v = next(((i, c) for i, c in rest if _num(c) is not None), rest[0])
        n = _num(v)
        figs.append({"label": label[:100], "value_text": _text(v)[:255], "value_numeric": n, "currency": _ccy(label, (head or {}).get(i), v) if n is not None else None,
                     "page": page, "line": r, "quote": " | ".join(_text(c) for _, c in cells)[:300]})
    return figs if kv >= 3 and 2 * kv >= wide and len(kinds) > 1 else None


def labelled(name: str, data: bytes) -> list[dict] | None:
    """Every labelled figure of a workbook whose sheets include a labelled summary; None when no sheet is one (a table of rows: the mapping path)."""
    if not name.lower().endswith(SHEETS):
        return None
    try:
        sheets = _grid(name, data)
    except Exception:  # noqa: BLE001 - a book that cannot be opened here takes the normal path, which names the problem
        return None
    out = [f for k, (_, rows) in enumerate(sheets, 1) for f in (sheet_figures(rows, k) or [])]
    return out or None


def pdf_figures(pairs: list[dict], picture: bool = False, agree: dict | None = None) -> list[dict]:
    """v3_pdf's label/value pairs (page, line, quote kept) as figures, each with a confidence (v3_review.score): read by rules from the text layer,
    checked that the value is in the quoted lines; a picture scan is low; `agree` ({label: bool}, v3_judge.field_agreement) lowers a disagreement."""
    from sh_platform.v1 import v3_review as R
    out = []
    for p in pairs:
        n = _num(p["value"])
        out.append({"label": p["label"][:100], "value_text": p["value"][:255], "value_numeric": n,
                    "currency": _ccy(p["label"], p["value"]) if n is not None else None, "page": p["page"], "line": p["line"], "quote": p["quote"],
                    "confidence": R.score(known=True, picture=picture, checks_ok=p["value"] in p["quote"], agree=(agree or {}).get(p["label"]))})
    return out


def write(plane, doc_id: str, figs: list[dict], layout: str = "") -> int:
    """The figures as extracted fields of the document, once (a document that already has fields keeps them). Copied as written from the
    sender's own file, so they are confirmed, as 2.0's reader confirms a high-confidence field; one below the signed `review_confidence_min`
    is held (v3_review) until a person confirms or corrects it."""
    from sh_platform.v1 import v3_review as R
    from sh_platform.v1 import wh
    from sh_platform.v1.e1_write import org_id
    ids, quotes = [], {}
    with wh.connect(plane) as con:
        have = con.execute("select count(*) from extracted_fields where document_id = ?", [doc_id]).fetchone()[0]
        if have:
            return int(have)
        org, at = org_id(plane), wh.now()
        for f in figs:
            conf = f.get("confidence", "high")
            st = R.status_for(plane, conf)
            fid = wh.new_id()
            con.execute("insert into extracted_fields (id, organization_id, document_id, field_name, value_text, value_numeric, currency, page_number,"
                        " line_number, confidence, review_status, reviewed_at) values (?,?,?,?,?,?,?,?,?,?,?,?)",
                        [fid, org, doc_id, f["label"], f["value_text"], f["value_numeric"], f["currency"], f["page"], f["line"], conf, st, at if st == "confirmed" else None])
            if st == "held":
                ids.append(fid)
                quotes[fid] = f.get("quote") or ""
    if ids:
        R.remember_quotes(plane, quotes)
        R.remember_layout(plane, ids, layout)
    return len(figs)


def _register(x, name: str, data: bytes, source: str | None) -> str | None:
    code = re.sub(r"\W+", "_", source or "upload").lower()[:24]
    try:
        return x.api("leo", "POST", "/api/v1/evidence/documents", {"filename": name, "data": base64.b64encode(data).decode(), "source_code": code}).get("id")
    except Exception:  # noqa: BLE001 - named to the person at load
        log.warning("a labelled workbook could not be registered as a document", exc_info=True)
        return None


# ---- hooks of v3_harbor ------------------------------------------------------------------------------------------------------------

def intake(x) -> None:
    """a_intake: a source workbook that is a labelled summary is registered as a document and kept out of the mapping path."""
    from sh_platform.v1 import v3_harbor as H
    st = x.th["labelled"] = {}
    books = [f for f in H.my_files(x, "source") if f["name"].lower().endswith(SHEETS)]
    if not books:
        return
    with H._skill(x, BOOK, "Reading the summary workbooks as labelled figures"):
        for f in books:
            data = H.blob(x.plane, f["name"])
            figs = labelled(f["name"], data)
            if figs:
                st[f["name"]] = {"doc_id": _register(x, f["name"], data, f.get("source")), "n": len(figs)}


def skip(x, name: str) -> bool:
    """True for a file this module loads (a_intake's spreadsheet loop leaves it alone)."""
    return name in (x.th.get("labelled") or {})


def d2_files(x) -> list[dict]:
    rows = [{"name": n, "read_as": f"labelled figures ({v['n']})", "goes_to": "extracted fields of the document", "state": "reused", "confidence": "high"}
            for n, v in (x.th.get("labelled") or {}).items()]
    rows += [{"name": n, "read_as": "labelled figures", "goes_to": "extracted fields of the document", "state": "reused", "confidence": "high"}
             for n, f in ((x.th.get("pdf") or {}).get("files") or {}).items() if f.get("status") == "figures"]
    return rows


def load(x) -> tuple[int, list[dict]]:
    """b_load: the labelled workbooks and the PDFs v3_pdf marked "figures". (loaded, set-aside lines)."""
    from sh_platform.v1 import v3_harbor as H
    from sh_platform.v1 import v3_pdf as V
    todo = [(n, v.get("doc_id"), BOOK) for n, v in (x.th.get("labelled") or {}).items()]
    todo += [(n, f.get("doc_id"), PDF) for n, f in ((x.th.get("pdf") or {}).get("files") or {}).items() if f.get("status") == "figures"]
    n_ok, aside = 0, []
    for name, doc_id, sid in todo:
        with H._skill(x, sid, f"Keeping the labelled figures of {name}, each with its place in the file"):
            try:
                data = H.blob(x.plane, name)
                pages = V.pages_of(name, data) if V.is_pdf(name) else None
                figs = pdf_figures(V.pairs_of(pages), V.is_picture(data)) if pages is not None else labelled(name, data) or []
                if not doc_id or not figs:
                    raise ValueError("it could not be kept as a document" if not doc_id else "no labelled figure could be read from it")
                write(x.plane, doc_id, figs, V.layout_of(pages)[0] if pages else "")
            except Exception as e:  # noqa: BLE001 - one file set aside, in one plain line; the rest carry on
                log.warning("labelled figures of %s not kept", name, exc_info=True)
                aside.append({"file": name, "title": f"{name} did not load", "text": f"I set it aside: {e}."[:200]})
                H.set_status(x.plane, name, "question")
                continue
        H.set_status(x.plane, name, "loaded")
        n_ok += 1
    return n_ok, aside
