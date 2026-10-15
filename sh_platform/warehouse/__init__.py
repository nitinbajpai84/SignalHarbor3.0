"""The company warehouse: one DuckDB file inside the plane (A-064), restricted marts in a separate file (A-065).

Paths come only from the signed-in company's plane, never from input. No route, no screen here.
"""
import csv
import datetime as dt
import logging
import shutil
import time
import uuid
from pathlib import Path

import duckdb

from sh_platform.plane import CompanyPlane, valid_module_id

log = logging.getLogger(__name__)

SCHEMAS = ("bronze", "silver", "gold")
LOAD_ID, LOADED_AT = "_load_id", "_loaded_at"


def _file(plane: CompanyPlane, name: str) -> Path:
    return plane.path("warehouse", name)


def _q(ident: str) -> str:
    return '"' + ident.replace('"', '""') + '"'


def open_company(plane: CompanyPlane) -> duckdb.DuckDBPyConnection:
    """Open warehouse/company.duckdb read-write; bronze, silver and gold exist on return."""
    plane.area("warehouse").mkdir(parents=True, exist_ok=True)
    con = duckdb.connect(str(_file(plane, "company.duckdb")))
    for s in SCHEMAS:
        con.execute(f"create schema if not exists {s}")
    return con


def open_restricted(plane: CompanyPlane) -> duckdb.DuckDBPyConnection:
    """Open warehouse/restricted.duckdb (restricted marts only; gold schema)."""
    plane.area("warehouse").mkdir(parents=True, exist_ok=True)
    con = duckdb.connect(str(_file(plane, "restricted.duckdb")))
    con.execute("create schema if not exists gold")
    return con


def open_readonly(plane: CompanyPlane) -> duckdb.DuckDBPyConnection:
    return duckdb.connect(str(_file(plane, "company.readonly.duckdb")), read_only=True)


def _read_tabular(path: Path) -> tuple[list[str], list[list[str]]]:
    """Header and text rows of a CSV or the first sheet of an XLSX (from bronze_loader / landing.read_tabular)."""
    if path.suffix.lower() == ".xlsx":
        from openpyxl import load_workbook
        wb = load_workbook(path, read_only=True, data_only=True)
        try:
            raw = [list(r) for r in wb.worksheets[0].iter_rows(values_only=True)]
        finally:
            wb.close()
        raw = [r for r in raw if any(v not in (None, "") for v in r)]

        def cell(v) -> str:
            if v is None:
                return ""
            if isinstance(v, dt.datetime):
                return v.date().isoformat() if v.time() == dt.time() else v.isoformat(sep=" ")
            if isinstance(v, dt.date):
                return v.isoformat()
            if isinstance(v, float) and v.is_integer():
                return str(int(v))
            return str(v)
        rows = [[cell(v) for v in r] for r in raw]
    elif path.suffix.lower() == ".csv":
        with path.open(newline="", encoding="utf-8") as fh:
            rows = list(csv.reader(fh))
    else:
        raise ValueError("only .csv and .xlsx files can be loaded")
    if not rows:
        raise ValueError("file is empty")
    return [str(h).strip() for h in rows[0]], rows[1:]


def load_bronze(plane: CompanyPlane, source: str, file: str) -> str:
    """Land a file from the plane's landing zone as bronze.<source>; returns the load id.

    Columns are kept as landed (text) plus _load_id and _loaded_at. A repeat load appends.
    """
    if not valid_module_id(source):
        raise ValueError("invalid source name")
    header, rows = _read_tabular(plane.path("landing", file))
    if not header or len(set(header)) != len(header) or LOAD_ID in header or LOADED_AT in header:
        raise ValueError("file header is not loadable")
    load_id, now = uuid.uuid4().hex, dt.datetime.now(dt.timezone.utc).replace(tzinfo=None)
    width = len(header)
    cols = ", ".join(f"{_q(h)} varchar" for h in header)
    con = open_company(plane)
    try:
        con.execute(f"create table if not exists bronze.{source} ({cols}, {LOAD_ID} varchar, {LOADED_AT} timestamp)")
        con.executemany(f"insert into bronze.{source} values ({', '.join(['?'] * (width + 2))})",
                        [(r + [""] * width)[:width] + [load_id, now] for r in rows])
    finally:
        con.close()
    log.info("bronze load %s: %s rows into bronze.%s", load_id, len(rows), source)
    return load_id


def refresh_readonly_copy(plane: CompanyPlane) -> Path:
    """Copy company.duckdb to company.readonly.duckdb (query tools read the copy, never the live file)."""
    con = open_company(plane)  # makes sure the file exists, then flushes it to disk
    try:
        con.execute("checkpoint")
    finally:
        con.close()
    src, dst = _file(plane, "company.duckdb"), _file(plane, "company.readonly.duckdb")
    tmp = dst.with_suffix(".tmp")
    shutil.copyfile(src, tmp)
    for i in range(10):  # Windows refuses the swap while a reader has the copy open; readers are short
        try:
            tmp.replace(dst)
            break
        except PermissionError:
            if i == 9:
                raise
            time.sleep(0.3)
    return dst


def counts_only(plane: CompanyPlane, table: str) -> dict[str, int]:
    """Row and column counts of schema.table. Never rows (egress guard 1)."""
    schema, _, name = table.partition(".")
    if schema not in SCHEMAS or not valid_module_id(name):
        raise ValueError("table must be bronze.x, silver.x or gold.x")
    con = open_company(plane)
    try:
        cols = con.execute("select count(*) from information_schema.columns where table_schema = ? and table_name = ?",
                           [schema, name]).fetchone()[0]
        if not cols:
            raise LookupError("unknown table")
        rows = con.execute(f"select count(*) from {schema}.{name}").fetchone()[0]
    finally:
        con.close()
    return {"rows": rows, "columns": cols}
