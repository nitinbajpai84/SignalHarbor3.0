"""B3 mapping sheet: draft from the signed SOW (rule-based when the model is off, Span's words with a stub model), edit a row, download
the workbook, and the refusals (a sponsor may not draft; an unsigned SOW cannot be mapped)."""
import io
import json
import shutil
from pathlib import Path
from types import SimpleNamespace

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from sh_platform.app import create_app
from sh_platform.jsonfile import read_json_strict, write_json_atomic
from sh_platform.settings import Settings
from tests.v1.seeded import SEED

ROOT = Path(__file__).resolve().parents[2]
ACCOUNTS = SEED / "_control" / "accounts.json"
pytestmark = pytest.mark.skipif(not ACCOUNTS.is_file(), reason="no seeded demo data")
PASSWORD = "b3-test-only-not-a-real-password"  # set on the tmp copy, so the test never reads the shared credentials file
IID = "int-b3"
LEELA, PRIYA = "leela.raman.neera@example.test", "priya.nair.neera@example.test"

SOW = {"id": "sow-b3", "intent_id": IID, "title": "Portfolio returns", "state": "signed", "version": 1,
       "kpis": [{"code": "FEE", "name": "Fee ratio", "formula": "fee_amount / market_value", "source_ids": ["custodian", "fees"]},
                {"code": "BMK", "name": "Benchmark relative", "formula": "return - benchmark_return", "source_ids": ["benchmark"]}],
       "sources": [{"id": "custodian", "name": "Custodian positions", "kind": "file", "sample_path": "landing/custodian/positions.csv", "fields": []},
                   {"id": "fees", "name": "Fee schedule", "kind": "file", "fields": []}],
       "catalogue": [
           {"table": "positions", "layer": "landing", "source_id": "custodian", "feeds_kpis": ["FEE"],
            "fields": [{"name": "market_value", "type": "decimal", "source_field": "Market Value"}]},
           {"table": "positions_clean", "layer": "silver", "source_id": "custodian", "feeds_kpis": ["FEE"],
            "fields": [{"name": "market_value", "type": "decimal", "source_field": "Market Value"},
                       {"name": "as_of", "type": "date", "source_field": "as_of"}]},
           {"table": "benchmark", "layer": "silver", "source_id": "benchmark", "feeds_kpis": ["BMK"],
            "fields": [{"name": "index_level", "type": "decimal", "source_field": "index_level"}]},
           {"table": "fee_ratio", "layer": "gold", "source_id": "fees", "feeds_kpis": ["FEE"],
            "fields": [{"name": "fee_amount", "type": "decimal", "source_field": ""}]}],
       "unclear": [], "governance": [{"rule": "Only the family office team reads positions.", "kind": "access", "applies_to": ["positions"]}],
       "signature": {"name": "Arjun Mehta", "role": "company_admin", "at": "2026-10-07"}}
SCAN = {"intent_id": IID, "at": "2026-10-07", "kpis": [{"code": "FEE", "status": "partly"}, {"code": "BMK", "status": "not"}], "agreed": None}


@pytest.fixture()
def world(tmp_path, monkeypatch):
    monkeypatch.setenv("SH_MODEL_PROVIDER", "off")
    data = tmp_path / "data"
    shutil.copytree(SEED / "_control", data / "_control")
    for c in ("neera", "fresh"):
        (data / c).mkdir()
    neera = data / "neera"
    write_json_atomic(neera / "intents" / IID / "sow.json", SOW)
    write_json_atomic(neera / "build" / "scan" / f"{IID}.json", SCAN)
    app = create_app(Settings(mode="dev", data_root=str(data), session_secret="fake-test-secret-not-real"))
    app.state.model_env = {"SH_MODEL_PROVIDER": "off"}
    for email in (LEELA, PRIYA):
        app.state.account_store.set_password(email, PASSWORD)
    return app, neera


def _client(world, email):
    c = TestClient(world[0])
    assert c.post("/api/auth/login", json={"email": email, "password": PASSWORD}).status_code == 200
    c.headers["X-CSRF-Token"] = c.cookies.get("sh_csrf", "")
    return c


def test_draft_edit_download(world):
    leela = _client(world, LEELA)
    r = leela.post(f"/api/v1/build/{IID}/mapping")
    assert r.status_code == 200, r.text
    got = r.json()
    assert got["fallback"] is True and got["message"]  # the model is off: labelled
    rows = {(x["target_table"], x["target_field"]): x for x in got["sheet"]["rows"]}
    assert got["sheet"]["version"] == 1 and len(rows) == 5
    land = rows[("positions", "market_value")]
    assert land["source_table"] == "positions.csv" and land["status"] == "mapped" and "Only the family office" in land["rule"]
    assert rows[("positions_clean", "market_value")]["transformation"].startswith("Rename Market Value to market_value")
    assert rows[("benchmark", "index_level")]["status"] == "unclear"  # benchmark is not in the SOW sources and BMK scanned "not"
    assert rows[("fee_ratio", "fee_amount")]["status"] == "unclear"  # no source field named
    assert read_json_strict(world[1] / "build" / "mapping" / f"{IID}.json")["rows"][0]["n"] == 1
    assert not (world[1] / "intents" / IID / "room.json").exists()  # no rooms: nothing is written

    n = land["n"]
    r = leela.put(f"/api/v1/build/{IID}/mapping/rows/{n}", json={"transformation": "Load as is, keep the file name.", "status": "mapped"})
    assert r.status_code == 200 and r.json()["row"]["transformation"] == "Load as is, keep the file name."
    assert leela.put(f"/api/v1/build/{IID}/mapping/rows/{n}", json={"status": "done"}).status_code == 409
    again = leela.post(f"/api/v1/build/{IID}/mapping").json()  # a redraft keeps the person's edit (rule 7)
    assert again["sheet"]["version"] == 2
    assert next(x for x in again["sheet"]["rows"] if x["n"] == n)["transformation"] == "Load as is, keep the file name."

    x = leela.get(f"/api/v1/build/{IID}/mapping.xlsx")
    assert x.status_code == 200 and "mapping-int-b3-v2.xlsx" in x.headers["content-disposition"]
    ws = load_workbook(io.BytesIO(x.content)).active
    assert ws["A1"].value == "No." and ws.max_row == 6 and ws.freeze_panes == "A2"


def test_span_proposes_with_a_model(world):
    app = world[0]
    answer = json.dumps({"positions_clean.as_of": "Read the custodian's as_of as the month-end date."})
    app.state.build_model_factory = lambda plane: SimpleNamespace(complete=lambda *a, **k: SimpleNamespace(text=answer, tool_calls=[]))
    got = _client(world, LEELA).post(f"/api/v1/build/{IID}/mapping").json()
    assert got["fallback"] is False
    row = next(x for x in got["sheet"]["rows"] if x["target_field"] == "as_of")
    assert row["transformation"] == "Read the custodian's as_of as the month-end date."


def test_refusals(world):
    priya = _client(world, PRIYA)
    r = priya.post(f"/api/v1/build/{IID}/mapping")
    assert r.status_code == 403  # the sponsor is not the architect
    assert priya.get(f"/api/v1/build/{IID}/mapping").status_code == 404  # nothing drafted yet
    sow = dict(SOW, state="draft", signature=None)
    write_json_atomic(world[1] / "intents" / IID / "sow.json", sow)
    r = _client(world, LEELA).post(f"/api/v1/build/{IID}/mapping")
    assert r.status_code == 409 and "not signed" in r.json()["message"]
