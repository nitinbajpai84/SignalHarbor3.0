"""E0c2: good load with locators, totals mismatch rejects everything, layout change stops, unknown fund held back, later file supersedes.
Runs against the real landing zone, profile, mapping records, instrument matching and writer (tmp company, synthetic E0d workbooks)."""
import io
import types
from pathlib import Path

import pytest
from openpyxl import load_workbook

from sh_platform.plane import PlaneStore
from sh_platform.profile.view import Viewer
from sh_platform.v1 import e0b_mappings, e0b_tools, e0c_landing, e0c_load, wh
from sh_platform.v1.e1_write import org_id

FIX = Path(__file__).resolve().parents[2] / "tests" / "v1" / "fixtures" / "e0d" / "workbooks"
FUNDS = ["Cedar Ridge Growth Fund LP", "Northgate Credit Opportunities", "Halden Real Assets Partners",
         "Marlow Global Equity Fund", "Pinecrest Infrastructure Fund II", "Ashby Venture Partners III"]
ENG = Viewer("Leela Raman", frozenset({"analyst"}), False, "acc-leela")


@pytest.fixture()
def env(tmp_path):
    wh.reset_cache()
    plane = PlaneStore(tmp_path / "data").create("acme")
    org = org_id(plane)
    with wh.connect(plane) as con:
        sid = e0c_load._source_system(con, org, "upload")
        for code in ("EVG-GROWTH", "EVG-CORE", "EVG-LEGACY"):
            con.execute("insert into portfolios (id, organization_id, legal_entity_id, portfolio_code, portfolio_name, base_currency) values (?,?,?,?,?,'USD')",
                        [wh.new_id(), org, wh.new_id(), code, code])

    def know(*names):  # an instrument and an approved alias for the fund name, as E2c would leave them
        with wh.connect(plane) as con:
            for n in names:
                iid = wh.new_id()
                con.execute("insert into instruments (id, organization_id, canonical_name, instrument_type, created_at) values (?,?,?,'fund',?)", [iid, org, n, wh.now()])
                con.execute("insert into instrument_identifiers (id, organization_id, instrument_id, source_system_id, identifier_type, identifier_value, mapping_status) "
                            "values (?,?,?,?,'name',?,'approved')", [wh.new_id(), org, iid, sid, n])

    def land(name, data=None, as_name=None):
        data = data or (FIX / f"{name}.xlsx").read_bytes()
        return e0c_landing.receive(plane, data, as_name or f"{name}.xlsx", "ops@evergreen.test")["file_id"]

    def approve(fid):
        rec = e0c_landing.get(plane, fid)
        prof = e0b_tools.profile_sheet(e0c_landing.original_path(plane, rec))
        mid = e0b_mappings.save_proposal(plane, {**e0b_tools.propose_mapping(plane, prof)["mapping"], "file_name": rec["name"]})
        e0b_mappings.approve(plane, mid, ENG)
        return mid

    def edited(name, fn):
        wb = load_workbook(FIX / f"{name}.xlsx")
        fn(wb.active)
        buf = io.BytesIO()
        wb.save(buf)
        return buf.getvalue()

    return types.SimpleNamespace(plane=plane, know=know, land=land, approve=approve, edited=edited)


def one(plane, sql, *a):
    with wh.connect(plane) as con:
        return con.execute(sql, list(a)).fetchone()[0]


def test_good_load_writes_batch_records_locators_and_rows(env):
    env.know(*FUNDS)
    fid = env.land("pos_2026_09_thousands")
    mid = env.approve(fid)
    r = e0c_load.load(env.plane, fid, ENG)
    assert r["status"] == "loaded" and r["rows_loaded"] == 6 and r["held_back"] == [] and r["mapping_id"] == mid
    assert one(env.plane, "select count(*) from position_snapshots") == 6
    assert one(env.plane, "select sum(source_market_value) from position_snapshots") == 107956000  # thousands applied
    assert one(env.plane, "select count(*) from source_records where recorded_at is not null and sheet_name = 'Positions' and row_number is not null "
                          "and cell_range like 'A%' and raw_payload_uri like 'file:%'") == 6
    assert one(env.plane, "select min(row_number) from source_records") == 5
    assert one(env.plane, "select count(*) from batch_control_totals where passed") == 2  # row count + rows written
    assert one(env.plane, "select count(*) from position_snapshots where recorded_at is not null and view_type = 'operational' and data_status = 'official'") == 6
    assert e0c_load.load(env.plane, fid, ENG)["status"] == "already_loaded"
    from sh_platform.v1 import e0a_catalogue
    assert e0a_catalogue.layer_of(env.plane, "bronze_upload_positions") == "bronze"  # D1: a load registers its bronze table (B4 names)


def test_stated_total_is_checked_and_passes(env):
    env.know(*FUNDS)
    fid = env.land("pos_2026_09_totals")
    env.approve(fid)
    r = e0c_load.load(env.plane, fid, ENG)
    assert r["status"] == "loaded" and len(r["control_totals"]) == 3 and all(m["passed"] for m in r["control_totals"])


def test_totals_mismatch_rejects_everything(env):
    env.know(*FUNDS)
    fid = env.land("cash_2026_09_bad_total")
    env.approve(fid)
    r = e0c_load.load(env.plane, fid, ENG)
    assert r["status"] == "rejected_totals"
    assert "Total of Amount" in r["message"] and "2,357,000" in r["message"] and "2,332,000" in r["message"]
    assert one(env.plane, "select count(*) from cash_movements") == 0
    assert one(env.plane, "select count(*) from source_records") == 0
    assert one(env.plane, "select validation_status from source_batches") == "rejected_totals"
    assert one(env.plane, "select count(*) from batch_control_totals where not passed") == 1


def test_layout_change_stops_and_asks_for_a_new_proposal(env):
    env.know(*FUNDS)
    first = env.land("pos_2026_09_custodian")
    env.approve(first)
    changed = env.edited("pos_2026_09_custodian", lambda ws: ws.cell(4, 6).__setattr__("value", "Mkt Val Renamed"))
    fid = env.land("x", changed, as_name="pos_2026_10_custodian.xlsx")  # same pattern, one header renamed
    r = e0c_load.load(env.plane, fid, ENG)
    assert r["status"] == "waiting_for_mapping" and r["reason"] == "layout_changed"
    assert e0b_mappings.get(env.plane, r["proposal_id"])["layout_change"] is True
    assert one(env.plane, "select count(*) from position_snapshots") == 0 and one(env.plane, "select count(*) from source_batches") == 0


def test_no_approved_mapping_waits(env):
    r = e0c_load.load(env.plane, env.land("pos_2026_09_thousands"), ENG)
    assert r["status"] == "waiting_for_mapping" and r["reason"] == "no_mapping" and r["proposal_id"].startswith("map-")


def test_unknown_fund_is_held_back_and_listed(env):
    env.know(*FUNDS[:-1])  # Ashby has no instrument yet
    fid = env.land("val_mgr_estimate_dmy")
    env.approve(fid)
    r = e0c_load.load(env.plane, fid, ENG)
    assert r["status"] == "loaded" and r["rows_loaded"] == 5 and len(r["held_back"]) == 1
    assert r["held_back"][0]["row"] == 9 and "Ashby" in r["held_back"][0]["reason"]
    assert one(env.plane, "select count(*) from valuation_observations") == 5
    assert one(env.plane, "select count(*) from source_records where quality_status = 'held_back'") == 1
    assert one(env.plane, "select max(valuation_date) from valuation_observations").isoformat() == "2026-09-30"  # day-first text date


def test_a_later_file_for_the_same_subject_supersedes(env):
    env.know(*FUNDS)
    fid = env.land("val_mgr_estimate_dmy")
    env.approve(fid)
    assert e0c_load.load(env.plane, fid, ENG)["status"] == "loaded"
    amended = env.edited("val_mgr_estimate_dmy", lambda ws: ws.cell(4, 3).__setattr__("value", 16200000))
    assert e0c_load.load(env.plane, env.land("x", amended, as_name="val_mgr_estimate_dmy.xlsx"), ENG)["status"] == "loaded"
    assert one(env.plane, "select count(*) from valuation_observations") == 12
    assert one(env.plane, "select count(*) from valuation_observations where supersedes_valuation_id is not null") == 6


def test_a_csv_is_profiled_and_loaded_like_a_sheet(env):
    env.know("Cedar Ridge Growth Fund LP", "Marlow Global Equity Fund")
    csv = ("Position Date,Fund Name,Portfolio,Units,Market Value (USD)\n"
           "30/09/2026,Cedar Ridge Growth Fund LP,EVG-GROWTH,100,1500000\n"
           "30/09/2026,Marlow Global Equity Fund,EVG-CORE,50,2500000\n").encode()
    fid = env.land("x", csv, as_name="pos_2026_09_plain.csv")
    prof = e0b_tools.profile_sheet(e0c_landing.original_path(env.plane, e0c_landing.get(env.plane, fid)))
    assert prof["sheet"] == "pos_2026_09_plain" and prof["row_count"] == 2 and prof["header_row"] == 1
    env.approve(fid)
    r = e0c_load.load(env.plane, fid, ENG)
    assert r["status"] == "loaded" and r["rows_loaded"] == 2
    assert one(env.plane, "select sum(source_market_value) from position_snapshots") == 4000000


def test_a_rejected_attempt_stays_as_history_and_a_retry_adds_a_batch(env):
    env.know(*FUNDS)
    fid = env.land("cash_2026_09_bad_total")
    env.approve(fid)
    assert e0c_load.load(env.plane, fid, ENG)["status"] == "rejected_totals"
    assert e0c_load.load(env.plane, fid, ENG)["status"] == "rejected_totals"
    assert one(env.plane, "select count(*) from source_batches") == 2
    assert one(env.plane, "select count(*) from batch_control_totals") == one(env.plane, "select count(*) from batch_control_totals bt join source_batches b on b.id = bt.source_batch_id")


def test_an_iso_text_date_needs_no_format():
    """UC5: '2026-07-31' in a text cell reads as a date without a format in the mapping; other text dates still need one."""
    from datetime import date as _date
    from sh_platform.v1 import e0c_load
    assert e0c_load._convert("2026-07-31", "position_date", None) == _date(2026, 7, 31)
    with pytest.raises(ValueError):
        e0c_load._convert("31/07/2026", "position_date", None)


def test_a_totals_block_under_the_data_is_read_as_group_totals_not_rows(tmp_path):
    """UC5: 'Portfolio totals (USD)' after a blank row made the load count 31 rows for 27 and refuse the file."""
    import openpyxl
    from sh_platform.v1 import e0c_load
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Positions"
    ws.append(["Portfolio", "Instrument", "Market value (USD)"])
    for row in (["P1", "A", 10], ["P1", "B", 5], ["P2", "C", 7]):
        ws.append(row)
    ws.append([]); ws.append(["Portfolio totals (USD)", None, None]); ws.append(["P1", None, 15]); ws.append(["P2", None, 7])
    path = tmp_path / "pos.xlsx"; wb.save(path)
    _, _, rows, totals = e0c_load._read(path, "Positions", 1, {})
    assert len(rows) == 3
    assert totals[e0c_load.GROUP_TOTALS] == {"P1": {"Market value (USD)": 15}, "P2": {"Market value (USD)": 7}}
    assert totals[e0c_load.GROUP_COLUMN] == "Portfolio"


def test_every_sheet_is_loaded_in_one_step(tmp_path, monkeypatch):
    """UC5: a workbook with Positions and Latest statements loads both sheets with '*', each with its own mapping; the worst status is the file's."""
    import openpyxl
    from sh_platform.v1 import e0c_landing as LND, e0c_load
    wb = openpyxl.Workbook(); wb.active.title = "Positions"; wb.create_sheet("Latest statements"); p = tmp_path / "two.xlsx"; wb.save(p)
    assert LND.sheet_names(p) == ["Positions", "Latest statements"]
    seen = []
    monkeypatch.setattr(LND, "get", lambda plane, fid: {"state": "new", "sha256": "x", "name": "two.xlsx"})
    monkeypatch.setattr(LND, "original_path", lambda plane, rec: p)
    monkeypatch.setattr(LND, "set_state", lambda plane, fid, state, **kw: {"state": state, **kw})
    monkeypatch.setattr(e0c_load, "load", lambda plane, fid, viewer, model=None, sheet=None: seen.append(sheet) or
                        {"status": "loaded" if sheet == "Positions" else "waiting_for_mapping", "message": sheet})
    out = LND.load(None, "f1", None, sheet=LND.ALL_SHEETS)
    assert seen == ["Positions", "Latest statements"] and out["state"] == "waiting_for_mapping"
    assert set(out["last_load"]["sheets"]) == {"Positions", "Latest statements"}


def test_the_day_a_file_arrived_can_be_given_at_upload():
    """UC5: an upload after the fact keeps the real arrival time, so a check 'as received by' that day sees the file."""
    from sh_platform.v1 import e0c_landing as LND
    assert LND._arrival("2026-08-04") == "2026-08-04T00:00:00+00:00"
    assert LND._arrival("2026-08-04T06:00:00Z") == "2026-08-04T06:00:00+00:00"
    assert LND._arrival(None) is None
    with pytest.raises(LND.Refused):
        LND._arrival("2999-01-01")
    with pytest.raises(LND.Refused):
        LND._arrival("4 Aug")
