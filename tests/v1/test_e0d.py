"""E0d: synthetic workbooks and the evaluation set."""
import hashlib
import json
import sys
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "tools"))
import eval_data_agent as ev  # noqa: E402
import make_e0d_workbooks as gen  # noqa: E402
import seed_e0_demo  # noqa: E402

FIX = gen.FIX


def _cells(d):
    out = {}
    for p in sorted((d / "workbooks").glob("*.xlsx")):
        wb = load_workbook(p)
        out[p.name] = [(ws.title, sorted(map(str, ws.merged_cells.ranges)), [list(r) for r in ws.iter_rows(values_only=True)])
                       for ws in wb]
    return out


def _md5(d):
    return {p.name: hashlib.md5(p.read_bytes()).hexdigest() for p in (d / "expected").glob("*.json")}


def test_generator_deterministic(tmp_path):
    a, b = tmp_path / "a", tmp_path / "b"
    for d in (a, b):
        gen.main(d)
    assert _cells(a) == _cells(b) and len(_cells(a)) >= 20
    assert _md5(a) == _md5(b) == _md5(FIX)  # the committed set is what the generator writes
    assert _cells(a) == _cells(FIX)


def test_expected_targets_exist_in_dbml():
    fields = ev.dbml_fields()
    assert "quantity" in fields["position_snapshots"]
    exps = list((FIX / "expected").glob("*.json"))
    assert len(exps) >= 20
    for p in exps:
        e = json.loads(p.read_text())
        assert (FIX / "workbooks" / e["workbook"]).is_file()
        for src, m in e["columns"].items():
            if m["target"] == "ignore":
                continue
            t, f = m["target"].split(".")
            assert f in fields.get(t, []), f"{p.name}: {src} -> {m['target']}"
        for f in e["control_totals"]["sums"]:
            t, c = f.split(".")
            assert c in fields[t]


def test_every_trap_present_and_layout_real():
    seen = set()
    for p in (FIX / "expected").glob("*.json"):
        e = json.loads(p.read_text())
        seen |= set(e["traps"])
        ws = load_workbook(FIX / "workbooks" / e["workbook"])[e["sheet"]]
        assert sorted(c.value for c in ws[e["header_row"]] if c.value) == sorted(e["columns"])
        if "merged_header" in e["traps"]:
            assert ws.merged_cells.ranges
        if "totals_row" in e["traps"]:
            assert str(ws.cell(e["exclude_rows"][0], 1).value).lower() == "total"
        if "in_progress" in e["traps"]:
            assert e["status"] == "in_progress"
    assert set(gen.TRAPS) <= seen


def test_stated_totals_mismatch_is_real():
    e = json.loads((FIX / "expected" / "cash_2026_09_bad_total.json").read_text())
    ws = load_workbook(FIX / "workbooks" / e["workbook"])[e["sheet"]]
    body = sum(ws.cell(r, 5).value for r in range(e["first_data_row"], e["last_data_row"] + 1))
    assert ws.cell(e["exclude_rows"][0], 5).value != body


def test_baseline_eval_runs():
    r = ev.score("evergreen-test", ev.baseline_profile, ev.baseline_mapping)
    assert r["files"] >= 20 and 0.0 <= r["mapping_accuracy"] <= 1.0 and r["invented_targets"] == 0
    assert r["totals_mismatch_caught"] == "1/1" and r["in_progress_caught"] == "3/3"
    assert ev.main(["evergreen-test", "--baseline"]) == 0


def test_seed_does_not_overwrite(tmp_path):
    n = len(list(seed_e0_demo.SRC.glob("*.xlsx")))
    assert seed_e0_demo.seed("co", tmp_path) == (n, 0)
    f = tmp_path / "co" / "landing" / "inbox" / "nav_admin_2026_08.xlsx"
    f.write_bytes(b"mine")
    assert seed_e0_demo.seed("co", tmp_path) == (0, n) and f.read_bytes() == b"mine"
