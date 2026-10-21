"""W1, the file-format feeds: MT940, camt.053, BAI2, MT535, MT536 read through the pack's readers, and the connectors that fetch them. The samples are written by
tools/gen_feeds.py from the six-month office, so every figure is checked against the office's own bank statement and custody files. Synthetic data only."""
from __future__ import annotations

import csv
import email.message
import importlib.util
import json
from datetime import date
from decimal import Decimal
from pathlib import Path

import openpyxl
import pytest

from sh_platform.v1 import v3_harbor as H
from sh_platform.v1 import v3_pack_readers as PK
from sh_platform.v1 import wh
from sh_platform.v1.connectors import base
from sh_platform.v1.connectors import files as CF

HERE = Path(__file__).resolve().parent
OFFICE, FEEDS = HERE / "fixtures" / "office6m", HERE / "fixtures" / "feeds"
_SPEC = importlib.util.spec_from_file_location("g_w3_intake", HERE / "test_w3_intake.py")
W3 = importlib.util.module_from_spec(_SPEC)
_SPEC.loader.exec_module(W3)

MT940 = {"HB-HOLD-USD": "mt940_HB-HOLD-USD_2026-06-30.sta", "HB-TRUST-USD": "mt940_HB-TRUST-USD_2026-06-30.sta"}
CAMT = {"HB-HOLD-SGD": "camt053_HB-HOLD-SGD_2026-06-30.xml"}
BAI = {"TCB-HOLD-USD": "bai2_TCB-HOLD-USD_2026-06-30.bai"}
M535, M536 = "mt535_EG-LIQUID_2026-06-30.txt", "mt536_EG-LIQUID_2026-06-30.txt"
ALL = sorted(p.name for p in FEEDS.glob("*") if p.is_file())
KIND = {**{n: "swift_mt940" for n in MT940.values()}, **{n: "iso_camt053" for n in CAMT.values()}, **{n: "bai2" for n in BAI.values()}, M535: "swift_mt535", M536: "swift_mt536"}


def blob(name: str) -> bytes:
    return (FEEDS / name).read_bytes()


def q(plane, sql, *a):
    return W3.rows(plane, sql, *a)


def load(tmp_path, names=(), feeds=ALL, extra=None):
    plane, x = W3.flat_thread(tmp_path, list(names), {**{n: blob(n) for n in feeds}, **(extra or {})})
    PK.intake(x)
    return plane, x


def drop_more(plane, x, files: dict[str, bytes]):
    reg = json.loads(H._fp(plane).read_text(encoding="utf-8"))
    for n, data in files.items():
        H.drop(plane, n, data)
        reg.append({"name": n, "arrived_at": "2026-10-13T08:00:00Z", "status": "recognised", "source": "fund flows", "thread": "t-w3"})
    H._fp(plane).write_text(json.dumps(reg), encoding="utf-8")
    x.hide = set()
    PK.intake(x)


@pytest.fixture(scope="module")
def loaded(tmp_path_factory):
    plane, x = load(tmp_path_factory.mktemp("feeds"))
    assert PK.load(x) == (len(ALL), [])
    return plane


# ---- what the office has: the figures the feeds must tie out to ------------------------------------------------------------------------

def statement(account: str):
    f = next(OFFICE.glob(f"20??-??/bank_statements/{account.replace('-', '_')}_2026-06.xlsx"))
    wb = openpyxl.load_workbook(f, read_only=True)
    tx = [(r[1], date.fromisoformat(r[0]), Decimal(str(r[5])), r[2]) for r in list(wb["Transactions"].iter_rows(values_only=True))[1:]]
    last = list(wb["Daily balances"].iter_rows(values_only=True))[-1]
    return tx, Decimal(str(last[5])), date.fromisoformat(last[0])


def feed_lines(plane, name: str):
    return q(plane, "select m.reference, m.effective_date, m.amount, m.movement_type, s.sheet_name, s.row_number, s.raw_payload_uri from cash_movements m "
                    "join source_records s on s.id = m.source_record_id where s.raw_payload_uri like ? order by m.effective_date, m.reference", f"%/{name}")


def closing(plane, account_number: str, day: date):
    return q(plane, "select p.quantity, p.source_market_value from position_snapshots p join instruments i on i.id = p.instrument_id join bank_accounts b on b.cash_instrument_id = i.id "
                    "where b.account_ref = ? and p.position_date = ?", account_number[-4:], day)


# ---- reading ------------------------------------------------------------------------------------------------------------------------------

def test_the_five_formats_are_recognised_from_content_whatever_the_file_is_called():
    for name in ALL:
        assert PK.kind_of("renamed.bin", blob(name))[0] == KIND[name]
        assert PK.kind_of(name, blob(name))[0] == KIND[name]
    for other in ("2026-10/deal_pipeline/deal_pipeline_2026-10-13.csv", "2026-10/custodian_positions/positions_20261013.csv", "2026-10/bank_statements/HB_HOLD_USD_2026-09.xlsx"):
        assert PK.kind_of(Path(other).name, (OFFICE / other).read_bytes())[0] not in KIND.values()


def test_every_feed_file_is_read_and_says_what_it_did(tmp_path):
    plane, x = load(tmp_path)
    assert {v["kind"] for v in x.th["pack"].values()} == set(KIND.values())
    assert PK.load(x) == (6, [])


@pytest.mark.parametrize("account,name", [*MT940.items(), *CAMT.items(), *BAI.items()])
def test_a_bank_statement_ties_out_to_the_offices_own_statement_line_for_line(loaded, account, name):
    tx, close, day = statement(account)
    got = feed_lines(loaded, name)
    assert sorted((r[0], r[1], r[2]) for r in got) == sorted((t[0], t[1], t[2]) for t in tx)
    assert sum(r[2] for r in got) == sum(t[2] for t in tx)
    number = {"HB-HOLD-USD": "1101", "HB-TRUST-USD": "1138", "HB-HOLD-SGD": "1104", "TCB-HOLD-USD": "1112"}[account]
    assert [tuple(map(Decimal, r)) for r in closing(loaded, number, day)] == [(close, close)]
    d = q(loaded, "select count(*), max(document_date) from documents where original_filename = ? and document_type = 'bank_statement'", name)[0]
    assert d == (1, day)


def test_lines_are_typed_from_the_banks_code(loaded):
    types = {t[0]: t[1] for r in MT940.values() for t in [(x[0], x[3]) for x in feed_lines(loaded, r)]}
    assert types["IN-HB-HOLD-USD-05"] == "interest" and types["BF-HB-HOLD-USD-06"] == "fee" and types["FXS-20260602-HOLD"] == "fx_conversion"
    assert types["CC-CALL-A-11"] == "payment" and any(v == "settlement" for v in types.values())
    camt = {x[0]: x[3] for x in feed_lines(loaded, CAMT["HB-HOLD-SGD"])}
    assert set(camt.values()) >= {"interest", "fee", "fx_conversion", "payment"}
    bai = {x[0]: x[3] for x in feed_lines(loaded, BAI["TCB-HOLD-USD"])}
    assert sorted(bai.values()) == ["fee", "interest", "interest"]


def test_custody_holdings_tie_out_to_the_custodians_positions_of_30_june(loaded):
    want = {r["security_name"]: (Decimal(r["quantity"]), Decimal(r["market_value"])) for r in
            csv.DictReader((OFFICE / "2026-06/custodian_positions/positions_20260630.csv").open(encoding="utf-8"))}
    got = q(loaded, "select i.canonical_name, p.quantity, p.source_market_value from position_snapshots p join instruments i on i.id = p.instrument_id "
                    "join source_records s on s.id = p.source_record_id where s.raw_payload_uri like ?", f"%/{M535}")
    assert len(want) == 20 and {r[0]: (Decimal(r[1]), Decimal(r[2])) for r in got} == want
    assert q(loaded, "select count(*) from documents where original_filename = ? and document_type = 'custody_statement'", M535)[0][0] == 1


def test_custody_settlements_tie_out_to_the_custodians_buys_and_sells_of_june(loaded):
    want = [r for r in csv.DictReader(next(OFFICE.glob("20??-??/custodian_transactions/transactions_2026-06.csv")).open(encoding="utf-8")) if r["type"] in ("buy", "sell")]
    got = {r[0]: r for r in q(loaded, "select t.external_trade_id, t.transaction_type, t.quantity_delta, t.cash_amount, t.settlement_date, t.trade_date, t.status, s.sheet_name, s.row_number "
                                      "from transactions t join source_records s on s.id = t.source_record_id where s.raw_payload_uri like ?", f"%/{M536}")}
    assert len(want) == 3 and len(got) == len(want)
    lines = blob(M536).decode().split("\n")
    for w in want:
        r = got[w["reference"].removeprefix("CT-")]
        sign = -1 if w["type"] == "sell" else 1
        assert (r[1], Decimal(r[2]), Decimal(r[3]), r[4], r[5], r[6]) == (w["type"], sign * Decimal(w["quantity"]), -sign * Decimal(w["amount"]), date.fromisoformat(w["settle_date"]),
                                                                          date.fromisoformat(w["txn_date"]), "settled")
        assert r[7] == ":16R:SETTRAN" and lines[r[8] - 1] == ":16R:SETTRAN" and w["reference"] in "\n".join(lines[r[8]:r[8] + 14])


# ---- receipts ------------------------------------------------------------------------------------------------------------------------------

def test_every_row_keeps_the_file_and_the_line_or_tag_it_came_from(loaded):
    for name in [*MT940.values(), *CAMT.values(), *BAI.values()]:
        lines = blob(name).decode().split("\n")
        rows = feed_lines(loaded, name)
        assert rows
        for ref, _, _, _, tag, row, uri in rows:
            assert uri.endswith(f"/{name}") and uri.startswith("2026-06/") and tag in {":61:", "Ntry", "16"}
            assert ref in "\n".join(lines[row - 1:row + 3]) and (tag == "Ntry") == lines[row - 1].lstrip().startswith("<Ntry") and (tag != ":61:" or lines[row - 1].startswith(":61:"))
    for ref in q(loaded, "select s.sheet_name, s.row_number, s.raw_payload_uri, i.identifier_value from position_snapshots p join source_records s on s.id = p.source_record_id "
                         "join instrument_identifiers i on i.instrument_id = p.instrument_id and i.identifier_type = 'isin' where s.raw_payload_uri like ?", f"%/{M535}"):
        lines = blob(M535).decode().split("\n")
        assert ref[0] == ":16R:FIN" and lines[ref[1] - 1] == ":16R:FIN" and lines[ref[1]] == f":35B:ISIN {ref[3]}"


def test_the_feeds_are_their_own_sources_so_they_can_be_reconciled_to_the_files(loaded):
    codes = {r[0] for r in q(loaded, "select source_code from source_systems")}
    assert set(KIND.values()) <= codes


# ---- the same file or the same lines twice ----------------------------------------------------------------------------------------------------

def counts(plane):
    return [q(plane, f"select count(*) from {t}")[0][0] for t in ("cash_movements", "position_snapshots", "transactions", "documents", "extracted_fields", "source_records")]


def test_dropping_a_file_twice_writes_nothing_twice(tmp_path):
    plane, x = load(tmp_path)
    assert PK.load(x) == (6, [])
    first = counts(plane)
    drop_more(plane, x, {f"copy_of_{n}": blob(n) for n in ALL})  # the same files again under other names
    assert PK.load(x)[1] == [] and counts(plane) == first


def test_a_changed_file_with_the_same_lines_writes_no_line_twice(tmp_path):
    plane, x = load(tmp_path)
    assert PK.load(x) == (6, [])
    first = counts(plane)
    again = {"mt940_HB-HOLD-USD_2026-06-30_resent.sta": blob(MT940["HB-HOLD-USD"]) + b"\n\n", "mt535_EG-LIQUID_2026-06-30_resent.txt": blob(M535) + b"\n",
             "mt536_EG-LIQUID_2026-06-30_resent.txt": blob(M536).replace(b":28E:1/ONLY", b":28E:2/ONLY")}
    drop_more(plane, x, again)
    assert PK.load(x)[1] == []
    after = counts(plane)
    assert after[0:3] == first[0:3]  # lines, positions, transactions: none again
    assert after[3] == first[3] + 3  # one document per file read (statement, custody statement, settlements), but no rows


def test_a_line_the_offices_statement_already_has_is_not_written_again_in_either_order(tmp_path):
    xlsx = "2026-07/bank_statements/HB_HOLD_USD_2026-06.xlsx"
    ref = MT940["HB-HOLD-USD"]
    # one drop: the statement and the feed together
    plane, x = load(tmp_path / "a", names=[xlsx], feeds=[ref])
    assert PK.load(x) == (2, [])
    tx, close, day = statement("HB-HOLD-USD")
    assert q(plane, "select count(*) from cash_movements where reference in (select reference from cash_movements where reference like 'IN-HB-HOLD-USD-%')")[0][0] == 2
    assert q(plane, "select count(*) from cash_movements where currency = 'USD' and effective_date between date '2026-06-01' and date '2026-06-30'")[0][0] >= len(tx)
    assert len(feed_lines(plane, ref)) == 0  # every line is the statement's, so none is written again; the feed adds its closing balance only if the statement has none
    # statement first, the feed in a later drop
    plane, x = load(tmp_path / "b", names=[xlsx], feeds=[])
    assert PK.load(x) == (1, [])
    n = q(plane, "select count(*) from cash_movements")[0][0]
    drop_more(plane, x, {ref: blob(ref)})
    assert PK.load(x)[1] == [] and q(plane, "select count(*) from cash_movements")[0][0] == n


def test_a_file_that_does_not_read_is_set_aside_and_the_others_load(tmp_path):
    bad = blob(MT940["HB-HOLD-USD"]).replace(b":61:260601C283,12", b":61:26060XC283,12")
    plane, x = load(tmp_path, feeds=[M535], extra={"mt940_HB-HOLD-USD_2026-06-29_bad.sta": bad})
    ok, aside = PK.load(x)
    assert ok == 1 and len(aside) == 1 and "mt940_HB-HOLD-USD_2026-06-29_bad.sta" in aside[0]["file"]
    assert q(plane, "select count(*) from position_snapshots")[0][0] == 20 and q(plane, "select count(*) from cash_movements")[0][0] == 0


def test_a_statement_for_an_account_the_office_does_not_have_is_said_so(tmp_path):
    other = blob(MT940["HB-HOLD-USD"]).replace(b":25:HB-4400-1101", b":25:XX-0000-0000")
    plane, x = load(tmp_path, feeds=[], extra={"mt940_XX_2026-06-30.sta": other})
    rep = PK.read_files(plane, [("mt940_XX_2026-06-30.sta", other, "swift_mt940", "mt940_XX_2026-06-30.sta")], PK._standing(plane))[0]
    assert rep["swift_mt940"]["rows"] == 0 and "not one of the office's accounts" in rep["swift_mt940"]["not_loaded"][0]


def test_a_statement_whose_lines_do_not_add_up_is_flagged(tmp_path):
    off = blob(MT940["HB-HOLD-USD"]).replace(b"C260531USD723709,90", b"C260531USD723709,91")
    plane, x = load(tmp_path, feeds=[], extra={"mt940_HB-HOLD-USD_2026-06-30_off.sta": off})
    rep = PK.read_files(plane, [("mt940_HB-HOLD-USD_2026-06-30_off.sta", off, "swift_mt940", "mt940_HB-HOLD-USD_2026-06-30_off.sta")], PK._standing(plane))[0]
    assert any("is not the closing balance" in f for f in rep["swift_mt940"]["flagged"])


# ---- the connectors -------------------------------------------------------------------------------------------------------------------------

def test_the_connectors_are_registered():
    assert {"sftp_folder", "imap_inbox", "mt940", "camt053", "mt535_536", "bai2"} <= set(base.REGISTRY)
    assert base.REGISTRY["sftp_folder"].env == ("SH_FEED_SFTP_DIR",)
    assert base.REGISTRY["imap_inbox"].env == ("SH_FEED_IMAP_HOST", "SH_FEED_IMAP_USER", "SH_FEED_IMAP_PASSWORD")


def test_a_connector_without_its_settings_names_what_is_missing_and_fetches_nothing(tmp_path, monkeypatch):
    plane = W3.PlaneStore(tmp_path / "data").create("c")
    for n in (*base.REGISTRY["imap_inbox"].env, "SH_FEED_SFTP_DIR"):
        monkeypatch.delenv(n, raising=False)
    s = base.run(plane, base.REGISTRY["imap_inbox"], date(2026, 6, 1))
    assert not s.ok and s.missing == ["SH_FEED_IMAP_HOST", "SH_FEED_IMAP_USER", "SH_FEED_IMAP_PASSWORD"] and "password" not in s.message.lower()
    assert base.run(plane, base.REGISTRY["mt940"], date(2026, 6, 1)).missing == ["SH_FEED_SFTP_DIR"]


def test_the_sftp_folder_and_each_format_connector_fetch_their_own_files(tmp_path, monkeypatch):
    box = tmp_path / "sftp"
    box.mkdir()
    for n in ALL:
        (box / n).write_bytes(blob(n))
    (box / "readme.txt").write_text("not a feed", encoding="utf-8")
    (box / "undated.sta").write_bytes(blob(MT940["HB-HOLD-USD"]))
    monkeypatch.setenv("SH_FEED_SFTP_DIR", str(box))
    got = lambda cid, since=date(2026, 1, 1): sorted(f.name for f in base.REGISTRY[cid].fetch(since))  # noqa: E731
    assert len(got("sftp_folder")) == len(ALL) + 2
    assert set(MT940.values()) < set(got("mt940")) and len(got("mt940")) == 3 and got("camt053") == sorted(CAMT.values()) and got("bai2") == sorted(BAI.values())
    assert got("mt535_536") == [M535, M536]
    assert got("sftp_folder", date(2099, 1, 1)) == []
    assert all(CF.DATED.search(n) for n in got("mt940"))  # a file with no date in its name lands under the day it was fetched
    plane = W3.PlaneStore(tmp_path / "data").create("c")
    s = base.run(plane, base.REGISTRY["bai2"], date(2026, 1, 1))
    assert s.ok and s.files == 1 and (plane.path("dropbox") / BAI["TCB-HOLD-USD"]).read_bytes() == blob(BAI["TCB-HOLD-USD"])


class FakeImap:
    log: list = []

    def __init__(self, host):
        FakeImap.log = [("connect", host)]
        self.mail = {}

    def login(self, user, password):
        FakeImap.log.append(("login", user))

    def select(self, folder, readonly=False):
        FakeImap.log.append(("select", folder, readonly))

    def search(self, charset, *crit):
        FakeImap.log.append(("search", *crit))
        return "OK", [b"1 2"]

    def fetch(self, num, what):
        return "OK", [(b"1 (RFC822)", MAILS[int(num)])]

    def logout(self):
        FakeImap.log.append(("logout",))


def _mail(subject: str, day: str, attach: dict[str, bytes]) -> bytes:
    m = email.message.EmailMessage()
    m["Subject"], m["From"], m["To"], m["Date"], m["Message-ID"] = subject, "statements@bank.example", "ops@office.example", f"{day} 08:00:00 +0000", f"<{subject}@x>"
    m.set_content("Your statement is attached.")
    for n, data in attach.items():
        m.add_attachment(data, maintype="application", subtype="octet-stream", filename=n)
    return m.as_bytes()


MAILS = {1: _mail("June statement", "Tue, 30 Jun 2026", {"stmt.sta": blob(MT940["HB-HOLD-USD"]), "stmt.xml": blob(CAMT["HB-HOLD-SGD"])}), 2: _mail("Hello", "Wed, 01 Jul 2026", {})}


def test_the_imap_inbox_fetches_mails_and_their_attachments_with_a_fake_client(tmp_path, monkeypatch):
    for n, v in (("SH_FEED_IMAP_HOST", "imap.test"), ("SH_FEED_IMAP_USER", "ops@office.example"), ("SH_FEED_IMAP_PASSWORD", "not-a-real-secret")):
        monkeypatch.setenv(n, v)
    got = CF.ImapInbox(client_factory=FakeImap).fetch(date(2026, 6, 15))
    assert ("search", "SINCE", "15-Jun-2026") in FakeImap.log and ("select", "INBOX", True) in FakeImap.log and FakeImap.log[-1] == ("logout",)
    names = [f.name for f in got]
    assert sum(n.endswith(".eml") for n in names) == 2 and "2026-06-30_stmt.sta" in names and "2026-06-30_stmt.xml" in names
    assert all("not-a-real-secret" not in f.source for f in got) and all(f.source.startswith("IMAP inbox ops@office.example") for f in got)
    by = {f.name: f for f in got}
    assert PK.kind_of("2026-06-30_stmt.sta", by["2026-06-30_stmt.sta"].data)[0] == "swift_mt940" and PK.kind_of("x.xml", by["2026-06-30_stmt.xml"].data)[0] == "iso_camt053"
    # and what is fetched goes through the same readers
    plane, x = load(tmp_path, feeds=[], extra={"2026-06-30_stmt.sta": by["2026-06-30_stmt.sta"].data})
    assert PK.load(x) == (1, [])
    assert q(plane, "select count(*) from cash_movements")[0][0] == 6
