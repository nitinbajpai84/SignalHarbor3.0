"""K10 Consume reading room: the Morning Note, Daily P&L and NAV, Market signals, the fund league, the reports library, the report PDF and the daily
email, checked against docs/test/10UC-ANSWER-KEY.md (P01-P07, P14, P16, P21-P22, P25) on the seeded Evergreen universe (synthetic). No model is called."""
import json
from datetime import date
from decimal import Decimal as D
from pathlib import Path

import pytest
import yaml
from fastapi import FastAPI
from fastapi.testclient import TestClient

from sh_platform.plane import PlaneStore
from sh_platform.v1 import e1_access, e5_brief, k10_note as N, k10_routes, wh
from tools import seed_10uc as S

AS_OF = date(2026, 10, 9)
Q4 = D("0.0001")
Q2 = D("0.01")
ALL_ROLES = {"company_admin", "sponsor", "analyst", "owner", "approver", "participant"}


@pytest.fixture(scope="module")
def world(tmp_path_factory):
    mp = pytest.MonkeyPatch()
    mp.setenv("SH_MODEL_PROVIDER", "off")
    wh.reset_cache()
    root = tmp_path_factory.mktemp("shk10")
    S.seed(root)
    plane = PlaneStore(root).plane_for("evergreen").ensure()
    ids, _ = S.EG.accounts(root)
    V = S.viewers(ids)
    d = V["Daniel Foo"]
    unsigned = N.signals(plane, d, "USD", AS_OF, "month")   # V6: nothing is compared until the head of data signs the fund-to-index mapping
    anita = V["Anita Rao"]
    N.sign_index_map(plane, anita, {f["fund_id"]: f["proposed_series_key"] for f in N.index_map(plane, anita)["funds"] if f["proposed_series_key"]})
    yield {"unsigned": unsigned, "plane": plane, "V": V, "d": d, "note": N.morning_note(plane, d, "USD", AS_OF), "sig": N.signals(plane, d, "USD", AS_OF, "month"),
           "league": N.league(plane, d, AS_OF), "lib": N.reports_library(plane, d), "pnl": N.pnl(plane, d, "USD", AS_OF),
           "today": N.today(plane, d, "USD", AS_OF), "health": N.health(plane, d, "USD", AS_OF)}
    mp.undo()


# ---- Morning Note -----------------------------------------------------------------------------------------------------------

def test_note_headline_and_tiles(world):
    n = world["note"]
    assert n["headline"] == "The office is worth USD 57.78 million, 180,000 more than September's report said."                       # P01, P02
    t = {x["key"]: x for x in n["tiles"]}
    assert (t["value"]["value"], t["value"]["note"], t["value"]["status"]) == (D("57780000"), "+180,000 since the report", "Revised")
    assert t["return"]["value"].quantize(Q4) == D("4.1667") and t["return"]["note"] == "Benchmark 2.71% · +1.46 pts"                  # P14
    assert (t["calls"]["value"], t["calls"]["note"]) == (D("1050000"), "450,000 confirmed · 600,000 not yet")                         # P21, P22
    assert t["limits"]["value"] == "2" and all(x["trail_id"] for x in n["tiles"])


def test_note_what_moved_it_and_who_is_waiting(world):
    n = world["note"]
    moved = {m["label"]: m for m in n["moved"]}
    assert moved["Evergreen Growth I changed since the report"]["amount"] == D("180000") and moved["The liquid book rose in September"]["amount"] == D("1648000")   # P06
    assert [m["amount"] for m in n["moved"][-2:]] == [D("132000"), D("130400")]
    assert n["needs_you"]["count"] == 9 and len(n["needs_you"]["items"]) == 3 and n["needs_you"]["href"] == "/centre/action"
    assert "Kestrel Ridge Macro Fund" in n["watch"][0]["text"] and n["built"]["text"].endswith("signed by Daniel Foo.")
    assert {m["name"] for m in n["markets"]} >= {"Global equity index", "US dollar in Singapore dollars"} and all(m["outside"] for m in n["markets"])


def test_note_in_singapore_dollars(world):
    n = N.morning_note(world["plane"], world["d"], "SGD", AS_OF)
    assert n["tiles"][0]["value"] == D("77494536.00") and "SGD 77.49 million" in n["headline"]                                         # P16


# ---- Daily P&L and NAV ------------------------------------------------------------------------------------------------------

def test_pnl_issued_against_known_now_in_both_currencies(world):
    p = world["pnl"]
    tot = next(l for l in p["issued"]["lines"] if l["label"] == "Total")
    assert (tot["issued"], tot["known_now"], tot["difference"]) == (D("57600000"), D("57780000"), D("180000"))
    assert (p["other_total"]["issued"], p["other_total"]["known_now"]) == (D("77253120.00"), D("77494536.00")) and p["rate"]["rate"] == D("1.3412")   # P16
    assert p["liquid_book"]["month"]["change"] == D("1648000") and not p["liquid_book"]["day"]["available"]
    assert "Not available yet" in p["liquid_book"]["day"]["message"]


def test_pnl_administrator_against_our_view(world):
    b = {x["portfolio"]: x for x in world["pnl"]["bridge"]}
    assert b["Entity B operating portfolio"]["difference"] == D("350000") and b["Evergreen Liquid portfolio"]["difference"] == 0       # P03, P04
    assert "Saltmarsh" in b["Entity B operating portfolio"]["why"] and b["Evergreen Liquid portfolio"]["why"] == "Nothing to explain."
    assert world["pnl"]["bridge_total"]["difference"] == D("730000") and world["pnl"]["total"] == D("103405000")
    assert {e["name"]: e["value"] for e in world["pnl"]["by_entity"]}["Evergreen Holdings"] == D("16580000")


# ---- Market signals ---------------------------------------------------------------------------------------------------------

def _row(sig, name):
    return next(r for r in sig["indexes"] if r["name"].startswith(name))


def test_index_moves_for_the_month(world):
    s = world["sig"]
    eq, bd, px = _row(s, "Global equity"), _row(s, "Global bond"), _row(s, "Private markets proxy")
    assert [x["periods"]["month"]["change_pct"].quantize(Q4) for x in (eq, bd, px)] == [D("3.5002"), D("0.4002"), D("3.7502")]          # P14
    assert (eq["periods"]["month"]["start_level"], eq["periods"]["month"]["end_level"]) == (D("1038.22"), D("1074.56"))
    assert eq["direction"] == "Rising" and eq["outside"] and eq["source"] == "Index provider" and px["source"] == "Signed rule"
    fx = _row(s, "US dollar")
    assert fx["periods"]["month"]["end_level"] == D("1.3412") and fx["latest"] == {"date": date(2026, 10, 8), "level": D("1.3388")}      # P16
    assert fx["source"] == "WMR 4pm fixing"


def test_a_market_with_no_feed_says_so(world):
    r = _row(world["sig"], "Regional equity")
    assert r["periods"]["month"]["change_pct"] is None and r["message"] == "Not available yet." and r["source"] is None


def test_periods_year_to_date_and_since_inception(world):
    ytd = N.signals(world["plane"], world["d"], "USD", AS_OF, "ytd")
    e = _row(ytd, "Global equity")["periods"]["ytd"]
    assert (e["start"], e["end"]) == (date(2025, 12, 31), date(2026, 9, 30)) and e["change_pct"].quantize(Q2) == D("8.89") and ytd["period_start"] == date(2025, 12, 31)
    inc = N.signals(world["plane"], world["d"], "USD", AS_OF, "inception")
    assert _row(inc, "Global equity")["periods"]["inception"]["change_pct"].quantize(Q2) == D("7.46") and inc["period_start"] == date(2025, 9, 30)
    with pytest.raises(ValueError):
        N.signals(world["plane"], world["d"], "USD", AS_OF, "decade")


def test_portfolios_against_their_benchmarks(world):
    by = {r["benchmark_label"]: r for r in world["sig"]["portfolios"] if r["portfolio"] == "Evergreen Liquid portfolio"}
    blend, sleeve = by["Office blend"], by["60/40 equity and bond"]
    assert [x.quantize(Q4) for x in (blend["return_pct"], blend["benchmark_pct"], blend["ahead_pts"])] == [D("4.1667"), D("2.7072"), D("1.4595")]   # P14
    assert [x.quantize(Q4) for x in (sleeve["benchmark_pct"], sleeve["ahead_pts"])] == [D("2.2602"), D("1.9065")] and blend["status"] == "Official"
    assert "version 2 is still being agreed" in world["sig"]["benchmark_note"] and "global equity index 42%" in world["sig"]["benchmark_note"]


def test_funds_against_their_index_have_no_invented_return(world):
    f = {x["fund"]: x for x in world["sig"]["funds"]}
    assert f["Kestrel Ridge Macro Fund"]["compared_with"] == "Global bond index" and f["Kestrel Ridge Equity Long/Short Fund"]["compared_with"] == "Global equity index"
    assert f["Evergreen Growth I"]["compared_with"].startswith("Private markets proxy") and f["Evergreen Income II"]["compared_with"] == "Global bond index"
    assert all(x["fund_return"] is None and x["alpha"] is None and "Not available yet" in x["message"] for x in f.values())


def test_no_fund_is_compared_until_the_mapping_is_signed_by_the_head_of_data(world):
    plane, V = world["plane"], world["V"]
    assert all(x["compared_with"] is None and not x["index_signed"] and x["message"] == N.NO_INDEX for x in world["unsigned"]["funds"])
    m = N.index_map(plane, V["Anita Rao"])
    assert m["signed"] and m["signed_by"] == "Anita Rao" and m["can_sign"] and N.index_map(plane, V["Leo Park"])["can_sign"] is False
    for who in ("Leo Park", "Daniel Foo", "Mei Chen"):   # only the head of data seat signs
        with pytest.raises(N.Refused) as e:
            N.sign_index_map(plane, V[who], {m["funds"][0]["fund_id"]: "IDX:GLBD:SYN"})
        assert e.value.status == 403
    with pytest.raises(N.Refused):
        N.sign_index_map(plane, V["Anita Rao"], {m["funds"][0]["fund_id"]: "IDX:NOT:REAL"})
    f = m["funds"][0]
    out = N.sign_index_map(plane, V["Anita Rao"], {f["fund_id"]: "IDX:GLBD:SYN"})   # a fund's index can be changed by a new signature
    assert out["version"] == m["version"] + 1 and next(x for x in out["funds"] if x["fund_id"] == f["fund_id"])["series_key"] == "IDX:GLBD:SYN"
    N.sign_index_map(plane, V["Anita Rao"], {f["fund_id"]: f["series_key"]})        # put it back for the tests that follow


def test_the_short_read_is_worded_from_the_figures(world):
    r = world["sig"]["read"]
    assert "returned 4.17% while global equity index rose 3.50%" in r and "Northwind Logistics (+132,000)" in r and "Silverpine Software (+130,400)" in r
    assert "US dollar rose against the Singapore dollar" in r


def test_manager_news_is_outside_material_with_source_and_date(world):
    news = world["sig"]["news"]
    k = next(n for n in news if n["manager"] == "Kestrel Ridge Advisors" and n["source"] == "Manager newsletter")
    assert k["date"] == "2026-09-28" and k["outside"] and "Priya Nair is sole portfolio manager of the Macro Fund" in k["text"]       # P25
    assert all(n["source"].startswith("Manager") and n["date"] for n in news) and news == sorted(news, key=lambda n: n["date"], reverse=True)


def test_a_part_that_fails_leaves_the_rest(world, monkeypatch):
    monkeypatch.setattr(N, "_news", lambda *a: 1 / 0)
    monkeypatch.setattr(N, "_funds", lambda *a: 1 / 0)
    s = N.signals(world["plane"], world["d"], "USD", AS_OF, "month")
    assert s["news"] == [] and s["funds"] == [] and s["indexes"] and s["portfolios"]


# ---- Fund league ------------------------------------------------------------------------------------------------------------

def test_league_ranks_nine_funds_and_flags_the_untrusted(world):
    lg = world["league"]
    assert len(lg["funds"]) == 9 and lg["managers"] == 6 and [f["rank"] for f in lg["funds"]] == list(range(1, 10))
    f = {x["fund"]: x for x in lg["funds"]}
    assert (f["Evergreen Growth I"]["value"], f["Evergreen Growth I"]["status"]) == ("10180000", "Revised") and "replaced the estimate (+180,000)" in f["Evergreen Growth I"]["note"]   # P07
    assert f["Cedarpoint Ventures III"]["status"] == "Stale" and f["Cedarpoint Ventures III"]["note"].startswith("Last valued 30 Jun (101 days)")
    assert f["Cedarpoint Opportunities II"]["status"] == "Late" and "192 days" in f["Cedarpoint Opportunities II"]["note"]
    assert f["Saltmarsh Infrastructure Fund II"]["status"] == "Unverified" and f["Kestrel Ridge Macro Fund"]["status"] == "Estimated"
    s = lg["summary"]
    assert (s["official"], s["revised"], s["estimated"], s["old"], s["total"], s["ahead"]) == (3, 1, 2, 3, 9, None)
    assert "192 days old" in lg["worst"]["text"]


def test_league_can_be_ranked_by_value(world):
    by_value = N.league(world["plane"], world["d"], AS_OF, "value")["funds"]
    vals = [D(f["value"]) for f in by_value]
    assert vals == sorted(vals, reverse=True) and by_value[0]["fund"] == "Evergreen Growth I"
    with pytest.raises(ValueError):
        N.league(world["plane"], world["d"], AS_OF, "luck")


# ---- Reports library and the PDF --------------------------------------------------------------------------------------------

def test_reports_library_lists_the_signed_report_and_the_live_views(world):
    rows = {r["name"]: r for r in world["lib"]["reports"]}
    r = rows["September investment report"]
    assert (r["state"], r["version"], r["signed_by"]) == ("Signed", 1, "Daniel Foo") and r["pdf"].endswith("/pdf") and r["summary"].startswith("USD and SGD. Issued 3 Oct")
    assert rows["Exposure and limits"]["summary"] == "2 limit cases open." and rows["Manager scorecard"]["state"] == "Live"
    assert world["lib"]["waiting"] == 0 and world["lib"]["pdf_available"]


def test_a_signed_version_exports_as_a_pdf(world):
    import fitz
    v = world["lib"]["reports"][0]["versions"][0]["id"]
    data, name = N.report_pdf(world["plane"], world["d"], v)
    assert data[:5] == b"%PDF-" and name == "september-investment-report-2026-09-30.pdf"
    text = "".join(p.get_text() for p in fitz.open(stream=data, filetype="pdf"))
    assert "57,600,000" in text and "57,780,000" in text and "signed by Daniel Foo" in text


def test_a_signed_version_exports_as_a_powerpoint(world, monkeypatch):
    import io

    from pptx import Presentation
    v = world["lib"]["reports"][0]["versions"][0]["id"]
    assert world["lib"]["pptx_available"] and world["lib"]["reports"][0]["pptx"].endswith("/pptx")
    data, name = N.report_pptx(world["plane"], world["d"], v)
    assert name == "september-investment-report-2026-09-30.pptx"
    text = " ".join(sh.text_frame.text if sh.has_text_frame else "|".join(c.text for r in sh.table.rows for c in r.cells) for sl in Presentation(io.BytesIO(data)).slides for sh in sl.shapes
                    if sh.has_text_frame or sh.has_table)
    assert "57,600,000" in text and "57,780,000" in text and "signed by Daniel Foo" in text
    monkeypatch.setattr(N, "pptx_available", lambda: False)
    with pytest.raises(N.Refused) as e:
        N.report_pptx(world["plane"], world["d"], v)
    assert e.value.status == 404 and "cannot be made" in e.value.message


def test_an_unsigned_version_and_a_missing_one_are_refused(world, monkeypatch):
    v = world["lib"]["reports"][0]["versions"][0]["id"]
    real = e5_brief.brief
    monkeypatch.setattr(e5_brief, "brief", lambda *a, **k: {**real(*a, **k), "version_status": "draft"})
    with pytest.raises(N.Refused) as e:
        N.report_pdf(world["plane"], world["d"], v)
    assert e.value.status == 409 and "waiting for a signature" in e.value.message
    monkeypatch.setattr(e5_brief, "brief", real)
    with pytest.raises(N.Refused) as e:
        N.report_pdf(world["plane"], world["d"], "00000000-0000-0000-0000-000000000000")
    assert e.value.status == 404
    monkeypatch.setattr(N, "pdf_available", lambda: False)
    with pytest.raises(N.Refused) as e:
        N.report_pdf(world["plane"], world["d"], v)
    assert e.value.status == 404 and "cannot be made" in e.value.message


# ---- the daily email --------------------------------------------------------------------------------------------------------

def test_email_carries_the_same_figures_and_escapes_everything(world):
    note = {**world["note"], "headline": "<script>alert(1)</script> The office", "company": 'Evergreen "Office" & Co'}
    e = N.render_email(note, "https://sh.example")
    assert e["subject"] == "Morning Note, 9 Oct: the office is worth USD 57.78 million"
    assert "<script" not in e["html"] and "&lt;script&gt;" in e["html"] and "&amp; Co" in e["html"]
    clean = N.render_email(world["note"])
    for must in ("USD 57.78 million", "4.17%", "+1.46 pts", "USD 1.05 million", "450,000 confirmed"):
        assert must in clean["html"] and must in clean["text"]
    assert "src=" not in clean["html"] and "<script" not in clean["html"] and 'href="/centre/action/cases/CASE-0006"' in clean["html"]
    assert 'href="https://sh.example/centre/consume"' in e["html"]


# ---- access and routes ------------------------------------------------------------------------------------------------------

def test_a_person_sees_only_their_entities(world, monkeypatch):
    plane, v = world["plane"], world["V"]["Mei Chen"]
    with wh.connect(plane, read_only=True) as con:
        pf = {str(r[0]) for r in con.execute("select p.id from portfolios p join legal_entities e on e.id = p.legal_entity_id where e.entity_code = 'EG-HOLD'").fetchall()}
        ent = {str(r[0]) for r in con.execute("select id from legal_entities where entity_code = 'EG-HOLD'").fetchall()}
    monkeypatch.setattr(e1_access, "visible_portfolios", lambda p, x: set(pf))
    monkeypatch.setattr(e1_access, "visible_entities", lambda p, x: set(ent))
    monkeypatch.setattr(e1_access, "level", lambda p, x, portfolio_id=None, **k: "read" if str(portfolio_id) in pf else "none")   # cases are filtered by the same access
    out = {"note": N.morning_note(plane, v, "USD", AS_OF), "pnl": N.pnl(plane, v, "USD", AS_OF), "sig": N.signals(plane, v, "USD", AS_OF, "month")}
    text = json.dumps(out, default=str)
    assert "Entity B" not in text and "22875000" not in text and "22525000" not in text
    assert [r["portfolio"] for r in out["sig"]["portfolios"]] == ["Evergreen Growth portfolio"]


def test_every_route_is_declared_for_all_six_roles():
    cfg = yaml.safe_load((Path(N.__file__).parent / "k10.permissions.yaml").read_text(encoding="utf-8"))
    declared = {f"GET /api/v1/consume/{p}" for p in ("today", "health", "note", "email-note", "pnl", "signals", "league", "reports", "index-map", "reports/{version_id}/pdf", "reports/{version_id}/pptx")}
    sign = "POST /api/v1/consume/index-map"   # V6: the one write; the head of data seat is checked again in k10_note
    assert set(cfg["routes"]) == declared | {sign} and all(set(v["roles"]) == ALL_ROLES for k, v in cfg["routes"].items() if k != sign)
    assert set(cfg["routes"][sign]["roles"]) == {"company_admin", "analyst"}
    assert {f"{next(iter(r.methods))} {r.path}" for r in k10_routes.router.routes} == declared | {sign}


def test_the_routes_send_plain_strings_and_refuse_bad_input(world, monkeypatch):
    monkeypatch.setattr(k10_routes, "viewer_of", lambda r: world["d"])
    monkeypatch.setattr(k10_routes, "_plane", lambda r: world["plane"])
    app = FastAPI()
    app.include_router(k10_routes.router)
    c = TestClient(app)
    r = c.get("/api/v1/consume/note?currency=SGD&as_of=2026-10-09")
    assert r.status_code == 200 and r.json()["tiles"][0]["value"] == "77494536.00" and "0E-" not in r.text
    assert c.get("/api/v1/consume/note?currency=EUR").status_code == 400 and c.get("/api/v1/consume/signals?period=decade").status_code == 400
    assert c.get("/api/v1/consume/league?rank=luck").status_code == 400 and c.get("/api/v1/consume/note?as_of=yesterday").status_code == 400
    assert c.get("/api/v1/consume/email-note?as_of=2026-10-09").json()["subject"].startswith("Morning Note, 9 Oct")
    for p in ("pnl", "signals", "league", "reports", "today", "health"):
        assert c.get(f"/api/v1/consume/{p}?as_of=2026-10-09").status_code == 200
    v = world["lib"]["reports"][0]["versions"][0]["id"]
    pdf = c.get(f"/api/v1/consume/reports/{v}/pdf")
    assert pdf.status_code == 200 and pdf.headers["content-type"] == "application/pdf" and "attachment" in pdf.headers["content-disposition"] and pdf.content[:5] == b"%PDF-"
    gone = c.get("/api/v1/consume/reports/00000000-0000-0000-0000-000000000000/pdf")
    assert gone.status_code == 404 and "could not find" in gone.json()["message"]


# ---- CONS1: Today, Financial health, where each outside line came from, and a faster room ------------------------------------

def test_today_is_one_sentence_four_numbers_and_what_moved(world):
    t = world["today"]
    assert t["headline"] == "The office is worth USD 103.41 million. The liquid book rose 1.65 million in September and beat its benchmark by 1.46 points."
    k = {x["key"]: x for x in t["tiles"]}
    assert (k["nav"]["value"], k["nav"]["note"]) == (D("103405000"), "4 portfolios · 3 entities")
    assert (k["month"]["label"], k["month"]["value"], k["month"]["note"]) == ("P&L, September", D("1648000"), "Liquid book 4.17%")
    assert (k["bench"]["kind"], k["bench"]["value"].quantize(Q2), k["bench"]["note"]) == ("points", D("1.46"), "Benchmark 2.71%")
    assert k["day"]["available"] is False and k["day"]["value"] is None and k["day"]["note"] == "Waits for the daily price file."   # never an invented figure
    assert all(x["trail_id"] and "sources" not in (x["where"] or {}) for x in t["tiles"])
    labels = [m["label"] for m in t["moved"]]
    assert labels[:3] == ["The liquid book rose in September", "Largest contributor: Northwind Logistics", "Next: Silverpine Software"]
    loose = [m for m in t["moved"] if m.get("unconfirmed")]
    assert {m["label"] for m in loose} >= {"Saltmarsh Infrastructure Fund II"} and all(m["status"] in ("Estimated", "Unverified") for m in loose)
    assert "needs_you" not in t and "reports" not in t   # Action has the cases; Reports has the reports


def test_today_names_the_source_of_the_markets_line(world):
    m = world["today"]["markets"]
    assert m["text"].startswith("Equity +3.50%, bond +0.40% in September, USD/SGD 1.33")
    assert m["sources"] and all(s["tag"] == "FILE" and s["label"] == "Market data file (demo data)" and s["date"] for s in m["sources"])


def test_today_in_singapore_dollars_and_with_a_part_failing(world, monkeypatch):
    assert N.today(world["plane"], world["d"], "SGD", AS_OF)["tiles"][0]["currency"] == "SGD"
    monkeypatch.setattr(N, "_markets_line", lambda *a: 1 / 0)
    assert N.today(world["plane"], world["d"], "USD", AS_OF)["markets"] is None
    with pytest.raises(ValueError):
        N.today(world["plane"], world["d"], "EUR", AS_OF)


def test_health_has_each_portfolio_entity_and_how_far_the_value_is_trusted(world):
    h = world["health"]
    rows = {r["portfolio"]: r for r in h["portfolios"]}
    liquid = rows["Evergreen Liquid portfolio"]
    assert (liquid["value"], liquid["pnl"], liquid["trust"]["status"]) == (D("41200000"), D("1648000"), "Official")
    assert liquid["return_pct"].quantize(Q2) == D("4.17") and liquid["ahead_pts"].quantize(Q2) == D("1.46") and liquid["benchmark_label"] == "Office blend"
    assert rows["Evergreen Growth portfolio"]["trust"] == {"status": "Revised", "detail": "Evergreen Growth I"}
    assert rows["Evergreen Funds portfolio"]["trust"]["status"] == "Estimated" and "Kestrel Ridge" in rows["Evergreen Funds portfolio"]["trust"]["detail"]
    assert rows["Entity B operating portfolio"]["trust"]["status"] == "Stale"
    assert h["total"]["value"] == D("103405000") == sum(D(r["value"]) for r in h["portfolios"]) and h["total"]["count"] == 4 and h["total"]["unexplained"] == 0
    assert h["total"]["pnl"] == sum(r["pnl"] for r in h["portfolios"])
    assert {e["entity"]: e["portfolios"] for e in h["entities"]} == {"Entity B": 1, "Evergreen Family Trust": 2, "Evergreen Holdings": 1}
    assert sum(e["value"] for e in h["entities"]) == h["total"]["value"]
    assert h["period"]["month"] == "September" and all(r["where"]["text"] and r["trail_id"] for r in h["portfolios"])


def test_health_fund_counts_match_the_league(world):
    f, s = world["health"]["funds"], world["league"]["summary"]
    assert (f["official"], f["revised"], f["estimated"], f["old"], f["total"]) == (s["official"], s["revised"], s["estimated"], s["old"], s["total"]) == (3, 1, 2, 3, 9)
    assert f["estimated_who"] == "Kestrel Ridge Advisors ×2" and "Cedarpoint Ventures" in f["old_who"]


def test_every_outside_series_names_its_source_date_and_whether_it_is_a_file_or_a_live_api(world):
    rows = {r["series_key"]: r for r in world["sig"]["indexes"]}
    assert rows["IDX:GLEQ:SYN"]["provenance"] == {"label": "Market data file (demo data)", "tag": "FILE", "date": date(2026, 9, 30), "demo": True}
    assert rows["FX:USDSGD:WMR_4PM"]["provenance"]["date"] == date(2026, 10, 8)
    assert rows["proxy"]["provenance"]["tag"] is None and "equity index" in rows["proxy"]["provenance"]["label"]   # calculated here, not an outside line
    assert world["sig"]["headlines"] == []   # no news API is connected in the seed


def test_a_live_api_series_and_its_headlines_are_tagged_live(world):
    from sh_platform.consume import market_store
    from sh_platform.v1.c5_connect import api as C5
    plane = world["plane"]
    C5.write_rows(plane, "fred", [{"series": "DGS10", "series_key": "YLD:US10Y:FRED", "label": "x", "type": "yield", "date": d, "value": D(v), "base": None, "quote": None}
                                  for d, v in ((date(2026, 8, 31), "4.00"), (date(2026, 9, 30), "4.20"), (date(2026, 10, 8), "4.10"))])
    market_store.add_news(plane, [{"id": "news-1", "title": "Rates held", "summary": "", "measures": [], "kind": "news", "live_api": True, "outside": True, "source": "Finnhub · Wire", "date": "2026-10-08"}])
    try:
        sig = N.signals(plane, world["d"], "USD", AS_OF, "month")
        y = next(r for r in sig["indexes"] if r["series_key"] == "YLD:US10Y:FRED")
        assert y["kind"] == "yield" and y["name"].startswith("US 10-year Treasury yield") and y["provenance"]["tag"] == "LIVE API" and y["provenance"]["label"] == "FRED (St. Louis Fed)"
        assert not y["provenance"]["demo"] and y["periods"]["month"]["change_abs"] == D("0.20")
        assert [(h["title"], h["source"], h["date"], h["tag"]) for h in sig["headlines"]] == [("Rates held", "Finnhub · Wire", "2026-10-08", "LIVE API")]
        assert "Treasury" not in N.today(plane, world["d"], "USD", AS_OF)["markets"]["text"]   # the one line stays equity, bonds and the exchange rate
    finally:
        with wh.connect(plane) as con:
            con.execute("delete from market_observations where series_key = 'YLD:US10Y:FRED'")
        market_store._path(plane).write_text("{}", encoding="utf-8")


def test_the_same_question_in_one_request_is_worked_out_once(world, monkeypatch):
    from sh_platform.v1 import scope_memo
    plane, d = world["plane"], world["d"]
    calls = []
    real = wh.connect
    monkeypatch.setattr(wh, "connect", lambda *a, **k: calls.append(1) or real(*a, **k))
    with scope_memo.scope():
        first = N.today(plane, d, "USD", AS_OF)
        n_first = len(calls)
        again = N.today(plane, d, "USD", AS_OF)
    assert again == first and len(calls) < n_first * 1.3   # the second pass reads almost nothing: the answers are reused
    calls.clear()
    N.today(plane, d, "USD", AS_OF)
    N.today(plane, d, "USD", AS_OF)
    assert len(calls) >= n_first * 1.8   # outside a scope nothing is kept, so tests and scripts see no change


def test_signing_inside_a_request_answers_with_what_was_signed(world):
    from sh_platform.v1 import scope_memo
    plane, anita = world["plane"], world["V"]["Anita Rao"]
    m = N.index_map(plane, anita)
    f = m["funds"][0]
    with scope_memo.scope():   # as the POST route runs it
        N.index_map(plane, anita)
        out = N.sign_index_map(plane, anita, {f["fund_id"]: "IDX:GLBD:SYN"})
        again = N.sign_index_map(plane, anita, {f["fund_id"]: f["series_key"]})   # a second signature reads the file afresh, never the earlier read
    assert out["version"] == m["version"] + 1 and again["version"] == m["version"] + 2 == N.index_map(plane, anita)["version"]


def test_a_caller_cannot_change_what_the_next_caller_sees():
    from sh_platform.v1 import scope_memo

    @scope_memo.memo
    def answer():
        return {"a": [1]}
    with scope_memo.scope():
        answer()["a"].append(2)
        assert answer() == {"a": [1]}


def test_the_room_keeps_answers_until_something_is_written(world, monkeypatch):
    from sh_platform.v1 import read_cache
    monkeypatch.setattr(k10_routes, "viewer_of", lambda r: world["d"])
    monkeypatch.setattr(k10_routes, "_plane", lambda r: world["plane"])
    runs = []
    real = N.today
    monkeypatch.setattr(N, "today", lambda *a, **k: runs.append(1) or real(*a, **k))
    app = FastAPI()
    app.include_router(k10_routes.router)

    @app.post("/api/anything")
    def write():
        return {"ok": True}
    app.add_middleware(read_cache.BumpOnWrite)
    c = TestClient(app)
    url = "/api/v1/consume/today?as_of=2026-10-09"
    a = c.get(url).json()
    assert c.get(url).json() == a and len(runs) == 1               # the second visit is served from what was kept
    c.get(url + "&currency=SGD")
    assert len(runs) == 2                                            # another question is its own answer
    c.post("/api/anything")                                          # a write anywhere
    c.get(url)
    assert len(runs) == 3                                            # drops everything kept
    monkeypatch.setattr(k10_routes, "viewer_of", lambda r: world["V"]["Mei Chen"])
    c.get(url)
    assert len(runs) == 4                                            # another viewer never gets this viewer's answer
