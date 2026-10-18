"""T1b: templates are document skeletons people choose; the three documents open laid out from them; North's patch fills the matching
section; a required grey section blocks signing; the process map template can be switched without losing content; export keeps layout."""
import io
import shutil
from pathlib import Path

import pytest
from fastapi.testclient import TestClient

from sh_platform.app import create_app
from sh_platform.jsonfile import write_json_atomic
from sh_platform.plane import PlaneStore
from sh_platform.profile.view import Viewer
from sh_platform.settings import Settings
from sh_platform.v1 import i2_north as N
from sh_platform.v1 import i4_sow as S
from sh_platform.v1 import i6_chat as C
from sh_platform.v1 import t1_layout as L
from sh_platform.v1 import t1_templates as T
from tests.v1.seeded import SEED

ROOT = Path(__file__).resolve().parents[2]
CREDS = SEED / "_control" / "seed_credentials.txt"
ARJUN = "arjun.mehta.neera@example.test"
pytestmark = pytest.mark.skipif(not CREDS.is_file(), reason="no seeded demo data")
ADMIN = Viewer("Arjun Mehta", frozenset({"company_admin"}), False, "acc-1")


@pytest.fixture(scope="module")
def world(tmp_path_factory):
    mp = pytest.MonkeyPatch()
    mp.setenv("SH_MODEL_PROVIDER", "off")
    data = tmp_path_factory.mktemp("t1b") / "data"
    shutil.copytree(SEED, data)
    app = create_app(Settings(mode="dev", data_root=str(data), session_secret="fake-test-secret-not-real"))
    app.state.model_env = {"SH_MODEL_PROVIDER": "off"}
    yield app, PlaneStore(data)
    mp.undo()


def client(app):
    pw = {l.split()[0]: l.split()[1] for l in CREDS.read_text().splitlines() if l.strip() and not l.startswith("#")}
    c = TestClient(app)
    assert c.post("/api/auth/login", json={"email": ARJUN, "password": pw[ARJUN]}).status_code == 200
    c.headers["X-CSRF-Token"] = c.cookies.get("sh_csrf", "")
    return c


def doc(lay, kind):
    return next(d for d in lay["documents"] if d["kind"] == kind)


def sec(lay, kind, key):
    return next(s for s in doc(lay, kind)["sections"] if s["key"] == key)


def test_two_templates_per_kind_are_seeded_for_every_company_and_once(world):
    _, store = world
    for company in ("neera", "fresh"):
        plane = store.plane_for(company)
        T.ensure(plane)
        before = sorted(p.name for p in plane.path("templates").glob("*.json"))
        T.ensure(plane)
        assert before == sorted(p.name for p in plane.path("templates").glob("*.json"))      # nothing written twice
        names = {k: sorted(t["name"] for t in T.listing(plane, ADMIN, k)) for k in T.KINDS}
        assert names["intent_of_work"] == ["Quick data request", "Reporting automation"]
        assert names["process_map"] == ["Step list", "Swim lane"] and names["definitions"] == ["KPI cards", "KPI glossary"]
        assert all(len(names[k]) == 2 for k in T.KINDS if k != "benchmarks")
        std = T.template_for(plane, "intent_of_work")
        assert std["name"] == "Reporting automation" and std["sections"] and T.goals_from(std) and T.required_keys(std)
        assert [f["key"] for f in std["fields"]] == [s["key"] for s in std["sections"]]       # fields are derived from the sections


def test_choosing_lays_out_the_three_documents_from_the_templates(world):
    _, store = world
    plane = store.plane_for("neera")
    iid = N.create(plane, ADMIN, "A one-off request")["id"]
    pin = T.choose(plane, iid, "quick-data-request")
    assert pin["companions"]["process_map"]["id"] == "step-list" and pin["companions"]["definitions"]["id"] == "kpi-cards"
    lay = C.view(plane, iid)["layout"]
    assert [d["kind"] for d in lay["documents"]] == ["intent_of_work", "process_map", "definitions"]
    assert [d["template"]["name"] for d in lay["documents"]] == ["Quick data request", "Step list", "KPI cards"]
    intent = doc(lay, "intent_of_work")
    assert intent["sections"][-1]["key"] == T.FREE_KEY and intent["sections"][-1]["heading"] == "Anything else"
    for d in lay["documents"]:
        for s in d["sections"]:
            assert s["heading"] and s["instruction"]                                           # every heading with its grey instruction
    grey = [s["key"] for s in intent["sections"] if not s["filled"]]
    assert {"kpis", "inputs", "rules"} <= set(grey) and sec(lay, "intent_of_work", "goal")["filled"]   # the first line fills the goal
    assert sec(lay, "definitions", "kpis")["value"]["columns"][:2] == ["name", "definition"]  # the empty table keeps its columns
    other = N.create(plane, ADMIN, "A monthly report")["id"]                                   # the default pins the standard and its companions
    assert T.choose(plane, other)["companions"]["process_map"]["id"] == "standard-process-map"
    assert doc(C.view(plane, other)["layout"], "intent_of_work")["template"]["name"] == "Reporting automation"


def test_norths_patch_fills_the_matching_section_and_grey_sections_block_signing(world):
    _, store = world
    plane = store.plane_for("neera")
    iid = N.create(plane, ADMIN, "A one-off request")["id"]
    T.choose(plane, iid, "quick-data-request")
    block = T.sign_block(plane, iid)
    assert block.startswith("The intent of work cannot be signed yet.") and "The figures wanted" in block and "Where the data is" in block
    sow = N.read_sow(plane, iid)
    C.merge(sow, {"goal": {"line": "Custody fees for September", "why": "The committee asks", "needed_by": "Friday"},
                  "kpis": [{"name": "Custody fees", "definition": "Fees paid to custodians", "owner": "Finance"}],
                  "sources": [{"name": "Finance invoices", "kind": "file", "holds": "Invoices", "owner": "Sam"}],
                  "governance": [{"rule": "Committee members only", "kind": "access"}],
                  "template_fields": {T.FREE_KEY: "Please keep it short"}}, None, {}, "North")
    write_json_atomic(N._path(plane, iid, "sow.json"), sow)
    lay = C.view(plane, iid)["layout"]
    assert sec(lay, "intent_of_work", "kpis")["filled"] and sec(lay, "intent_of_work", "kpis")["value"] == ["Custody fees: Fees paid to custodians, Finance"]
    assert sec(lay, "intent_of_work", "frequency")["value"] == "Friday" and sec(lay, "intent_of_work", "comments")["value"] == "Please keep it short"
    assert sec(lay, "intent_of_work", "inputs")["from"] == ["Finance invoices"]                # read from a file, marked as such
    assert sec(lay, "definitions", "kpis")["value"]["rows"][0]["name"] == "Custody fees"      # the same content in the definitions document
    assert T.sign_block(plane, iid) is None                                                    # nothing required is grey now


def test_a_required_section_of_a_companion_is_named(world):
    _, store = world
    plane = store.plane_for("fresh")
    admin = Viewer("Fresh Owner", frozenset({"company_admin"}), False, "acc-9")
    iid = N.create(plane, admin, "Another request")["id"]
    T.choose(plane, iid, "quick-data-request")
    sow = N.read_sow(plane, iid)
    C.merge(sow, {"kpis": [{"name": "Fees", "definition": "Fees"}], "sources": [{"name": "Ledger", "kind": "system"}],
                  "governance": [{"rule": "Admins only"}], "goal": {"needed_by": "Friday"}}, None, {}, "North")
    write_json_atomic(N._path(plane, iid, "sow.json"), sow)
    assert T.sign_block(plane, iid) is None
    T.switch(plane, iid, "process_map", "standard-process-map")                                # the swim lane needs its steps
    assert "Process map: Lanes and steps" in T.sign_block(plane, iid)


def test_switching_the_process_map_template_keeps_the_content(world):
    app, store = world
    plane = store.plane_for("neera")
    iid = N.create(plane, ADMIN, "A monthly report")["id"]
    T.choose(plane, iid)
    sow = N.read_sow(plane, iid)
    C.merge(sow, {}, [{"name": "Pull file", "who": "Operations", "does": "Downloads it", "system": "Portal"},
                      {"name": "Reconcile", "who": "Finance", "does": "In a spreadsheet", "system": "Ledger"}], {}, "North")
    write_json_atomic(N._path(plane, iid, "sow.json"), sow)
    before = C.view(plane, iid)["layout"]
    assert doc(before, "process_map")["template"]["name"] == "Swim lane" and len(sec(before, "process_map", "steps")["value"]["rows"]) == 2
    c = client(app)
    assert c.put(f"/api/v1/intents/{iid}/template/process_map", json={"template_id": "step-list"}).status_code == 200
    after = C.view(plane, iid)["layout"]
    assert doc(after, "process_map")["template"]["name"] == "Step list" and sec(after, "process_map", "steps")["view"] == "steps"
    assert sec(after, "process_map", "steps")["value"]["rows"] == sec(before, "process_map", "steps")["value"]["rows"]
    assert doc(after, "process_map")["map"]["steps"][0]["name"] == "Pull file"
    assert c.put(f"/api/v1/intents/{iid}/template/process_map", json={"template_id": "kpi-cards"}).status_code == 404     # not a process map template
    assert c.put(f"/api/v1/intents/{iid}/template/benchmarks", json={"template_id": "standard-benchmarks"}).status_code == 422


def test_export_keeps_the_templates_layout_and_headings(world):
    app, store = world
    plane = store.plane_for("neera")
    iid = N.create(plane, ADMIN, "A one-off request")["id"]
    T.choose(plane, iid, "quick-data-request")
    c = client(app)
    md = c.get(f"/api/v1/intents/{iid}/sow/download?format=md")
    assert md.status_code == 200 and "### The figures wanted" in md.text and "### Where the data is" in md.text and "## Process map (Step list)" in md.text
    assert "_One line: what you need and why it matters now._" not in md.text                 # the goal is filled, so no grey instruction
    assert "_Each figure, in your words._" in md.text                                         # an empty section keeps its grey instruction
    if S.docx_available():
        import docx
        r = c.get(f"/api/v1/intents/{iid}/sow/download?format=docx")
        heads = [p.text for p in docx.Document(io.BytesIO(r.content)).paragraphs if p.style.name.startswith("Heading")]
        assert "The figures wanted" in heads and "Anything else" in heads
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(c.get(f"/api/v1/intents/{iid}/sow/download?format=xlsx").content))
    assert wb.sheetnames == ["Intent of work", "Process map", "Definitions"]
    assert any(row[0] == "The figures wanted" for row in wb["Intent of work"].iter_rows(values_only=True))
    old = N.create(plane, ADMIN, "Began before templates")["id"]                               # no pin: the plain eight sections, as before
    assert "## 1. Goal" in c.get(f"/api/v1/intents/{old}/sow/download?format=md").text
    assert L.layout(plane, old) is not None


def test_signed_intent_and_retired_template_cannot_be_pinned_or_switched(world):
    app, store = world
    plane = store.plane_for("neera")
    iid = N.create(plane, ADMIN, "A signed one")["id"]
    T.choose(plane, iid)
    write_json_atomic(N._path(plane, iid, "sow.json"), {**N.read_sow(plane, iid), "state": "signed"})
    with pytest.raises(N.Refused) as e:
        T.switch(plane, iid, "process_map", "step-list")
    assert e.value.status == 409
    T.retire(plane, ADMIN, "step-list")
    assert (T._read(T._p(plane, "settings.json")) or {}).get("seed") == T.SEED_VERSION  # retire keeps the seed key
    with pytest.raises(N.Refused):
        T.choose(plane, N.create(plane, ADMIN, "Another")["id"], "step-list")


def test_filled_sections_carry_edits_and_a_correction_replaces_the_rule(world):
    _, store = world
    plane = store.plane_for("neera")
    iid = N.create(plane, ADMIN, "A one-off request")["id"]
    T.choose(plane, iid, "quick-data-request")
    sow = N.read_sow(plane, iid)
    C.merge(sow, {"kpis": [{"name": "Custody fees", "definition": "Fees paid"}],
                  "governance": [{"rule": "Only the CEO team, Nitin Guard and Yakub may view the report", "kind": "access"}]}, None, {}, "North")
    last = {"governance": [dict(r) for r in sow["governance"]]}
    C.merge(sow, {"governance": [{"rule": "Only the CEO team, Nitin Guard, Yakub and Gerd may view the report", "kind": "access"}]}, None, last, "North")
    assert [r["rule"] for r in sow["governance"]] == ["Only the CEO team, Nitin Guard, Yakub and Gerd may view the report"]
    C.merge(sow, {"governance": [{"rule": "Keep it seven years", "kind": "retention"}]}, None, {}, "North")
    assert len(sow["governance"]) == 2                                                         # a different rule is still a new line
    write_json_atomic(N._path(plane, iid, "sow.json"), sow)
    lay = C.view(plane, iid)["layout"]
    for kind, key in (("intent_of_work", "kpis"), ("definitions", "kpis")):
        s = sec(lay, kind, key)
        assert s["filled"] and any(e["section"] == "kpis" and e["index"] == 0 and e["field"] == "definition" for e in s["edits"])
    rules = sec(lay, "intent_of_work", "rules")
    assert rules["filled"] and [(e["section"], e["field"], e["index"]) for e in rules["edits"]] == [("governance", "rule", 0), ("governance", "rule", 1)]
