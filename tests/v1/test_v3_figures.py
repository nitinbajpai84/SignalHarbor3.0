"""Harbor 3.0 figures (sh_platform/v1/v3_figures.py): each figure of the brief ends in a filled ready table or in a check line that names it.
Real: the 2.0 design rules (b10_design), the intake reader (e0c_load), the warehouse. Stubbed: the HTTP calls of a chain (a recording `api`), the model."""
import json
from types import SimpleNamespace

import pytest

from sh_platform.plane import PlaneStore
from sh_platform.v1 import b10_design as D
from sh_platform.v1 import e0c_load, e1_write
from sh_platform.v1 import v3_figures as F
from sh_platform.v1 import v3_harbor as H
from sh_platform.v1 import wh
from tools import seed_evergreen as EG

NW = {"name": "Net worth", "definition": "Sum of positions value plus cash for each portfolio", "unit": "USD", "period": "monthly"}
BREAKS = {"name": "Breaks over 5,000", "definition": "Cash items above the threshold", "unit": "GBP", "period": "daily"}
GOLD = [{"name": "gold_p_holdings_monthly", "purpose": "Value of each fund", "columns": [{"name": "value", "type": "DECIMAL(18,2)"}], "depends_on": ["valuation_observations", "instruments"], "kpis": []},
        {"name": "gold_p_net_worth_entity_monthly", "purpose": "Net worth by entity", "columns": [{"name": "net_worth", "type": "DECIMAL(18,2)"}], "depends_on": ["gold_p_holdings_monthly"], "kpis": []},
        {"name": "gold_p_exceptions", "purpose": "Exceptions", "columns": [{"name": "title", "type": "VARCHAR"}], "depends_on": ["exceptions"], "kpis": []}]


@pytest.fixture()
def plane(tmp_path):
    EG.company(tmp_path)
    p = PlaneStore(str(tmp_path)).plane_for("evergreen").ensure()
    wh.use(p)
    return p


def test_a_figure_is_covered_by_a_table_that_lists_it_names_a_column_for_it_or_shares_its_words():
    assert [g["name"] for g in F.covering(NW, GOLD)] == ["gold_p_net_worth_entity_monthly"]
    assert F.covering(BREAKS, GOLD) == []  # "exceptions" is not "breaks": a figure with no table is not matched to the nearest one
    assert F.covers({"name": "Allocation drift"}, {"name": "gold_p_x", "kpis": ["allocation drift"], "columns": []})
    assert F.covers({"name": "Allocation drift"}, {"name": "gold_p_x", "columns": [{"name": "allocation_drift"}]})


def test_a_ready_table_is_dead_only_when_every_fact_table_it_reads_is_empty():
    assert F.sources_of(GOLD[1], GOLD) == {"valuation_observations", "instruments"}  # through the other gold table
    rows = {"valuation_observations": 0, "position_snapshots": 29, "instruments": 33}  # master data is there before any file
    assert F.dead(GOLD[1], GOLD, rows)
    assert not F.dead(GOLD[1], GOLD, {**rows, "valuation_observations": 3})
    assert not F.dead(GOLD[2], GOLD, rows)  # reads no fact table: not judged dead before it is built


def test_what_a_figure_is_split_by_comes_from_its_words_or_the_goal_that_names_it():
    assert F.dims_of({"name": "Cash by entity", "definition": ""}, "") == ["entity"]
    assert F.dims_of(NW, "gives me net worth by portfolio on a board") == ["portfolio"]  # "for each portfolio" in its definition
    assert F.dims_of({"name": "Allocation drift", "definition": "Difference from target"}, "gives me net worth by portfolio") == []  # the goal names another figure


def test_a_design_that_has_no_column_for_the_figure_does_not_hold_it():
    assert F.has_column("Net worth", ["month", "portfolio", "net_worth"])
    assert not F.has_column("Time-weighted return net of fees", ["month", "portfolio", "net_worth"])  # net worth is not the return on it
    assert not F.has_column("Allocation drift", ["month", "net_worth"])


def test_a_column_of_nulls_is_a_placeholder_not_a_design():
    sql = "create table g as select m, sum(v) as net_worth, cast(null as numeric(18,2)) as \"Allocation drift\" from t group by m"
    assert F.null_columns(sql) == {"Allocation drift"}
    assert F.null_columns("create table g as select m, sum(v) as net_worth from t group by m") == set()


def test_the_question_names_the_month_the_figure_and_the_dimension():
    assert F.parse_when("What was it at 30 September 2026?") == (2026, 9, 30)
    assert F.parse_when("net worth in September 2026") == (2026, 9, None)
    assert F.parse_when("on 8 October?") == (None, 10, 8)
    assert F.parse_when("2026-08") == (2026, 8, None) and F.parse_when("what are we worth") is None
    ks = [NW, {"name": "Cash balance per entity"}]
    assert F.kpi_for("What was the Evergreen Liquid portfolio worth at 30 September 2026", ks) is NW
    assert F.kpi_for("Which holiday is next", ks) is None
    cash = [{"name": "Cash balance per entity", "definition": "Administrator cash figure for each legal entity"}, {"name": "Breaks over 5,000", "definition": "Cash items above the threshold flagged"}]
    assert F.kpi_for("Which entity has a cash break on 8 October and by how much?", cash)["name"] == "Breaks over 5,000"  # "break" is only in one figure
    assert F.kpi_for("What is our cash by entity today?", cash)["name"] == "Cash balance per entity"
    names = {"EG-LIQUID": "Evergreen Liquid portfolio", "EG-GROWTH": "Evergreen Growth portfolio", "EG-FUNDS": "Evergreen Funds portfolio"}
    q = F._tokens("What was the Evergreen Liquid portfolio worth")
    assert F.match_values(list(names), names, q) == ["EG-LIQUID"]  # the name's words, not the code's
    assert F.match_values(list(names), names, F._tokens("What are the Evergreen portfolios worth")) == []  # all alike: no filter


def _ready(plane, rows):
    with wh.connect(plane) as con:
        con.execute("create table gold_p_position_snapshots_monthly_by_portfolio (month date, portfolio varchar, net_worth decimal(18,2))")
        con.executemany("insert into gold_p_position_snapshots_monthly_by_portfolio values (?, ?, ?)", rows)
    return {"name": "gold_p_position_snapshots_monthly_by_portfolio",
            "columns": [{"name": "month", "type": "DATE"}, {"name": "portfolio", "type": "VARCHAR"}, {"name": "net_worth", "type": "DECIMAL(18,2)"}], "rows": len(rows)}


def test_the_answer_is_read_from_the_ready_table_for_the_month_and_portfolio_asked(plane):
    t = _ready(plane, [("2026-08-01", "EG-LIQUID", 40100000), ("2026-09-01", "EG-LIQUID", 41200000), ("2026-09-01", "EG-GROWTH", 16400000)])
    names = {"EG-LIQUID": "Evergreen Liquid portfolio", "EG-GROWTH": "Evergreen Growth portfolio"}
    got = F._read(plane, t, NW, "What was the Evergreen Liquid portfolio worth at 30 September 2026?", names)
    assert got["figure"]["value"] == "41,200,000.00" and "September 2026" in got["answer"] and "Evergreen Liquid portfolio" in got["answer"]
    both = F._read(plane, t, NW, "What is our net worth in September 2026?", names)  # no portfolio named: each one, and the total
    assert both["rows"] == [{"label": "Evergreen Liquid portfolio", "value": "41,200,000.00"}, {"label": "Evergreen Growth portfolio", "value": "16,400,000.00"}] and "Total 57,600,000.00" in both["answer"]
    later = F._read(plane, t, NW, "What is the Evergreen Liquid portfolio worth in November 2026?", names)  # a month the table does not have: said, not guessed
    assert later["figure"] is None and later["answer"].startswith("Not available yet") and "September 2026" in later["answer"]


def test_a_ready_table_that_keeps_the_masters_id_is_read_by_name_and_a_run_id_is_not_a_split(plane):
    ids = {"43709bc3-b8f0-485f-a204-4130bb032b7d": "Evergreen Liquid portfolio", "f606dc49-4b26-4bcf-bd44-7f5a5f617076": "Evergreen Growth portfolio"}
    with wh.connect(plane) as con:
        con.execute("create table gold_p_metrics (run_id varchar, portfolio_id uuid, month timestamp, net_worth decimal(38,8))")
        con.executemany("insert into gold_p_metrics values ('run-1', ?, '2026-09-01', ?)", [(k, v) for k, v in zip(ids, (41200000, 16400000))])
    t = {"name": "gold_p_metrics", "rows": 2, "columns": [{"name": "run_id", "type": "VARCHAR"}, {"name": "portfolio_id", "type": "UUID"}, {"name": "month", "type": "TIMESTAMP"},
                                                            {"name": "net_worth", "type": "DECIMAL(38,8)"}]}
    got = F._read(plane, t, NW, "What was the Evergreen Liquid portfolio worth at 30 September 2026?", ids)
    assert got["figure"]["value"] == "41,200,000" and "run-1" not in got["answer"] and "Evergreen Liquid portfolio" in got["answer"]


def test_a_table_with_no_split_by_what_the_question_asks_says_so_instead_of_giving_a_total(plane):
    t = _ready(plane, [("2026-09-01", "EG-LIQUID", 41200000)]) if False else None
    with wh.connect(plane) as con:
        con.execute("create table gold_p_cash_by_day (day date, administrator varchar, cash_variance decimal(18,2))")
        con.execute("insert into gold_p_cash_by_day values ('2026-10-08', 'desk', -33200), ('2026-10-07', 'desk', 100)")
    t = {"name": "gold_p_cash_by_day", "rows": 2, "columns": [{"name": "day", "type": "DATE"}, {"name": "administrator", "type": "VARCHAR"}, {"name": "cash_variance", "type": "DECIMAL(18,2)"}]}
    k = {"name": "Cash variance"}
    no = F._read(plane, t, k, "Which entity has a cash variance on 8 October?", {})
    assert no["figure"] is None and "no split by entity" in no["answer"]
    yes = F._read(plane, t, k, "What is the cash variance on 8 October?", {})  # no year: the 8th of October the table has
    assert yes["figure"]["value"] == "-33,200.00"
    assert F._read(plane, t, k, "What is the cash variance on 9 October?", {})["answer"].startswith("Not available yet")


def test_an_office_with_no_figure_asked_for_leaves_the_question_to_lumen(plane):
    assert F.figure_answer(plane, SimpleNamespace(), "hello", None, None) is None
    assert F.figure_answer(plane, SimpleNamespace(), "hello", "prj-1", {"id": "t-1", "iid": None}) is None


# ---- the design step and the pack -----------------------------------------------------------------------------------------------------------

def _x(plane, pack_state="approved", gold=None, kpis=None, rows=None, design=None):
    calls: list = []

    def api(who, method, path, body=None, **kw):
        calls.append((who, method, path.split("/architecture")[-1], body))
        if method == "GET" and path.endswith("/pack"):
            return {"pack": {"state": pack_state, "target": {"gold": gold if gold is not None else GOLD}}, "versions": [{"version": 1}, {"version": 2}]}
        return {}
    x = SimpleNamespace(plane=plane, th={"id": "t-1", "pid": "prj-1", "iid": "ini-1"}, api=api, app=None, account_id=None)
    return x, calls


def test_a_figure_with_no_ready_table_gets_one_designed_and_the_approved_pack_becomes_a_new_version(plane, monkeypatch):
    new = {"name": "gold_p_position_snapshots_monthly_by_portfolio", "purpose": "Net worth", "sql": "create table gold_p_position_snapshots_monthly_by_portfolio as select cast(date_trunc('month', f.position_date) as date) as month, sum(f.source_market_value) as net_worth from position_snapshots as f group by 1",
           "depends_on": ["position_snapshots"], "kpis": ["Net worth"], "col_defs": {"net_worth": "Sum of positions"}, "grain": "Month", "keys": ["month"]}
    monkeypatch.setattr(F, "kpis_of", lambda plane, iid: [NW])
    monkeypatch.setattr(F, "design", lambda x, todo, filled: ([new], {}))
    monkeypatch.setattr(F, "table_rows", lambda plane, copy=False: {"valuation_observations": 0, "position_snapshots": 29} if not copy else {"gold_p_net_worth_entity_monthly": 0})
    x, calls = _x(plane)
    monkeypatch.setattr(H, "b_run", lambda x: calls.append(("leo", "RUN", "round for the new version", None)))
    assert F.cover(x, "loaded") == "1 ready tables added"
    verbs = [(w, m, p) for w, m, p, _ in calls if m != "GET"]
    assert verbs[0] == ("leo", "PUT", "/pack/edit") and ("leo", "POST", "/pack/submit") in verbs and ("anita", "POST", "/pack/approve") in verbs
    assert verbs.index(("leo", "RUN", "round for the new version")) > verbs.index(("anita", "POST", "/pack/approve"))  # the new version gets its own round before anyone goes live
    added = calls[1][3]
    assert added["section"] == "gold_add" and added["key"] == new["name"] and added["changes"]["definitions"]["month"].startswith("The first day")  # every column defined: nothing for a person to settle
    assert x.th["figures"][0]["table"] == "gold_p_net_worth_entity_monthly"  # the old table still lists the figure; the new one is the one with rows


def test_a_draft_pack_is_edited_in_place_and_a_figure_that_is_covered_costs_no_design(plane, monkeypatch):
    monkeypatch.setattr(F, "kpis_of", lambda plane, iid: [NW])
    monkeypatch.setattr(F, "table_rows", lambda plane, copy=False: {})
    monkeypatch.setattr(F, "design", lambda *a: pytest.fail("a covered figure is not designed again"))
    built = [GOLD[0], {**GOLD[1], "purpose": "Recipe net_worth: Net worth from the pack's recipe"}, GOLD[2]]  # a figure with a recipe is covered once its table is the recipe's
    x, calls = _x(plane, pack_state="draft", gold=built)
    assert F.cover(x, "draft") == "all 1 figures have a ready table"
    monkeypatch.setattr(F, "design", lambda x, todo, filled: ([], {"Breaks over 5,000": "it is a count under a condition"}))
    monkeypatch.setattr(F, "kpis_of", lambda plane, iid: [BREAKS])
    assert F.cover(x, "draft") == "0 ready tables added; no table for Breaks over 5,000"
    assert [c for c in calls if c[1] != "GET"] == []  # nothing was written, and nothing was approved


def test_a_design_the_plan_check_refuses_is_named_with_its_figure_and_the_chain_goes_on(plane, monkeypatch):
    new = {"name": "gold_p_bad", "sql": "create table gold_p_bad as select 1", "depends_on": [], "kpis": ["Net worth"]}
    monkeypatch.setattr(F, "kpis_of", lambda plane, iid: [NW])
    monkeypatch.setattr(F, "design", lambda x, todo, filled: ([new], {}))
    monkeypatch.setattr(F, "table_rows", lambda plane, copy=False: {})
    x, _ = _x(plane, gold=[])

    def refuse(who, method, path, body=None, **kw):
        if method == "PUT":
            raise H.Q("The plan check refused gold_p_bad.")
        return {"pack": {"state": "draft", "target": {"gold": []}}}
    x.api = refuse
    assert F.cover(x, "draft").startswith("0 ready tables added")
    assert "The plan check refused" in x.th["figure_notes"]["Net worth"]


def test_decision_three_names_each_figure_with_no_numbers_and_why(plane, monkeypatch):
    from sh_platform.v1 import b9_verify as V
    with wh.connect(plane) as con:
        con.execute("create table gold_p_net_worth_entity_monthly (net_worth decimal(18,2))")
        con.execute("create table gold_p_position_snapshots_monthly_by_portfolio (month date, net_worth decimal(18,2))")
        con.execute("insert into gold_p_position_snapshots_monthly_by_portfolio values ('2026-09-01', 41200000)")
    gold = [*GOLD, {"name": "gold_p_position_snapshots_monthly_by_portfolio", "purpose": "Net worth", "depends_on": ["position_snapshots"],
                    "columns": [{"name": "month", "type": "DATE"}, {"name": "net_worth", "type": "DECIMAL(18,2)"}], "kpis": []}]
    x, _ = _x(plane, gold=gold)
    x.th["exceptions"] = [{"title": "Evergreen_Office_Investment_Report_2026-09.pdf did not load"}]
    x.th["figure_notes"] = {"Breaks over 5,000": "Span could not place it on a table that holds data."}
    x.api = lambda who, method, path, body=None, **kw: {"pack": {"target": {"gold": gold}}}
    monkeypatch.setattr(V, "_copy", lambda plane: wh.connect(plane, read_only=True))  # no trial-run copy in this office: the warehouse itself
    monkeypatch.setattr(F, "kpis_of", lambda plane, iid: [NW, BREAKS, {"name": "Holdings value"}])
    F.check(x)
    by = {f["kpi"]: f for f in x.th["figures"]}
    assert by["Net worth"]["ok"] and by["Net worth"]["table"] == "gold_p_position_snapshots_monthly_by_portfolio" and by["Net worth"]["rows"] == 1
    assert not by["Breaks over 5,000"]["ok"] and by["Breaks over 5,000"]["reason"].startswith("Span could not place it")
    assert not by["Holdings value"]["ok"] and "did not build" in by["Holdings value"]["reason"]
    lines = F.d3_lines(x.th)
    assert lines[0] == {"label": "Net worth", "value": "1 rows", "ok": True, "note": "Ready table position snapshots monthly by portfolio has 1 rows."}
    assert lines[1]["value"] == "Not available yet" and lines[1]["ok"] is False  # never a silent pass


def test_a_table_of_rows_without_a_value_for_the_figure_is_not_a_figure(plane, monkeypatch):
    from sh_platform.v1 import b9_verify as V
    with wh.connect(plane) as con:
        con.execute("create table gold_p_position_snapshots_monthly_by_portfolio (month date, net_worth decimal(18,2))")
        con.execute("insert into gold_p_position_snapshots_monthly_by_portfolio values ('2026-09-01', null), ('2026-10-01', null)")
    gold = [{"name": "gold_p_position_snapshots_monthly_by_portfolio", "purpose": "Net worth", "depends_on": ["position_snapshots"],
             "columns": [{"name": "month", "type": "DATE"}, {"name": "net_worth", "type": "DECIMAL(18,2)"}], "kpis": []}]
    x, _ = _x(plane, gold=gold)
    x.api = lambda who, method, path, body=None, **kw: {"pack": {"target": {"gold": gold}}}
    monkeypatch.setattr(V, "_copy", lambda plane: wh.connect(plane, read_only=True))
    monkeypatch.setattr(F, "kpis_of", lambda plane, iid: [NW])
    F.check(x)
    assert x.th["figures"][0]["ok"] is False and "2 rows but none holds a value for Net worth" in x.th["figures"][0]["reason"]


def test_install_puts_the_figure_step_after_the_run_and_the_lines_into_decision_three():
    names = [s[2].__name__ for s in H.CHAINS["B"][0]]
    assert names.index("b_figures") == names.index("b_run") + 1
    assert H.SKILLS_OF["b_figures"] == (F.ADDS_UP,)
    assert H.DECISIONS["decision3"].__name__ == "decision3"


# ---- loading, honestly ---------------------------------------------------------------------------------------------------------------------

def _loader(answers):
    seen: list = []

    def api(who, method, path, body=None, **kw):
        seen.append((who, path))
        return answers.pop(0) if path.endswith("/load") else {"approved": ["Atlas Harbour Shipping"]} if "approve-exact" in path else {}
    return SimpleNamespace(api=api), seen


def test_a_file_the_load_held_back_or_rejected_is_not_called_loaded():
    x, _ = _loader([{"state": "rejected_totals", "last_load": {"status": "rejected_totals", "message": "Row count: the file says 29 but the rows add up to 31. Nothing was loaded."}}])
    assert F.load_file(x, "a.xlsx", "f1") == (False, "Row count: the file says 29 but the rows add up to 31. Nothing was loaded.")
    x, _ = _loader([{"state": "nothing_loaded", "last_load": {"status": "nothing_loaded", "message": "Nothing was loaded.", "held_back": [{"row": 2, "reason": "side 'sell usd / buy sgd' is not buy or sell"}]}}])
    ok, why = F.load_file(x, "t.csv", "f2")
    assert not ok and "Row 2: side 'sell usd / buy sgd' is not buy or sell." in why


def test_fund_names_that_are_exactly_the_instruments_own_are_approved_and_the_file_is_loaded_again():
    held = {"state": "nothing_loaded", "last_load": {"status": "nothing_loaded", "needs_approval": {"funds": [{"name": "Atlas Harbour Shipping"}]}}}
    x, seen = _loader([held, {"state": "loaded", "last_load": {"status": "loaded", "message": "Loaded 29 of 29 rows."}}])
    assert F.load_file(x, "admin.xlsx", "f1") == (True, "Loaded 29 of 29 rows.")
    assert [p.rsplit("/", 1)[-1] for _, p in seen] == ["load", "approve-exact", "load"]


def test_a_second_layout_proposes_a_mapping_that_is_approved_as_the_head_of_data_and_the_file_loaded_again():
    waiting = {"state": "waiting_for_mapping", "last_load": {"status": "waiting_for_mapping", "proposal_id": "map-9"}}
    x, seen = _loader([waiting, {"state": "loaded", "last_load": {"status": "loaded", "message": "ok"}}])
    assert F.load_file(x, "admin.xlsx", "f1")[0]
    assert ("anita", "/api/v1/model/mappings/map-9/approve") in seen


def test_the_table_ends_at_its_first_blank_row_so_notes_below_it_are_not_rows(tmp_path):
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.append(["Statement"]), ws.append(["Portfolio", "Market Value"]), ws.append(["EG-A", 100]), ws.append(["EG-B", 200]), ws.append([None])
    ws.append(["Rows: 2. Generated by the administrator."]), ws.append(["Synthetic data."])
    ws.append([None]), ws.append(["Total", 300])  # a totals row after a blank still counts as the totals row
    p = tmp_path / "s.xlsx"
    wb.save(p)
    cols = {"Portfolio": {"table": "portfolios", "field": "portfolio_code", "transform": None}, "Market Value": {"table": "position_snapshots", "field": "source_market_value", "transform": None}}
    _, _, rows, totals = e0c_load._read(p, "Sheet", 2, cols)
    assert [r[1]["Portfolio"] for r in rows] == ["EG-A", "EG-B"] and totals is None or totals["Market Value"] == 300


from tests.v1.test_e0c2 import ENG, env, one  # noqa: E402,F401 - the real landing zone, mapping records, instrument matching and writer on a tmp company


def test_a_row_with_no_date_is_named_and_held_back_not_the_whole_file(env):
    from openpyxl import load_workbook
    env.know(*[n for n in __import__("tests.v1.test_e0c2", fromlist=["FUNDS"]).FUNDS])
    import io
    wb = load_workbook(__import__("tests.v1.test_e0c2", fromlist=["FIX"]).FIX / "pos_2026_09_thousands.xlsx")
    ws = wb.active
    ws.cell(6, 1).value = None  # "value not received": no date on the second position
    ws.append([None])
    ws.append(["Rows: 6. Page 1 of 1. Synthetic data."])  # a footer below the table
    buf = io.BytesIO()
    wb.save(buf)
    fid = env.land("pos_footer", buf.getvalue(), "pos_footer.xlsx")
    env.approve(fid)
    r = e0c_load.load(env.plane, fid, ENG)
    assert r["status"] == "loaded" and r["rows_loaded"] == 5
    assert [h["reason"] for h in r["held_back"]] == ["no position date in the row"]  # the footer is not a row, the undated row is named
    assert one(env.plane, "select count(*) from position_snapshots") == 5


def test_a_table_that_holds_the_loaded_data_outranks_an_equal_fit(plane):
    words = D._toks("Net worth Sum of positions value plus cash")
    plain_rank = [r["table"] for r in D.rank_tables(plane, words)]
    with D.prefer({"position_snapshots"}):
        top = D.rank_tables(plane, words)[0]["table"]
    assert top == "position_snapshots" and plain_rank  # the preference changes the order only inside the block
    assert [r["table"] for r in D.rank_tables(plane, words)] == plain_rank


def test_two_sources_that_fit_a_file_name_equally_are_told_apart_by_what_the_file_holds(tmp_path, monkeypatch):
    from openpyxl import Workbook
    from sh_platform.v1 import i4_sow
    wb = Workbook()
    ws = wb.active
    ws.title = "Positions"
    ws.append(["Month-end positions and valuations"]), ws.append(["Portfolio", "Instrument", "Quantity", "Market Value"])
    (tmp_path / "dropbox").mkdir()
    wb.save(tmp_path / "dropbox" / "administrator_export_2026-09.xlsx")
    plane = SimpleNamespace(path=lambda *p: tmp_path.joinpath(*p))
    monkeypatch.setattr(i4_sow, "read", lambda plane, iid: {"sources": [{"name": "Administrator positions file", "holds": "positions and cash balances"},
                                                                         {"name": "Trading export", "holds": "trade details for reconciliation"}]})
    got = F.pick_source(plane, "administrator_export_2026-09.xlsx", ["Trading export", "Administrator positions file"], {"iid": "ini-1"})
    assert got == "Administrator positions file"
    assert F.pick_source(plane, "administrator_export_2026-09.xlsx", ["Trading export", "Trading export two"], {"iid": "ini-1"}) is None  # no clear winner: the person is asked


# ---- a total repeated under every entity, a change table, an account no portfolio has -------------------------------------------------------

def test_a_join_on_the_organisation_alone_is_refused_and_a_join_through_a_key_is_not():
    from sh_platform.v1 import b6_span as S
    bad = "create table g as select le.id, sum(cm.amount) from cash_movements cm join legal_entities le on cm.organization_id = le.organization_id group by le.id"
    good = "select le.id, sum(cm.amount) from cash_movements cm join portfolios p on p.id = cm.portfolio_id join legal_entities le on le.id = p.legal_entity_id group by le.id"
    assert "organization_id alone" in S.org_only_join(bad) and S.org_only_join(good) is None


def test_the_same_total_under_every_entity_is_not_a_split_and_the_answer_says_what_did_not_load(plane):
    with wh.connect(plane) as con:
        con.execute("create table gold_p_cash (entity_id uuid, entity_name varchar, day date, cash_variance decimal(18,2))")
        for e, n in (("43709bc3-b8f0-485f-a204-4130bb032b7d", "Entity B"), ("f606dc49-4b26-4bcf-bd44-7f5a5f617076", "Evergreen Holdings")):
            con.executemany("insert into gold_p_cash values (?, ?, ?, ?)", [(e, n, "2026-10-07", 10000), (e, n, "2026-10-08", -33200)])
    t = {"name": "gold_p_cash", "rows": 4, "columns": [{"name": "entity_id", "type": "UUID"}, {"name": "entity_name", "type": "VARCHAR"}, {"name": "day", "type": "DATE"},
                                                     {"name": "cash_variance", "type": "DECIMAL(18,2)"}]}
    th = {"exceptions": [{"title": "bank.xlsx did not load", "text": "TCB-1: portfolio not approved yet"}]}
    got = F._read(plane, t, {"name": "Cash variance"}, "Which entity has a cash variance on 8 October?", {}, th)
    assert got["figure"] is None and "same figure for every entity" in got["answer"] and "bank.xlsx: TCB-1" in got["answer"] and "33,200" not in got["answer"]
    with wh.connect(plane) as con:  # the same table with figures that differ by entity is read as before, one label per entity (the id and its name are one)
        con.execute("update gold_p_cash set cash_variance = 12400 where entity_name = 'Entity B' and day = '2026-10-08'")
    ok = F._read(plane, t, {"name": "Cash variance"}, "Which entity has a cash variance on 8 October?", {}, th)
    assert ok["rows"] == [{"label": "Entity B", "value": "12,400.00"}, {"label": "Evergreen Holdings", "value": "-33,200.00"}] or {r["label"] for r in ok["rows"]} == {"Entity B", "Evergreen Holdings"}


def _change_table(plane, rows):
    with wh.connect(plane) as con:
        con.execute("create table gold_p_kpi_changes (portfolio varchar, reporting_month date, original_value decimal(28,2), revised_value decimal(28,2), delta decimal(28,2))")
        rows and con.executemany("insert into gold_p_kpi_changes values (?, ?, ?, ?, ?)", rows)
    cols = [{"name": "portfolio", "type": "VARCHAR"}, {"name": "reporting_month", "type": "DATE"}, {"name": "original_value", "type": "DECIMAL(28,2)"},
            {"name": "revised_value", "type": "DECIMAL(28,2)"}, {"name": "delta", "type": "DECIMAL(28,2)"}]
    return {"gold_p_kpi_changes": {"name": "gold_p_kpi_changes", "rows": len(rows), "columns": cols}}


def test_why_did_it_move_reads_original_revised_and_delta_from_the_table_for_the_fund_named(plane):
    live = _change_table(plane, [("Fund A", "2026-09-01", 10000000, 10180000, 180000), ("Fund B", "2026-09-01", 5000000, 5000000, 0)])
    gold = [{"name": "gold_p_kpi_changes"}]
    assert F.CHANGE_Q.search("Why did Fund A move since the issued September report?") and not F.CHANGE_Q.search("What was the net worth in September?")
    got = F._change_answer(plane, "Why did Fund A move since the issued September report?", "prj", {}, gold, live)
    assert "10,000,000.00" in got["answer"] and "10,180,000.00" in got["answer"] and "180,000.00" in got["answer"] and "Fund B" not in got["answer"]
    assert F._change_answer(plane, "x", "prj", {}, [{"name": "gold_p_cash"}], {}) is None  # no such table: the question is left to Lumen


def test_an_empty_change_table_falls_back_to_the_two_valuations_on_record_and_never_invents_the_order(plane):
    live = _change_table(plane, [])
    with wh.connect(plane) as con:
        org = e1_write.org_id(plane)
        inst = wh.new_id()
        con.execute("insert into instruments (id, organization_id, canonical_name, instrument_type, created_at) values (?, ?, 'Evergreen Growth I', 'fund', now())", [inst, org])
        for v in (10000000, 10180000):
            con.execute("insert into valuation_observations (id, organization_id, instrument_id, valuation_date, received_at, recorded_at, value_amount, currency, valuation_type, approval_status) "
                        "values (?, ?, ?, '2026-09-30', now(), now(), ?, 'USD', 'manager_estimate', 'unapproved')", [wh.new_id(), org, inst, v])
    got = F._change_answer(plane, "Why did Evergreen Growth I move since the issued September report?", "prj", {}, [{"name": "gold_p_kpi_changes"}], live)
    assert "10,000,000 and 10,180,000" in got["answer"] and "differ by 180,000" in got["answer"] and "do not say which came later" in got["answer"]
    assert F._change_answer(plane, "Why did it move in March?", "prj", {}, [{"name": "gold_p_kpi_changes"}], live)["answer"].startswith("Not available yet")


def test_an_account_number_no_portfolio_has_is_asked_as_one_plain_question_and_the_answer_ties_it_to_the_entity(plane):
    from openpyxl import Workbook
    org = e1_write.org_id(plane)
    with wh.connect(plane) as con:
        ent = wh.new_id()
        con.execute("insert into legal_entities (id, organization_id, entity_code, entity_name, base_currency) values (?, ?, 'EG-B', 'Entity B', 'USD')", [ent, org])
        con.execute("insert into portfolios (id, organization_id, legal_entity_id, portfolio_code, portfolio_name, base_currency) values (?, ?, ?, 'EG-B-OPS', 'Entity B operating', 'USD')", [wh.new_id(), org, ent])
    wb = Workbook()
    ws = wb.active
    ws.append(["Value date", "Account number", "Currency", "Net movement"])
    ws.append(["2026-09-10", "TCB-7730-210958", "USD", 6000]), ws.append(["2026-09-11", "EG-B-OPS", "USD", 5])
    plane.path("dropbox", "bank.xlsx").parent.mkdir(parents=True, exist_ok=True)
    wb.save(plane.path("dropbox", "bank.xlsx"))
    x = SimpleNamespace(plane=plane, th={"maps": {"m1": {"name": "bank.xlsx"}}}, api=lambda *a, **k: {})
    assert F.account_questions(x) == [{"id": "acct:TCB-7730-210958", "text": "Which entity owns account TCB-7730-210958?", "options": ["Entity B"], "assumed": None}]  # the known code is not asked
    assert "tied to Entity B" in F.tie_account(x, "TCB-7730-210958", "Entity B")
    assert F.account_questions(x) == []  # the account is a portfolio of the entity now, so the file can load
    with wh.connect(plane, read_only=True) as con:
        assert con.execute("select e.entity_name from portfolios p join legal_entities e on e.id = p.legal_entity_id where p.portfolio_code = 'TCB-7730-210958'").fetchone() == ("Entity B",)
