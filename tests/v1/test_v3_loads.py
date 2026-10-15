"""Harbor 3.0 loads (c1 fixes): an FX fixing file loads as market observations of type fx; a report's summary workbook and a report PDF are kept as
labelled figures (the document's extracted fields), never as rows that need a fund; a file with no home is set aside in one line and the rest go on;
SQL calling a function the warehouse lacks is refused by the plan check; a failed run on the copy names its failing step. Synthetic data only."""
import io
import json
import types
from datetime import datetime, timezone

import pytest
from openpyxl import Workbook

from sh_platform.plane import PlaneStore
from sh_platform.profile.view import Viewer
from sh_platform.v1 import b6_plan, e0b_mappings, e0b_tools, e0c_landing, e0c_load, v3_loads as L, v3_pdf as V, wh
from sh_platform.v1.b11_builder import run_trouble
from sh_platform.v1.e1_write import org_id
from sh_platform.v1.i2_north import Refused

ENG = Viewer("Leela Raman", frozenset({"analyst"}), False, "acc-leela")
FX = "Date,Pair,Fixing,Rate (SGD per 1 USD),Source\n2026-09-29,USD/SGD,WMR 4pm,1.3405,Synthetic\n2026-09-30,USD/SGD,WMR 4pm,1.3412,Synthetic\n"
SUMMARY = [["Example Office"], ["Monthly Investment Report - 30 September 2026"], [],
           ["Report code", "report_2026-09"], ["Period end", "2026-09-30"], ["Prepared by", "A. Analyst"], ["USD/SGD fixing (WMR 4pm)", 1.3412],
           ["Version", "As issued"], [], ["Line", "Value (USD)", "Status", "Basis"], ["Growth Fund I", 10000000, "estimated", "Estimate"],
           ["Liquid book", 41200000, "official", "Administrator"], ["Total", 51200000, "estimated", "Sum"], [], ["Total (SGD)", 68669440, None, "Total x 1.3412"]]


def book(rows) -> bytes:
    wb = Workbook()
    for r in rows:
        wb.active.append(r)
    wb.create_sheet("Positions").append(["Security", "Name", "Quantity", "Price", "Value (USD)"])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


@pytest.fixture()
def plane(tmp_path):
    wh.reset_cache()
    return PlaneStore(tmp_path / "data").create("acme")


# ---- FX: the model's market observations ----------------------------------------------------------------------------------------

def test_fx_file_is_read_as_market_data_not_cash():
    prof = {"columns": ["Date", "Pair", "Fixing", "Rate (SGD per 1 USD)", "Source"], "sheet": "fx_usdsgd"}
    assert e0b_tools.guess_purpose(prof) == "market_data"
    got = e0b_tools._rules({**prof, "purpose": "market_data", "column_info": [{"name": c} for c in prof["columns"]]})
    assert got["Pair"] == "market_observations.series_key" and got["Rate (SGD per 1 USD)"] == "market_observations.value"
    assert got["Date"] == "market_observations.observed_at" and got["Fixing"] == e0b_tools.IGNORE


def test_fx_file_loads_as_fx_observations_with_both_currencies(plane):
    fid = e0c_landing.receive(plane, FX.encode(), "fx_usdsgd_wmr4pm.csv", "ops@example.test")["file_id"]
    prof = e0b_tools.profile_sheet(e0c_landing.original_path(plane, e0c_landing.get(plane, fid)))
    e0b_mappings.approve(plane, e0b_mappings.save_proposal(plane, {**e0b_tools.propose_mapping(plane, prof)["mapping"], "file_name": "fx_usdsgd_wmr4pm.csv"}), ENG)
    r = e0c_load.load(plane, fid, ENG)
    assert r["status"] == "loaded" and r["rows_loaded"] == 2 and r["observation_table"] == "market_observations"
    with wh.connect(plane) as con:
        rows = con.execute("select series_key, observation_type, base_currency, quote_currency, value from market_observations order by observed_at").fetchall()
    assert [tuple(x[:4]) for x in rows] == [("USD/SGD", "fx", "USD", "SGD")] * 2 and float(rows[-1][4]) == 1.3412


# ---- labelled figures: a report's summary workbook and a report PDF ---------------------------------------------------------------

def test_summary_workbook_is_read_as_labelled_figures_with_currency_and_row():
    figs = {f["label"]: f for f in L.labelled("report.xlsx", book(SUMMARY))}
    assert figs["Period end"]["value_text"] == "2026-09-30" and figs["Period end"]["value_numeric"] is None
    assert figs["Total"]["value_numeric"] == 51200000 and figs["Total"]["currency"] == "USD" and figs["Total"]["line"] == 13
    assert figs["Total (SGD)"]["currency"] == "SGD" and figs["USD/SGD fixing (WMR 4pm)"]["currency"] is None
    assert "Security" not in figs  # the positions sheet is a table of rows, not figures


def test_tables_of_rows_are_not_labelled_figures():
    assert L.labelled("fx.csv", FX.encode()) is None
    two_cols = [["Fund", "Amount"], *[[f"Fund {i}", 1000 * i] for i in range(1, 8)]]  # a list of funds and amounts stays a table
    assert L.sheet_figures(two_cols) is None


def test_a_three_column_pdf_table_pairs_each_amount_with_its_line():
    run = [(1, "Line"), (2, "Value (USD)"), (3, "Status"), (4, "Growth Fund I"), (5, "10,000,000.00"), (6, "estimated"),
           (7, "Liquid book"), (8, "41,200,000.00"), (9, "official"), (10, "Total (USD)"), (11, "51,200,000.00")]
    got = [(lb[1], v[1]) for lb, v in V._block(run)]
    assert got == [("Growth Fund I", "10,000,000.00"), ("Liquid book", "41,200,000.00"), ("Total (USD)", "51,200,000.00")]
    alternating = [(1, "Fund"), (2, "Growth Fund I"), (3, "Units"), (4, "1,000"), (5, "Value"), (6, "USD 5,000")]
    assert [(lb[1], v[1]) for lb, v in V._block(alternating)] == [("Fund", "Growth Fund I"), ("Units", "1,000"), ("Value", "USD 5,000")]  # parity still wins


def test_figures_are_written_once_as_confirmed_extracted_fields(plane):
    figs = L.labelled("report.xlsx", book(SUMMARY))
    did = wh.new_id()
    assert L.write(plane, did, figs) == len(figs) and L.write(plane, did, figs) == len(figs)
    with wh.connect(plane) as con:
        n, total, ccy, st = con.execute("select count(*), max(value_numeric) filter (where field_name = 'Total'), max(currency) filter (where field_name = 'Total'),"
                                        " min(review_status) from extracted_fields where document_id = ?", [did]).fetchone()
    assert n == len(figs) and total == 51200000 and ccy == "USD" and st == "confirmed" and org_id(plane)


# ---- Harbor: the loads hook, set aside in one line, never "none would load" while one file loads ------------------------------------

class X:
    def __init__(self, plane, files):
        self.plane, self.files, self.calls = plane, files, []
        self.th = {"id": "t-1", "labelled": {}, "pdf": {"files": {}}}

    def api(self, who, method, path, body=None, **k):
        self.calls.append(path)
        return {"id": wh.new_id()} if path == "/api/v1/evidence/documents" else {}


def test_load_keeps_good_files_and_sets_aside_one_with_a_plain_line(plane, monkeypatch):
    from sh_platform.v1 import v3_harbor as H
    files = {"report.xlsx": book(SUMMARY), "notes.xlsx": book([["Just a title"]])}
    monkeypatch.setattr(H, "blob", lambda p, n: files[n])
    monkeypatch.setattr(H, "my_files", lambda x, kind=None: [{"name": n, "source": "Issued reports"} for n in files])
    monkeypatch.setattr(H, "set_status", lambda p, n, s: x.status.__setitem__(n, s))
    monkeypatch.setattr(H, "save", lambda p, th: None)
    x = X(plane, files)
    x.status = {}
    L.intake(x)
    assert list(x.th["labelled"]) == ["report.xlsx"] and L.skip(x, "report.xlsx") and not L.skip(x, "notes.xlsx")
    assert L.d2_files(x)[0]["goes_to"] == "extracted fields of the document"
    x.th["labelled"]["broken.xlsx"] = {"doc_id": None, "n": 1}
    files["broken.xlsx"] = files["report.xlsx"]
    n, aside = L.load(x)
    assert n == 1 and x.status == {"report.xlsx": "loaded", "broken.xlsx": "question"}
    assert aside == [{"file": "broken.xlsx", "title": "broken.xlsx did not load", "text": "I set it aside: it could not be kept as a document."}]


# ---- the run on the copy: a function the warehouse lacks, and the failing step named -----------------------------------------------

def test_plan_check_refuses_a_function_the_warehouse_does_not_have():
    t = {"name": "gold_x", "layer": "gold", "depends_on": ["documents"],
         "sql": "create table gold_x as select uuid_generate_v4() as exception_id, d.id from documents d"}
    with pytest.raises(Refused) as e:
        b6_plan._check_sql(t, {"documents": {"id": "uuid"}}, {"documents"})
    assert "uuid_generate_v4" in e.value.message and "DuckDB" in e.value.message
    ok = dict(t, sql="create table gold_x as select gen_random_uuid() as exception_id, d.id from documents d")
    assert [c["name"] for c in b6_plan._check_sql(ok, {"documents": {"id": "uuid"}}, {"documents"})[0]] == ["exception_id", "id"]


def test_a_failed_run_names_its_step_in_plain_words():
    cur = {"run": {"status": "failed", "error": "gold_x_exceptions could not be built: Catalog Error: Scalar Function with name uuid_generate_v4 does not exist!",
                   "steps": [{"label": "Rows as they arrive.", "status": "filled by file loads"}, {"label": "Lists each late statement.", "status": "failed"}]}}
    assert run_trouble(cur) == ("The run on the copy stopped at the step that builds gold_x_exceptions (Lists each late statement): it calls uuid_generate_v4, "
                                "which the warehouse does not have. Send the design back to the architect to correct that step.")
    assert run_trouble({"run": {"status": "ok"}}) is None and run_trouble(None) is None


# ---- c2 bank statements: an opening line holds back by name, Debit/Credit columns become one signed amount ---------------------------

def _bank_plane(plane, rows, header):
    org = org_id(plane)
    with wh.connect(plane) as con:
        ent = wh.new_id()
        con.execute("insert into legal_entities (id, organization_id, entity_code, entity_name, base_currency) values (?, ?, 'EG-B', 'Entity B', 'USD')", [ent, org])
        con.execute("insert into portfolios (id, organization_id, legal_entity_id, portfolio_code, portfolio_name, base_currency) values (?, ?, ?, 'TCB-1', 'Entity B operating', 'USD')", [wh.new_id(), org, ent])
    csv = ",".join(header) + "\n" + "\n".join(",".join(map(str, r)) for r in rows) + "\n"
    fid = e0c_landing.receive(plane, csv.encode(), "bank.csv", "ops@example.test")["file_id"]
    prof = e0b_tools.profile_sheet(e0c_landing.original_path(plane, e0c_landing.get(plane, fid)))
    e0b_mappings.approve(plane, e0b_mappings.save_proposal(plane, {**e0b_tools.propose_mapping(plane, prof)["mapping"], "file_name": "bank.csv"}), ENG)
    r = e0c_load.load(plane, fid, ENG)
    return r


def test_an_opening_balance_line_is_held_back_by_name_not_a_constraint_failure(plane):
    r = _bank_plane(plane, [["2026-09-09", "TCB-1", "USD", ""], ["2026-09-10", "TCB-1", "USD", 6000]], ["Value date", "Account", "Currency", "Net movement"])
    assert r["status"] == "loaded" and r["rows_loaded"] == 1
    assert "no amount in the row" in r["held_back"][0]["reason"]


def test_debit_and_credit_columns_load_as_one_signed_amount(plane):
    r = _bank_plane(plane, [["2026-09-10", "TCB-1", "USD", 0, 250.5], ["2026-09-11", "TCB-1", "USD", 40, 0]], ["Value date", "Account", "Currency", "Debit", "Credit"])
    assert r["status"] == "loaded" and r["rows_loaded"] == 2
    with wh.connect(plane) as con:
        assert [float(a) for (a,) in con.execute("select amount from cash_movements order by effective_date").fetchall()] == [250.5, -40.0]


# ---- c3 bank statement workbook: the movements come from the transactions sheet, balances are kept out ------------------------------

def _statement(path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Daily balances"
    for r in [["Bank (synthetic)"], ["Account number: TCB-1"], [], ["Value date", "Account number", "Currency", "Net movement", "Closing ledger balance"],
              ["2026-10-07", "TCB-1", "USD", 0, 879200], ["2026-10-08", "TCB-1", "USD", -11200, 868000]]:
        ws.append(r)
    ws = wb.create_sheet("Transactions")
    for r in [["Value date", "Reference", "Type", "Amount", "Currency"],  # no heading: the account is named on the balances sheet beside it
              ["2026-10-08", "FEE-1", "fee", -1600, "USD"], ["2026-10-08", "WIRE-1", "wire", -9600, "USD"]]:
        ws.append(r)
    wb.save(path)


def test_statement_workbook_loads_its_transactions_not_its_balances(plane, tmp_path):
    org = org_id(plane)
    with wh.connect(plane) as con:
        ent = wh.new_id()
        con.execute("insert into legal_entities (id, organization_id, entity_code, entity_name, base_currency) values (?, ?, 'EG-B', 'Entity B', 'USD')", [ent, org])
        con.execute("insert into portfolios (id, organization_id, legal_entity_id, portfolio_code, portfolio_name, base_currency) values (?, ?, ?, 'TCB-1', 'Acct', 'USD')", [wh.new_id(), org, ent])
    _statement(tmp_path / "bank_statement.xlsx")
    fid = e0c_landing.receive(plane, (tmp_path / "bank_statement.xlsx").read_bytes(), "bank_statement.xlsx", "ops@example.test")["file_id"]
    path = e0c_landing.original_path(plane, e0c_landing.get(plane, fid))
    assert e0c_load._movement_sheet(plane, path, None) == ("Transactions", None)
    sheet, why = e0c_load._movement_sheet(plane, path, "Daily balances")
    assert "no table for balances" in why and "Transactions" in why
    prof = e0b_tools.profile_sheet(path, "Transactions")
    e0b_mappings.approve(plane, e0b_mappings.save_proposal(plane, {**e0b_tools.propose_mapping(plane, prof)["mapping"], "file_name": "bank_statement.xlsx"}), ENG)
    r = e0c_load.load(plane, fid, ENG)  # no sheet named: the transactions are loaded, on the account the sheet's heading names
    assert r["status"] == "loaded" and r["rows_loaded"] == 2, r
    assert e0c_load.load(plane, fid, ENG, sheet="Daily balances")["status"] == "nothing_loaded"
    with wh.connect(plane) as con:
        assert [float(a) for (a,) in con.execute("select amount from cash_movements order by amount").fetchall()] == [-9600.0, -1600.0]


# ---- c3 risk_officer: a table that a proposed model change adds, and nothing fills, is not "loaded" ---------------------------------

def test_a_gold_table_on_a_proposed_silver_table_that_nothing_fills_is_refused(plane):
    from sh_platform.v1 import b6_span as S
    sql = "create table gold_m as select portfolio_id, concentration_percent from silver_conc"
    assert S.reads_nothing_loaded(plane, {"tables": []}, sql).startswith("it reads only silver_conc")  # not in the warehouse yet, no plan table builds it
    with wh.connect(plane) as con:
        con.execute("create table silver_conc (portfolio_id uuid, concentration_percent decimal(18,2))")
    assert S.reads_nothing_loaded(plane, {"tables": []}, sql)  # in the warehouse, no rows
    plan = {"tables": [{"name": "silver_conc", "layer": "silver", "sql": "insert into silver_conc select portfolio_id, 1 from silver_other"}]}
    assert S.reads_nothing_loaded(plane, plan, sql)  # the plan's silver reads another empty table: the chain is empty too
    plan["tables"][0]["sql"] = "insert into silver_conc select portfolio_id, 1 from position_snapshots"
    assert S.reads_nothing_loaded(plane, plan, sql) is None  # built from a table a load fills


def test_a_break_between_two_sources_is_not_placed_by_the_rules_as_a_plain_total():
    from sh_platform.v1 import b10_design as D
    assert D._COMPARE.search("Cash break amount") and D._COMPARE.search("Administrator against the bank") and not D._COMPARE.search("Net worth by portfolio")


def test_the_design_ask_always_lets_a_hundred_through_for_a_percentage():
    from sh_platform.v1 import b10_design as D
    assert "100" in D._intent_numbers({"kpis": [{"name": "Concentration"}]}) and "5,000" in D._intent_numbers({"goal": "over 5,000"})
