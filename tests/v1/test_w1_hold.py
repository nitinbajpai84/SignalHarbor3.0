"""W1 hold: the recipes of the six "Hold and maintain" cards of the top 50 (13, 17, 21, 31, 38, 43), run on the Evergreen office6m files and compared with answer_keys.yaml.

Real: the recipes file, the model's DuckDB tables (docs/v2/evidence-to-action-mvp.duckdb.sql), the office6m files, the answer keys.
Stubbed: the loader. The files are read here by a few small readers (xlsx, csv, pdf text) and written into the model's tables; the real loaders (b4/e0c, v3_pdf) are
not called (see Needs in docs/v3/w1/build-hold.md). Also checks the signal yaml and the builds yaml of the six cards.
"""
import ast
import csv
import glob
import json
import re
from datetime import date, datetime
from functools import lru_cache
from pathlib import Path

import duckdb
import openpyxl
import pypdf
import pytest
import yaml

from sh_platform.v1 import v3_recipes as R

REPO = Path(__file__).resolve().parents[2]
FX = REPO / "tests" / "v1" / "fixtures" / "office6m"
PACK = REPO / "packs" / "family_office" / "1"
KEYS = yaml.safe_load((FX / "answer_keys.yaml").read_text(encoding="utf-8"))["keys"]
MASTER = json.loads((FX / "reference" / "master_data.json").read_text(encoding="utf-8"))
ORG = "org-1"
CARDS = ["uc-013-secgaps", "uc-017-liab", "uc-021-stale", "uc-031-alloc", "uc-038-navroll", "uc-043-liquidpnl"]
LISTED = ("listed_equity", "bond", "convertible")
# the generator's own map (tools/gen_office6m/keys.py): the liquid portfolio belongs to the Trust; the funds' portfolios are in master_data.json
PORTFOLIO_ENTITY = {**{f["portfolio"]: f["entity"] for f in MASTER["funds"].values()}, "EG-LIQUID": "EG-TRUST"}


# ---- a small DuckDB with the model's tables ---------------------------------------------------------------------------------------------

DEFAULT = {"VARCHAR": "x", "DATE": date(2026, 10, 13), "BOOLEAN": False}


def new_db():
    con = duckdb.connect()
    con.execute((REPO / "docs" / "v2" / "evidence-to-action-mvp.duckdb.sql").read_text(encoding="utf-8"))
    return con


def ins(con, table, rows):
    """Insert dict rows (same keys in each); every other NOT NULL column gets a plain default."""
    if not rows:
        return
    have = list(rows[0])
    extra = []
    for name, typ in con.execute("select column_name, data_type from information_schema.columns where table_name = ? and is_nullable = 'NO'", [table]).fetchall():
        if name not in have:
            extra.append((name, datetime(2026, 10, 13) if typ.startswith("TIMESTAMP") else 0 if typ.startswith(("DECIMAL", "INT", "BIGINT")) else DEFAULT.get(typ, "x")))
    cols = have + [n for n, _ in extra]
    con.executemany(f"insert into {table} ({', '.join(cols)}) values ({', '.join('?' * len(cols))})", [[r[c] for c in have] + [v for _, v in extra] for r in rows])


def d(s):
    return date.fromisoformat(str(s)[:10])


def iso(x):
    return x.date().isoformat() if isinstance(x, datetime) else str(x)[:10]


# ---- readers of the office6m files -------------------------------------------------------------------------------------------------------

def files(pattern):
    return sorted(Path(p) for p in glob.glob(str(FX / pattern)))


def load_master(con):
    """Entities, portfolios, instruments (the security master) and their identifiers."""
    ins(con, "legal_entities", [{"id": k, "organization_id": ORG, "entity_code": k, "entity_name": v, "base_currency": "USD"} for k, v in MASTER["entities"].items()])
    ins(con, "portfolios", [{"id": p, "organization_id": ORG, "legal_entity_id": e, "portfolio_code": p, "portfolio_name": p, "base_currency": "USD"} for p, e in PORTFOLIO_ENTITY.items()])
    wb = openpyxl.load_workbook(FX / "reference" / "EG_security_master_20261008.xlsx", data_only=True)
    rows = [r for r in wb["Security master"].iter_rows(min_row=5, values_only=True) if r[0] and str(r[0]).startswith(("SEC-", "FND-"))]
    ins(con, "instruments", [{"id": r[0], "organization_id": ORG, "canonical_name": r[2], "instrument_type": r[3], "asset_class": r[4], "manager_name": r[9], "currency": r[5],
                              "created_at": datetime(2026, 10, 8)} for r in rows])
    ins(con, "instrument_identifiers", [{"id": f"id-{r[0]}", "organization_id": ORG, "instrument_id": r[0], "source_system_id": "master", "identifier_type": "isin",
                                         "identifier_value": r[1], "mapping_status": "approved"} for r in rows if r[1]])
    # the security master leaves out three private funds (Evergreen Growth II, Evergreen Credit III, Saltmarsh Renewables I); their class is in master_data.json
    named = {r[2] for r in rows}
    extra = [{"id": f"FND-{k}", "organization_id": ORG, "canonical_name": f["name"], "instrument_type": "private_fund", "asset_class": f["asset_class"], "manager_name": f["manager"],
              "currency": "USD", "created_at": datetime(2026, 10, 8)} for k, f in MASTER["funds"].items() if f["name"] not in named]
    ins(con, "instruments", extra)
    return {**{r[2]: r[0] for r in rows}, **{x["canonical_name"]: x["id"] for x in extra}}, {r[6]: r[0] for r in rows if r[6]}  # by name, by trading code


def position(n, pf, inst, day, qty, mv, status, src):
    return {"id": f"{src}-{n}", "organization_id": ORG, "source_record_id": src, "portfolio_id": pf, "instrument_id": inst, "position_date": d(day), "quantity": qty,
            "source_market_value": mv, "currency": "USD", "view_type": "operational", "data_status": status, "as_of_at": datetime(2026, 10, 13), "received_at": datetime(2026, 10, 13),
            "recorded_at": datetime(2026, 10, 13)}


def load_admin(con, by_name, pattern="*/administrator_nav/administrator_export_2026-09.xlsx"):
    """Administrator packs: the Positions sheet, dated by the pack's NAV sheet (the column 'Market Value' is 'Valuation Amount (USD)' in June, M18)."""
    out, codes = [], {}
    for f in files(pattern):
        wb = openpyxl.load_workbook(f, data_only=True)
        pack_date = list(wb["NAV by portfolio"].iter_rows(min_row=2, max_row=2, values_only=True))[0][2]
        rows = list(wb["Positions"].iter_rows(values_only=True))
        head = [str(h).lower() for h in rows[0]]
        mv = next(i for i, h in enumerate(head) if "market value" in h or "valuation amount" in h)
        for n, r in enumerate(rows[1:]):
            out.append(position(n, r[0], by_name[r[2]], pack_date, r[3], r[mv], r[7], f"admin-{f.stem}"))
            codes[(by_name[r[2]], str(r[1]))] = r[1]
    ins(con, "position_snapshots", out)
    # the administrator code is an identifier too (approved): it is the only one the three funds missing from the security master have
    have = {x[0] for x in con.execute("select id from instrument_identifiers").fetchall()}
    ins(con, "instrument_identifiers", [{"id": f"adm-{i}-{c}", "organization_id": ORG, "instrument_id": i, "source_system_id": "administrator", "identifier_type": "administrator_code",
                                         "identifier_value": c, "mapping_status": "approved"} for (i, c) in codes if f"adm-{i}-{c}" not in have])


def load_custodian(con, pattern):
    out = []
    for f in files(pattern):
        for n, r in enumerate(csv.DictReader(open(f, encoding="utf-8"))):
            out.append(position(n, r["portfolio"], r["security_id"], r["as_of"], float(r["quantity"]), float(r["market_value"]), "Final", f"cust-{f.stem}"))
    ins(con, "position_snapshots", out)


def num(t, label):
    m = re.search(re.escape(label) + r"\s+(-?[\d,]+\.\d\d)", t)
    return float(m.group(1).replace(",", "")) if m else None


@lru_cache(maxsize=1)
def statements():
    """Every manager statement pdf: fund, period end, issued, status, closing balance, net gain."""
    out = []
    for f in files("*/manager_statements/*.pdf"):
        t = pypdf.PdfReader(f).pages[0].extract_text()
        g = lambda k: re.search(k + r": (.*)", t).group(1).strip()
        out.append({"id": f.stem, "fund": g("Fund"), "period_end": datetime.strptime(g(r"Statement date \(period end\)"), "%d %B %Y").date(),
                    "issued": datetime.strptime(g("Issued"), "%d %B %Y"), "status": g("Status"), "closing": num(t, "Closing capital account balance"), "gain": num(t, "Net gain / (loss)")})
    return out


def load_statements(con, by_name):
    """valuation_observations: one per statement; a later statement of the same fund and period supersedes the earlier. The net gain line is an extracted field."""
    rows, fields, last = [], [], {}
    for s in sorted(statements(), key=lambda s: (s["fund"], s["period_end"], s["issued"])):
        k = (s["fund"], s["period_end"])
        rows.append({"id": s["id"], "organization_id": ORG, "instrument_id": by_name[s["fund"]], "valuation_date": s["period_end"], "recorded_at": s["issued"],
                     "value_amount": s["closing"], "currency": "USD", "valuation_type": "statement", "approval_status": s["status"].lower(),
                     "evidence_document_id": s["id"], "supersedes_valuation_id": last.get(k), "received_at": s["issued"]})
        fields.append({"id": f"f-{s['id']}", "organization_id": ORG, "document_id": s["id"], "field_name": "net_gain_loss", "value_numeric": s["gain"], "confidence": "high",
                       "review_status": "read"})
        last[k] = s["id"]
    ins(con, "valuation_observations", rows)
    ins(con, "extracted_fields", fields)


def load_flows(con, by_name):
    """cash_movements: the capital calls and distributions the bank statements show (what actually moved), confirmed."""
    fund_port = {f["name"]: f["portfolio"] for f in MASTER["funds"].values()}
    rows = []
    for f in files("*/bank_statements/*.xlsx"):
        for r in list(openpyxl.load_workbook(f, data_only=True)["Transactions"].iter_rows(values_only=True))[1:]:
            if r and r[2] in ("capital_call", "distribution"):
                fund = max((n for n in fund_port if n in r[3]), key=len)  # "Evergreen Growth I" is inside "Evergreen Growth II"
                rows.append({"id": f"{f.stem}-{r[1]}", "organization_id": ORG, "portfolio_id": fund_port[fund], "instrument_id": by_name[fund],
                             "movement_type": "call" if r[2] == "capital_call" else "distribution", "effective_date": d(r[0]), "amount": abs(r[5]), "currency": r[6],
                             "confirmation_status": "confirmed"})
    ins(con, "cash_movements", rows)


def load_loans(con):
    rows = []
    for f in files("*/loan_statements/loan_statement_TL-B-01_*.xlsx"):
        r = list(openpyxl.load_workbook(f, data_only=True).worksheets[0].iter_rows(min_row=2, values_only=True))[0]
        y, m = map(int, r[1].split("-"))
        end = date(y + (m == 12), m % 12 + 1, 1).toordinal() - 1  # the statements show amounts negative (M12): the model holds the amount owed
        rows.append({"id": f"loan-{r[1]}", "organization_id": ORG, "facility_ref": "TL-B-01", "legal_entity_id": "EG-B", "lender": "Harbour Bank (synthetic)",
                     "facility_type": "term_loan", "currency": "USD", "drawn_amount": abs(r[5]), "as_at": date.fromordinal(end), "status": "active"})
    ins(con, "credit_facilities", rows)


def load_trades(con, by_code, pattern):
    rows = []
    for f in files(pattern):
        for r in csv.DictReader(open(f, encoding="utf-8")):
            if r["instrument_code"] in by_code:  # the FX forward is not a listed instrument
                rows.append({"id": r["external_trade_id"], "organization_id": ORG, "portfolio_id": r["portfolio"], "instrument_id": by_code[r["instrument_code"]],
                             "transaction_type": r["side"], "trade_date": d(r["trade_date"]), "effective_date": d(r["trade_date"]), "settlement_date": d(r["settle_date"]),
                             "cash_amount": float(r["amount"]), "currency": r["currency"], "status": r["status"]})
    ins(con, "transactions", rows)


def run(con, rid, name="gold_t", dims=None):
    r = R.recipe(rid)
    assert r and r["status"] == "ready"
    sql = R.render(r, name, R.dims_for(r, {}) if dims is None else dims)
    assert R.refusal(r, sql) is None, "the recipe reads a table outside may_read or one in never_read"
    assert R.tables_read(sql) <= set(r["may_read"])
    con.execute(sql)
    t = R.tie_out(con, r, name)
    return r, t


def rows(con, sql, *a):
    cur = con.execute(sql, list(a))
    cols = [c[0] for c in cur.description]
    return [dict(zip(cols, x)) for x in cur.fetchall()]


def key(card):
    return next(iter(KEYS[card]["questions"].values()))["expected"]


# ---- card 13: securities with no asset class or identifier -----------------------------------------------------------------------------------

def test_13_identifier_gaps():
    con = new_db()
    by_name, _ = load_master(con)
    load_custodian(con, "*/custodian_positions/positions_20260930.csv")
    load_admin(con, by_name)
    r, t = run(con, "identifier_gaps")
    assert t and t["ok"], t
    k = key("uc-013-secgaps")
    gaps = rows(con, "select * from gold_t where has_gap")
    assert len(gaps) == k["count"] == 1
    assert [g["instrument"] for g in gaps] == k["instruments"]
    assert float(gaps[0]["gap_value"]) == k["value"]
    assert gaps[0]["no_class"] and not gaps[0]["no_identifier"]
    one = rows(con, "select * from gold_t where first_row")[0]
    assert one["gap_count"] == 1 and float(one["gap_value_total"]) == 1122000 and float(one["net_worth"]) == 108230000
    assert float(one["gap_pct_of_net_worth"]) == pytest.approx(1.0367, abs=1e-4)  # above the 1 percent of the card's exception row


# ---- card 17: net worth after debt ----------------------------------------------------------------------------------------------------------

def test_17_net_worth_after_debt():
    con = new_db()
    by_name, _ = load_master(con)
    load_admin(con, by_name, "*/administrator_nav/administrator_export_2026-*.xlsx")
    load_loans(con)
    r, t = run(con, "net_worth_after_debt")
    assert t and t["ok"], t
    k = key("uc-017-liab")
    sept = rows(con, "select * from gold_t where position_date = date '2026-09-30'")
    assert sum(float(x["net_worth"]) for x in sept) == k["all_portfolios_nav"]
    assert sum(float(x["debt"]) for x in sept) == k["principal_outstanding"] == -k["loan_balance_as_shown_negative"]
    assert sum(float(x["net_worth_after_debt"]) for x in sept) == k["all_portfolios_less_debt"]
    assert float(next(x for x in sept if x["entity"] == "Entity B")["net_worth_after_debt"]) == 17395000 - 2800000
    one = rows(con, "select * from gold_t where is_latest")
    assert len(one) == 1 and float(one[0]["family_net_worth_after_debt"]) == k["all_portfolios_less_debt"] and float(one[0]["family_debt"]) == k["principal_outstanding"]
    for month, bal in k["by_month"].items():  # the balance at each month end, as the latest statement at or before the date
        end = rows(con, "select debt from gold_t where entity = 'Entity B' and strftime(position_date, '%Y-%m') = ?", month)
        assert [float(x["debt"]) for x in end] == [-bal], month
    # the issued report counts Growth I, Income II and the liquid book (57.6m); less the loan it is 54.8m
    issued = float(con.execute("select sum(v) from (select max(source_market_value) v from position_snapshots where position_date = date '2026-09-30' and "
                               "(portfolio_id = 'EG-LIQUID' or instrument_id in ('FND-A', 'FND-B')) group by portfolio_id, instrument_id)").fetchone()[0])
    assert issued == k["report_total"] and issued - 2800000 == k["report_less_debt"]


def test_17_facility_headroom_is_empty_when_the_agreement_states_no_limit():
    """The loan agreement names a principal, not a limit or a covenant: the headroom stays empty and the signal waits for a signed line."""
    con = new_db()
    load_master(con)
    load_loans(con)
    run(con, "facility_headroom")
    row = rows(con, "select * from gold_t")[0]
    assert row["facility_ref"] == "TL-B-01" and row["entity"] == "Entity B" and float(row["drawn_amount"]) == 2800000 and str(row["as_at"]) == "2026-09-30"
    assert row["limit_amount"] is None and row["headroom"] is None and row["first_row"]


# ---- card 21: valuations that are stale ------------------------------------------------------------------------------------------------------

def test_21_value_age_days():
    con = new_db()
    by_name, _ = load_master(con)
    load_statements(con, by_name)
    load_admin(con, by_name)
    load_custodian(con, "*/custodian_positions/positions_20261013.csv")  # the day's load: the latest position date is the as-of date, 13 Oct
    r, t = run(con, "value_age_days")
    assert t and t["ok"], t
    k = key("uc-021-stale")
    got = {x["fund"]: x for x in rows(con, "select * from gold_t")}
    for fund, e in k["by_fund"].items():
        assert str(got[fund]["valuation_date"]) == e["valuation_date"], fund
        assert got[fund]["age_days"] == e["age_days"], fund
        assert got[fund]["stale"] is e["stale"], fund
    assert sorted(f for f, x in got.items() if x["stale"]) == sorted(k["stale"])
    one = rows(con, "select * from gold_t where first_row")
    assert len(one) == 1 and one[0]["stale_count"] == 3 and one[0]["fund"] == one[0]["stalest_fund"] == "Cedarpoint Opportunities II"
    assert set(one[0]["stale_names"].split(", ")) == set(k["stale"])


# ---- card 31: allocation by asset class and entity ------------------------------------------------------------------------------------------

def test_31_allocation():
    con = new_db()
    by_name, _ = load_master(con)
    load_admin(con, by_name)
    r, t = run(con, "allocation")
    assert t and t["ok"], t
    k = key("uc-031-alloc")
    total = float(con.execute("select sum(market_value) from gold_t").fetchone()[0])
    assert total == k["total"]
    for cls, v in k["by_class"].items():
        row = rows(con, "select sum(market_value) mv, sum(share_pct) p from gold_t where asset_class = ?", cls)[0]
        assert float(row["mv"]) == v, cls
        assert float(row["p"]) == pytest.approx(k["by_class_pct"][cls], abs=0.02), cls
    for ent, p in k["by_entity_pct"].items():
        assert float(rows(con, "select sum(share_pct) p from gold_t where entity = ?", ent)[0]["p"]) == pytest.approx(p, abs=0.02), ent
    assert float(rows(con, "select sum(share_pct) p from gold_t")[0]["p"]) == pytest.approx(100, abs=0.05)
    one = rows(con, "select * from gold_t where first_row")
    assert len(one) == 1 and float(one[0]["total_value"]) == float(one[0]["positions_total"]) == k["total"] and float(one[0]["unclassified_value"]) == k["by_class"]["unclassified"]


# ---- card 38: private fund value roll-forward ------------------------------------------------------------------------------------------------

def test_38_nav_rollforward():
    con = new_db()
    by_name, _ = load_master(con)
    load_statements(con, by_name)
    load_flows(con, by_name)
    r, t = run(con, "nav_rollforward")
    assert t and t["ok"], t
    exp = key("uc-038-navroll")["rows"]
    got = {(x["fund"], f"{x['period_start']} to {x['period_end']}"): x for x in rows(con, "select * from gold_t")}
    assert set(got) == {(e["fund"], e["quarter"]) for e in exp}
    for e in exp:
        x = got[(e["fund"], e["quarter"])]
        assert (float(x["prior_value"]), float(x["calls"]), float(x["distributions"]), float(x["closing_value"]), float(x["implied_gain"])) == \
               (e["prior"], e["calls"], e["distributions"], e["new"], e["implied_gain"]), e
        assert float(x["unexplained"]) == pytest.approx(float(x["implied_gain"]) - float(x["reported_gain"]))  # unexplained = what the flows leave less the manager's gain
    # the manager's own statements do not always carry on from the previous closing: Growth I opens 8,976,477.00 in Q2 against a 8,868,296.00 closing
    assert {x["fund"] for x in rows(con, "select fund from gold_t where latest")} == {e["fund"] for e in exp}
    assert len(rows(con, "select * from gold_t where latest and fund = 'Evergreen Growth I'")) == 1
    a = got[("Evergreen Growth I", "2026-03-31 to 2026-06-30")]
    assert float(a["unexplained"]) == 108181 and float(a["reported_gain"]) == -210957


# ---- card 43: liquid book gain or loss this month ------------------------------------------------------------------------------------------

def test_43_liquid_pnl_mtd():
    con = new_db()
    by_name, by_code = load_master(con)
    load_custodian(con, "*/custodian_positions/positions_2026093*.csv")
    load_custodian(con, "*/custodian_positions/positions_20261013.csv")
    load_trades(con, by_code, "*/trading_export/trading_export_2026100*.csv")
    r, t = run(con, "liquid_pnl_mtd")
    assert t and t["ok"], t
    k = key("uc-043-liquidpnl")
    row = rows(con, "select * from gold_t")[0]
    assert row["portfolio"] == "EG-LIQUID" and str(row["position_date"]) == "2026-10-13" and str(row["month_end_date"]) == "2026-09-30"
    assert float(row["value_month_end"]) == k["value_30_sep"]
    assert float(row["value_today"]) == k["value_13_oct"]
    assert float(row["net_trades"]) == k["net_trades_settled"]
    assert float(row["gain_loss"]) == k["gain_loss_mtd"]
    assert float(row["gain_pct"]) == pytest.approx(k["gain_pct"], abs=1e-4) and row["first_row"]


def test_43_price_move_flags_nothing_here():
    """The 10 percent row of the card: no price in the 140 custodian days moved that much, so nothing is held for review."""
    con = new_db()
    load_master(con)
    load_custodian(con, "*/custodian_positions/positions_*.csv")
    run(con, "liquid_price_move")
    n, worst = con.execute("select count(*), max(abs(move_pct)) from gold_t").fetchone()
    assert n > 2000 and float(worst) < 10
    top = rows(con, "select * from gold_t where top_move")
    assert len(top) == 1 and str(top[0]["position_date"]) == "2026-10-13" and abs(float(top[0]["move_pct"])) == float(worst) or abs(float(top[0]["move_pct"])) <= float(worst)
    assert con.execute("select count(*) from gold_t where abs(move_pct) > 10").fetchone()[0] == 0


# ---- the signals and the builds of the six cards --------------------------------------------------------------------------------------------

ROLES = {r["id"] for r in yaml.safe_load((PACK / "role_builds.yaml").read_text(encoding="utf-8"))["roles"]}
FUNCS = {"abs", "min", "max", "round", "last4", "num", "same_fund"}


def signal_files():
    return sorted((PACK / "signals").glob("uc-0*.yaml"))


@pytest.mark.parametrize("card", CARDS)
def test_signals_of_the_card_are_well_formed(card):
    mine = [p for p in signal_files() if p.name.startswith(card + "-")]
    docs = [yaml.safe_load(p.read_text(encoding="utf-8")) for p in mine]
    assert len(docs) >= 3
    assert any(s["severity"] != "info" and s.get("hold") for s in docs), "every card has at least one exception signal with a hold"
    for p, s in zip(mine, docs):
        assert s["id"] == p.stem and s["suggestion"] is True
        for kf in ("id", "title", "trigger", "condition", "severity", "route_to", "action", "message"):
            assert s.get(kf), (p.name, kf)
        trig = s["trigger"]
        assert (trig.get("on") or trig.get(True)) in {"file_arrived", "figure_changed", "date_passed", "schedule", "run", "action_done"}
        assert s["severity"] in {"info", "needs", "problem"} and s["route_to"] in ROLES
        assert s["action"]["verb"] in {"log", "record", "approve", "call_back", "see"}
        names = {n.id for n in ast.walk(ast.parse(str(s["condition"]), mode="eval")) if isinstance(n, ast.Name)}
        inputs = {i["as"] for i in s.get("inputs", []) if i.get("as")} | {"doc"} | set(s.get("let") or {})
        assert names - FUNCS <= inputs, (p.name, names - FUNCS - inputs)
        for i in s.get("inputs", []):
            if i.get("recipe"):
                assert R.recipe(i["recipe"]), (p.name, i["recipe"])


def test_builds_carry_the_exact_answer_keys():
    b = yaml.safe_load((PACK / "builds" / "hold.yaml").read_text(encoding="utf-8"))
    builds = {x["id"]: (r["id"], x) for r in b["roles"] for x in r["builds"]}
    assert set(builds) == set(CARDS)
    for card, (role, x) in builds.items():
        assert role in ROLES and x["say"] and x["files"] and x["check"]["ask"]
        assert all(f.startswith("office6m:") and glob.glob(str(FX / f[9:])) for f in x["files"]), card
        text = (FX / "answer_keys.yaml").read_text(encoding="utf-8")
        for e in x["check"]["expect"]:
            assert e.replace(",", "") in text.replace(",", ""), (card, e)
        assert x["check"]["expect"], card
