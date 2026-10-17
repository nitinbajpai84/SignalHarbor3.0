"""Wave 1, domain office (cards 15, 30, 50 and the signals of 9): each recipe of recipes/office.yaml run over a small DuckDB made from the model's own DDL and the office6m
fixture files, and the numbers compared with answer_keys.yaml (never typed here). The model tables are loaded by small readers in this file, because the real loaders do not
read allowance schedules, vendor contracts or fee invoices yet (docs/v3/w1/build-office.md, Needs)."""
import ast
import csv
import glob
import re
from decimal import Decimal
from pathlib import Path

import duckdb
import openpyxl
import pypdf
import pytest
import yaml

from sh_platform.v1 import v3_recipes as R

REPO = Path(__file__).resolve().parents[2]
OFFICE = REPO / "tests" / "v1" / "fixtures" / "office6m"
PACK = REPO / "packs" / "family_office" / "1"
DDL = REPO / "docs" / "v2" / "evidence-to-action-mvp.duckdb.sql"
KEYS = yaml.safe_load((OFFICE / "answer_keys.yaml").read_text(encoding="utf-8"))["keys"]
ORG = "org-eg"


def _key(card: str) -> dict:
    return next(iter(KEYS[card]["questions"].values()))["expected"]


def _recipe(rid: str) -> dict:
    return next(r for r in yaml.safe_load((PACK / "recipes" / "office.yaml").read_text(encoding="utf-8"))["recipes"] if r["id"] == rid)


@pytest.fixture()
def con():
    c = duckdb.connect()
    c.execute(DDL.read_text(encoding="utf-8"))
    return c


def _build(con, rid: str) -> str:
    """Run the recipe's SQL on the copy, check it only reads its allowed tables, run its tie-out. -> the ready table's name."""
    r, name = _recipe(rid), "gold_" + rid
    assert r.get("status") != "stub" and r["sql"]
    sql = R.render(r, name)
    assert R.refusal(r, sql) is None, R.refusal(r, sql)
    con.execute(sql)
    t = R.tie_out(con, r, name)
    assert t and t["ok"], t
    return name


def _rows(con, name: str, order: str) -> list[dict]:
    cur = con.execute(f"select * from {name} order by {order}")
    cols = [d[0] for d in cur.description]
    return [dict(zip(cols, row)) for row in cur.fetchall()]


def _pdf(path) -> str:
    return "\n".join(p.extract_text() for p in pypdf.PdfReader(str(path)).pages)


# ---- the card 30 data: vendor contracts read from the four contract PDFs --------------------------------------------------------------

def _load_vendors(con, recorded: str = "2026-10-13 09:00:00") -> None:
    master = {r["vendor"].lower(): r for r in csv.DictReader((OFFICE / "reference" / "vendor_master.csv").open(encoding="utf-8"))}
    for i, f in enumerate(sorted((OFFICE / "reference").glob("vendor_contract_*.pdf"))):
        text = _pdf(f)
        lines = text.splitlines()
        name = master[lines[0].strip().lower()]["vendor"]
        cost = re.search(r"Annual cost: ([A-Z]{3}) ([\d,]+\.\d\d)", text)
        notice = re.search(r"(\d+) days' notice", text)
        con.execute("insert into vendor_contracts (id, organization_id, source_record_id, received_at, recorded_at, vendor_name, service, next_renewal_date, auto_renews, notice_days, "
                    "annual_cost, currency, status) values (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, 'Active')",
                    [f"vc{i}", ORG, f"sr{i}", recorded, recorded, name, lines[2].strip(), re.search(r"Renewal date: (\S+)", text).group(1), "renews" in text.lower(),
                     int(notice.group(1)) if notice else None, Decimal(cost.group(2).replace(",", "")), cost.group(1)])


def test_renewals_vendors_match_the_key(con):
    _load_vendors(con)
    got = _rows(con, _build(con, "renewals_vendors"), "next_renewal_date")
    key = _key("uc-030-vendors")
    want = sorted(key["renewing_within_90_days"], key=lambda x: x["renewal"])
    assert len(got) == key["count"] == 2
    for g, w in zip(got, want):
        assert (g["vendor_name"], g["currency"], float(g["annual_cost"]), str(g["next_renewal_date"]), g["days_to_renewal"]) == \
               (w["vendor"], w["currency"], w["annual_cost"], w["renewal"], w["days_to_renewal"])
    quill, orchard = got
    assert quill["days_to_notice_deadline"] == 48 and quill["notice_deadline"] is None and quill["auto_renews"] == 1  # no notice period stated: the renewal date is the clock
    assert orchard["notice_days"] == 60 and str(orchard["notice_deadline"]) == "2026-11-01" and orchard["days_to_notice_deadline"] == 19
    assert all(g["decision_open"] == 1 for g in got)


def test_renewals_vendors_status_and_supersede(con):
    _load_vendors(con)
    con.execute("update vendor_contracts set status = 'Renewal approved' where vendor_name = 'Orchard IT Services'")
    con.execute("update vendor_contracts set status = 'Ended' where vendor_name = 'Quill Research Subscriptions'")
    got = _rows(con, _build(con, "renewals_vendors"), "vendor_name")
    assert [(g["vendor_name"], g["decision_open"]) for g in got] == [("Orchard IT Services", 0)]


# ---- the card 15 data: allowance schedules and the Trust's bank statements ----------------------------------------------------------------

def _load_trust(con, months: list[str], statements: list[Path]) -> None:
    con.execute("insert into legal_entities (id, organization_id, entity_code, entity_name, base_currency) values ('le-trust', ?, 'EG-TRUST', 'Evergreen Family Trust', 'USD')", [ORG])
    con.execute("insert into portfolios (id, organization_id, legal_entity_id, portfolio_code, portfolio_name, base_currency) values ('pf-trust', ?, 'le-trust', 'EG-TRUST-CASH', "
                "'Evergreen Family Trust cash', 'USD')", [ORG])
    n = 0
    for m in months:  # one schedule file is one set of rows, valid for its month
        ws = openpyxl.load_workbook(OFFICE / m / "family_allowances" / f"allowance_schedule_{m}.xlsx").active
        for member, rel, amount, due, _paying in list(ws.iter_rows(values_only=True))[1:]:
            n += 1
            first = due[:8] + "01"
            con.execute("insert into family_members (id, organization_id, source_record_id, received_at, recorded_at, member_name, relationship, legal_entity_id, allowance_amount, "
                        "allowance_currency, allowance_frequency, allowance_from, allowance_to, paid_from_portfolio_id) values (?, ?, 'sr', ?, ?, ?, ?, 'le-trust', ?, 'USD', 'monthly', "
                        "cast(? as date), last_day(cast(? as date)), 'pf-trust')", [f"fm{n}", ORG, due, due, member, rel, amount, first, first])
    seen = set()
    for s in statements:
        for value_date, ref, kind, _d, _cp, amount, ccy in list(openpyxl.load_workbook(s)["Transactions"].iter_rows(values_only=True))[1:]:
            if kind == "allowance" and ref not in seen:
                seen.add(ref)
                con.execute("insert into cash_movements (id, organization_id, source_record_id, portfolio_id, movement_type, effective_date, received_at, recorded_at, amount, currency, "
                            "confirmation_status) values (?, ?, ?, 'pf-trust', 'allowance', ?, ?, ?, ?, ?, 'confirmed')", [ref, ORG, ref, value_date, value_date, value_date, amount, ccy])


def _owed(rows: list[dict]) -> dict:
    return {(r["member_name"], str(r["month"])[:7]): float(r["owed"]) for r in rows}


def test_allowance_due_matches_the_key(con):
    """The inputs of the key: the September and October schedules and the September Trust statement."""
    _load_trust(con, ["2026-09", "2026-10"], [OFFICE / "2026-10" / "bank_statements" / "HB_TRUST_USD_2026-09.xlsx"])
    rows = _rows(con, _build(con, "allowance_due"), "month, member_name")
    key = _key("uc-015-allowances")
    owed = _owed(rows)
    for member, amount in key["due_october"].items():
        assert owed[(member, "2026-10")] == amount
    earlier = [r for r in rows if r["earlier"] == 1 and r["owed"] > 0]
    assert [(r["member_name"], str(r["month"])[:7], float(r["owed"])) for r in earlier] == [(u["member"], u["month"], u["amount"]) for u in key["unpaid_earlier"]]
    assert float(sum(r["owed"] for r in rows)) == key["total_owed"] == 28000
    assert {r["entity"] for r in rows} == {"Evergreen Family Trust"}


def test_allowance_due_over_all_history_only_the_planted_gap(con):
    """Every schedule April to October and every Trust USD statement April to September: the only unpaid line before October is Ethan Chen's September."""
    stm = sorted(Path(p) for p in glob.glob(str(OFFICE / "*" / "bank_statements" / "HB_TRUST_USD_2026-??.xlsx")))
    assert len(stm) >= 6
    _load_trust(con, ["2026-04", "2026-05", "2026-06", "2026-07", "2026-08", "2026-09", "2026-10"], stm)
    owed = {k: v for k, v in _owed(_rows(con, _build(con, "allowance_due"), "month, member_name")).items() if v}
    assert owed == {("Ethan Chen", "2026-09"): 5000.0, ("Lin Chen", "2026-10"): 8000.0, ("Maya Chen", "2026-10"): 6000.0, ("Ethan Chen", "2026-10"): 5000.0, ("Wen Chen", "2026-10"): 4000.0}


def test_trust_cash_cannot_cover_the_allowances_after_the_buffer():
    """The cash half of the card: the Trust's USD at 13 October (both banks, read from the daily balances) against the signed buffer."""
    cash = Decimal("0")
    for bank in ("HB", "TCB"):
        ws = openpyxl.load_workbook(OFFICE / "2026-10" / "bank_statements" / f"{bank}_TRUST_USD_2026-10_mtd.xlsx")["Daily balances"]
        cash += Decimal(str(list(ws.iter_rows(values_only=True))[-1][5]))
    assert float(cash) == _key("uc-015-allowances")["trust_usd_cash"] == 213895.52
    rules = {r["rule"]: r["value"] for r in csv.DictReader((OFFICE / "reference" / "office_rules.csv").open(encoding="utf-8"))}
    assert float(cash) - float(rules["cash_buffer_usd"]) < _key("uc-015-allowances")["total_owed"]  # the signal holds the payment list


# ---- the card 50 data: fee invoices and the fee schedule read from PDF ----------------------------------------------------------------------

def _vendor(line: str) -> str:
    return re.sub(r"\s*\(synthetic\)", "", line.strip(), flags=re.I).title()


def _doc(con, did: str, dtype: str, title: str, dated=None) -> None:
    con.execute("insert into documents (id, organization_id, source_system_id, access_group_id, document_type, title, original_uri, content_hash, document_date, received_at, recorded_at, "
                "document_status, processing_status) values (?, ?, 'ss', 'ag', ?, ?, ?, ?, ?, now(), now(), 'Final', 'Done')", [did, ORG, dtype, title, title, did, dated])


def _field(con, did: str, name: str, text=None, num=None, ccy=None) -> None:
    con.execute("insert into extracted_fields (id, organization_id, document_id, field_name, value_text, value_numeric, currency, confidence, review_status) values (?, ?, ?, ?, ?, ?, ?, 'high', 'Accepted')",
                [f"{did}-{name}-{text or ''}", ORG, did, name, text, num, ccy])


def _load_fee_invoices(con) -> None:
    from datetime import datetime
    for f in sorted(glob.glob(str(OFFICE / "*" / "invoices" / "*FEE-*.pdf"))):
        t = _pdf(f)
        did = re.search(r"Invoice no: (\S+)", t).group(1)
        _doc(con, did, "invoice", did, datetime.strptime(re.search(r"Invoice date: (\d+ \w+ \d{4})", t).group(1), "%d %B %Y").date())
        _field(con, did, "invoice_no", text=did)
        _field(con, did, "vendor", text=_vendor(t.splitlines()[0]))
        _field(con, did, "fee_base", num=Decimal(re.search(r"Fee base \(USD\): ([\d,\.]+)", t).group(1).replace(",", "")), ccy="USD")
        _field(con, did, "fee_rate", num=Decimal(re.search(r"Fee rate per year: ([\d\.]+)%", t).group(1)) / 100)
        _field(con, did, "amount_due", num=Decimal(re.search(r"Amount due: USD ([\d,\.]+)", t).group(1).replace(",", "")), ccy="USD")
    _doc(con, "fee-schedule", "fee_schedule", "Fee schedule on file")
    for line in _pdf(OFFICE / "reference" / "fee_schedule.pdf").splitlines():
        m = re.match(r".*\((.+)\): ([\d\.]+)% a year .*billed (quarterly)", line)
        if m:
            _field(con, "fee-schedule", "service_fee_rate", text=m.group(1), num=Decimal(m.group(2)) / 100)
            _field(con, "fee-schedule", "service_fee_periods", text=m.group(1), num=Decimal(4))


def test_invoice_variance_matches_the_key(con):
    _load_fee_invoices(con)
    got = {r["invoice_no"]: r for r in _rows(con, _build(con, "invoice_variance"), "invoice_no")}
    key = _key("uc-050-feeinv")
    assert set(got) == {i["id"] for i in key["invoices"]}
    for w in key["invoices"]:
        g = got[w["id"]]
        assert (float(g["billed"]), float(g["expected"]), float(g["variance"]), float(g["rate_billed_pct"])) == (w["billed"], w["expected"], w["variance"], w["rate_billed_pct"])
    assert [i for i, g in got.items() if g["variance"] > 0] == key["overcharged"] == ["FEE-CUST-Q3"]
    assert float(got["FEE-CUST-Q3"]["rate_schedule_pct"]) == 0.05  # billed at 0.06, the schedule says 0.05


def test_invoice_without_a_schedule_rate_stays_in_with_no_variance(con):
    _load_fee_invoices(con)
    con.execute("delete from extracted_fields where field_name like 'service_fee_%' and value_text = 'Anchorage Custody'")
    got = {r["invoice_no"]: r for r in _rows(con, _build(con, "invoice_variance"), "invoice_no")}
    assert got["FEE-CUST-Q3"]["variance"] is None and got["FEE-CUST-Q3"]["expected"] is None and got["FEE-ADMIN-Q3"]["variance"] == 0


# ---- card 9 (uc9-sop): no table holds a process, so the key is checked on the narrative itself ---------------------------------------------

def test_sop_steps_and_owners_match_the_key():
    text = (OFFICE / "reference" / "month_end_process_narrative.md").read_text(encoding="utf-8")
    steps = re.findall(r"### Step \d+\. (.+?)\n\n- Who: (.+?)\n- System: (.+?)\n- Manual spreadsheet or hand step: (\w+)\n- Control: (.+?)\n- Known exception: (.+?)\n", text)
    key = _key("uc9-sop")
    assert len(steps) == key["steps"] == 6
    assert {s[0]: s[1] for s in steps} == key["owners"]
    assert not [s[0] for s in steps if not (s[1] and s[2] and s[4])]  # no step without an owner, a system and a check: the exception signal stays quiet


# ---- the signals, the builds and the figures they point at ---------------------------------------------------------------------------------

CARDS = {"uc9-sop": "uc9-sop-", "uc-015-allowances": "uc-015-allowances-", "uc-030-vendors": "uc-030-vendors-", "uc-050-feeinv": "uc-050-feeinv-"}
SIGNALS = {p.stem: yaml.safe_load(p.read_text(encoding="utf-8")) for pre in CARDS.values() for p in sorted((PACK / "signals").glob(pre + "*.yaml"))}
for _s in SIGNALS.values():  # YAML reads a bare `on` as True; the engine renames it the same way
    _s["trigger"] = {("on" if k is True else k): v for k, v in _s["trigger"].items()}
ROLES = {r["id"] for r in yaml.safe_load((PACK / "role_builds.yaml").read_text(encoding="utf-8"))["roles"]}
FUNCS = {"abs", "min", "max", "round", "last4", "num", "same_fund"}


def _names(expr: str) -> set[str]:
    return {n.id for n in ast.walk(ast.parse(str(expr), mode="eval")) if isinstance(n, ast.Name)} - FUNCS


@pytest.mark.parametrize("card,prefix", CARDS.items())
def test_each_card_has_signals_with_an_exception(card, prefix):
    own = {i: s for i, s in SIGNALS.items() if i.startswith(prefix)}
    assert len(own) >= 3
    assert [i for i, s in own.items() if s.get("hold")], "no exception signal with a hold"


@pytest.mark.parametrize("sid", SIGNALS)
def test_signal_is_in_the_engine_format(sid):
    s = SIGNALS[sid]
    assert s["id"] == sid and s["suggestion"] is True and s["title"]
    assert s["trigger"]["on"] in {"file_arrived", "figure_changed", "date_passed", "schedule", "run", "action_done"}
    assert s["route_to"] in ROLES
    assert s["severity"] in {"info", "needs", "problem"}
    assert s["action"]["verb"] in {"log", "record", "approve", "call_back", "see"} and s["action"]["options"]
    known = {"doc"} | {i["as"] for i in s.get("inputs", []) if "as" in i} | set(s.get("let") or {})
    used = _names(s["condition"]) | {n for e in (s.get("let") or {}).values() for n in _names(e)}
    assert used <= known, used - known
    for ref in re.findall(r"\{([a-z_]+)[\.:\}]", s["message"]):
        assert ref in known, ref
    if s.get("hold"):
        assert s["severity"] == "problem" and s["hold"]["blocks"] and s["hold"]["until"]
    if s["trigger"]["on"] == "action_done":
        assert s["trigger"]["signal"] in SIGNALS


def test_figure_inputs_name_a_recipe_that_exists():
    here = {r["id"] for r in yaml.safe_load((PACK / "recipes" / "office.yaml").read_text(encoding="utf-8"))["recipes"]}
    other = {r["id"] for f in (PACK / "recipes").glob("*.yaml") for r in (yaml.safe_load(f.read_text(encoding="utf-8")) or {}).get("recipes", [])}
    for s in SIGNALS.values():
        for i in s.get("inputs", []):
            if i["kind"] == "figure" and i.get("recipe"):
                assert i["recipe"] in other
    assert {"renewals_vendors", "allowance_due", "invoice_variance"} <= here


def test_builds_exist_and_their_numbers_are_the_key():
    roles = yaml.safe_load((PACK / "builds" / "office.yaml").read_text(encoding="utf-8"))["roles"]
    builds = {b["id"]: (r["id"], b) for r in roles for b in r["builds"]}
    assert set(builds) == {"uc-015-allowances", "uc-030-vendors", "uc-050-feeinv"}  # uc9-sop keeps its build in role_builds.yaml
    for bid, (role, b) in builds.items():
        assert role in ROLES and b["say"] and b["check"]["ask"] and b["check"]["expect"]
        assert b["check"]["ask"].rstrip("?") == next(iter(KEYS[bid]["questions"])).rstrip("?")
        for f in b["files"]:
            assert f.startswith("office6m:") and glob.glob(str(OFFICE / f[len("office6m:"):])), f
        flat = yaml.safe_dump(KEYS[bid]["questions"]).replace(",", "")
        for e in b["check"]["expect"]:
            assert e.replace(",", "") in flat, (bid, e)
    assert not [b for r in yaml.safe_load((PACK / "role_builds.yaml").read_text(encoding="utf-8"))["roles"] for b in r["builds"] if b["id"] in builds]
