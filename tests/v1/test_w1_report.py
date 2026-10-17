"""W1 Report and govern: the recipes of packs/family_office/1/recipes/report.yaml proved on the synthetic Evergreen office (tests/v1/fixtures/office6m).

The office files are loaded into an in-memory DuckDB with the company model (docs/v2/evidence-to-action-mvp.duckdb.sql) by `load_office` below: a small loader
written for this test, because the platform loader does not read the office6m files yet (see docs/v3/w1/build-report.md, Needs). Each recipe's SQL is then
run exactly as v3_recipes renders it and its numbers are compared with tests/v1/fixtures/office6m/answer_keys.yaml. The signal yaml and the persona builds
are checked for their shape. Synthetic data only."""
from __future__ import annotations

import csv
import glob
import json
import re
import uuid
from datetime import date, datetime, timedelta
from pathlib import Path

import duckdb
import openpyxl
import pytest
import yaml

ROOT = Path(__file__).resolve().parents[2]
OFFICE = ROOT / "tests" / "v1" / "fixtures" / "office6m"
PACK = ROOT / "packs" / "family_office" / "1"
DDL = ROOT / "docs" / "v2" / "evidence-to-action-mvp.duckdb.sql"
KEYS = yaml.safe_load((OFFICE / "answer_keys.yaml").read_text(encoding="utf-8"))["keys"]
ORG = "org-evergreen"
SRC = "src-office6m"
NOW = "2026-10-13 12:00:00+00"


# ---- the recipes, rendered the way v3_recipes renders them -------------------------------------------------------------------------------

def _recipes() -> dict[str, dict]:
    return {r["id"]: r for r in yaml.safe_load((PACK / "recipes" / "report.yaml").read_text(encoding="utf-8"))["recipes"]}


def _snippets() -> dict[str, str]:
    return yaml.safe_load((PACK / "recipes" / "_shared.yaml").read_text(encoding="utf-8"))["snippets"]


def sql_of(rid: str, name: str | None = None) -> str:
    r, sn = _recipes()[rid], _snippets()
    return " ".join(r["sql"].replace("{dims}", "").replace("{name}", name or f"t_{rid}").replace("{positions}", sn["positions"]).split())


def build(con, rid: str) -> str:
    name = f"t_{rid}"
    con.execute(f"drop table if exists {name}")
    con.execute(sql_of(rid, name))
    return name


def rows(con, name: str, where: str = "true", order: str = "1") -> list[dict]:
    cur = con.execute(f"select * from {name} where {where} order by {order}")
    cols = [c[0] for c in cur.description]
    return [dict(zip(cols, r)) for r in cur.fetchall()]


# ---- the loader: office6m files -> company model -----------------------------------------------------------------------------------------

class Loader:
    def __init__(self, con):
        self.con, self.n = con, 0

    def uid(self, p: str) -> str:
        self.n += 1
        return f"{p}-{self.n:06d}"

    def ins(self, table: str, row: dict) -> None:
        base = {"id": self.uid(table[:4]), "organization_id": ORG}
        row = {**base, **row}
        cols = ", ".join(row)
        self.con.execute(f"insert into {table} ({cols}) values ({', '.join('?' * len(row))})", list(row.values()))


def _xl(path: Path):
    return openpyxl.load_workbook(path, data_only=True)


def _month_end(d: str) -> date:
    y, m = int(d[:4]), int(d[5:7])
    return (date(y + (m == 12), m % 12 + 1, 1)) - timedelta(days=1)


def _day(s: str) -> str:
    """One custodian row writes its date day first (10/07/2026, planted schema drift); the settle date beside it is ISO."""
    return datetime.strptime(s, "%d/%m/%Y").date().isoformat() if "/" in s else s


def _pdf_text(path: Path) -> str:
    import pypdf
    return "\n".join(p.extract_text() or "" for p in pypdf.PdfReader(str(path)).pages)


def load_office(con) -> None:
    con.execute("set TimeZone = 'UTC'")
    con.execute(DDL.read_text(encoding="utf-8"))
    L = Loader(con)
    ref = OFFICE / "reference"
    stamp = {"received_at": NOW, "recorded_at": NOW}

    # portfolios and instruments: the security master (reference/EG_security_master_*.xlsx)
    pf_names = {"EG-LIQUID": "Evergreen Liquid Book"}  # the issued report calls the liquid portfolio this
    pid, iid, ibyname = {}, {}, {}
    sm = list(_xl(next(ref.glob("EG_security_master_*.xlsx"))).worksheets[0].iter_rows(values_only=True))
    for r in sm[4:]:
        if not r[0] or not r[2]:
            continue
        code, name, typ, cls, ccy, mgr, pcode = r[0], r[2], r[3], r[4], r[5], r[9], r[10]
        iid[code] = L.uid("ins")
        ibyname[name] = iid[code]
        L.ins("instruments", {"id": iid[code], "canonical_name": name, "instrument_type": typ, "asset_class": cls, "manager_name": mgr, "currency": ccy, "created_at": NOW})
    for code in ("EG-LIQUID", "EG-GROWTH", "EG-PRIVATE", "EG-B-OPS", "EG-FUNDS"):
        pid[code] = L.uid("pf")
        L.ins("portfolios", {"id": pid[code], "legal_entity_id": "le-1", "portfolio_code": code, "portfolio_name": pf_names.get(code, code), "base_currency": "USD"})

    # administrator exports: one pack per month end (Positions sheet; the NAV sheet holds the 'As of' date)
    for f in sorted(glob.glob(str(OFFICE / "*" / "administrator_nav" / "administrator_export_*.xlsx"))):
        wb = _xl(Path(f))
        asof = date.fromisoformat(str(list(wb["NAV by portfolio"].iter_rows(values_only=True))[1][2])[:10])
        for r in list(wb["Positions"].iter_rows(values_only=True))[1:]:
            if not r[0] or r[5] is None:  # 'Not received': the administrator sent no value
                continue
            if r[2] not in ibyname:  # a fund the security master does not list (F, G, K): an instrument with no class and no manager
                ibyname[r[2]] = L.uid("ins")
                L.ins("instruments", {"id": ibyname[r[2]], "canonical_name": r[2], "instrument_type": "private_fund", "currency": r[6], "created_at": NOW})
            L.ins("position_snapshots", {"source_record_id": SRC, "portfolio_id": pid[r[0]], "instrument_id": ibyname[r[2]], "position_date": asof,
                                         "as_of_at": f"{r[8]} 00:00:00+00", "quantity": r[3], "source_market_value": r[5], "currency": r[6],
                                         "view_type": "accounting", "data_status": r[7], **stamp})

    # custodian transactions of the liquid book (buys and sells; income lines too)
    seen = set()
    for f in sorted(glob.glob(str(OFFICE / "*" / "custodian_transactions" / "*.csv"))):
        for r in csv.DictReader(open(f, encoding="utf-8")):
            if r["reference"] in seen or not r["security_id"]:
                continue
            seen.add(r["reference"])
            amt, kind = float(r["amount"]), r["type"]
            L.ins("transactions", {"source_record_id": SRC, "portfolio_id": pid["EG-LIQUID"], "instrument_id": iid[r["security_id"]], "external_trade_id": r["reference"],
                                   "transaction_type": kind, "trade_date": _day(r["txn_date"]), "effective_date": _day(r["txn_date"]), "settlement_date": r["settle_date"],
                                   "quantity_delta": (float(r["quantity"]) * (1 if kind == "buy" else -1)) if r["quantity"] else None,
                                   "cash_amount": -amt if kind == "buy" else amt, "currency": r["currency"], "recorded_at": NOW, "status": "settled"})

    # the benchmark workbook: definitions, components, and the March month-end level (the daily files start on 1 April)
    wb = _xl(next(ref.glob("EG-OFFICE-BLEND_*.xlsx")))
    series, level_rows = {}, list(wb["Index Levels"].iter_rows(values_only=True))
    for ver, sheet in ((1, "Components v1"), (2, "Components v2 (In Progress)")):
        rs = list(wb[sheet].iter_rows(values_only=True))
        status = str(rs[1][0]).split(":")[1].strip()
        eff = date.fromisoformat(str(rs[2][1])[:10])
        did = L.uid("bd")
        L.ins("benchmark_definitions", {"id": did, "benchmark_code": "EG-OFFICE-BLEND", "benchmark_name": "Evergreen office blend", "version_number": ver, "reporting_currency": "USD",
                                        "effective_from": eff, "calculation_method": str(rs[5][1]), "approval_status": "approved" if status == "Approved" else "in_progress",
                                        "approved_by": rs[3][1] if status == "Approved" else None})
        for c in rs[8:11]:
            series[c[0]] = c[1]
            L.ins("benchmark_components", {"benchmark_definition_id": did, "market_series_key": c[1], "component_weight": c[2], "spread_bps": c[3], "weight_effective_from": eff,
                                           "treatment_note": c[4]})
    head = level_rows[3]
    for r in level_rows[4:]:
        if r[0] and str(r[0]) == "2026-03-31":
            for col in (1, 2):
                L.ins("market_observations", {"source_record_id": SRC, "series_key": series[head[col]], "observation_type": "index_level", "observed_at": f"{r[0]} 00:00:00+00",
                                              "value": r[col], "quality_status": "official", **stamp})

    # daily index levels and the WMR USD/SGD fixing
    for f in sorted(glob.glob(str(OFFICE / "*" / "index_levels" / "index_levels_*.csv"))):
        for r in csv.DictReader(open(f, encoding="utf-8")):
            for col, key in series.items():
                L.ins("market_observations", {"source_record_id": SRC, "series_key": key, "observation_type": "index_level", "observed_at": f"{r['Date']} 00:00:00+00",
                                              "value": float(r[col]), "quality_status": "official", **stamp})
    for f in sorted(glob.glob(str(OFFICE / "*" / "fx_rates" / "fx_usdsgd_*.csv"))):
        for r in csv.DictReader(open(f, encoding="utf-8")):
            L.ins("market_observations", {"source_record_id": SRC, "series_key": "FX:USDSGD:WMR4PM", "observation_type": "fx", "observed_at": f"{r['Date']} 16:00:00+00",
                                          "value": float(r["Rate (SGD per 1 USD)"]), "base_currency": "USD", "quote_currency": "SGD", "quality_status": "official", **stamp})

    # issued reports (xlsx): one document per month, one extracted field per line (assumption A3)
    for f in sorted(glob.glob(str(OFFICE / "*" / "investment_reports" / "Evergreen_Office_Investment_Report_*.xlsx"))):
        month = re.search(r"_(\d{4}-\d{2})\.xlsx", f)[1]
        did, me = L.uid("doc"), _month_end(month)
        L.ins("documents", {"id": did, "source_system_id": "sys-1", "access_group_id": "ag-1", "document_type": "investment_report", "title": Path(f).name, "original_filename": Path(f).name,
                            "original_uri": f, "content_hash": uuid.uuid5(uuid.NAMESPACE_URL, f).hex, "document_date": me, "valuation_date": me,
                            "received_at": f"{(me + timedelta(days=3)).isoformat()} 10:00:00+00", "recorded_at": NOW, "document_status": "issued", "processing_status": "done"})
        for n, r in enumerate(list(_xl(Path(f))["Report"].iter_rows(values_only=True))[1:], 2):
            L.ins("extracted_fields", {"document_id": did, "field_name": r[0], "value_text": r[2], "value_numeric": r[1], "currency": "SGD" if "SGD" in r[0] else "USD", "page_number": 1,
                                       "line_number": n, "confidence": "high", "review_status": "accepted"})

    # the auditor's sample (audit/audit_request_*.xlsx): one field per column and item (the item number is the line)
    f = next(glob.iglob(str(OFFICE / "*" / "audit" / "audit_request_*.xlsx")))
    did = L.uid("doc")
    L.ins("documents", {"id": did, "source_system_id": "sys-1", "access_group_id": "ag-1", "document_type": "audit_request", "title": Path(f).name, "original_filename": Path(f).name,
                        "original_uri": f, "content_hash": uuid.uuid5(uuid.NAMESPACE_URL, f).hex, "document_date": "2026-09-22", "received_at": NOW, "recorded_at": NOW,
                        "document_status": "received", "processing_status": "done"})
    for r in list(_xl(Path(f)).worksheets[0].iter_rows(values_only=True))[1:]:
        for name, kind, val in (("report_month", "text", r[1]), ("line", "text", r[2]), ("figure", "num", r[3])):
            L.ins("extracted_fields", {"document_id": did, "field_name": name, "value_text": val if kind == "text" else None, "value_numeric": val if kind == "num" else None,
                                       "page_number": 1, "line_number": r[0], "confidence": "high", "review_status": "accepted"})

    # Fund A: the estimate and the amended valuation that replaces it (manager statements, closing balance)
    est, amd = [next(glob.iglob(str(OFFICE / "*" / "manager_statements" / f"northgate_A_2026-09-30_{k}.pdf"))) for k in ("estimate", "amended")]
    val = {}
    for k, f in (("estimate", est), ("amended", amd)):
        val[k] = float(re.search(r"Closing capital account balance\s+([\d,]+\.\d\d)", _pdf_text(Path(f)))[1].replace(",", ""))
    for k, eid in (("estimate", "val-A-est"), ("amended", "val-A-amd")):
        L.ins("valuation_observations", {"id": eid, "instrument_id": ibyname["Evergreen Growth I"], "valuation_date": "2026-09-30", "value_amount": val[k], "currency": "USD",
                                         "valuation_type": "manager_estimate" if k == "estimate" else "manager_final", "approval_status": "official",
                                         "supersedes_valuation_id": "val-A-est" if k == "amended" else None,
                                         "received_at": "2026-10-02 09:00:00+00" if k == "estimate" else "2026-10-07 09:00:00+00", "recorded_at": NOW})

    # the signed rules: the allocation policy v2026.1 (limits section 3, reporting rules section 6)
    policy = _pdf_text(ref / "Evergreen_Investment_Allocation_Policy_v2026.1.pdf")
    for must in ("0%\n5%", "60%\n90%", "0%\n55%", "20%\n60%", "100,000 USD or 0.5%"):
        assert must in policy, must  # the numbers below are the policy's
    rules = [("limit_single_name_liquid", {"portfolio": "EG-LIQUID", "dimension": "single_name", "bucket": "any single name", "min_pct": 0, "max_pct": 5}),
             ("limit_listed_equity_liquid", {"portfolio": "EG-LIQUID", "dimension": "asset_class", "bucket": "listed equity", "min_pct": 60, "max_pct": 90}),
             ("limit_manager_entity_b", {"portfolio": "EG-B-OPS", "dimension": "manager", "bucket": "Cedarpoint Ventures", "min_pct": 0, "max_pct": 55}),
             ("limit_hedge_funds", {"portfolio": "EG-FUNDS", "dimension": "asset_class", "bucket": "hedge fund", "min_pct": 20, "max_pct": 60}),
             ("cio_materiality", {"amount": 100000, "pct": 0.5})]
    for code, params in rules:
        rid = L.uid("rule")
        L.ins("business_rules", {"id": rid, "rule_code": code, "rule_type": "policy" if code.startswith("limit") else "threshold", "owner_user_id": "u-sofia"})
        L.ins("rule_versions", {"business_rule_id": rid, "version_number": 1, "definition_text": code, "parameters": json.dumps(params), "effective_from": "2026-01-01",
                                "approval_status": "approved", "approved_by": "Daniel Foo", "approved_at": "2026-01-01 00:00:00+00"})
        if code == "limit_single_name_liquid":  # the question of card 48: what if the limit were 6? (proposed_limit in the answer key); a proposal, not a signed rule
            L.ins("rule_versions", {"business_rule_id": rid, "version_number": 2, "definition_text": code, "parameters": json.dumps({**params, "max_pct": 6}), "effective_from": "2026-10-13",
                                    "approval_status": "proposed"})

    # Two rules the key assumes and the office files do not carry. They stand in for rules the office has still to sign (docs/v3/w1/build-report.md, Needs 4 and 8); the
    # recipes read them like any signed rule, and without them the recipes say so (tests below remove them and look). Not in the policy text:
    #   single_name_exempt  the key and the committee notes leave the sovereign treasury note out of the 5 percent single-name limit
    #   benchmark_link      the key compares the liquid book with the office blend benchmark
    for code, params in (("single_name_exemption_treasury", {"portfolio": "EG-LIQUID", "dimension": "single_name_exempt", "instrument": "Harbour Treasury Note 2029 (synthetic sovereign)"}),
                         ("benchmark_link", {"portfolio": "EG-LIQUID", "benchmark_code": "EG-OFFICE-BLEND"})):
        rid = L.uid("rule")
        L.ins("business_rules", {"id": rid, "rule_code": code, "rule_type": "policy" if code.startswith("single") else "mapping", "owner_user_id": "u-sofia"})
        L.ins("rule_versions", {"business_rule_id": rid, "version_number": 1, "definition_text": code, "parameters": json.dumps(params), "effective_from": "2026-01-01",
                                "approval_status": "approved", "approved_by": "Daniel Foo", "approved_at": "2026-01-01 00:00:00+00"})

    # the pipeline (fund teaser: minimum commitment) and the capital call notice of 13 October (cash_movements, noticed)
    teaser = _pdf_text(next(OFFICE.glob("*/prospects/fund_teaser_Harborlight_Ventures_IV.pdf")))
    L.ins("deal_pipeline", {"source_record_id": SRC, "deal_name": "Harborlight Ventures IV", "manager_name": "Harborlight Capital Partners", "strategy": "venture",
                            "target_commitment": float(re.search(r"Minimum commitment: USD ([\d,]+)", teaser)[1].replace(",", "")), "currency": "USD", "stage": "under review",
                            "stage_on": "2026-09-30", **stamp})
    notice = _pdf_text(next(OFFICE.glob("2026-10/notices/*call7.pdf")))
    L.ins("cash_movements", {"source_record_id": SRC, "portfolio_id": pid["EG-PRIVATE"], "instrument_id": ibyname["Cedarpoint Ventures III"], "movement_type": "call",
                             "effective_date": "2026-10-23", "due_date": "2026-10-23", "notice_kind": "capital_call", "amount": -float(re.search(r"Amount to be paid by the investor\s+USD ([\d,]+)", notice)[1].replace(",", "")),
                             "currency": "USD", "confirmation_status": "noticed", **stamp})


@pytest.fixture(scope="module")
def con():
    c = duckdb.connect()
    load_office(c)
    yield c
    c.close()


# ---- helpers for the checks --------------------------------------------------------------------------------------------------------------

def tie(con, rid: str) -> tuple[float, float]:
    t = _recipes()[rid]["tie_out"]["sql"]
    t = t.replace("{name}", f"t_{rid}").replace("{positions_check}", _snippets()["positions_check"])
    a, b = con.execute(" ".join(t.split())).fetchone()
    return float(a or 0), float(b or 0)


def num(v) -> float:
    return float(v)


def key(card: str, n: int = 0) -> dict:
    return list(KEYS[card]["questions"].values())[n]["expected"]


MINE = ("concentration", "change_since_report", "time_weighted_return", "attribution", "single_name_breach", "drift_vs_target", "limit_changes",
        "report_tie_out", "report_ready", "report_total_fx", "trace_coverage", "pack_items")


@pytest.fixture(scope="module")
def built(con):
    return {rid: build(con, rid) for rid in MINE}


# ---- the recipes as a set ----------------------------------------------------------------------------------------------------------------

def test_every_recipe_is_built_reads_only_what_it_may_and_ties_out(con, built):
    from sh_platform.v1 import v3_recipes as R
    rs = _recipes()
    ready = [r for r in rs.values() if r.get("status") != "stub"]
    assert {r["id"] for r in ready} == set(MINE)
    assert rs["networth_walk"]["status"] == "stub"  # card 75 is not in the top 50
    for r in ready:
        sql = sql_of(r["id"])
        cols = {c[0] for c in con.execute(f"describe t_{r['id']}").fetchall()}
        assert r.get("may_read") and r.get("never_read") is not None and r.get("formula") and r.get("grain") and r["sql"] and r.get("names"), r["id"]
        assert R.refusal(r, sql) is None, (r["id"], R.refusal(r, sql))  # A8: its SQL reads only the tables it names
        assert not R.tables_read(sql) & set(r["never_read"]), r["id"]
        assert "--" not in sql and "{" not in sql, r["id"]  # one line of SQL, every placeholder filled
        assert r["measure"] in cols and set(r.get("columns") or {}) <= cols, r["id"]
        if r.get("tie_out"):
            a, b = tie(con, r["id"])
            assert abs(a - b) <= 0.01, (r["id"], r["tie_out"]["label"], a, b)


# ---- card 12, uc5-limits -----------------------------------------------------------------------------------------------------------------

def test_uc5_limits_the_largest_name_over_the_limit_is_5_75_and_a_holding_has_no_class(con, built):
    k = key("uc5-limits")
    r = rows(con, built["single_name_breach"], "holding = 'Harborview Semiconductor'")[0]
    assert (num(r["concentration_pct"]), num(r["limit_pct"]), num(r["exposure"]), r["status"]) == (k["pct_of_portfolio"], k["limit_pct"], k["value"], "real breach")
    assert rows(con, built["single_name_breach"], "not exempt", "exposure desc")[0]["holding"] == k["holding"]  # the largest name that is not the sovereign note, as the key says
    assert [x["holding"] for x in rows(con, built["single_name_breach"], "no_asset_class")] == ["Pacific Rim Convertible 2029 note"]
    assert [x["holding"] for x in rows(con, built["single_name_breach"], "exempt")] == ["Harbour Treasury Note 2029 (synthetic sovereign)"]
    assert num(rows(con, built["concentration"], "issuer = 'Harborview Semiconductor' and position_date = '2026-09-30'")[0]["concentration_pct"]) == 5.75  # the shared recipe agrees


def test_uc5_limits_data_finding_other_names_are_also_over_five_percent(con, built):
    """The key and the committee notes name only the largest name. By the same data six more names are above 5 percent at 30 September (reported under Needs 6): the
    key's 'real_breaches: 1' counts the largest name only, and no signed rule or committee decision in the files excuses the other six, so the pack says 7."""
    over = rows(con, built["single_name_breach"], "over_limit", "concentration_pct desc")
    assert len(over) == 7 and over[0]["holding"] == "Harborview Semiconductor" and all(o["status"] == "real breach" for o in over)


def test_a_holding_is_exempt_only_when_a_signed_rule_says_so(con, built):
    """The sovereign note is exempt through the signed exemption rule (stand-in row in the loader), not through its name. Without the row it is a breach like the rest."""
    exempt = rows(con, built["single_name_breach"], "exempt")
    assert [(x["holding"], x["exemption_rule"], x["status"]) for x in exempt] == [("Harbour Treasury Note 2029 (synthetic sovereign)", "single_name_exemption_treasury", "exempt")]
    con.execute("update rule_versions set approval_status = 'proposed' where definition_text = 'single_name_exemption_treasury'")  # not signed
    try:
        n = build(con, "single_name_breach")
        sov = rows(con, n, "holding ilike '%sovereign%'")[0]
        assert sov["exempt"] is False and sov["status"] == "real breach" and num(sov["concentration_pct"]) == 7.08 and len(rows(con, n, "over_limit")) == 8
        lc = rows(con, build(con, "limit_changes"), "version_number = 1 and not is_opening", "month_end")
        assert {r["largest_holding"] for r in lc} == {"Harbour Treasury Note 2029 (synthetic sovereign)"} and all(r["over_limit"] == 1 for r in lc)  # the figures follow the rule on record
        pack = {x["section"]: x for x in rows(con, build(con, "pack_items"))}
        assert pack["Breaches"]["items"] == 8 and "Harbour Treasury Note 2029" in pack["Breaches"]["status"]
    finally:
        con.execute("update rule_versions set approval_status = 'approved' where definition_text = 'single_name_exemption_treasury'")
        for rid in ("single_name_breach", "limit_changes", "pack_items"):
            build(con, rid)


def test_uc5_limits_a_breach_on_a_value_older_than_90_days_is_possibly_stale_not_real(con):
    hvs = "instrument_id = (select id from instruments where canonical_name = 'Harborview Semiconductor') and position_date = '2026-09-30'"
    con.execute(f"update position_snapshots set as_of_at = '2026-06-01 00:00:00+00' where {hvs}")
    try:
        r = rows(con, build(con, "single_name_breach"), "holding = 'Harborview Semiconductor'")[0]
        assert r["value_age_days"] == 121 and r["status"] == "possibly stale"
    finally:
        con.execute(f"update position_snapshots set as_of_at = '2026-09-30 00:00:00+00' where {hvs}")
        build(con, "single_name_breach")


def test_uc5_limits_other_limits_are_the_key_percentages(con, built):
    k = key("uc5-limits", 1)
    got = {r["rule_code"]: r for r in rows(con, built["drift_vs_target"])}
    assert num(got["limit_listed_equity_liquid"]["allocation_pct"]) == k["listed_equity_pct"]
    assert num(got["limit_hedge_funds"]["allocation_pct"]) == k["hedge_funds_pct_of_eg_funds"]
    assert num(got["limit_manager_entity_b"]["allocation_pct"]) == k["cedarpoint_pct_of_eg_b_ops"]


# ---- card 14, morning-networth-note ----------------------------------------------------------------------------------------------------

def test_morning_note_total_in_usd_and_sgd_at_the_wmr_fixing(con, built):
    k = key("morning-networth-note")
    r = rows(con, built["report_total_fx"])[0]
    assert (num(r["total_usd"]), num(r["total_sgd"]), num(r["usd_sgd_fixing"]), r["issued_on"].isoformat()) == (k["total_usd"], k["total_sgd"], k["fixing"], k["issued"])
    assert num(r["total_sgd"]) == num(r["total_sgd_reported"]) and r["total_status"] == "estimated"


# ---- card 20, uc1-why-nav-moved ---------------------------------------------------------------------------------------------------------

def test_uc1_fund_a_moved_by_180000_over_the_materiality_line(con, built):
    k = key("uc1-why-nav-moved")
    r = rows(con, built["change_since_report"])
    assert len(r) == 1 and (num(r[0]["original_value"]), num(r[0]["amended_value"]), num(r[0]["change_since_report"])) == (k["issued"], k["amended"], k["change"])
    mat = json.loads(con.execute("select parameters from rule_versions v join business_rules r on r.id = v.business_rule_id where r.rule_code = 'cio_materiality'").fetchone()[0])
    assert (abs(k["change"]) > mat["amount"] or 100 * abs(k["change"]) / k["issued"] > mat["pct"]) is k["over_materiality"]
    assert round(100 * k["change"] / k["issued"], 1) == k["change_pct"]


# ---- card 23, uc-023-tieout ----------------------------------------------------------------------------------------------------------------

def test_uc023_the_report_differs_from_the_administrator_on_three_lines(con, built):
    k = key("uc-023-tieout")
    diffs = rows(con, built["report_tie_out"], "difference <> 0", "report_date")
    assert [(d["report_month"], d["line"], num(d["report_value"]), num(d["administrator_value"]), num(d["difference"])) for d in diffs] == [
        (x["month"], x["line"], x["report"], x["administrator"], x["difference"]) for x in k["differences"]]
    assert con.execute(f"select count(distinct report_month) from {built['report_tie_out']}").fetchone()[0] == k["months_checked"]
    assert all(not d["explained_by_amendment"] for d in diffs)  # none of the three is an amendment: they are the planted breaks
    assert {r["report_month"] for r in rows(con, built["report_tie_out"], "lines_off = 0")} == {"2026-07", "2026-08", "2026-09"}  # those months are ready to sign


# ---- card 24, uc-024-audit ----------------------------------------------------------------------------------------------------------------

def test_uc024_seven_of_eight_sampled_figures_trace_to_a_report_line(con, built):
    k = key("uc-024-audit")
    r = rows(con, built["trace_coverage"], "true", "item")
    assert len(r) == k["sampled"] and sum(1 for x in r if x["traced"]) == k["traced"] and num(r[0]["share_traced_pct"]) == k["share_traced_pct"]
    bad = [x for x in r if not x["traced"]]
    assert [(x["item"], x["report_month"], x["line"], num(x["figure"])) for x in bad] == [(u["item"], u["month"], u["line"], u["figure"]) for u in k["untraceable"]]
    assert all(x["source_file"] and x["source_page"] == 1 for x in r if x["traced"]) and bad[0]["untraced_items"] == 1


# ---- card 32, uc-032-driftrep -------------------------------------------------------------------------------------------------------------

def test_uc032_no_bucket_is_outside_its_policy_band_at_30_september(con, built):
    k = key("uc-032-driftrep")
    got = {r["rule_code"]: r for r in rows(con, built["drift_vs_target"])}
    for row in k["rows"]:
        code = {"listed_equity (EG-LIQUID)": "limit_listed_equity_liquid", "hedge_fund (EG-FUNDS)": "limit_hedge_funds"}[row["rule"]]
        assert num(got[code]["allocation_pct"]) == row["pct"] and got[code]["outside_band"] is row["outside"] and got[code]["target_pct"] is None
    assert [r["rule_code"] for r in rows(con, built["drift_vs_target"], "outside_band")] == k["outside"] == []


def test_uc032_a_bucket_outside_its_band_on_an_old_value_is_marked_stale(con):
    rule = "definition_text = 'limit_listed_equity_liquid' and version_number = 1"
    hvs = "position_date = '2026-09-30' and instrument_id = (select id from instruments where canonical_name = 'Harborview Semiconductor')"
    con.execute(f"update rule_versions set parameters = replace(cast(parameters as varchar), '\"max_pct\": 90', '\"max_pct\": 75') where {rule}")
    try:
        r = rows(con, build(con, "drift_vs_target"), "rule_code = 'limit_listed_equity_liquid'")[0]
        assert r["outside_band"] is True and r["status"] == "outside band" and r["possibly_stale"] is False
        con.execute(f"update position_snapshots set as_of_at = '2026-05-01 00:00:00+00' where {hvs}")
        r = rows(con, build(con, "drift_vs_target"), "rule_code = 'limit_listed_equity_liquid'")[0]
        assert r["status"] == "outside band, possibly stale"
    finally:
        con.execute(f"update rule_versions set parameters = replace(cast(parameters as varchar), '\"max_pct\": 75', '\"max_pct\": 90') where {rule}")
        con.execute(f"update position_snapshots set as_of_at = '2026-09-30 00:00:00+00' where {hvs}")
        build(con, "drift_vs_target")


# ---- card 33, uc-033-breachesc ------------------------------------------------------------------------------------------------------------

def test_uc033_cost_to_cure_is_324210_53(con, built):
    k = key("uc-033-breachesc")
    r = rows(con, built["single_name_breach"], "holding = 'Harborview Semiconductor'")[0]
    assert (num(r["exposure"]), num(r["portfolio_total"]), num(r["limit_pct"]), num(r["cost_to_cure"])) == (k["exposure"], k["total"], k["limit_pct"], k["cost_to_cure"])
    assert abs((num(r["exposure"]) - num(r["cost_to_cure"])) / (num(r["portfolio_total"]) - num(r["cost_to_cure"])) - 0.05) < 1e-6  # selling it puts the name at the limit of what is left


# ---- card 42, uc4-vs-benchmark ------------------------------------------------------------------------------------------------------------

def test_uc4_september_portfolio_against_the_blended_benchmark(con, built):
    k = key("uc4-vs-benchmark")
    r = rows(con, built["time_weighted_return"], "portfolio = 'Evergreen Liquid Book' and period_kind = 'month' and period_label = '2026-09'")[0]
    assert (num(r["portfolio_return_pct"]), num(r["benchmark_return_pct"]), num(r["active_pct"])) == (k["portfolio_return_pct"], k["benchmark_return_pct"], k["active_pct"])
    assert (num(r["start_value"]), num(r["end_value"]), num(r["net_flow"])) == (k["liquid_aug"], k["liquid_sep"], 0)


def test_uc4_a_missing_index_level_means_no_benchmark_return_and_names_the_series(con):
    where = "series_key = 'IDX:GLBD:SYN' and cast(observed_at at time zone 'UTC' as date) >= '2026-09-01'"
    con.execute(f"create temp table lost as select * from market_observations where {where}")
    con.execute(f"delete from market_observations where {where}")
    try:
        r = rows(con, build(con, "time_weighted_return"), "portfolio = 'Evergreen Liquid Book' and period_label = '2026-09'")[0]
        assert r["benchmark_return_pct"] is None and r["active_pct"] is None and r["months_without_benchmark"] == 1 and r["benchmark_gap"] == "IDX:GLBD:SYN"
        assert num(r["portfolio_return_pct"]) == 4.1667  # the portfolio's own return does not depend on the benchmark
    finally:
        con.execute("insert into market_observations select * from lost")
        con.execute("drop table lost")
        build(con, "time_weighted_return")


# ---- card 44, uc-044-twr ----------------------------------------------------------------------------------------------------------------------

def test_uc044_monthly_quarter_and_since_april_returns_match_the_key(con, built):
    k = key("uc-044-twr")
    t, liquid = built["time_weighted_return"], "portfolio = 'Evergreen Liquid Book'"
    months = rows(con, t, f"{liquid} and period_kind = 'month'", "period_label")
    assert [round(num(m["portfolio_return_pct"]), 4) for m in months] == k["monthly_pct"]
    q = {r["period_label"]: r for r in rows(con, t, f"{liquid} and period_kind = 'quarter'")}
    assert (num(q["2026Q2"]["portfolio_return_pct"]), num(q["2026Q3"]["portfolio_return_pct"])) == (k["q2_apr_jun_pct"], k["q3_jul_sep_pct"]) and q["2026Q2"]["complete"] and q["2026Q3"]["complete"]
    since = rows(con, t, f"{liquid} and period_kind = 'all'")[0]
    assert (num(since["portfolio_return_pct"]), num(since["benchmark_return_pct"]), since["months"]) == (k["since_1_apr_pct"], k["benchmark_since_1_apr_pct"], 6)
    ytd, y12 = rows(con, t, f"{liquid} and period_kind = 'ytd'")[0], rows(con, t, f"{liquid} and period_kind = 'last_12_months'")[0]
    assert ytd["complete"] is False and y12["complete"] is False and y12["months"] == 6  # the data starts on 1 April, so the year to date and 12 months say they are short


def test_uc044_a_flow_in_the_month_is_counted_at_half_weight(con, built):
    r = rows(con, built["time_weighted_return"], "portfolio = 'Evergreen Liquid Book' and period_label = '2026-04'")[0]
    assert num(r["net_flow"]) == 775970 and num(r["portfolio_return_pct"]) == -0.2944
    s, e, f = num(r["start_value"]), num(r["end_value"]), num(r["net_flow"])
    assert abs(100 * (e - s - f) / (s + f / 2) - -0.2944) < 0.00005  # not the plain (end - start) / start, which would be +1.66


# ---- card 45, uc-045-reportpack ----------------------------------------------------------------------------------------------------------

def test_uc045_three_of_five_sections_are_ready(con, built):
    k = key("uc-045-reportpack")
    r = rows(con, built["report_ready"], "true", "section")
    assert len(r) == k["sections"] and sorted(x["section"] for x in r if x["ready"]) == sorted(k["ready"]) and sorted(x["section"] for x in r if not x["ready"]) == sorted(k["not_ready"])
    assert num(r[0]["share_ready_pct"]) == k["share_ready_pct"] and all(x["ties"] for x in r)
    assert {x["section"]: x["reason"] for x in r if not x["ready"]} == {"Evergreen Growth I": "rests on an estimate", "Total": "rests on an estimate"}


# ---- card 46, uc-046-attrib ----------------------------------------------------------------------------------------------------------------

def test_uc046_effects_add_to_the_active_return_for_september(con, built):
    k = key("uc-046-attrib")
    r = {x["asset_class"]: x for x in rows(con, built["attribution"], "portfolio = 'Evergreen Liquid Book' and month_end = '2026-09-30'")}
    assert set(r) == {"listed_equity", "bond", "unclassified"}
    for cls, want in (("listed_equity", k["effects_pct"]["equity"]), ("bond", k["effects_pct"]["bond"]), ("unclassified", k["effects_pct"]["unclassified"])):
        for part in ("allocation", "selection", "interaction"):
            assert abs(num(r[cls][part + "_pct"]) - want[part]) <= 0.0001, (cls, part, r[cls][part + "_pct"], want[part])  # the key rounds the index returns first; 0.0001 is its rounding
    x = r["bond"]
    assert (num(x["portfolio_return_pct"]), num(x["benchmark_return_pct"]), num(x["active_pct"]), num(x["sum_of_effects_pct"]), x["adds_up"]) == (
        k["portfolio_pct"], k["liquid_benchmark_pct"], k["active_pct"], k["sum_of_effects_pct"], k["adds_up"])


def test_uc046_a_month_with_flows_adds_up_too(con, built):
    """April to July have flows: the weights are average-capital weights, so the effects still add to the active return (this was a flagged break before)."""
    at = rows(con, built["attribution"])
    assert {x["portfolio"] for x in at} == {"Evergreen Liquid Book"} and all(x["adds_up"] for x in at)
    assert {x["month_end"].isoformat() for x in at} == {"2026-04-30", "2026-05-31", "2026-06-30", "2026-07-31", "2026-08-31", "2026-09-30"}
    april = {x["asset_class"]: x for x in rows(con, built["attribution"], "month_end = '2026-04-30'")}
    r = list(april.values())[0]
    assert abs(num(r["sum_of_effects_pct"]) - num(r["active_pct"])) <= 0.0001 and num(r["portfolio_return_pct"]) == -0.2944  # the same portfolio return as the return recipe


# ---- card 47, uc-047-icpack ----------------------------------------------------------------------------------------------------------------

def test_uc047_the_pack_has_five_sections_with_owners_and_the_open_items(con, built):
    k = key("uc-047-icpack")
    r = {x["section"]: x for x in rows(con, built["pack_items"])}
    assert sorted(r) == sorted(s["section"] for s in k["sections"])
    assert r["Report"]["amendments_open"] == k["amendments_open"] and num(r["Report"]["amount"]) == 57600000 and r["Report"]["status"] == "1 line(s) estimated"
    assert (r["Report"]["owner_role"], r["Attribution"]["owner_role"], r["Breaches"]["owner_role"], r["Proposals"]["owner_role"]) == ("cio", "investment_analyst", "risk_officer", "cfo")
    assert r["Performance"]["status"] == r["Attribution"]["status"] == "inputs ready"
    p = r["Proposals"]  # call 7 noticed and the Harborlight ticket: two amounts, never added (a call is owed, a ticket is a possible commitment)
    assert p["items"] == 2 and (num(p["call_amount"]), num(p["deal_amount"]), p["amount"]) == (1250000, 2000000, None) and "1250000" in p["status"] and "2000000" in p["status"]
    # the key's 'short USD 270,000' is the figure cash_after_calls (cash domain), not produced by this recipe
    assert r["Breaches"]["items"] == 7 and "Harborview Semiconductor at 5.75" in r["Breaches"]["status"]  # the key counts the largest only (Needs 6)


# ---- card 48, uc-048-policyupd -------------------------------------------------------------------------------------------------------------

def test_uc048_largest_name_by_month_end_and_breaches_at_5_and_6_percent(con, built):
    k = key("uc-048-policyupd")
    t = built["limit_changes"]
    cur = rows(con, t, "version_number = 1 and not is_opening", "month_end")  # the first position date (31 March) is the opening position: shown in the table, not counted
    assert {r["month_end"].isoformat(): num(r["largest_name_pct"]) for r in cur} == k["max_single_name_pct_by_month_end"]
    assert len(rows(con, t, "version_number = 1")) == 7 and [r["month_end"].isoformat() for r in rows(con, t, "is_opening")] == ["2026-03-31"] * 2  # March is there, flagged
    assert cur[0]["months_tested"] == 6 and cur[0]["breaches_in_period"] == k["breaches_at_5_pct"] and cur[0]["approval_status"] == "approved" and num(cur[0]["limit_pct"]) == k["current_limit"]
    prop = rows(con, t, "version_number = 2 and not is_opening", "month_end")
    assert prop[0]["breaches_in_period"] == k["breaches_at_6_pct"] and prop[0]["months_tested"] == 6 and prop[0]["approval_status"] == "proposed" and num(prop[0]["limit_pct"]) == k["proposed_limit"]
    assert [r["largest_holding"] for r in cur[1:]] == ["Harborview Semiconductor"] * 5  # April's largest is a bond


def test_uc048_the_proposed_limit_never_changes_the_signed_checks(con, built):
    assert num(rows(con, built["single_name_breach"], "holding = 'Harborview Semiconductor'")[0]["limit_pct"]) == 5
    assert not rows(con, built["drift_vs_target"], "max_pct = 6")  # only signed (approved) versions are tested elsewhere


# ---- the signals ---------------------------------------------------------------------------------------------------------------------------

CARDS = ["uc5-limits", "morning-networth-note", "uc1-why-nav-moved", "uc-023-tieout", "uc-024-audit", "uc-032-driftrep", "uc-033-breachesc", "uc4-vs-benchmark", "uc-044-twr",
         "uc-045-reportpack", "uc-046-attrib", "uc-047-icpack", "uc-048-policyupd"]


def signals(card: str) -> dict[str, dict]:
    out = {}
    for p in sorted((PACK / "signals").glob(f"{card}-*.yaml")):
        s = yaml.safe_load(p.read_text(encoding="utf-8"))
        s["trigger"] = {("on" if k is True else k): v for k, v in s["trigger"].items()}  # YAML 1.1 reads a bare `on` as true
        out[s["id"]] = s
    return out


def fire(sig: dict, doc: dict, rules: dict | None = None):
    """Evaluate a signal the way the engine does (v3_signals.evaluate): every declared field must be there, then the lets, the condition, the message.
    -> the message when it fires, False when it does not, None when an input is missing (skipped, never guessed)."""
    from sh_platform.v1 import v3_signals as SG
    funcs = {"abs": abs, "min": min, "max": max, "round": round, "last4": SG.last4, "num": SG.num}
    declared = next(i["fields"] for i in sig["inputs"] if i["kind"] == "doc")
    ns = {"doc": {}}
    for f in declared:
        if doc.get(f) in (None, ""):
            return None
        ns["doc"][f] = SG._dec(doc[f])
    ns.update(rules or {})
    try:
        for n, e in (sig.get("let") or {}).items():
            ns[n] = SG.calc(e, ns, funcs)
        hit = bool(SG.calc(sig["condition"], ns, funcs))
    except SG.Missing:
        return None
    return SG.render(sig["message"], ns) if hit else False


def test_every_card_has_three_signals_one_of_them_an_exception_with_a_hold_and_a_role_that_exists():
    from sh_platform.v1 import v3_signals as SG
    role_ids = {r["id"] for r in yaml.safe_load((PACK / "role_builds.yaml").read_text(encoding="utf-8"))["roles"]}
    roles = {r: r for r in role_ids}
    seen = set()
    for card in CARDS:
        sg = signals(card)
        assert len(sg) == 3, (card, list(sg))
        for sid, s in sg.items():
            assert SG.check(s, roles) is None, (sid, SG.check(s, roles))
            assert s["id"].startswith(card + "-") and s.get("suggestion") is True and s["route_to"] in role_ids and s["id"] not in seen, sid
            seen.add(s["id"])
            declared = set(next(i["fields"] for i in s["inputs"] if i["kind"] == "doc"))
            used = set(re.findall(r"doc\.([A-Za-z_]\w*)", " ".join([s["condition"], s["message"], s["subject"], str((s.get("hold") or {}).get("scope", ""))])))
            assert used == declared, (sid, used ^ declared)  # a name the condition or message uses is declared, and nothing declared is unused
        exc = [s for s in sg.values() if s["severity"] == "problem"]
        assert len(exc) == 1 and exc[0].get("hold", {}).get("blocks") and exc[0]["hold"].get("until"), card  # one exception signal per card, and it holds an action
        assert all(not s.get("hold") for s in sg.values() if s is not exc[0]), card


def test_signals_of_uc5_limits_fire_on_the_office_data(con, built):
    sg = signals("uc5-limits")
    sn = rows(con, built["single_name_breach"], "true", "holding")
    msgs = [m for m in (fire(sg["uc5-limits-breach"], r) for r in sn) if m]
    assert any("Harborview Semiconductor is 5.75 percent of Evergreen Liquid Book (2,368,000 of 41,200,000) against the 5 percent single-name limit" in m for m in msgs) and len(msgs) == 7
    assert [m for m in (fire(sg["uc5-limits-no-class"], r) for r in sn) if m][0].startswith("Pacific Rim Convertible 2029 note (1,122,000 in Evergreen Liquid Book) has no asset class")
    assert not [m for m in (fire(sg["uc5-limits-possibly-stale"], r) for r in sn) if m]  # every value is current at 30 September
    stale = fire(sg["uc5-limits-possibly-stale"], {**[r for r in sn if r["holding"] == "Harborview Semiconductor"][0], "status": "possibly stale", "value_age_days": 121})
    assert stale and "121 days old" in stale


def test_signals_of_the_morning_note_and_uc1(con, built):
    sg = signals("morning-networth-note")
    r = rows(con, built["report_total_fx"])[0]
    assert "USD 57,600,000 (SGD 77,253,120 at 1.3412) in the report of 2026-09-30, issued on 2026-10-03, status estimated" in fire(sg["morning-networth-note-today"], {**r, "days_since_issue": 10})
    assert fire(sg["morning-networth-note-report-old"], {**r, "days_since_issue": 10}) is False and "40 days ago" in fire(sg["morning-networth-note-report-old"], {**r, "days_since_issue": 40})
    call = {"entity": "Evergreen Holdings", "currency": "USD", "call_amount": 1250000, "due_in_days": 7, "shortfall": 270000}
    assert "short USD 270,000" in fire(sg["morning-networth-note-call-uncovered"], call) and fire(sg["morning-networth-note-call-uncovered"], {**call, "due_in_days": 10}) is False
    sg = signals("uc1-why-nav-moved")
    ch = rows(con, built["change_since_report"])[0]
    mat = {"mat": {"amount": 100000, "pct": 0.5}}
    assert "10,000,000 in the issued report and is now 10,180,000: a change of 180,000" in fire(sg["uc1-why-nav-moved-amended"], ch)
    assert fire(sg["uc1-why-nav-moved-bridge"], ch) and "over the materiality line of 100,000 or 0.5%" in fire(sg["uc1-why-nav-moved-over-materiality"], ch, mat)
    assert fire(sg["uc1-why-nav-moved-over-materiality"], {**ch, "change_since_report": 50000}, mat) is False


def test_signals_of_the_tie_out_and_the_audit_sample(con, built):
    sg = signals("uc-023-tieout")
    tt = rows(con, built["report_tie_out"], "true", "report_date, line")
    differs = [m for m in (fire(sg["uc-023-tieout-differs"], r) for r in tt) if m]
    assert len(differs) == 3 and "In the 2026-05 report Evergreen Growth I is 8,976,477 and the administrator holds 9,087,343: a difference of -110,866" in differs[1]
    assert len([m for m in (fire(sg["uc-023-tieout-unexplained"], r) for r in tt) if m]) == 3
    ties = [m for m in (fire(sg["uc-023-tieout-ties"], r) for r in tt) if m]
    assert len(ties) == 9 and all("All 3 lines" in m for m in ties)  # three months of three lines (the engine fires once per month: the subject is the month)
    assert fire(sg["uc-023-tieout-unexplained"], {**[r for r in tt if r["difference"] != 0][1], "explained_by_amendment": 1}) is False  # an amendment explains it
    sg = signals("uc-024-audit")
    tc = rows(con, built["trace_coverage"], "true", "item")
    assert len([m for m in (fire(sg["uc-024-audit-traced"], r) for r in tc) if m]) == 7 and fire(sg["uc-024-audit-pack-ready"], tc[0]) is False
    un = [m for m in (fire(sg["uc-024-audit-untraced"], r) for r in tc) if m]
    assert len(un) == 1 and "Item 8" in un[0] and "1 of 8 items are untraced; the traced share is 87.50%" in un[0]


def test_signals_of_drift_and_the_breach_note(con, built):
    sg = signals("uc-032-driftrep")
    dr = rows(con, built["drift_vs_target"])
    assert not [m for m in (fire(sg["uc-032-driftrep-outside"], {**r, "was_outside": 0}) for r in dr) if m]  # nothing is outside at 30 September
    over = {**[r for r in dr if r["rule_code"] == "limit_listed_equity_liquid"][0], "allocation_pct": 91, "status": "outside band", "outside_band": True, "was_outside": 0}
    assert "91.00 percent of Evergreen Liquid Book against a band of 60 to 90 percent" in fire(sg["uc-032-driftrep-outside"], over)
    assert fire(sg["uc-032-driftrep-back-inside"], {**dr[0], "was_outside": 1, "outside_band": 0}) and fire(sg["uc-032-driftrep-back-inside"], {**dr[0], "was_outside": 0, "outside_band": 0}) is False
    assert fire(sg["uc-032-driftrep-stale"], {**over, "status": "outside band, possibly stale", "oldest_value_days": 183}) and fire(sg["uc-032-driftrep-stale"], over) is False
    sg = signals("uc-033-breachesc")
    sn = rows(con, built["single_name_breach"], "holding = 'Harborview Semiconductor'")[0]
    assert "Selling 324,210.53 brings it back to the limit" in fire(sg["uc-033-breachesc-cost-to-cure"], sn)
    waiver = {"committee": "Investment committee", "decided_on": "2026-10-20", "title": "Harborview trim", "subject_id": "Harborview Semiconductor", "outcome": "waiver", "expires_on": "2026-12-31", "recorded_by": "Raj Patel"}
    assert "set it to expire on 2026-12-31" in fire(sg["uc-033-breachesc-waiver-recorded"], waiver) and fire(sg["uc-033-breachesc-waiver-recorded"], {**waiver, "outcome": "cured"}) is False
    old = {"holding": "Harborview Semiconductor", "expires_on": "2026-12-31", "days_past_expiry": 3, "breach_open": 1}
    assert fire(sg["uc-033-breachesc-waiver-expired"], old) and fire(sg["uc-033-breachesc-waiver-expired"], {**old, "breach_open": 0}) is False


def test_signals_of_performance_attribution_and_the_report_pack(con, built):
    sg = signals("uc4-vs-benchmark")
    tw = rows(con, built["time_weighted_return"], "portfolio = 'Evergreen Liquid Book'", "period_kind, to_date")
    sept = [r for r in tw if r["period_label"] == "2026-09"][0]
    assert "returned 4.1667 percent in 2026-09 against 2.7072 percent for the EG-OFFICE-BLEND benchmark: 1.4595 points active" in fire(sg["uc4-vs-benchmark-month"], sept)
    assert fire(sg["uc4-vs-benchmark-index-missing"], sept) is None  # no gap in the data: benchmark_gap is empty, so the signal is skipped, never guessed
    gap = {**sept, "benchmark_return_pct": None, "active_pct": None, "months_without_benchmark": 1, "benchmark_gap": "IDX:GLBD:SYN"}
    assert "IDX:GLBD:SYN has no level for the month ending 2026-09-30" in fire(sg["uc4-vs-benchmark-index-missing"], gap) and fire(sg["uc4-vs-benchmark-month"], gap) is None
    assert "version 2" in fire(sg["uc4-vs-benchmark-definition-changed"], {"benchmark_code": "EG-OFFICE-BLEND", "version_number": 2, "previous_version": 1, "approval_status": "in_progress", "effective_from": "2026-10-01"})
    sg = signals("uc-044-twr")
    q = [m for m in (fire(sg["uc-044-twr-periods"], r) for r in tw) if m]
    assert len(q) == 2 and "returned 3.5148 percent in 2026Q3 against 5.0052 percent" in q[1]
    assert len([m for m in (fire(sg["uc-044-twr-flow-timing"], r) for r in tw) if m]) == 5  # April to August had a flow; September had none
    q3 = [r for r in tw if r["period_label"] == "2026Q3"][0]
    br = {**q3, "stated_return_pct": 3.51}  # the linked figure is the recipe's own column; the stated one is what the report being checked says
    assert fire(sg["uc-044-twr-no-link"], br) and fire(sg["uc-044-twr-no-link"], {**br, "stated_return_pct": 3.5148}) is False
    sg = signals("uc-045-reportpack")
    rr = rows(con, built["report_ready"])
    assert len([m for m in (fire(sg["uc-045-reportpack-drafted"], r) for r in rr) if m]) == 1 and not [m for m in (fire(sg["uc-045-reportpack-all-green"], r) for r in rr) if m]
    assert not [m for m in (fire(sg["uc-045-reportpack-no-source"], r) for r in rr) if m]
    assert fire(sg["uc-045-reportpack-all-green"], {**rr[0], "section": "Total", "share_ready_pct": 100})
    assert "has no source figure" in fire(sg["uc-045-reportpack-no-source"], {**rr[0], "source_value": None, "has_source": 0})
    sg = signals("uc-046-attrib")
    at = rows(con, built["attribution"], "true", "month_end, asset_class")
    assert len([m for m in (fire(sg["uc-046-attrib-effects"], r) for r in at) if m]) == 6 * 3 and not [m for m in (fire(sg["uc-046-attrib-does-not-add"], r) for r in at) if m]  # an exception signal: silent on good data
    assert fire(sg["uc-046-attrib-does-not-add"], {**at[0], "adds_up": 0, "sum_of_effects_pct": 1.0, "active_pct": 1.2})


def test_signals_of_the_committee_pack_and_the_policy_update(con, built):
    sg = signals("uc-047-icpack")
    pk = rows(con, built["pack_items"])
    asm = [m for m in (fire(sg["uc-047-icpack-assembled"], r) for r in pk) if m]
    assert len(asm) == 1 and "The committee pack as at 2026-09-30 has five sections" in asm[0]
    sec = {"section": "Breaches", "owner_role": "risk_officer", "status": "7 real breach(es) open", "done": 0, "days_to_committee": 2}
    assert "owner is the risk_officer" in fire(sg["uc-047-icpack-open-late"], sec) and fire(sg["uc-047-icpack-open-late"], {**sec, "done": 1}) is False
    assert fire(sg["uc-047-icpack-section-done"], {**sec, "done": 1}) and fire(sg["uc-047-icpack-section-done"], sec) is False
    sg = signals("uc-048-policyupd")
    lc = rows(con, built["limit_changes"], "true", "version_number, month_end")
    rev = [m for m in (fire(sg["uc-048-policyupd-review"], r) for r in lc) if m]
    assert len(rev) == 1 and "limit of 6 percent for Evergreen Liquid Book would have been breached in 1 of 6 month ends" in rev[0]
    assert not [m for m in (fire(sg["uc-048-policyupd-leaves-breach"], r) for r in lc) if m]  # at 6 percent the September name (5.75) is inside
    tight = {**[r for r in lc if r["version_number"] == 2 and r["is_latest"] == 1][0], "limit_pct": 5.5, "over_limit": 1}
    assert "Harborview Semiconductor (2,368,000, 5.75 percent) would still be in breach" in fire(sg["uc-048-policyupd-leaves-breach"], tight)
    assert fire(sg["uc-048-policyupd-approved"], {"rule_code": "limit_single_name_liquid", "version_number": 2, "limit_pct": 6, "decided_on": "2026-10-20", "effective_from": "2026-10-21", "outcome": "approved"})


# ---- the persona builds -----------------------------------------------------------------------------------------------------------------------

def test_persona_builds_exist_once_with_the_keys_files_and_numbers():
    base = yaml.safe_load((PACK / "role_builds.yaml").read_text(encoding="utf-8"))["roles"]
    mine = yaml.safe_load((PACK / "builds" / "report.yaml").read_text(encoding="utf-8"))["roles"]
    old = {b["id"] for r in base for b in r["builds"]}
    new = {b["id"]: (r["id"], b) for r in mine for b in r["builds"]}
    assert set(new) == {"uc-023-tieout", "uc-024-audit", "uc-032-driftrep", "uc-033-breachesc", "uc-044-twr", "uc-045-reportpack", "uc-046-attrib", "uc-047-icpack", "uc-048-policyupd"}
    assert not set(new) & old  # a card that already has a role_builds build keeps it, nothing is duplicated
    assert {"uc5-limits", "morning-networth-note", "uc1-why-nav-moved", "uc4-vs-benchmark"} <= old
    owners = {c["id"]: c["owner"] for c in yaml.safe_load((PACK / "use_cases_100.yaml").read_text(encoding="utf-8"))["use_cases"]}
    for cid, (role, b) in new.items():
        assert owners[cid] == role and b["say"].strip() and b["check"]["ask"].endswith("?") and b["check"]["expect"] and b["skills"] and b["dashboard"]
        assert b["files"] == ["office6m:" + f for f in KEYS[cid]["input_files"]]
        assert all(glob.glob(str(OFFICE / f[len("office6m:"):])) for f in b["files"]), cid  # every file pattern finds files
    for cid, expect in (("uc-032-driftrep", ["80.61", "54.55", "26.73"]), ("uc-033-breachesc", ["324,210.53"]), ("uc-044-twr", ["0.2958", "3.5148", "3.821", "5.7563"]), ("uc-024-audit", ["87.5"])):
        assert set(expect) <= set(new[cid][1]["check"]["expect"])


def test_every_top_50_report_card_has_signals_and_a_build():
    cards = [c for c in yaml.safe_load((PACK / "use_cases_100.yaml").read_text(encoding="utf-8"))["use_cases"] if c["domain"] == "Report and govern" and c.get("priority", 999) <= 50]
    assert sorted(c["id"] for c in cards) == sorted(CARDS)
    base = {b["id"] for r in yaml.safe_load((PACK / "role_builds.yaml").read_text(encoding="utf-8"))["roles"] for b in r["builds"]}
    mine = {b["id"] for r in yaml.safe_load((PACK / "builds" / "report.yaml").read_text(encoding="utf-8"))["roles"] for b in r["builds"]}
    for c in cards:
        assert c["id"] in base | mine and len(signals(c["id"])) == len(c["signals"]) == 3, c["id"]


# ---- the return recipe is right for every portfolio, not only the keyed one ----------------------------------------------------------------

def test_a_portfolio_with_carried_forward_or_estimated_values_has_no_return_and_says_why(con, built):
    tw = built["time_weighted_return"]
    others = rows(con, tw, "portfolio <> 'Evergreen Liquid Book'")
    assert others and all(r["portfolio_return_pct"] is None and r["months_not_usable"] == r["months"] for r in others)  # no 0.0000 and no start-value gap
    assert {r["data_note"] for r in others} == {"a value is carried forward or estimated"}
    assert all(r["benchmark_linked"] is False and r["benchmark_return_pct"] is None for r in others)  # no signed link to a benchmark: no benchmark return
    liquid = rows(con, tw, "portfolio = 'Evergreen Liquid Book' and period_kind = 'month'")
    assert all(r["benchmark_linked"] and r["months_not_usable"] == 0 and r["data_note"] is None for r in liquid)


def test_a_holding_change_with_no_trade_on_record_is_not_a_return(con):
    """A holding is missing at the end of September and no trade explains it (there are none in September): that month has no return, not a start-value gap."""
    iid = con.execute("select id from instruments where canonical_name = 'Harborview Semiconductor'").fetchone()[0]
    con.execute("create temp table held_snap as select * from position_snapshots where instrument_id = ? and position_date = '2026-09-30'", [iid])
    con.execute("delete from position_snapshots where instrument_id = ? and position_date = '2026-09-30'", [iid])
    try:
        r = rows(con, build(con, "time_weighted_return"), "portfolio = 'Evergreen Liquid Book' and period_label = '2026-09'")[0]
        assert r["portfolio_return_pct"] is None and r["data_note"] == "the holdings changed with no trade on record"
    finally:
        con.execute("insert into position_snapshots select * from held_snap")
        con.execute("drop table held_snap")
        build(con, "time_weighted_return")


def test_without_a_signed_benchmark_link_there_is_no_benchmark_return(con):
    con.execute("update rule_versions set approval_status = 'proposed' where definition_text = 'benchmark_link'")
    try:
        r = rows(con, build(con, "time_weighted_return"), "portfolio = 'Evergreen Liquid Book' and period_label = '2026-09'")[0]
        assert num(r["portfolio_return_pct"]) == 4.1667 and r["benchmark_return_pct"] is None and r["benchmark_linked"] is False
    finally:
        con.execute("update rule_versions set approval_status = 'approved' where definition_text = 'benchmark_link'")
        build(con, "time_weighted_return")


def test_income_is_shown_beside_the_return_not_inside_it(con, built):
    r = rows(con, built["time_weighted_return"], "portfolio = 'Evergreen Liquid Book' and period_label = '2026-04'")[0]
    assert num(r["income"]) > 100000 and num(r["portfolio_return_pct"]) == -0.2944  # the key's method (income is not in the return); the amount is visible


def test_a_trace_matches_a_line_with_a_note_in_brackets_when_the_figure_is_the_same(con):
    """The planted item 8 reads 'Evergreen Growth I (as sampled)'; a sample item with that note and the report's own figure still traces."""
    f = con.execute("select id from extracted_fields where field_name = 'line' and value_text like '%(as sampled)' ").fetchone()[0]
    fig = con.execute("select value_numeric from extracted_fields where document_id = (select document_id from extracted_fields where id = ?) and line_number = 8 and field_name = 'figure'", [f]).fetchone()
    doc = con.execute("select document_id, line_number from extracted_fields where id = ?", [f]).fetchone()
    real = con.execute("select f.value_numeric from extracted_fields f join documents d on d.id = f.document_id where d.document_type = 'investment_report' and d.document_date = '2026-07-31' and f.field_name = 'Evergreen Growth I'").fetchone()[0]
    con.execute("update extracted_fields set value_numeric = ? where document_id = ? and line_number = ? and field_name = 'figure'", [real, doc[0], doc[1]])
    try:
        r = rows(con, build(con, "trace_coverage"), "item = 8")[0]
        assert r["traced"] is True and r["share_traced_pct"] == 100
    finally:
        con.execute("update extracted_fields set value_numeric = ? where document_id = ? and line_number = ? and field_name = 'figure'", [fig[0], doc[0], doc[1]])
        build(con, "trace_coverage")
