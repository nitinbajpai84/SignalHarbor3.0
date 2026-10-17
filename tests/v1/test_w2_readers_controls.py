"""Readers of the controls group (v3_pack_readers_controls): expense ledger and budget, payroll drafts, SOPs, access register, KYC register, onboarding file, bank
forms, family access, glossary, break register, load log, inbox and new feed. Real readers on the six-month office; expected numbers come from answer_keys.yaml or
from the files themselves, never from this file. Synthetic data only."""
from __future__ import annotations

import importlib.util
import json
from datetime import date
from decimal import Decimal
from pathlib import Path

import openpyxl

from sh_platform.v1 import v3_pack_readers as PK
from sh_platform.v1 import v3_pack_readers_controls as C

HERE = Path(__file__).parent


def _load(name: str):
    spec = importlib.util.spec_from_file_location(name, HERE / f"{name}.py")
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


W2 = _load("test_w2_loader")
W3 = _load("test_w3_intake")
key, OFFICE = W2.key, W2.OFFICE
office = W2.office  # the module fixture: the office loaded through the real readers
TODAY = date(2026, 10, 13)

KINDS = ["expense_ledger", "office_budget", "payroll_drafts", "sops", "access_register", "kyc_register", "onboarding", "bank_forms", "family_access", "glossary", "break_register",
         "load_log", "inbox", "new_feed"]
RECOGNISED = {  # one file of each kind that content alone tells
    "expense_ledger": "2026-10/expense_ledger/expense_ledger_2026-04_to_2026-09.xlsx", "office_budget": "2026-04/office_budget/office_budget_2026.xlsx",
    "payroll_drafts": "2026-10/payroll_drafts/payroll_run_2026-10_draft.xlsx", "sops": "2026-10/sops/sop_payroll-release.md", "access_register": "2026-10/access_register/access_register_2026-10-13.csv",
    "kyc_register": "2026-10/kyc_register/kyc_register_2026-10-13.csv", "onboarding": "2026-10/onboarding/onboarding_checklist_Harborlight.csv",
    "bank_forms": "2026-10/bank_forms/account_opening_form_TCB_TRUST_EUR.pdf", "family_access": "2026-10/family_access/family_view_grants_2026-10-13.csv",
    "glossary": "2026-10/glossary/office_words_2026-10-01.csv", "break_register": "2026-10/break_register/break_register_2026-10-13.csv", "load_log": "2026-10/load_log/load_log_2026-10-13.csv"}
ALSO = {"kyc_pack_Evergreen_Holdings_2026.pdf": ("2026-10/kyc_register/kyc_pack_Evergreen_Holdings_2026.pdf", "kyc_register"),
        "kyc_pack_Harborlight_Capital_Partners.pdf": ("2026-10/onboarding/kyc_pack_Harborlight_Capital_Partners.pdf", "onboarding"),
        "subscription_document_Harborlight_Ventures_IV_draft.pdf": ("2026-10/onboarding/subscription_document_Harborlight_Ventures_IV_draft.pdf", "onboarding"),
        "signing_policy_2026.pdf": ("2026-10/bank_forms/signing_policy_2026.pdf", "bank_forms"), "family_access_policy_2026.pdf": ("2026-10/family_access/family_access_policy_2026.pdf", "family_access"),
        "word_review_2026-10-08.md": ("2026-10/glossary/word_review_2026-10-08.md", "glossary")}


def fields(office, dtype: str, name: str) -> list[dict]:
    return office.rows("select f.* from extracted_fields f join documents d on d.id = f.document_id where d.document_type = ? and f.field_name = ? order by f.line_number", dtype, name)


def bdays(a: date, b: date) -> int:
    from datetime import timedelta
    return sum(1 for i in range(1, (b - a).days + 1) if (a + timedelta(days=i)).weekday() < 5)


# ---- recognition and the report ---------------------------------------------------------------------------------------------------------

def test_each_kind_is_recognised_from_its_content_not_its_name():
    for kind, rel in RECOGNISED.items():
        assert PK.kind_of("renamed" + Path(rel).suffix, (OFFICE / rel).read_bytes()) == (kind, "renamed" + Path(rel).suffix), kind
    for name, (rel, kind) in ALSO.items():
        assert PK.kind_of("x" + Path(rel).suffix, (OFFICE / rel).read_bytes())[0] == kind, name
    assert PK.kind_of("fund_flows.csv", W3.GENERIC) is None  # a table no reader knows stays on the generic path
    assert PK.kind_of("2026-10-02_joiner_leaver.eml", b"From: a\nSubject: b\n\nhi") is None  # mail stays a document
    # the older kinds are still themselves
    assert PK.kind_of("p.xlsx", (OFFICE / "2026-09/payroll/payroll_2026-09.xlsx").read_bytes())[0] == "payroll"


def test_every_kind_is_read_and_says_what_it_did(office):
    for kind in KINDS:
        r = office.report[kind]
        assert not r["not_loaded"] and r["files"] == len([p for p in OFFICE.glob(f"20??-??/{kind}/*") if p.is_file() and (p.suffix != ".eml" or kind != "inbox")]), (kind, r)
        assert r["rows"] > 0, kind
    assert all(office.report[k]["files"] for k in KINDS)


# ---- expenses, budget, approver limits (cards 73, 89, 90) -------------------------------------------------------------------------------------------

def test_expense_ledger_rows_equal_the_file_and_the_key(office):
    k = key("uc-089-expapp")
    rows = office.rows("select * from office_expenses")
    assert len(rows) == k["payments"] == 151 and all(r["source_record_id"] and r["source_document_id"] and r["approver_user_id"] for r in rows)
    by = {}
    for r in rows:
        by[r["approver_name"]] = by.get(r["approver_name"], 0) + 1
    assert by == k["approvals_by_person"]
    wb = openpyxl.load_workbook(next(OFFICE.glob("*/expense_ledger/*.xlsx")), data_only=True)
    for ccy in ("SGD", "USD"):  # nothing lost or changed on the way: totals per currency equal the file's
        want = sum(Decimal(str(r[5])) for r in wb["Ledger"].iter_rows(min_row=2, values_only=True) if r[4] == ccy)
        assert sum(r["amount"] for r in rows if r["currency"] == ccy) == want
    assert {r["legal_entity_id"] is not None for r in rows} == {True}


def test_payments_over_the_approvers_limit_are_the_two_of_the_key(office):
    k = key("uc-089-expapp")
    flagged = office.report["expense_ledger"]["flagged"]
    assert sorted(f.split(":")[0] for f in flagged if "whose limit" in f) == sorted(o["reference"] for o in k["over_limit"])
    rules = office.rows("select b.rule_code, v.parameters, v.approval_status, v.approved_by from business_rules b join rule_versions v on v.business_rule_id = b.id where b.rule_code like 'approver_limit_%' and v.version_number = 1")
    assert {r["rule_code"]: json.loads(r["parameters"])["limit"] for r in rules} == {f"approver_limit_{C.slug(n)}": float(v) for n, v in k["limits"].items()}
    assert {r["approval_status"] for r in rules} == {"proposed"} and all(r["approved_by"] is None for r in rules)  # the reader writes them unsigned (v1)
    # the office signed them as v2 beside the reader's v1 (A48): the reader never touches a signed rule
    assert office.one("select count(*) from business_rules b join rule_versions v on v.business_rule_id = b.id where b.rule_code like 'approver_limit_%' "
                      "and v.version_number = 2 and v.approval_status = 'approved'") == len(rules)


def test_the_budget_is_written_as_proposed_rules_that_give_the_keys_half_year(office):
    k = key("uc-090-offcost")["by_category"]
    rules = {r["rule_code"]: r for r in office.rows("select b.rule_code, v.parameters, v.approval_status from business_rules b join rule_versions v on v.business_rule_id = b.id where b.rule_code like 'office_budget_%' and v.version_number = 1")}
    assert len(rules) == len(k) == 13
    for cat, v in k.items():
        r = rules[f"office_budget_{C.slug(cat)}"]
        assert json.loads(r["parameters"])["annual"] / 2 == v["budget_half_year"] and r["approval_status"] == "proposed", cat
    spent = {}
    for r in office.rows("select category, currency, amount from office_expenses where currency = 'USD'"):
        spent[r["category"]] = spent.get(r["category"], 0) + r["amount"]
    assert spent["Grants"] == Decimal(175000) == Decimal(str(k["Grants"]["actual"]))  # Grants are all in USD: the ledger total is the key's actual


# ---- payroll (card 81) -----------------------------------------------------------------------------------------------------------------

def test_payroll_draft_and_the_paid_runs_give_the_keys_figures_and_the_two_flags(office):
    k = key("uc-081-payroll")
    d = office.one("select sum(net_pay) from staff_payroll where run_status = 'draft'"), office.one("select sum(employer_contribution) from staff_payroll where run_status = 'draft'")
    assert (float(d[0]), float(d[1])) == (k["draft_net"], k["draft_employer_cpf"])
    sept = office.rows("select * from staff_payroll where pay_period = date '2026-09-01'")
    assert float(sum(r["net_pay"] for r in sept)) == k["september_net"] and all(r["run_status"] == "paid" for r in sept)
    assert office.one("select count(distinct pay_period) from staff_payroll where run_status = 'paid'") == 6
    flags = office.report["payroll_drafts"]["flagged"]
    assert any("Grace Ong gross changed" in f and "+20.0 percent" in f for f in flags) and any("Jun Tan" in f and "not on the staff register" in f for f in flags) and len(flags) == len(k["flags"])
    jun = office.rows("select * from staff_payroll where employee_name = 'Jun Tan'")
    assert len(jun) == 1 and jun[0]["employment_kind"] == "contractor" and jun[0]["run_status"] == "draft"
    # pay is personal data: the source document is in the personal access group
    assert office.one("select count(*) from staff_payroll s join documents d on d.id = s.source_document_id join access_groups g on g.id = d.access_group_id where g.group_name = 'personal'") == 64


# ---- SOPs (cards 95, 96) -------------------------------------------------------------------------------------------------------------------

def test_sops_give_owner_review_date_runs_and_steps_that_age_to_the_key(office):
    k = {s["sop"]: s for s in key("uc-095-sopreview")["sops"]}
    docs = office.rows("select id, title from documents where document_type = 'sop'")
    assert {d["title"] for d in docs} == set(k)
    for d in docs:
        f = {r["field_name"]: r for r in office.rows("select * from extracted_fields where document_id = ?", d["id"]) if r["field_name"] in ("owner", "last_reviewed_on", "runs_per_month")}
        age = (TODAY - date.fromisoformat(f["last_reviewed_on"]["value_text"])).days
        assert (f["owner"]["value_text"], f["last_reviewed_on"]["value_text"], age, age > 365) == (k[d["title"]]["owner"], k[d["title"]]["last_reviewed"], k[d["title"]]["age_days"], k[d["title"]]["due"])
    steps = office.rows("select f.value_text, f.value_numeric from extracted_fields f join documents d on d.id = f.document_id where d.title = 'Bank statement chase' and f.field_name = 'step' order by f.line_number")
    assert [(s["value_text"], float(s["value_numeric"])) for s in steps] == [("List expected statements", 20.0), ("Chase late or missing ones", 40.0), ("Log the receipt", 10.0)]
    r = key("uc-096-automation")["ranking"][0]
    assert sum(float(s["value_numeric"]) for s in steps) * r["runs_per_month"] == r["manual_minutes_per_month"]  # 70 minutes x 4 runs: the key's 280


# ---- access (cards 91, 92, 100) -------------------------------------------------------------------------------------------------------------

def test_access_register_gives_the_unapproved_the_unused_and_the_leavers_grants(office):
    k91, k92 = key("uc-091-jml"), key("uc-092-accessrev")
    assert office.one("select count(*) from system_access_grants") == k92["grants_reviewed"] == 38
    none = office.rows("select person_name, system_name, access_level from system_access_grants where approved_by_name is null order by person_name")
    assert [(r["person_name"], r["system_name"], r["access_level"]) for r in none] == sorted((g["person"], g["system"], g["permission"]) for g in k91["granted_without_approver"])
    idle = office.rows("select person_name, system_name, last_used_on from system_access_grants where last_used_on < date '2026-10-13' - 90 and person_name <> 'J. Ng' order by last_used_on")
    assert [(r["person_name"], r["system_name"], str(r["last_used_on"]), (TODAY - r["last_used_on"]).days) for r in idle] == \
        [(g["person"], g["system"], g["last_used"], g["days_unused"]) for g in sorted(k92["not_used_in_90_days"], key=lambda g: g["last_used"])]
    leaver = office.rows("select system_name from system_access_grants where person_name = 'J. Ng' order by system_name")
    assert [r["system_name"] for r in leaver] == sorted(k91["leaver_access_to_remove"])
    assert [r["value_text"] for r in fields(office, "email", "leaver")] == [k91["leaver"]]
    joiner = {r["value_text"] for r in fields(office, "email", "joiner_access")}
    held = {r["system_name"] for r in office.rows("select system_name from system_access_grants where person_name = 'Wei Tan'")}
    assert sorted(held - joiner) == k91["joiner_access_not_requested"] and [r["value_text"] for r in fields(office, "email", "joiner")] == [k91["joiner"]]
    assert any("approved their own" in f for f in office.report["access_register"]["flagged"])  # a self-approved grant is said, not hidden


def test_family_access_over_grant_and_the_request(office):
    k = key("uc-100-famaccess")
    level = {"preview": "summary"}
    got = office.rows("select g.group_name, s.scope_value, s.access_level from access_group_scopes s join access_groups g on g.id = s.access_group_id where g.group_name like 'family:%'")
    assert {(r["group_name"][7:], r["scope_value"], level.get(r["access_level"], r["access_level"])) for r in got} == \
        {(m, e, lvl) for m, v in k["by_member"].items() for e, lvl in v["granted"].items() for e in [office.one("select entity_code from legal_entities where entity_name = ?", e)]}
    rules = {r["rule_code"]: json.loads(r["parameters"])["level"] for r in office.rows("select b.rule_code, v.parameters from business_rules b join rule_versions v on v.business_rule_id = b.id where b.rule_code like 'family_access_%'")}
    assert len(rules) == 12 and rules["family_access_wen_chen_eg_b"] == "none" and rules["family_access_lin_chen_eg_hold"] == "full"
    assert [f.split(" is granted")[0] for f in office.report["family_access"]["flagged"]] == k["over_grants"]
    req = {r["field_name"]: r["value_text"] for r in office.rows("select f.* from extracted_fields f join documents d on d.id = f.document_id where d.document_type = 'email' and f.field_name in ('requester', 'requested_entity', 'requested_level')")}
    assert (req["requester"], req["requested_entity"], req["requested_level"]) == (k["request"]["member"], "EG-HOLD", k["request"]["level"])


# ---- KYC, onboarding, bank forms (cards 93, 94, 97) -------------------------------------------------------------------------------------------

def test_kyc_register_gives_the_expiring_and_the_expired_and_holds_no_identity_numbers(office):
    k = key("uc-097-kycrefresh")
    rows = office.rows("select * from kyc_records where subject_kind <> 'manager'")
    assert len(rows) == k["documents"] == 12  # the pack's five lines are the register's: not written twice
    soon = sorted((str(r["expires_on"]), (r["expires_on"] - TODAY).days) for r in rows if TODAY <= r["expires_on"] <= date(2027, 1, 11))
    assert soon == sorted((e["expires"], e["days_left"]) for e in k["expiring_in_90_days"])
    gone = [r for r in rows if r["expires_on"] < TODAY]
    assert [(str(r["expires_on"]), (r["expires_on"] - TODAY).days, r["subject_name"], r["status"]) for r in gone] == [(e["expires"], e["days_left"], "Daniel Foo", "expired") for e in k["expired"]]
    cps = {r["value_text"] for r in office.rows("select f.value_text from extracted_fields f where f.field_name = 'counterparty'")}
    assert {"Harbour Bank (synthetic)", "Cedarpoint Ventures"} <= cps
    assert office.one("select count(*) from kyc_records k join documents d on d.id = k.source_document_id join access_groups g on g.id = d.access_group_id where g.group_name = 'personal'") == 16
    assert not [c for c in office.rows("select * from information_schema.columns where table_name = 'kyc_records'") if "number" in c["column_name"]]


def test_onboarding_checklist_counts_and_the_unsigned_subscription(office):
    k = key("uc-093-onbmgr")
    st = {r["line_number"]: r["value_text"] for r in fields(office, "onboarding_checklist", "checklist_status")}
    item = {r["line_number"]: r["value_text"] for r in fields(office, "onboarding_checklist", "checklist_item")}
    assert len(st) == k["items"] and sum(v == "complete" for v in st.values()) == k["complete"]
    assert [item[i] for i, v in st.items() if v != "complete"] == k["missing"]
    assert [r["value_text"] for r in fields(office, "subscription_document", "signature_status")] == ["not signed"]
    docs = office.rows("select document_kind, status from kyc_records where subject_kind = 'manager' order by document_kind")
    assert len(docs) == 4 and {d["status"] for d in docs} == {"provided"}


def test_bank_form_signatories_policy_and_the_pending_account(office):
    k = key("uc-094-onbbank")
    sig = [r["value_text"] for r in fields(office, "account_opening_form", "signatory")]
    assert sig == k["signatories"] and k["missing_signatory"] not in sig
    roles = [r["value_text"] for r in fields(office, "account_opening_form", "signatory_role")]
    assert not [r for r in roles if "Head of Finance" in r]  # the Head of Finance is not on the form
    pol = office.rows("select v.parameters, v.approval_status from business_rules b join rule_versions v on v.business_rule_id = b.id where b.rule_code = 'signing_policy' and v.version_number = 1")
    assert len(pol) == 1 and json.loads(pol[0]["parameters"]) == {"min_signatories": 2, "required_role": "Head of Finance"} and pol[0]["approval_status"] == "proposed"
    acct = office.rows("select * from bank_accounts where status = 'pending'")
    assert len(acct) == 1 and acct[0]["currency"] == "EUR" and acct[0]["legal_entity_id"] and acct[0]["source_document_id"]
    assert office.one("select count(*) from bank_accounts where status = 'pending'") == 1


# ---- glossary, breaks, load log, inbox, new feed (cards 98, 59, 61, 60, 99) ---------------------------------------------------------------------

def test_glossary_proposals_against_the_list(office):
    k = key("uc-098-wordlist")
    listed = {r["value_text"].lower() for r in fields(office, "glossary", "term")} | {r["value_text"].lower() for r in fields(office, "glossary", "alias")}
    terms = [r["value_text"] for r in fields(office, "glossary", "term")]
    assert len(terms) == k["list_size"]
    prop = [r["value_text"].lower() for r in office.rows("select f.* from extracted_fields f where f.field_name = 'proposed_term'")]
    assert len(prop) == k["proposals_made"] and sorted(set(prop)) == k["proposed"]
    assert sorted(t for t in set(prop) if t in listed) == k["already_in_list"] and sorted(t for t in set(prop) if t not in listed) == k["to_add"]


def test_break_register_gives_every_break_with_its_id_entity_and_age(office):
    k = key("uc-059-breakage")
    rows = office.rows("select e.external_ref, e.description, e.detected_at, e.status, l.entity_name from exceptions e join legal_entities l on l.id = e.legal_entity_id where e.exception_type = 'break' order by 1")
    want = {b["id"]: b for b in k["all_open"]}
    assert [r["external_ref"] for r in rows] == sorted(want) and {r["status"] for r in rows} == {"open"}
    for r in rows:
        b = want[r["external_ref"]]
        assert (r["entity_name"], r["description"].split(";")[0], bdays(r["detected_at"].date(), TODAY)) == (b["entity"], b["type"], b["age_business_days"])
    older = sorted(r["external_ref"] for r in rows if r["description"].startswith("cash break") and bdays(r["detected_at"].date(), TODAY) > 5)
    assert older == sorted(b["id"] for b in k["cash_breaks_over_5_days"])


def test_load_log_marks_the_batches_and_gives_the_keys_score_where_every_file_arrived(office):
    k = {s["source"]: s for s in key("uc-061-dqscore")["by_source"]}
    rows = office.rows("""select json_extract_string(b.validation_summary, '$.load_log_source') as source, count(*) loads,
                          sum(case when b.arrived_on_time and t.passed and not b.manual_fix then 1 else 0 end) passed,
                          sum(case when not t.passed then 1 else 0 end) not_added, sum(case when b.manual_fix then 1 else 0 end) fixes, sum(case when not b.arrived_on_time then 1 else 0 end) late
                          from source_batches b join batch_control_totals t on t.source_batch_id = b.id where json_extract_string(b.validation_summary, '$.load_log_source') is not null group by 1""")
    got = {r["source"]: r for r in rows}
    for src in ("administrator_cash", "administrator_nav", "custodian_transactions", "notices", "fx_confirmations", "invoices", "trading_export", "commitments", "custodian_positions", "investment_reports"):
        r, e = got[src], k[src]
        assert (r["loads"], r["passed"], r["not_added"], r["fixes"], r["late"]) == (e["loads"], e["passed"], e["did_not_add_up"], e["manual_fixes"], e["late_or_missing"]), src
        assert round(100 * r["passed"] / r["loads"], 1) == e["score_pct"], src
    # a statement that never arrived has no batch: it is said in the report, and the late ones are on their batches
    never = [f for f in office.report["load_log"]["flagged"] if "never arrived" in f]
    for src in ("bank_statements", "manager_statements"):
        miss = len([f for f in never if f.startswith(src + ":")])
        assert got[src]["late"] + miss == k[src]["late_or_missing"] and got[src]["loads"] + miss == k[src]["loads"], src
    assert office.one("select count(*) from source_batches where arrived_on_time is not null") == office.one("select count(*) from batch_control_totals")


def test_inbox_files_are_read_by_content(office):
    k = key("uc-060-docinbox")
    docs = {d["original_filename"]: d for d in office.rows("select * from documents where original_filename in ('scan_0412.pdf', 'account_extract.xlsx', 'INV-OIT-2610.pdf', 'IMG_2210.pdf', 'data_export.xlsx') and document_type in ('inbox_file', 'unrecognised')")}
    assert len(docs) == k["files"] == 5
    act = {d["original_filename"]: {r["field_name"]: r["value_text"] for r in office.rows("select * from extracted_fields where document_id = ?", d["id"])} for d in docs.values()}
    for r in k["recognised"]:
        assert act[r["file"]]["belongs_in"] == r["belongs_in"] and docs[r["file"]]["document_status"] == ("duplicate" if "duplicate" in r["action"] else "final"), r["file"]
    for r in k["unrecognised"]:
        assert docs[r["file"]]["document_type"] == "unrecognised" and act[r["file"]]["belongs_in"] == "unknown" and act[r["file"]]["inbox_action"].split()[0] == r["action"].split()[0], r["file"]


def test_new_feed_gets_a_proposed_mapping_with_a_confidence_mark_per_column_and_loads_no_rows(office):
    k = key("uc-099-newfeed")
    m = office.rows("select * from sheet_mappings where file_pattern like 'pinnacle%'")
    assert len(m) == 1 and m[0]["mapping_status"] == "proposed" and m[0]["approved_by"] is None
    cols = office.rows("select * from sheet_mapping_columns where sheet_mapping_id = ?", m[0]["id"])
    assert len(cols) == k["columns"] and sum(c["confidence"] == "high" for c in cols) == k["mapped_with_confidence"]
    assert sorted(c["source_column"] for c in cols if c["confidence"] == "doubt") == sorted(k["in_doubt"])
    assert round(100 * k["mapped_with_confidence"] / k["columns"]) == k["share_mapped_pct"]
    assert json.loads(office.one("select validation_summary from source_batches where sheet_mapping_id = ?", m[0]["id"]))["rows_in_sample"] == k["rows_in_sample"]
    assert office.one("select count(*) from position_snapshots s join source_records r on r.id = s.source_record_id where r.raw_payload_uri like '%pinnacle_custody%'") == 0  # no position from the feed until the mapping is approved


# ---- the way a row is written: append-only, superseding ------------------------------------------------------------------------------------

def test_a_row_is_written_once_and_a_change_supersedes_it():
    class Stub:
        def __init__(self):
            self.rows, self.n = {}, 0
            self.con = type("Con", (), {"execute": staticmethod(lambda *a, **k: (_ for _ in ()).throw(RuntimeError("no table")))})()

        def add(self, table, **kw):
            self.n += 1
            kw["id"] = f"id{self.n}"
            self.rows.setdefault(table, []).append(kw)
            return kw["id"]
    o = Stub()
    rec = dict(spent_on=date(2026, 4, 1), category="Premises", payee_name="Marina", legal_entity_id="e1", currency="SGD", amount=Decimal("18500.00"), reference="MOS-2604", approver_name="Grace Ong")
    first = C._put(o, "office_expenses", rec)
    assert first == "id1" and C._put(o, "office_expenses", {**rec, "amount": Decimal("18500")}) is None  # the same row, however it is written
    second = C._put(o, "office_expenses", {**rec, "amount": Decimal("19000")})
    assert o.rows["office_expenses"][1]["supersedes_office_expense_id"] == "id1" and second == "id2"
    assert C._put(o, "office_expenses", {**rec, "reference": "MOS-2605"}) == "id3" and "supersedes_office_expense_id" not in o.rows["office_expenses"][2]


def test_a_changed_ledger_from_a_flat_dropbox_supersedes_what_changed_and_flags_its_mess(tmp_path):
    plane, x = W3.flat_thread(tmp_path, [RECOGNISED["expense_ledger"], RECOGNISED["payroll_drafts"], RECOGNISED["access_register"]])
    PK.intake(x)
    assert {v["kind"] for v in x.th["pack"].values()} == {"expense_ledger", "payroll_drafts", "access_register"}
    assert PK.load(x) == (3, [])
    q = lambda sql: W3.rows(plane, sql)  # noqa: E731
    assert q("select count(*) from office_expenses")[0][0] == 151 and q("select count(*) from staff_payroll")[0][0] == 10 and q("select count(*) from system_access_grants")[0][0] == 38
    # a new ledger: one amount changed, one line twice, one line with no approver
    src = next(OFFICE.glob("*/expense_ledger/*.xlsx"))
    wb = openpyxl.load_workbook(src)
    ws = wb["Ledger"]
    ref = ws.cell(2, 8).value
    ws.cell(2, 6).value = 8500
    ws.append(list(c.value for c in ws[3]))  # the third line again
    ws.cell(4, 9).value = None
    out = tmp_path / "expense_ledger_2026-04_to_2026-10.xlsx"
    wb.save(out)
    data = out.read_bytes()
    assert PK.kind_of(out.name, data)[0] == "expense_ledger"
    W3.H.drop(plane, out.name, data)
    items = [(out.name, data, "expense_ledger", out.name)]
    report, aside, done = PK.read_files(plane, items, PK._standing(plane), "t-w3")
    assert aside == [] and done == [out.name]
    assert q("select count(*) from office_expenses")[0][0] == 153  # two changed lines (the amount, the lost approver) are one new row each; the other 149 are not written again; the line twice is counted once
    new = q(f"select amount, supersedes_office_expense_id is not null from office_expenses where reference = '{ref}'")
    assert sorted((float(a), s) for a, s in new) == [(8000.0, False), (8500.0, True)]
    flags = report["expense_ledger"]["flagged"]
    assert any("appears twice" in f for f in flags) and any("no approver" in f for f in flags)
