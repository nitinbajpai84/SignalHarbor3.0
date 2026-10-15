"""Administrator monthly exports (xlsx), Oct 2025 to Sep 2026. Deterministic, figures from docs/test/10uc_universe.json."""
import json, sys
from decimal import Decimal as D
from pathlib import Path
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
ROOT = Path(__file__).resolve().parents[2]
U = json.load(open(ROOT / "docs/test/10uc_universe.json"))
OUT = ROOT / "tests/v1/fixtures/10uc/admin_exports"
ADMIN = "Harbourline Fund Services (synthetic)"
RENAME_MONTH = "2026-03-31"          # planted: layout drift in this month only
COLS = ["Portfolio", "Admin Code", "Instrument", "Quantity", "Price", "Market Value", "Currency", "Valuation Status", "Valuation Date"]
RENAMED = {"Market Value": "Valuation Amount (USD)"}
PORT = {p["portfolio_code"]: p for p in U["portfolios"]}
FUND = {f["key"]: f for f in U["funds"]}
SEC = {s["key"]: s for s in U["securities"]}
NAV = {(r["portfolio"], r["nav_period_end"]): r for r in U["administrator_nav_monthly"]}
VAL = {(v["fund"], v["month_end"]): v for v in U["valuations_monthly"]}
MONTHS = list(U["liquid_book_month_end"])

def status(sid, month):
    """Administrator wording from the statement the value came from."""
    if sid is None: return "Not received"
    _, _, period, kind = sid.split("-", 3)[0], None, sid.split("-", 2)[2][:10], sid.rsplit("-", 1)[1]
    if period != month: return "Carried forward"
    return {"final": "Final", "estimate": "Estimate", "amended": "Amended"}[kind]

def rows_for(month):
    out = []
    for pc, p in PORT.items():
        if pc == "EG-LIQUID":
            for pos in U["liquid_book_month_end"][month]["positions"]:
                s = SEC[pos["security"]]
                out.append([pc, int(s["identifiers"]["administrator"]["administrator_code"]), s["canonical_name"], D(pos["quantity"]), D(pos["price"]), D(pos["value"]), "USD", "Priced", month])
        else:
            for k in p["holdings"]:
                v = VAL[k, month]; f = FUND[k]; val = v["administrator_value"]
                out.append([pc, int(f["identifiers"]["administrator"]["administrator_code"]), f["canonical_name"],
                            D(1) if val else None, D(val) if val else None, D(val) if val else None, f["currency"],
                            status(v["administrator_from"], month),
                            v["administrator_from"].split("-", 2)[2][:10] if val else None])  # period end of the statement the value came from
    return out

def build(month):
    wb = Workbook(); ws = wb.active; ws.title = "Positions"
    ws["A1"] = ADMIN; ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = "Month-end positions and valuations - Evergreen Office"
    ws["A3"] = f"Period end: {month}"
    nr = NAV["EG-LIQUID", month]["received_at"]
    ws["A4"] = f"Valued as at {nr.replace('T', ' ').replace(':00Z', ' UTC')} (2nd business day after period end). Currency: USD."
    cols = [RENAMED.get(c, c) if month == RENAME_MONTH else c for c in COLS]
    for i, c in enumerate(cols, 1):
        cell = ws.cell(6, i, c); cell.font = Font(bold=True); cell.fill = PatternFill("solid", fgColor="D9E1F2")
    r = 7
    for row in rows_for(month):
        for i, v in enumerate(row, 1):
            ws.cell(r, i, float(v) if isinstance(v, D) else v)
        ws.cell(r, 4).number_format = "#,##0.0000" if row[1] >= 5000 and False else "#,##0"
        ws.cell(r, 5).number_format = "#,##0.0000"; ws.cell(r, 6).number_format = "#,##0.00"; r += 1
    ws.cell(r + 1, 1, f"Rows: {r - 7}. Prices: equities per share, bonds per 1.00 nominal, private funds at reported holding value. Generated {ADMIN}, page 1 of 1.")
    ws.cell(r + 2, 1, "Synthetic data. Not an official statement.")
    for col, w in zip("ABCDEFGHI", (13, 11, 36, 12, 12, 18, 10, 17, 14)): ws.column_dimensions[col].width = w
    n = wb.create_sheet("Approved NAV")
    n["A1"] = ADMIN; n["A1"].font = Font(bold=True, size=14)
    n["A2"] = f"Approved NAV per portfolio - period end {month}"
    for i, c in enumerate(["Portfolio", "Portfolio Name", "NAV Period End", "NAV Amount", "Currency", "NAV Status", "Approved At (UTC)"], 1):
        n.cell(4, i, c).font = Font(bold=True)
    for j, pc in enumerate(PORT):
        x = NAV[pc, month]
        for i, v in enumerate([pc, PORT[pc]["portfolio_name"], month, float(D(x["nav_amount"])), x["currency"], "Approved", x["received_at"]], 1):
            n.cell(5 + j, i, v)
        n.cell(5 + j, 4).number_format = "#,##0.00"
    n.cell(10, 1, "NAV is the sum of investment holdings; cash is reported separately. Synthetic data.")
    for col, w in zip("ABCDEFG", (13, 32, 15, 18, 10, 12, 22)): n.column_dimensions[col].width = w
    path = OUT / f"administrator_export_{month[:7]}.xlsx"
    wb.save(path); return path

def verify(path, month):
    wb = load_workbook(path); ws = wb["Positions"]
    hdr = [c.value for c in ws[6] if c.value]
    want = [RENAMED.get(c, c) if month == RENAME_MONTH else c for c in COLS]
    assert hdr == want, (hdr, want)
    got = [[c.value for c in r[:9]] for r in ws.iter_rows(min_row=7, max_row=6 + len(rows_for(month)))]
    for g, e in zip(got, rows_for(month)):
        assert g[:3] == e[:3] and g[6:] == e[6:]
        for i in (3, 4, 5): assert (g[i] is None and e[i] is None) or D(str(g[i])) == e[i], (g, e)
        if g[3] is not None: assert D(str(g[3])) * D(str(g[4])) == D(str(g[5])), g
    for pc in PORT:
        s = sum((D(str(g[5])) for g in got if g[0] == pc and g[5] is not None), D(0))
        assert s == D(NAV[pc, month]["nav_amount"]), (pc, month, s)
    nv = {r[0].value: D(str(r[3].value)) for r in wb["Approved NAV"].iter_rows(min_row=5, max_row=8)}
    assert nv == {pc: D(NAV[pc, month]["nav_amount"]) for pc in PORT}
    return len(got)

if __name__ == "__main__":
    OUT.mkdir(parents=True, exist_ok=True)
    total = 0
    for m in MONTHS:
        total += verify(build(m), m)
    print(len(MONTHS), "files,", total, "position rows, verified")
