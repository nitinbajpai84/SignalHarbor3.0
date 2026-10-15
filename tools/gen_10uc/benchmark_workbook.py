"""Deterministic generator: benchmark specification workbook from docs/test/10uc_universe.json (synthetic)."""
import json, pathlib
from decimal import Decimal as D
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.properties import WorksheetProperties

ROOT = pathlib.Path(__file__).resolve().parents[2]
U = json.load(open(ROOT / "docs/test/10uc_universe.json"))
OUT = ROOT / "tests/v1/fixtures/10uc/benchmark_workbook"
FN = OUT / "EG-OFFICE-BLEND_Benchmark_Specification_2026-10.xlsx"
B = U["benchmark"]; NAMES = U["index_levels"]["series"]; ANCH = U["index_levels"]["month_end_anchors"]
HB = Font(bold=True, color="FFFFFF"); HF = PatternFill("solid", fgColor="1F3A5F")

def head(ws, title, sub):
    ws["A1"] = title; ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = sub; ws["A2"].font = Font(italic=True, color="666666")
    ws.oddFooter.left.text = "Evergreen Office - synthetic data"; ws.oddFooter.right.text = "Page &P of &N"

def table(ws, r, cols, rows):
    for j, c in enumerate(cols, 1):
        x = ws.cell(r, j, c); x.font = HB; x.fill = HF
    for i, row in enumerate(rows, r + 1):
        for j, v in enumerate(row, 1): ws.cell(i, j, v)
    for j in range(1, len(cols) + 1): ws.column_dimensions[chr(64 + j)].width = 34 if j in (1, 2, 6) else 18
    return r + len(rows)

def comp_sheet(wb, v):
    ws = wb.create_sheet(f"Components v{v['version_number']}" + (" (In Progress)" if v["approval_status"] == "in_progress" else ""))
    head(ws, f"{B['benchmark_name']} ({B['benchmark_code']}) - Version {v['version_number']}",
         f"Status: {'In Progress' if v['approval_status']=='in_progress' else 'Approved'}")
    ws["A3"], ws["B3"] = "Effective from", v["effective_from"]
    ws["A4"], ws["B4"] = "Approved by", v["approved_by"] or "(not approved)"
    ws["A5"], ws["B5"] = "Reporting currency", B["reporting_currency"]
    ws["A6"], ws["B6"] = "Calculation method", v["calculation_method"]
    rows = [[NAMES[c["market_series_key"]], c["market_series_key"], float(c["component_weight"]),
             int(c["spread_bps"]) if c["spread_bps"] else 0, c["treatment_note"]] for c in v["components"]]
    end = table(ws, 8, ["Component", "Series key", "Weight", "Spread (bps p.a.)", "Treatment note"], rows)
    ws.cell(end + 1, 1, "Total weight"); ws.cell(end + 1, 3, float(sum(D(c["component_weight"]) for c in v["components"])))
    ws.cell(end + 3, 1, "Formula note: component return = index monthly return + spread_bps / 12 / 100 (300 bps a year = 25 bps a month = 0.25 %). Blend = sum(weight x component return).")
    if v["approval_status"] == "in_progress":
        ws.cell(end + 4, 1, "IN PROGRESS - component sheet still being filled in. Proposed 1 Oct 2026 by Sofia Reyes (Risk). Not for certified reporting.").font = Font(bold=True, color="C00000")
        ws.sheet_properties.tabColor = "FFC000"
    return ws

def build():
    OUT.mkdir(parents=True, exist_ok=True)
    wb = Workbook(); ws = wb.active; ws.title = "Cover"
    head(ws, f"{B['benchmark_name']} - benchmark specification", "Evergreen Office, Risk and Finance. Synthetic data.")
    for i, (a, b) in enumerate([("Benchmark code", B["benchmark_code"]), ("Reporting currency", B["reporting_currency"]),
        ("Current approved version", "1 (effective 2026-01-01, approved by Daniel Foo)"),
        ("Next version", "2 - In Progress (effective 2026-10-01, proposed by Sofia Reyes)"),
        ("Private-markets spread", "300 bps a year over the Global Equity Index (synthetic)"),
        ("Liquid sleeve", "60% Global Equity Index (synthetic) / 40% Global Bond Index (synthetic)"),
        ("Sheets", "Components v1; Components v2 (In Progress); Index Levels; Sep 2026 Calc")], 4):
        ws.cell(i, 1, a).font = Font(bold=True); ws.cell(i, 2, b)
    ws.column_dimensions["A"].width = 30; ws.column_dimensions["B"].width = 80
    for v in B["versions"]: comp_sheet(wb, v)
    ws = wb.create_sheet("Index Levels"); head(ws, "Month-end index levels", "Synthetic series, base 1000.00 at 2025-09-30")
    ks = list(NAMES); rows = [[d] + [float(ANCH[k][d]) for k in ks] for d in ANCH[ks[0]]]
    table(ws, 4, ["Month end"] + [NAMES[k] for k in ks], rows)
    s = B["september_2026"]; ws = wb.create_sheet("Sep 2026 Calc"); head(ws, "September 2026 benchmark return (v1)", "Percent; v1 weights 0.42 / 0.28 / 0.30")
    table(ws, 4, ["Line", "Return %", "Note"], [
        ["Equity index return", float(s["equity_index_return_pct"]), "1074.56 / 1038.22 - 1"],
        ["Bond index return", float(s["bond_index_return_pct"]), "1041.22 / 1037.07 - 1"],
        ["Private-markets sleeve return", float(s["private_sleeve_return_pct"]), "equity + 0.25 (300 bps a year / 12)"],
        ["Blend v1 return", float(s["blend_v1_return_pct"]), "0.42 x equity + 0.28 x bond + 0.30 x private sleeve"],
        ["60/40 liquid sleeve return", float(s["sixty_forty_return_pct"]), "0.60 x equity + 0.40 x bond"]])
    wb.save(FN)

def verify():
    wb = load_workbook(FN)
    assert any("In Progress" in n for n in wb.sheetnames)
    for v in B["versions"]:
        ws = wb[[n for n in wb.sheetnames if n.startswith(f"Components v{v['version_number']}")][0]]
        for i, c in enumerate(v["components"], 9):
            assert str(ws.cell(i, 2).value) == c["market_series_key"] and D(str(ws.cell(i, 3).value)) == D(c["component_weight"])
            assert ws.cell(i, 4).value == (int(c["spread_bps"]) if c["spread_bps"] else 0)
        assert ws["B3"].value == v["effective_from"]
    ws = wb["Sep 2026 Calc"]; s = B["september_2026"]
    for r, k in zip(range(5, 10), ["equity_index_return_pct", "bond_index_return_pct", "private_sleeve_return_pct", "blend_v1_return_pct", "sixty_forty_return_pct"]):
        assert D(str(ws.cell(r, 2).value)) == D(s[k]), k
    e, b = (D(ANCH[k]["2026-09-30"]) / D(ANCH[k]["2026-08-31"]) - 1 for k in NAMES)
    assert round(e * 100, 4) == D(s["equity_index_return_pct"]) and round(b * 100, 4) == D(s["bond_index_return_pct"])
    assert round(D("0.42") * e * 100 + D("0.28") * b * 100 + D("0.30") * (e * 100 + D("0.25")), 4) == D(s["blend_v1_return_pct"])

    print("ok", FN.name)

if __name__ == "__main__":
    build(); verify()
