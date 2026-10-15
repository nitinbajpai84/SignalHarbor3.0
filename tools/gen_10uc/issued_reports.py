"""Generate issued monthly investment report packs (Jul, Aug, Sep 2026) from the universe.
Deterministic: no randomness; timestamps pinned. Run: python tools/gen_10uc/issued_reports.py"""
import datetime as dt
import json
from decimal import Decimal as D
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas

ROOT = Path(__file__).resolve().parents[2]
U = json.load(open(ROOT / "docs/test/10uc_universe.json", encoding="utf-8"))
OUT = ROOT / "tests/v1/fixtures/10uc/issued_reports"
FIXED = dt.datetime(2026, 1, 1)
# Issue day per period. Sep is fixed by the universe (report_issued); Jul and Aug use the same
# 10:00 cutoff / 10:30 issue pattern on the day the administrator NAV arrived (Jul) or the 3rd (Aug).
ISSUE = {"2026-07-31": "2026-08-04", "2026-08-31": "2026-09-03", "2026-09-30": "2026-10-03"}
FUND = {f["key"]: f["canonical_name"] for f in U["funds"]}
SEC = {s["key"]: s["canonical_name"] for s in U["securities"]}
ST = {s["statement_id"]: s for s in U["statements"]}
VAL = {(v["fund"], v["month_end"]): v for v in U["valuations_monthly"]}
ADM = {(r["portfolio"], r["nav_period_end"]): r for r in U["administrator_nav_monthly"]}
RI = U["report_issued"]


def pack(me):
    """Figures as issued: the administrator_value, i.e. what was known at the NAV before the cutoff."""
    day = ISSUE[me]
    lines = []
    for k in ("A", "B"):
        v = VAL[(k, me)]
        src = ST[v["administrator_from"]]
        if src["period_end"] != me:
            lines.append((FUND[k], D(v["administrator_value"]), "carried",
                          f"Carried forward from the {src['period_end']} statement; the {me} statement was not yet received at cutoff"))
        else:
            lines.append((FUND[k], D(v["administrator_value"]), src["status"],
                          f"{src['kind'].capitalize()} statement received {src['received_at'][:10]}"))
    lb = U["liquid_book_month_end"][me]
    lines.append(("Liquid book", D(lb["total"]), "official", "Administrator pricing, 20 securities"))
    total = sum(l[1] for l in lines)
    st = {l[2] for l in lines}
    tstat = "estimated" if "estimated" in st else "carried" if "carried" in st else "official"
    fx = D(U["fx_fixings"]["month_end"][me])
    p = dict(me=me, day=day, cutoff=f"{day}T10:00:00Z", issued=f"{day}T10:30:00Z", lines=lines, total=total,
             tstat=tstat, fx=fx, sgd=(total * fx).quantize(D("0.01")), pos=lb["positions"],
             nav_g=D(ADM[("EG-GROWTH", me)]["nav_amount"]), nav_l=D(ADM[("EG-LIQUID", me)]["nav_amount"]),
             nav_rx=ADM[("EG-GROWTH", me)]["received_at"])
    if me == RI["period_end"]:  # cross-check against the universe's own issued report
        assert [f"{l[1]:.2f}" for l in lines] == [x["value"] for x in RI["lines"][:3]]
        assert f"{total:.2f}" == RI["lines"][3]["value"] and f"{p['sgd']:.2f}" == RI["total_sgd"]; assert str(fx) == RI["fx_rate_used"]
        assert p["cutoff"] == RI["cutoff_at"] and p["issued"] == RI["issued_at"]
    return p


def stem(p): return f"Evergreen_Office_Investment_Report_{p['me'][:7]}"
def fmt(x): return f"{x:,.2f}"
def long_date(me): return f"{int(me[8:])} {dt.date.fromisoformat(me):%B %Y}"


def write_xlsx(p):
    wb = Workbook()
    wb.properties.creator = "Evergreen Office - Operations"; wb.properties.created = FIXED; wb.properties.modified = FIXED
    hdr = PatternFill("solid", fgColor="1F3A2E"); bold = Font(bold=True); white = Font(bold=True, color="FFFFFF")
    ws = wb.active; ws.title = "Summary"
    ws["A1"] = "Evergreen Office"; ws["A1"].font = Font(bold=True, size=16)
    ws["A2"] = f"Monthly Investment Report - {long_date(p['me'])}"; ws["A2"].font = Font(bold=True, size=12)
    meta = [("Report code", f"investment_report_{p['me'][:7]}"), ("Period end", p["me"]), ("Data cutoff (UTC)", p["cutoff"]),
            ("Issued (UTC)", p["issued"]), ("Prepared by", "Mei Chen, Operations analyst"), ("Approved by", RI["approved_by"]),
            ("Scope", "EG-GROWTH, EG-LIQUID"), ("Reporting currencies", "USD, SGD"),
            ("USD/SGD fixing (WMR 4pm)", float(p["fx"])), ("Version", "As issued - not restated")]
    for i, (k, v) in enumerate(meta, 4):
        ws.cell(i, 1, k).font = bold; ws.cell(i, 2, v)
    r = 4 + len(meta) + 1
    for c, h in enumerate(["Line", "Value (USD)", "Status", "Basis"], 1):
        x = ws.cell(r, c, h); x.font = white; x.fill = hdr
    for i, l in enumerate(p["lines"] + [("Total", p["total"], p["tstat"], "Sum of the lines above")], r + 1):
        ws.cell(i, 1, l[0]); ws.cell(i, 2, float(l[1])).number_format = "#,##0.00"; ws.cell(i, 3, l[2]); ws.cell(i, 4, l[3])
        if l[0] == "Total":
            for c in range(1, 5): ws.cell(i, c).font = bold
    i += 2
    ws.cell(i, 1, "Total (SGD)").font = bold; ws.cell(i, 2, float(p["sgd"])).number_format = "#,##0.00"
    ws.cell(i, 4, f"Total USD x {p['fx']}")
    ws.cell(i + 2, 1, "Figures are as known at the data cutoff above. Estimated and carried figures are marked in the Status column.").font = Font(italic=True)
    ws.cell(i + 3, 1, "Evergreen Office - synthetic data for testing only. Page 1 of 1").font = Font(italic=True, size=8)
    for col, w in zip("ABCD", (28, 20, 14, 90)): ws.column_dimensions[col].width = w
    h = wb.create_sheet("Liquid book")
    h["A1"] = f"Evergreen Liquid Book - positions at {p['me']}"; h["A1"].font = Font(bold=True, size=12)
    for c, t in enumerate(["Security", "Name", "Quantity", "Price", "Value (USD)"], 1):
        x = h.cell(3, c, t); x.font = white; x.fill = hdr
    for i, q in enumerate(p["pos"], 4):
        h.cell(i, 1, q["security"]); h.cell(i, 2, SEC[q["security"]])
        h.cell(i, 3, int(q["quantity"])).number_format = "#,##0"
        h.cell(i, 4, float(q["price"])).number_format = "0.0000" if D(q["price"]) < 2 else "0.00"
        h.cell(i, 5, float(q["value"])).number_format = "#,##0.00"
    t = 4 + len(p["pos"]); h.cell(t, 1, "Total").font = bold
    h.cell(t, 5, f"=SUM(E4:E{t-1})").number_format = "#,##0.00"; h.cell(t, 5).font = bold
    h.cell(t + 2, 1, "Bond prices per 1.00 nominal; convertible per unit; equities per share.").font = Font(italic=True, size=9)
    for col, w in zip("ABCDE", (12, 30, 14, 10, 18)): h.column_dimensions[col].width = w
    n = wb.create_sheet("Notes")
    for i, s in enumerate([
            f"Administrator NAV used: EG-GROWTH {fmt(p['nav_g'])}, EG-LIQUID {fmt(p['nav_l'])} (received {p['nav_rx'][:10]} 06:00 UTC).",
            "Fund lines use the value known at the cutoff; a manager estimate is permitted in the investment report only.",
            "Cash is held at entity level and is not included in portfolio NAV."], 1):
        n.cell(i, 1, s)
    n.column_dimensions["A"].width = 120
    wb.save(OUT / (stem(p) + ".xlsx"))


def write_pdf(p):
    c = canvas.Canvas(str(OUT / (stem(p) + ".pdf")), pagesize=A4, invariant=1)
    c.setTitle(f"Evergreen Office Investment Report {p['me'][:7]}"); c.setAuthor("Evergreen Office")
    W, H = A4; y = H - 60
    c.setFillColorRGB(.12, .23, .18); c.rect(0, H - 40, W, 40, fill=1, stroke=0); c.setFillColorRGB(1, 1, 1)
    c.setFont("Helvetica-Bold", 14); c.drawString(40, H - 26, "Evergreen Office"); c.setFont("Helvetica", 9)
    c.drawRightString(W - 40, H - 26, "Monthly Investment Report - one page summary")
    c.setFillColorRGB(0, 0, 0); y -= 20; c.setFont("Helvetica-Bold", 16); c.drawString(40, y, f"Period ending {long_date(p['me'])}")
    y -= 22; c.setFont("Helvetica", 9)
    for t in (f"Data cutoff: {p['cutoff'][:10]} {p['cutoff'][11:16]} UTC    Issued: {p['issued'][:10]} {p['issued'][11:16]} UTC",
              f"Prepared by Mei Chen (Operations analyst)    Approved by {RI['approved_by']}",
              "Scope: EG-GROWTH, EG-LIQUID    Reporting currencies: USD, SGD"):
        c.drawString(40, y, t); y -= 13
    y -= 14; x0, xv, xs = 40, 380, 470
    c.setFillColorRGB(.12, .23, .18); c.rect(36, y - 4, W - 72, 16, fill=1, stroke=0); c.setFillColorRGB(1, 1, 1); c.setFont("Helvetica-Bold", 9)
    c.drawString(x0, y, "Line"); c.drawRightString(xv, y, "Value (USD)"); c.drawString(xs, y, "Status")
    c.setFillColorRGB(0, 0, 0); y -= 18
    for l in p["lines"]:
        c.setFont("Helvetica", 10); c.drawString(x0, y, l[0]); c.drawRightString(xv, y, fmt(l[1])); c.drawString(xs, y, l[2]); y -= 16
    c.line(36, y + 10, W - 36, y + 10); c.setFont("Helvetica-Bold", 10)
    c.drawString(x0, y, "Total (USD)"); c.drawRightString(xv, y, fmt(p["total"])); c.drawString(xs, y, p["tstat"]); y -= 16
    c.drawString(x0, y, f"Total (SGD) at {p['fx']}"); c.drawRightString(xv, y, fmt(p["sgd"])); y -= 26
    c.drawString(40, y, "Notes"); y -= 14; c.setFont("Helvetica", 9)
    notes = [f"- {l[0]}: {l[3]}." for l in p["lines"]] + [
        f"- USD/SGD converted at the WMR 4pm fixing of {p['me']} ({p['fx']}).",
        f"- Administrator NAV received {p['nav_rx'][:10]} 06:00 UTC: EG-GROWTH {fmt(p['nav_g'])}; EG-LIQUID {fmt(p['nav_l'])}.",
        "- Figures are as known at the data cutoff. Estimated and carried values are marked in the Status column."]
    for t in notes:
        c.drawString(40, y, t); y -= 12
    c.setFont("Helvetica", 7.5); c.line(36, 50, W - 36, 50)
    c.drawString(40, 38, "Evergreen Office - synthetic data for testing only. Internal use; not investment advice.")
    c.drawRightString(W - 40, 38, "Page 1 of 1")
    c.drawString(40, 28, f"Issued version of investment_report_{p['me'][:7]}; the full pack is the accompanying workbook.")
    c.showPage(); c.save()


def verify(p):
    ws = load_workbook(OUT / (stem(p) + ".xlsx"))["Summary"]
    rows = {r[0].value: r for r in ws.iter_rows(min_row=1, max_row=40) if r[0].value}
    for l in p["lines"]:
        assert D(str(rows[l[0]][1].value)) == l[1] and rows[l[0]][2].value == l[2], l
    assert D(str(rows["Total"][1].value)) == p["total"] and D(str(rows["Total (SGD)"][1].value)) == p["sgd"]
    hs = load_workbook(OUT / (stem(p) + ".xlsx"))["Liquid book"]
    assert sum(D(str(hs.cell(i, 5).value)) for i in range(4, 24)) == p["lines"][2][1]
    for i, q in enumerate(p["pos"], 4):
        assert D(str(hs.cell(i, 3).value)) * D(str(hs.cell(i, 4).value)) == D(q["value"]), q
    pdf = open(OUT / (stem(p) + ".pdf"), "rb").read()
    assert pdf.startswith(b"%PDF") and pdf.count(b"/Type /Page\n") == 1, "one page"
    if p["me"] == RI["period_end"]:  # as issued: no trace of the amended 10,180,000
        assert all(l[1] != D("10180000") for l in p["lines"])


FACTS = {
    "2026-07-31": "Fund A and B are administrator values carried from the 30 Jun statements (the Jul statements arrived 5 Aug, after the 4 Aug cutoff); known now for Jul is higher (A 9,442,491.00, B 6,008,189.00). Gives a second 'issued vs known now' difference beyond Sep.",
    "2026-08-31": "Control month: A and B final statements received 2 Sep, before the cutoff; issued equals known now.",
    "2026-09-30": "P01 (A 10,000,000 estimated, B 6,400,000, Liquid book 41,200,000, total 57,600,000, SGD 77,253,120.00 at 1.3412, cutoff 3 Oct 10:00, issued 10:30). Pre-revision baseline for P02 (A amended to 10,180,000 on 7 Oct) and P03. Carries no sign of the revision.",
}

if __name__ == "__main__":
    OUT.mkdir(parents=True, exist_ok=True)
    md = ["# Issued reports (synthetic)", "",
          "Generated by `tools/gen_10uc/issued_reports.py` from `docs/test/10uc_universe.json`. Deterministic. Evergreen Office, fictional. Each pack is one xlsx (Summary, Liquid book, Notes) and one one-page PDF summary, as issued at the cutoff.", "",
          "| File | Holds | Planted facts / role |", "|---|---|---|"]
    for me in ISSUE:
        p = pack(me); write_xlsx(p); write_pdf(p); verify(p)
        md.append(f"| {stem(p)}.xlsx | Summary lines, SGD total, 20 liquid book positions; total {fmt(p['total'])} USD, cutoff {p['cutoff']} | {FACTS[me]} |")
        md.append(f"| {stem(p)}.pdf | One-page summary of the same figures, footer with page number | same as the workbook |")
        print(stem(p), fmt(p["total"]), fmt(p["sgd"]))
    md += ["", "Issue dates: Jul 4 Aug 2026, Aug 3 Sep 2026, Sep 3 Oct 2026 (a Saturday, as in the universe). Only the Sep cutoff and issue times are fixed by the universe; Jul and Aug follow the same 10:00 / 10:30 UTC pattern."]
    (OUT / "MANIFEST.md").write_text("\n".join(md) + "\n", encoding="utf-8")
