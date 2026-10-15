"""Security master + alias map (xlsx) for the Evergreen Office 10UC universe. Deterministic; synthetic only."""
import json, pathlib
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment

ROOT = pathlib.Path(__file__).resolve().parents[2]
U = json.load(open(ROOT / "docs/test/10uc_universe.json"))
OUT = ROOT / "tests/v1/fixtures/10uc/security_master"
FN = "EG_security_master_20261008.xlsx"
HDR = ["Internal ID", "ISIN / synthetic code", "Instrument name", "Instrument type", "Asset class", "Currency",
       "Trading system code", "Administrator code", "Manager ID", "Manager", "Portfolio", "Quantity 30 Sep 2026", "Status"]


def isin(seed):  # synthetic ISIN: unassigned country code ZZ + 9 digits + Luhn check
    body = "ZZ" + f"{(sum(ord(c) * (i + 7) for i, c in enumerate(seed)) * 7919) % 10**9:09d}"
    digits = "".join(str(int(c, 36)) for c in body)
    tot = sum(sum(divmod(int(d) * (2 if i % 2 == 0 else 1), 10)) for i, d in enumerate(reversed(digits)))
    return body + str((10 - tot % 10) % 10)


def rows():
    mg = {m["id"]: m["name"] for m in U["managers"]}
    r = []
    for s in U["securities"]:
        i = s["identifiers"]
        r.append([i["internal"]["security_id"], isin(s["key"]), s["canonical_name"], s["instrument_type"],
                  s["asset_class"] or "", s["currency"], i["trading_export"]["ticker"],
                  i["administrator"]["administrator_code"], "", "", s["portfolio"], int(s["quantity_30_sep"]),
                  "Active" if s["asset_class"] else "Active - asset class missing"])
    for f in U["funds"]:
        i = f["identifiers"]
        if "administrator" not in i:
            continue
        r.append(["FND-" + f["key"], isin("FND" + f["key"]), f["canonical_name"], f["instrument_type"],
                  f["asset_class"], f["currency"], "", i["administrator"]["administrator_code"], f["manager"],
                  mg[f["manager"]], f["portfolio"], int(f["units"]) if f["units"] else "", "Active"])
    return r


def aliases():
    a = []
    for s in U["securities"]:
        a.append([s["canonical_name"], "SEC-" + s["key"], "Internal canonical name", "Master"])
        a.append([s["identifiers"]["trading_export"]["ticker"], "SEC-" + s["key"], "Trading system ticker", "Trading export"])
        a.append([s["identifiers"]["administrator"]["administrator_code"], "SEC-" + s["key"], "Administrator code", "Administrator"])
    for f in U["funds"]:
        i = f["identifiers"]
        if "administrator" not in i:
            continue
        a.append([f["canonical_name"], "FND-" + f["key"], "Internal canonical name", "Master"])
        a.append([i["administrator"]["administrator_code"], "FND-" + f["key"], "Administrator code", "Administrator"])
        for n in i["manager_email"]["manager_fund_name"]:
            if n != f["canonical_name"]:
                a.append([n, "FND-" + f["key"], "Manager naming variant", "Manager statement / email"])
    return a


def build():
    OUT.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active; ws.title = "Security master"
    bold = Font(bold=True, color="FFFFFF"); fill = PatternFill("solid", fgColor="2F4F4F")
    def sheet(ws, title, hdr, data, widths):
        ws["A1"] = "Evergreen Office - " + title; ws["A1"].font = Font(bold=True, size=14)
        ws["A2"] = "Maintained by Data team (Anita Rao, Leo Park). As of 8 Oct 2026. Fictional data for testing."
        for c, h in enumerate(hdr, 1):
            x = ws.cell(4, c, h); x.font = bold; x.fill = fill; x.alignment = Alignment(wrap_text=True)
        for r, row in enumerate(data, 5):
            for c, v in enumerate(row, 1):
                ws.cell(r, c, v)
        ws.cell(5 + len(data) + 1, 1, f"Rows: {len(data)}   |   Page 1 of 1   |   Confidential - synthetic test data")
        for c, w in enumerate(widths, 1):
            ws.column_dimensions[chr(64 + c)].width = w
        ws.freeze_panes = "A5"; ws.oddFooter.center.text = "Page &P of &N"
    sheet(ws, "Security master", HDR, rows(), [12, 16, 44, 16, 16, 9, 14, 14, 15, 24, 12, 16, 28])
    sheet(wb.create_sheet("Alias map"), "Alias map", ["Alias", "Internal ID", "Alias type", "Seen in"], aliases(), [44, 12, 26, 28])
    wb.save(OUT / FN)


def verify():
    wb = load_workbook(OUT / FN)
    m = [r for r in wb["Security master"].iter_rows(min_row=5, values_only=True) if r[0] and str(r[0]).startswith(("SEC-", "FND-"))]
    assert len(m) == 29, len(m)
    by = {r[0]: r for r in m}
    for s in U["securities"]:
        r = by["SEC-" + s["key"]]
        assert r[2] == s["canonical_name"] and r[11] == int(s["quantity_30_sep"]) and r[6] == s["key"] + ".SYN" or r[6] == s["identifiers"]["trading_export"]["ticker"]
        assert r[7] == s["identifiers"]["administrator"]["administrator_code"]
        assert (r[4] or None) == s["asset_class"]
    assert [r[0] for r in m if not r[4]] == ["SEC-PRC29"]
    for f in U["funds"]:
        if f["key"] != "LIQ":
            assert by["FND-" + f["key"]][7] == f["identifiers"]["administrator"]["administrator_code"]
    al = [r for r in wb["Alias map"].iter_rows(min_row=5, values_only=True) if r[1]]
    amap = {}
    for a, i, *_ in al: amap.setdefault(a, set()).add(i)
    assert amap["KR Macro Fund Ltd - Class B USD"] == {"FND-H1"} == amap["Kestrel Ridge Macro Fund"]
    assert amap["Fund B - Evergreen Liquid Book"] == {"FND-B"}
    assert len({r[1] for r in m}) == 29
    return len(m), len(al)


if __name__ == "__main__":
    build(); print(verify())
