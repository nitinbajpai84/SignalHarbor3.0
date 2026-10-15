"""E0d: write ~20 messy synthetic workbooks for the fictional "Evergreen Office" plus an expected-mapping JSON for each.
Deterministic (fixed seed, fixed document properties). Usage: python tools/make_e0d_workbooks.py [out_dir]
All names and figures are invented."""
import datetime as dt
import json
import random
import sys
from pathlib import Path

from openpyxl import Workbook

ROOT = Path(__file__).resolve().parents[1]
FIX = ROOT / "tests" / "v1" / "fixtures" / "e0d"
FIXED = dt.datetime(2026, 10, 1, 9, 0, 0)
TRAPS = ["same_fund_two_names", "thousands", "day_month_swap", "merged_header", "totals_row", "renamed_column",
         "in_progress", "totals_mismatch"]

# (code, canonical name, alias used by another sender, portfolio)
FUNDS = [("EVG-F01", "Cedar Ridge Growth Fund LP", "Cedar Ridge Growth Fd", "EVG-GROWTH"),
         ("EVG-F02", "Northgate Credit Opportunities", "Northgate Credit Opps", "EVG-CORE"),
         ("EVG-F03", "Halden Real Assets Partners", "Halden RA Partners", "EVG-CORE"),
         ("EVG-F04", "Marlow Global Equity Fund", "Marlow Global Equity", "EVG-GROWTH"),
         ("EVG-F05", "Pinecrest Infrastructure Fund II", "Pinecrest Infra Fund 2", "EVG-LEGACY"),
         ("EVG-F06", "Ashby Venture Partners III", "Ashby Ventures III", "EVG-LEGACY")]
MONTHS = {"2026-08": dt.date(2026, 8, 31), "2026-09": dt.date(2026, 9, 30)}


def build(out_dir: Path):
    rng = random.Random(20260930)
    fig = {}
    for code, *_ in FUNDS:  # two consistent months per fund
        units = rng.randrange(20, 400) * 1000
        aug = round(units * rng.uniform(8, 140))
        fig[code] = {"units": units, "2026-08": aug, "2026-09": round(aug * (1 + rng.uniform(-0.03, 0.045)))}
    exp = {}

    def make(name, traps, domain, sheet, cols, rows, *, title=None, band=None, total=None, stated_delta=0, in_progress=False,
             sum_fields=(), note=None):
        """cols: [(header, target, transform|None)]; rows: raw cell lists. total: label for a totals row.
        band: [(first_col, last_col, text)] merged row above the header."""
        wb = Workbook()
        wb.properties.creator = "Evergreen Office (synthetic)"
        wb.properties.created = wb.properties.modified = FIXED
        ws = wb.active
        ws.title = sheet
        r = 1
        for t in (title or []):
            ws.cell(r, 1, t)
            r += 1
        if title:
            r += 1
        if band:
            for a, b, text in band:
                ws.cell(r, a, text)
                ws.merge_cells(start_row=r, start_column=a, end_row=r, end_column=b)
            r += 1
        header_row = r
        for c, (h, _, _) in enumerate(cols, 1):
            ws.cell(r, c, h)
        first = r + 1
        for row in rows:
            r += 1
            for c, v in enumerate(row, 1):
                ws.cell(r, c, v)
        totals_row = None
        ctl = {"row_count": len(rows), "sums": {}}
        idx = lambda f: next(i for i, (_, t, _) in enumerate(cols) if t == f)
        for f in sum_fields:
            mult = (cols[idx(f)][2] or {}).get("multiply", 1)
            ctl["sums"][f] = round(sum(x[idx(f)] for x in rows) * mult, 2)
        if total:
            r += 1
            totals_row = r
            ws.cell(r, 1, total)
            for k, f in enumerate(sum_fields):
                ws.cell(r, idx(f) + 1, sum(x[idx(f)] for x in rows) + (stated_delta if k == 0 else 0))
            ctl["stated_totals_row"] = totals_row
            ctl["stated_matches_rows"] = stated_delta == 0
            if stated_delta:
                ctl["stated_difference"] = {sum_fields[0]: stated_delta}
        if note:
            ws.cell(r + 2, 1, note)
        wb.save(out_dir / "workbooks" / f"{name}.xlsx")
        exp[name] = {
            "workbook": f"{name}.xlsx", "domain": domain, "traps": traps or ["clean"], "sheet": sheet,
            "header_row": header_row, "first_data_row": first, "last_data_row": first + len(rows) - 1,
            "exclude_rows": [totals_row] if totals_row else [],
            "status": "in_progress" if in_progress else "final",
            "columns": {h: ({"target": t, **({"transform": tf} if tf else {})}) for h, t, tf in cols},
            "control_totals": ctl}

    iso = lambda m: MONTHS[m]
    nm = {c: n for c, n, *_ in FUNDS}
    al = {c: a for c, _, a, _ in FUNDS}
    pf = {c: p for c, *_, p in FUNDS}
    P, V, C, N, B, M = ("position_snapshots.", "valuation_observations.", "cash_movements.", "nav_snapshots.",
                        "benchmark_components.", "market_observations.")
    FN = "instruments.canonical_name"

    def pos_rows(m, names=nm, scale=1):
        return [[iso(m), names[c], c, pf[c], fig[c]["units"], fig[c][m] if scale == 1 else round(fig[c][m] / scale), "USD"]
                for c, *_ in FUNDS]

    pos_cols = lambda mv, tf=None: [("Position Date", P + "position_date", None), ("Fund Name", FN, None),
        ("Fund ID", "instrument_identifiers.identifier_value", None), ("Portfolio", "portfolios.portfolio_code", None),
        ("Units", P + "quantity", None), (mv, P + "source_market_value", tf), ("Currency", P + "currency", None)]
    T = lambda s: ["Custodian statement (synthetic)", s]
    SM = [P + "source_market_value"]
    make("pos_2026_08_custodian", [], "positions", "Positions", pos_cols("Market Value"), pos_rows("2026-08"),
         title=T("As at 31 Aug 2026"), sum_fields=SM)
    make("pos_2026_09_custodian", ["renamed_column"], "positions", "Positions", pos_cols("Mkt Value USD"), pos_rows("2026-09"),
         title=T("As at 30 Sep 2026"), sum_fields=SM, note="Column 'Market Value' was renamed 'Mkt Value USD' from August.")
    make("pos_2026_09_altnames", ["same_fund_two_names"], "positions", "Positions", pos_cols("Market Value"),
         pos_rows("2026-09", al), title=T("As at 30 Sep 2026 (administrator names)"), sum_fields=SM)
    make("pos_2026_09_thousands", ["thousands"], "positions", "Positions", pos_cols("Market Value (USD 000)", {"multiply": 1000}),
         pos_rows("2026-09", nm, 1000), title=T("As at 30 Sep 2026 - values in USD thousands"), sum_fields=SM)
    make("pos_2026_09_totals", ["totals_row"], "positions", "Positions", pos_cols("Market Value"), pos_rows("2026-09"),
         total="Total", sum_fields=SM, title=T("As at 30 Sep 2026"))
    make("pos_2026_09_merged", ["merged_header"], "positions", "Positions", pos_cols("Market Value"), pos_rows("2026-09"),
         band=[(1, 4, "Holding"), (5, 7, "Valuation")], title=T("As at 30 Sep 2026"), sum_fields=SM)

    # valuations: text dates. The sheet holds a day above 12, so the order is resolvable.
    vcols = lambda fmt: [("Valuation Date", V + "valuation_date", {"date_format": fmt} if fmt else None), (  # noqa
        "Fund", FN, None), ("Estimated Value", V + "value_amount", None), ("Ccy", V + "currency", None), ("Type", V + "valuation_type", None)]
    vrows = lambda m, fmt, names=nm, typ="manager_estimate": [
        [iso(m) if fmt is None else iso(m).strftime(fmt), names[c], fig[c][m], "USD", typ] for c, *_ in FUNDS]
    VA = [V + "value_amount"]
    make("val_mgr_estimate_dmy", ["day_month_swap"], "valuations", "Estimates", vcols("%d/%m/%Y"), vrows("2026-09", "%d/%m/%Y"),
         title=["Manager estimates (synthetic)"], sum_fields=VA)
    make("val_mgr_estimate_mdy", ["day_month_swap"], "valuations", "Estimates", vcols("%m/%d/%Y"), vrows("2026-08", "%m/%d/%Y"),
         title=["Manager estimates, US style (synthetic)"], sum_fields=VA)
    make("val_mgr_final_aug", [], "valuations", "Final", vcols(None), vrows("2026-08", None, nm, "manager_final"), sum_fields=VA)
    make("val_admin_in_progress", ["in_progress", "same_fund_two_names"], "valuations", "Valuations - In Progress", vcols(None),
         vrows("2026-09", None, al, "administrator"), title=["DRAFT - still being reviewed by the administrator"],
         in_progress=True, sum_fields=VA)

    # cash movements
    cc = lambda amt: [("Date", C + "effective_date", None), ("Fund", FN, None), ("Portfolio", "portfolios.portfolio_code", None),
        ("Type", C + "movement_type", None), (amt, C + "amount", {"multiply": 1000} if "000" in amt else None),
        ("Currency", C + "currency", None)]
    kinds = ["call", "distribution", "call", "redemption", "distribution", "call"]
    flows = lambda mo, k: [[dt.date(2026, mo, 5 + i * 3), nm[c], pf[c], kinds[i], rng.randrange(50, 900) * k, "USD"]
                           for i, (c, *_) in enumerate(FUNDS)]
    CA = [C + "amount"]
    make("cash_2026_08", [], "cash", "Cash", cc("Amount"), flows(8, 1000), sum_fields=CA)
    make("cash_2026_09_thousands", ["thousands", "totals_row"], "cash", "Cash", cc("Amount (USD 000)"), flows(9, 1),
         total="Total", sum_fields=CA, title=["Cash movements - USD thousands"])
    make("cash_2026_09_bad_total", ["totals_mismatch", "totals_row"], "cash", "Cash", cc("Amount"), flows(9, 1000),
         total="Total", stated_delta=25000, sum_fields=CA, title=["Cash movements (stated total is wrong)"])

    # NAV
    ncols = lambda a: [("Period End", N + "nav_period_end", None), ("Portfolio", "portfolios.portfolio_code", None),
        (a, N + "nav_amount", None), ("Currency", N + "currency", None), ("Status", N + "nav_status", None)]
    navs = lambda m, s: [[iso(m), p, sum(fig[c][m] for c, *_ in FUNDS if pf[c] == p), "USD", s]
                         for p in ("EVG-CORE", "EVG-GROWTH", "EVG-LEGACY")]
    NA = [N + "nav_amount"]
    make("nav_admin_2026_08", [], "nav", "NAV", ncols("NAV"), navs("2026-08", "approved_official"), sum_fields=NA)
    make("nav_admin_2026_09", ["renamed_column"], "nav", "NAV", ncols("Net Asset Value"), navs("2026-09", "approved_official"),
         sum_fields=NA, note="'NAV' became 'Net Asset Value' from September.")
    make("nav_draft_in_progress", ["in_progress"], "nav", "NAV (In Progress)", ncols("NAV"), navs("2026-09", "draft"),
         in_progress=True, title=["In Progress - not for use"], sum_fields=NA)

    # capital calls: cash movements of type call
    kc = [("Notice Date", C + "effective_date", None), ("Fund", FN, None), ("Call Amount", C + "amount", None),
          ("Ccy", C + "currency", None), ("Due", "ignore", None)]
    kr = lambda: [[dt.date(2026, 9, 1 + i), nm[c], rng.randrange(1, 9) * 250000, "USD", dt.date(2026, 9, 15 + i)]
                  for i, (c, *_) in enumerate(FUNDS[:4])]
    make("calls_2026_q3", [], "capital_calls", "Calls", kc, kr(), sum_fields=CA)
    make("calls_2026_q3_merged", ["merged_header", "totals_row"], "capital_calls", "Calls", kc, kr(),
         band=[(1, 3, "Notice"), (4, 5, "Payment")], total="Total", sum_fields=CA)

    # benchmark and market series
    bc = [("Index", B + "market_series_key", None), ("Weight %", B + "component_weight", {"multiply": 0.01}),
          ("Spread (bps)", B + "spread_bps", None), ("Effective From", B + "weight_effective_from", None)]
    make("benchmark_blend_2026", ["totals_row"], "benchmark", "Benchmark", bc,
         [["GLOBAL-EQ-TR", 60, 0, dt.date(2026, 7, 1)], ["AGG-BOND-TR", 30, 0, dt.date(2026, 7, 1)],
          ["CASH-PLUS", 10, 300, dt.date(2026, 7, 1)]], total="Total", sum_fields=[B + "component_weight"],
         title=["Policy benchmark 2026H2"])
    make("benchmark_in_progress", ["in_progress"], "benchmark", "Benchmark - In Progress", bc,
         [["GLOBAL-EQ-TR", 70, 0, dt.date(2026, 10, 1)], ["AGG-BOND-TR", 30, 0, dt.date(2026, 10, 1)]], in_progress=True,
         title=["Proposed weights - In Progress"], sum_fields=[B + "component_weight"])
    make("market_index_levels_sep", [], "market", "Levels",
         [("Obs Date", M + "observed_at", None), ("Series", M + "series_key", None), ("Level", M + "value", None),
          ("Kind", M + "observation_type", None), ("Base Ccy", M + "base_currency", None)],
         [[dt.date(2026, 9, 1 + i), "GLOBAL-EQ-TR", round(1000 * (1 + 0.004 * i), 2), "index_level", "USD"] for i in range(5)])
    for name, e in exp.items():
        (out_dir / "expected" / f"{name}.json").write_text(json.dumps(e, indent=2, sort_keys=True) + "\n", encoding="utf-8")
    return sorted(exp)


def main(out: Path = FIX):
    for sub in ("workbooks", "expected"):
        (out / sub).mkdir(parents=True, exist_ok=True)
    print(f"wrote {len(build(out))} workbooks to {out / 'workbooks'}")


if __name__ == "__main__":
    main(Path(sys.argv[1]) if len(sys.argv) > 1 else FIX)
