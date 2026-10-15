"""Rebuild the UC8 synthetic source files: trades against the administrator (liquid and OTC).

Synthetic only: Evergreen Office, Evergreen Family Trust (EG-TRUST), portfolio EG-LIQUID.
Writes docs/demo/uc8/sources/<YYYY-MM>/ and prints the expected reconciliation per month
(the same figures are written into docs/demo/uc8/ANSWER-KEY.md).

Usage: python tools/make_uc8_data.py [out_dir]
"""
import calendar
import csv
import datetime as dt
import sys
from decimal import Decimal as D, ROUND_HALF_UP
from pathlib import Path

import openpyxl
from openpyxl.styles import Font
from reportlab.pdfgen import canvas

ROOT = Path(__file__).resolve().parents[1]
OUT = Path(sys.argv[1]) if len(sys.argv) > 1 else ROOT / "docs" / "demo" / "uc8" / "sources"
ALLOWANCE = D("5000.00")
BROKER = "Synthetic Broker One"
BANK = "Harbour Bank (synthetic)"
ADMIN = "Northcott Fund Administration (synthetic)"
ENTITY = "Evergreen Family Trust"
PORTFOLIO = "EG-LIQUID"

# key: (security name, instrument type, administrator code, trading ticker, counterparty)
INSTR = {
    "MRF": ("Meridian Foods", "listed_equity", "5004", "MRF.SYN", BROKER),
    "NWL": ("Northwind Logistics", "listed_equity", "5002", "NWL.SYN", BROKER),
    "BSU": ("Bay Street Utilities", "listed_equity", "5003", "BSU.SYN", BROKER),
    "CRH": ("Coastal Rail Holdings", "listed_equity", "5005", "CRH.SYN", BROKER),
    "HVS": ("Harborview Semiconductor", "listed_equity", "5001", "HVS.SYN", BROKER),
    "CRB31": ("Coastal Rail 5.1% 2031 bond", "bond", "5017", "CRB31.SYN", BROKER),
}
# month-end USD/SGD spot used to express SGD differences in USD (universe: 30 Sep spot 1.3412)
SPOT = {"2026-07": None, "2026-08": None, "2026-09": D("1.3412"), "2026-10": D("1.3388")}
FX_VALUE = {"FXF-2026-001": "2026-12-15", "FXF-2026-002": "2027-01-05"}


def money(x, dp=2):
    return str(D(x).quantize(D(1).scaleb(-dp), ROUND_HALF_UP))


def add_bd(day, n):
    while n:
        day += dt.timedelta(days=1)
        if day.weekday() < 5:
            n -= 1
    return day


def isin(root, code):
    """Synthetic ISIN: country QZ (ISO user-assigned), 9-char code, Luhn check digit."""
    body = "QZ" + (root + "000000")[:6] + f"{int(code) % 1000:03d}"
    digits = "".join(str(int(c, 36)) for c in body)
    total = 0
    for i, ch in enumerate(reversed(digits)):
        d = int(ch)
        if i % 2 == 0:
            d = d * 2 - 9 if d * 2 > 9 else d * 2
        total += d
    return body + str((10 - total % 10) % 10)


def fx(sell_amt, rate, buy_amt, value):
    return dict(sell_ccy="USD", sell_amt=sell_amt, buy_ccy="SGD", buy_amt=buy_amt, rate=rate, value=value)


def trade(tid, td, key, side, qty, price, gross=None, ccy="USD", adm="same", fxd=None):
    """One trading-system trade. adm: 'same' (administrator books it as we do), None (not booked),
    or a dict with overrides for td, qty, price, gross."""
    gross = gross or money(D(qty) * D(price))
    return dict(id=tid, td=td, key=key, side=side, qty=qty, price=price, gross=gross, ccy=ccy, fx=fxd,
                adm=dict(td=td, qty=qty, price=price, gross=gross) if adm == "same" else adm)


MONTHS = {
    "2026-07": dict(trades=[
        trade("T-20260706-001", "2026-07-06", "MRF", "buy", "15000", "31.20"),
        trade("T-20260714-002", "2026-07-14", "BSU", "sell", "5000", "45.60"),
        trade("T-20260721-003", "2026-07-21", "CRB31", "buy", "500000", "0.9915"),
        trade("T-20260728-004", "2026-07-28", "CRH", "buy", "3000", "72.10"),
    ], amendments=[]),
    "2026-08": dict(trades=[
        trade("T-20260804-001", "2026-08-04", "MRF", "buy", "8000", "31.00"),
        trade("T-20260811-002", "2026-08-11", "NWL", "buy", "4000", "38.90"),
        trade("T-20260818-003", "2026-08-18", "CRH", "buy", "2000", "73.40"),
        # administrator booked it one business day late: same quantity, price and amount
        trade("T-20260827-004", "2026-08-27", "BSU", "sell", "8000", "46.83",
              adm=dict(td="2026-08-28", qty="8000", price="46.83", gross="374640.00")),
    ], amendments=[]),
    "2026-09": dict(trades=[
        trade("T-20260903-001", "2026-09-03", "MRF", "buy", "10000", "32.10"),
        trade("T-20260910-002", "2026-09-10", "BSU", "sell", "6000", "47.30"),
        trade("T-20260915-F01", "2026-09-15", "FXF-2026-001", "sell USD / buy SGD", "2000000.00", "1.3400",
              gross="2680000.00", ccy="SGD", fxd=fx("2000000.00", "1.3400", "2680000.00", FX_VALUE["FXF-2026-001"])),
        # administrator books 10,200 instead of 12,000
        trade("T-20260917-004", "2026-09-17", "NWL", "buy", "12000", "42.15",
              adm=dict(td="2026-09-17", qty="10200", price="42.15", gross="429930.00")),
        trade("T-20260922-005", "2026-09-22", "CRH", "buy", "4000", "75.60"),
        # FX forward: administrator booked the confirmation rate 1.3452, trading export 1.3425
        trade("T-20260929-F02", "2026-09-29", "FXF-2026-002", "sell USD / buy SGD", "3000000.00", "1.3425",
              gross="4027500.00", ccy="SGD", fxd=fx("3000000.00", "1.3425", "4027500.00", FX_VALUE["FXF-2026-002"]),
              adm=dict(td="2026-09-29", qty="3000000.00", price="1.3452", gross="4035600.00")),
        trade("T-20260929-007", "2026-09-29", "CRB31", "buy", "500000", "0.9913"),
        # sell with no administrator booking
        trade("T-20260930-008", "2026-09-30", "HVS", "sell", "5000", "118.40", adm=None),
    ], amendments=[]),
    "2026-10": dict(trades=[
        trade("T-20261001-009", "2026-10-01", "MRF", "buy", "20000", "33.10"),
        trade("T-20261002-010", "2026-10-02", "BSU", "sell", "10000", "48.20"),
        trade("T-20261007-011", "2026-10-07", "CRB31", "buy", "1000000", "0.9910"),
    ], amendments=[
        # administrator corrected the September Northwind quantity to the trade (6 Oct)
        dict(ref="T-20260917-004", td="2026-09-17", qty="12000", price="42.15", gross="505800.00", ccy="USD",
             booked="2026-10-06", comment="Quantity corrected from 10,200 after trade check"),
        # Harbour Bank reissued the FX confirmation at the traded rate 1.3425
        dict(ref="T-20260929-F02", td="2026-09-29", qty="3000000.00", price="1.3425", gross="4027500.00", ccy="SGD",
             booked="2026-10-06", comment="Confirmation reissued by Harbour Bank at 1.3425"),
    ]),
}
# September trades still open at 30 Sep (the October cycle checks them against the amendments)
PRIOR_OPEN = {"2026-10": ["T-20260917-004", "T-20260929-F02", "T-20260930-008"]}


def business_days(month):
    y, m = map(int, month.split("-"))
    last = calendar.monthrange(y, m)[1]
    return [dt.date(y, m, d) for d in range(1, last + 1) if dt.date(y, m, d).weekday() < 5]


def trade_csv_rows(trades):
    for t in trades:
        if t["fx"]:
            f = t["fx"]
            yield [t["id"], t["td"], f["value"], PORTFOLIO, t["side"], t["key"], "", f"USD/SGD forward to {f['value']}",
                   "fx_forward", t["qty"], t["price"], t["gross"], t["ccy"], BANK, "executed", "OMS-SYN",
                   f["sell_ccy"], f["sell_amt"], f["buy_ccy"], f["buy_amt"], f["rate"], f["value"]]
        else:
            name, typ, code, ticker, cpty = INSTR[t["key"]]
            settle = add_bd(dt.date.fromisoformat(t["td"]), 2).isoformat()
            yield [t["id"], t["td"], settle, PORTFOLIO, t["side"], ticker, isin(t["key"], code), name, typ,
                   t["qty"], t["price"], t["gross"], "USD", cpty, "executed", "OMS-SYN"] + [""] * 6


CSV_HEAD = ["external_trade_id", "trade_date", "settle_date", "portfolio", "side", "instrument_code", "isin",
            "instrument_name", "instrument_type", "quantity", "price", "amount", "currency", "counterparty",
            "status", "source_system", "fx_sell_currency", "fx_sell_amount", "fx_buy_currency", "fx_buy_amount",
            "fx_contract_rate", "fx_value_date"]


def write_trading_export(month_dir, month, trades):
    per_day = {}
    for t in trades:
        per_day.setdefault(t["td"], []).append(t)
    for day in business_days(month):
        path = month_dir / f"trading_export_{day:%Y%m%d}.csv"
        with path.open("w", newline="", encoding="utf-8") as f:
            w = csv.writer(f)
            w.writerow(CSV_HEAD)
            w.writerows(trade_csv_rows(per_day.get(day.isoformat(), [])))


def write_admin(month_dir, month, trades, amendments):
    y, m = map(int, month.split("-"))
    last = calendar.monthrange(y, m)[1]
    arrived = add_bd(dt.date(y, m, last), 2)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Trades"
    ws.append([ADMIN])
    ws.cell(1, 1).font = Font(bold=True, size=12)
    ws.append([f"Trade confirmation report: {ENTITY}, portfolio {PORTFOLIO}"])
    ws.append([f"Period: 1 {calendar.month_name[m]} {y} to {last} {calendar.month_name[m]} {y}"])
    ws.append([f"Run date: {arrived:%d %b %Y}"])
    ws.append([])
    head = ["Client Reference", "Row Type", "Trade Date", "Settle Date", "Portfolio", "Administrator Code",
            "Security", "Buy/Sell", "Quantity", "Price", "Gross Amount", "Currency", "Counterparty", "Status",
            "Booked On", "Comment"]
    ws.append(head)
    for c in ws[6]:
        c.font = Font(bold=True)
    usd_total, sgd_total, n = D(0), D(0), 0
    for t in trades:
        a = t["adm"]
        if a is None:
            continue
        if t["fx"]:
            name, code, side, cpty = f"USD/SGD forward to {t['fx']['value']}", "FX", "sell USD / buy SGD", BANK
            settle = t["fx"]["value"]
        else:
            name, _, code, _, cpty = INSTR[t["key"]]
            side = t["side"]
            settle = add_bd(dt.date.fromisoformat(a["td"]), 2).isoformat()
        if a["td"] != t["td"]:
            comment = "Booked one business day after trade date"
        elif t["fx"] and a["price"] != t["price"]:
            comment = "Rate per confirmation"
        elif a["qty"] != t["qty"]:
            comment = "Quantity per administrator register"
        else:
            comment = ""
        ws.append([t["id"], "Trade", a["td"], settle, PORTFOLIO, code, name, side, a["qty"], a["price"], a["gross"],
                   t["ccy"], cpty, "Booked", a["td"], comment])
        if t["ccy"] == "USD":
            usd_total += D(a["gross"])
        else:
            sgd_total += D(a["gross"])
        n += 1
    for am in amendments:
        ws.append([am["ref"], "Amendment", am["td"], "", PORTFOLIO, "", "", "", am["qty"], am["price"], am["gross"],
                   am["ccy"], "", "Amended", am["booked"], am["comment"]])
    ws.append([])
    ws.append(["Total gross amount, USD trades", "", "", "", "", "", "", "", "", "", money(usd_total)])
    ws.append(["Total gross amount, SGD trades", "", "", "", "", "", "", "", "", "", money(sgd_total)])
    ws.append(["Trade rows", "", "", "", "", "", "", "", "", "", n])
    for col, width in zip("ABCDEFGHIJKLMNOP", [18, 12, 12, 12, 12, 16, 30, 16, 14, 10, 16, 10, 26, 12, 12, 44]):
        ws.column_dimensions[col].width = width
    wb.save(month_dir / f"administrator_trade_confirmations_{month}.xlsx")


def write_confirmation(path, ref, lines):
    c = canvas.Canvas(str(path), pagesize=(595, 842))
    c.setFont("Helvetica-Bold", 13)
    c.drawString(60, 790, f"OTC FX forward confirmation: {ref} (synthetic)")
    c.setFont("Helvetica", 10.5)
    y = 760
    for label, value in lines:
        c.drawString(60, y, f"{label}: {value}")
        y -= 20
    c.drawString(60, y - 20, "Synthetic document for the Signal Harbor demo. Not a real confirmation.")
    c.save()


def write_confirmations(month_dir, month, trades, amendments):
    odir = month_dir / "otc_confirms"
    odir.mkdir(exist_ok=True)
    for t in trades:
        if not t["fx"]:
            continue
        f = t["fx"]
        a = t["adm"] or dict(price=f["rate"], gross=f["buy_amt"])
        lines = [("Confirmation reference", t["key"]), ("Counterparty", BANK), ("Client", ENTITY),
                 ("Portfolio", PORTFOLIO), ("Trade date", t["td"]), ("Value date", f["value"]),
                 ("Sell currency", f["sell_ccy"]), ("Sell amount", f["sell_amt"]),
                 ("Buy currency", f["buy_ccy"]), ("Buy amount", a["gross"]), ("Contract rate", a["price"])]
        write_confirmation(odir / f"{t['key']}_confirmation_{t['td']}.pdf", t["key"], lines)
    for am in amendments:
        if am["ref"] == "T-20260929-F02":
            lines = [("Confirmation reference", "FXF-2026-002"), ("Counterparty", BANK), ("Client", ENTITY),
                     ("Portfolio", PORTFOLIO), ("Trade date", am["td"]), ("Value date", FX_VALUE["FXF-2026-002"]),
                     ("Sell currency", "USD"), ("Sell amount", "3000000.00"), ("Buy currency", "SGD"),
                     ("Buy amount", am["gross"]), ("Contract rate", am["price"]), ("Status", "Reissued, replaces 1.3452")]
            write_confirmation(odir / f"FXF-2026-002_confirmation_amended_{am['booked']}.pdf", "FXF-2026-002", lines)


def reconcile(month, trades, amendments, prior_open):
    """Expected result the product should show for this cycle."""
    spot = SPOT[month]
    out = dict(compared=len(trades), matched=0, breaks=[], cleared=[], carried=[])

    def usd(amount, ccy):
        return D(amount) if ccy == "USD" else (D(amount) / spot).quantize(D("0.01"), ROUND_HALF_UP)

    for t in trades:
        a = t["adm"]
        if a is None:
            out["breaks"].append(dict(id=t["id"], type="missing", key=t["key"], side=t["side"],
                                      net=usd(t["gross"], t["ccy"]), detail=f"export {t['qty']} at {t['price']} = {t['gross']}; not at administrator"))
        elif a["qty"] != t["qty"]:
            net = abs(D(t["gross"]) - D(a["gross"]))
            out["breaks"].append(dict(id=t["id"], type="quantity", key=t["key"], side=t["side"], net=usd(net, t["ccy"]),
                                      detail=f"export {t['qty']} at {t['price']} = {t['gross']}; admin {a['qty']} = {a['gross']}"))
        elif a["price"] != t["price"]:
            net = abs(D(t["gross"]) - D(a["gross"]))
            out["breaks"].append(dict(id=t["id"], type="price", key=t["key"], side=t["side"], net=usd(net, t["ccy"]),
                                      detail=f"export rate {t['price']} = {t['gross']} SGD; admin rate {a['price']} = {a['gross']} SGD"))
        elif a["td"] != t["td"]:
            out["breaks"].append(dict(id=t["id"], type="timing", key=t["key"], side=t["side"], net=D("0.00"),
                                      detail=f"export {t['td']}, admin {a['td']}; same quantity, price and amount"))
        else:
            out["matched"] += 1
    for ref in prior_open:
        am = next((x for x in amendments if x["ref"] == ref), None)
        (out["carried"] if am is None else out["cleared"]).append(ref)
    return out


def main():
    for month, spec in MONTHS.items():
        d = OUT / month
        d.mkdir(parents=True, exist_ok=True)
        write_trading_export(d, month, spec["trades"])
        write_admin(d, month, spec["trades"], spec["amendments"])
        write_confirmations(d, month, spec["trades"], spec["amendments"])
        res = reconcile(month, spec["trades"], spec["amendments"], PRIOR_OPEN.get(month, []))
        total = sum((b["net"] for b in res["breaks"]), D(0))
        print(f"== {month}: compared {res['compared']}, matched {res['matched']}, breaks {len(res['breaks'])}, "
              f"net break USD {total}")
        for b in res["breaks"]:
            flag = "over" if b["net"] > ALLOWANCE else "within"
            print(f"   {b['id']} {b['key']} {b['type']} net USD {b['net']} ({flag} allowance) {b['detail']}")
        if res["cleared"] or res["carried"]:
            print(f"   prior items cleared: {res['cleared']}; carried forward: {res['carried']}")


if __name__ == "__main__":
    main()
